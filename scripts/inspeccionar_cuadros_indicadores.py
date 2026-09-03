"""Lista fórmulas y valores de los cuadros anuales de Indicadores."""
import sys
from openpyxl import load_workbook


archivo = sys.argv[1]
libro_f = load_workbook(archivo, data_only=False, read_only=True)
libro_v = load_workbook(archivo, data_only=True, read_only=True)
hoja_f = libro_f["Indicadores"]
hoja_v = libro_v["Indicadores"]

print("FORMULAS MENSUALES CON HISTORICO")
for fila in hoja_f.iter_rows(min_row=1, max_row=70, max_col=18):
    for celda in fila:
        valor = celda.value
        texto = getattr(valor, "text", valor)
        if "Historico!" in str(texto or ""):
            print(celda.coordinate, "valor=", hoja_v[celda.coordinate].value, "formula=", texto)

for row in hoja_f.iter_rows(min_row=1, max_row=250, max_col=120):
    for celda in row:
        if str(celda.value or "").strip().casefold() == "año":
            fila, columna = celda.row, celda.column
            encabezados = [hoja_f.cell(fila, c).value for c in range(columna, columna + 4)]
            if [str(x or "").strip().casefold() for x in encabezados[:3]] != ["año", "objetivo", "anual"]:
                continue
            print(f"TABLA {celda.coordinate}: {encabezados}")
            for f in range(fila + 1, fila + 10):
                valores = [hoja_v.cell(f, c).value for c in range(columna, columna + 4)]
                formulas = [hoja_f.cell(f, c).value for c in range(columna, columna + 4)]
                print(f"  fila {f}: valores={valores} formulas={formulas}")
