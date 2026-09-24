"""Concilia el KPI Total Facturado con la fórmula de la hoja RESUMEN."""

import argparse
import json
import os
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app, db
from app.models import Facturacion2026, HistorialCambio


def texto(valor):
    return " ".join(str(valor or "").strip().lower().split())


def grupo(valor):
    clave = texto(valor)
    if clave.startswith("personal"):
        return "personal"
    if clave.startswith("santander"):
        return "santander"
    if clave in ("banco credicoop", "credicoop"):
        return "credicoop"
    if clave.startswith("naturgy"):
        return "naturgy"
    return {"bbva seguros": "bbva", "galicia seguros": "galicia"}.get(clave, clave)


def numero(valor):
    # SUMIFS ignora importes almacenados como texto.
    return float(valor) if isinstance(valor, (int, float)) and valor is not None else 0.0


def mes(valor):
    if isinstance(valor, (date, datetime)):
        return valor.strftime("%Y-%m")
    return str(valor or "")[:7]


def leer_objetivos(archivo):
    hoja = load_workbook(archivo, data_only=True, read_only=True)["2026"]
    objetivos = defaultdict(lambda: [0.0] * 8)
    meses = set()
    for fila in hoja.iter_rows(min_row=3, max_col=31, values_only=True):
        periodo = mes(fila[9])
        if not periodo.startswith("2026-"):
            continue
        # Algunas filas de Next Gen tienen cero en Cuenta/Nom Diras y el
        # servicio válido en Cuenta.
        nombre = fila[4] if fila[4] not in (None, "", 0) else fila[6]
        clave = (periodo, grupo(nombre))
        if not clave[1]:
            continue
        meses.add(periodo)
        for indice, columna in enumerate((27, 28, 29, 19, 20)):
            objetivos[clave][indice] += numero(fila[columna])
        objetivos[clave][5] += numero(fila[21])
        objetivos[clave][6] += numero(fila[24])
        objetivos[clave][7] += numero(fila[25])
    return objetivos, meses


def main(archivo, aplicar=False):
    objetivos, meses_archivo = leer_objetivos(archivo)
    app = create_app()
    with app.app_context():
        registros = Facturacion2026.query.filter(
            Facturacion2026.mes.in_(sorted(meses_archivo))
        ).order_by(Facturacion2026.id).all()
        por_grupo = defaultdict(list)
        for registro in registros:
            por_grupo[(registro.mes, grupo(registro.cliente))].append(registro)

        faltantes = sorted(
            clave for clave, valores in objetivos.items()
            if clave not in por_grupo and abs(sum(valores)) > 0.0001
        )
        if faltantes:
            raise RuntimeError(f"No hay registros en la app para: {faltantes}")

        cambios = []
        for clave, deseado in sorted(objetivos.items()):
            if clave not in por_grupo:
                continue
            propios = por_grupo[clave]
            actual = [
                sum(float(r.facturado_horas_control or 0) for r in propios),
                sum(float(r.variable_productivo_control or 0) for r in propios),
                sum(float(r.penalizaciones_bonos_control or 0) for r in propios),
                sum(float(r.netx_gen or 0) for r in propios),
                sum(float(r.otros or 0) for r in propios),
            ]
            deltas = [deseado[i] - actual[i] for i in range(5)]
            total_actual = sum(float(r.total_dashboard or 0) for r in propios)
            delta_total = deseado[5] - total_actual
            objetivo_horas_actual = sum(float(r.objetivo_facturacion_horas or 0) for r in propios)
            objetivo_variable_actual = sum(float(r.variable_objetivo or 0) for r in propios)
            delta_objetivo_horas = deseado[6] - objetivo_horas_actual
            delta_objetivo_variable = deseado[7] - objetivo_variable_actual
            if not any(abs(delta) > 0.0001 for delta in (
                *deltas, delta_total, delta_objetivo_horas, delta_objetivo_variable
            )):
                continue
            registro = next((r for r in propios if not r.es_next_gen), propios[0])
            antes = {
                "id": registro.id,
                "clave": clave,
                "actual": actual,
                "objetivo": deseado,
                "total_actual": total_actual,
                "objetivo_horas_actual": objetivo_horas_actual,
                "objetivo_variable_actual": objetivo_variable_actual,
            }
            registro.control_facturado_horas = float(registro.facturado_horas_control or 0) + deltas[0]
            registro.control_variable_productivo = float(registro.variable_productivo_control or 0) + deltas[1]
            registro.control_penalizaciones_bonos = float(registro.penalizaciones_bonos_control or 0) + deltas[2]
            registro.netx_gen = float(registro.netx_gen or 0) + deltas[3]
            registro.otros = float(registro.otros or 0) + deltas[4]
            if abs(delta_total) > 0.0001:
                registro.total_facturado_manual = float(registro.total_dashboard or 0) + delta_total
            if abs(delta_objetivo_horas) > 0.0001:
                registro_objetivo = next((
                    r for r in propios
                    if float(r.horas_objetivo or 0) != 0 and not r.objetivo_separar_ajuste_vh
                ), None)
                if registro_objetivo is None:
                    raise RuntimeError(f"No se puede ajustar Objetivo Horas para {clave}")
                registro_objetivo.valor_hora_objetivo = (
                    float(registro_objetivo.valor_hora_objetivo_calculo or 0)
                    + delta_objetivo_horas / float(registro_objetivo.horas_objetivo)
                )
            if abs(delta_objetivo_variable) > 0.0001:
                registro.variable_objetivo = float(registro.variable_objetivo or 0) + delta_objetivo_variable
            cambios.append(antes)

        total_excel = sum(valores[5] for valores in objetivos.values())
        total_app = sum(float(r.total_dashboard or 0) for r in registros)
        resultado = {
            "archivo": str(archivo),
            "grupos_ajustados": len(cambios),
            "total_excel": round(total_excel, 2),
            "total_app": round(total_app, 2),
            "diferencia": round(total_app - total_excel, 2),
        }
        if aplicar:
            historial = HistorialCambio(
                accion="reprocesamiento",
                entidad="facturacion_2026",
                entidad_id="2026",
                resumen=f"Conciliación Total Facturado desde Excel: {len(cambios)} grupos",
                antes=json.dumps(cambios, ensure_ascii=False, default=str),
                despues=json.dumps(resultado, ensure_ascii=False, default=str),
            )
            db.session.add(historial)
            db.session.commit()
            resultado["history_id"] = historial.id
        else:
            db.session.rollback()
        print(resultado)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--archivo", required=True, type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    main(args.archivo, args.apply)
