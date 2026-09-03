"""Auditoría de solo lectura de fórmulas, referencias y errores cacheados del libro."""
from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


REFERENCIA_HOJA = re.compile(r"(?:'([^']+)'|([A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9 _.-]+))!")


def celdas_reales(hoja):
    return hoja._cells.values()  # evita recorrer las 1.048.576 filas formateadas


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path)
    args = parser.parse_args()
    formulas = load_workbook(args.archivo, data_only=False, read_only=False)
    valores = load_workbook(args.archivo, data_only=True, read_only=False)
    hojas = set(formulas.sheetnames)
    total_formulas = 0
    problemas = []
    print("hoja|formulas|errores_cache|patrones")
    for hoja in formulas.worksheets:
        celdas_formula = [celda for celda in celdas_reales(hoja) if celda.data_type == "f"]
        total_formulas += len(celdas_formula)
        errores = []
        patrones = Counter()
        for celda in celdas_formula:
            formula = str(celda.value or "")
            cache = valores[hoja.title][celda.coordinate].value
            if isinstance(cache, str) and cache.startswith("#"):
                errores.append((celda.coordinate, cache))
            if "#REF!" in formula.upper():
                problemas.append((hoja.title, celda.coordinate, "REFERENCIA_ROTA", formula))
            if "[" in formula and "]" in formula:
                problemas.append((hoja.title, celda.coordinate, "VINCULO_EXTERNO", formula))
            for coincidencia in REFERENCIA_HOJA.finditer(formula):
                referida = (coincidencia.group(1) or coincidencia.group(2) or "").strip()
                if referida and referida not in hojas:
                    problemas.append((hoja.title, celda.coordinate, "HOJA_INEXISTENTE", referida))
            patron = re.sub(r"(\$?[A-Z]{1,3}\$?)\d+", r"\1#", formula.upper())
            patrones[patron] += 1
        print(f"{hoja.title}|{len(celdas_formula)}|{len(errores)}|{len(patrones)}")
        for coordenada, error in errores[:20]:
            problemas.append((hoja.title, coordenada, "ERROR_CACHE", error))
        if len(errores) > 20:
            problemas.append((hoja.title, "-", "ERROR_CACHE_ADICIONALES", len(errores) - 20))
    print(f"TOTAL_FORMULAS|{total_formulas}")
    print(f"TOTAL_PROBLEMAS|{len(problemas)}")
    for problema in problemas:
        print("PROBLEMA|" + "|".join(map(str, problema)))


if __name__ == "__main__":
    main()
