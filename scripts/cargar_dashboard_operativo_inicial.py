"""Carga la hoja Dashboard de un libro indicado explícitamente."""

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import create_app, db
from app.models import DashboardOperativo
from app.routes import DASHBOARD_COLUMNAS


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path, help="Libro Excel fuente")
    parser.add_argument("--sheet", default="Dashboard")
    parser.add_argument("--apply", action="store_true", help="Confirma la escritura en DATABASE_URL")
    args = parser.parse_args()
    if not args.archivo.is_file():
        parser.error(f"No existe el archivo: {args.archivo}")

    workbook = openpyxl.load_workbook(args.archivo, data_only=True, read_only=True)
    if args.sheet not in workbook.sheetnames:
        parser.error(f"No existe la hoja {args.sheet!r}")
    app = create_app()
    with app.app_context():
        count = 0
        for row in workbook[args.sheet].iter_rows(min_row=3, values_only=True):
            values = list(row)
            if len(values) < 41:
                values += [None] * (41 - len(values))
            fecha = values[5]
            cliente = str(values[6] or "").strip()
            campania = str(values[7] or "").strip()
            if not isinstance(fecha, (datetime, date)) or not cliente or not campania:
                continue
            month = fecha.strftime("%Y-%m")
            record = DashboardOperativo.query.filter_by(
                mes=month, cliente=cliente, campania=campania
            ).first()
            if not record:
                record = DashboardOperativo(
                    year=fecha.year, mes=month, cliente=cliente, campania=campania,
                    bajas_manual=float(values[40] or 0),
                )
                db.session.add(record)
            data = {}
            for index, key in enumerate(DASHBOARD_COLUMNAS):
                value = values[index + 1]
                data[key] = value.isoformat() if isinstance(value, (datetime, date)) else value
            data.update({"mes": month, "cliente": cliente, "campania": campania})
            data.pop("agentes_baja", None)
            record.datos = json.dumps(data, ensure_ascii=False)
            count += 1
        if args.apply:
            db.session.commit()
            print(f"APLICADO: Dashboard operativo, {count} fila(s).")
        else:
            db.session.rollback()
            print(f"SIMULADO: Dashboard operativo, {count} fila(s). Use --apply para guardar.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
