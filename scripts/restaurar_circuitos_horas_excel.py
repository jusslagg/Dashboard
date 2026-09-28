"""Restablece qué filas dependen del calendario y cuáles de horas auxiliares."""
import argparse
from collections import defaultdict
from pathlib import Path
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import create_app, db
from app.models import ProyeccionMatriz
from app.routes import (aplicar_calculo_proyeccion, fechas_feriadas_activas,
                        normalizar_header, registrar_historial, snapshot_proyeccion)
from scripts.sincronizar_proyectados_excel import (
    YEAR, claves_horas_auxiliares, identidad_canonica, mes, numero, texto,
)


def restaurar(archivo):
    valores = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    formulas = openpyxl.load_workbook(archivo, data_only=False, read_only=True)
    claves_auxiliares = claves_horas_auxiliares(formulas)
    grupos = defaultdict(list)
    nombres = {}
    for fila in valores['Matriz'].iter_rows(min_row=3, values_only=True):
        cliente, campania = identidad_canonica(fila[2], fila[3])
        clave_mes = mes(fila[5])
        if not cliente or not campania or not clave_mes or not clave_mes.startswith(str(YEAR)):
            continue
        clave = (normalizar_header(cliente), normalizar_header(campania), clave_mes)
        grupos[clave].append(fila)
        nombres[clave] = (cliente, campania)

    registros = {
        (normalizar_header(item.cliente), normalizar_header(item.campania), item.mes): item
        for item in ProyeccionMatriz.query.filter_by(year=YEAR).all()
    }
    antes, despues, faltantes = [], [], []
    for clave in sorted(claves_auxiliares):
        registro = registros.get(clave)
        filas = grupos.get(clave)
        if not registro or not filas:
            faltantes.append(clave)
            continue
        total_horas = sum(numero(fila[10]) for fila in filas)
        total_proyectadas = sum(numero(fila[11]) for fila in filas)
        cumplimiento = total_proyectadas / total_horas * 100 if total_horas else 100
        total_dotacion = sum(numero(fila[9]) for fila in filas)
        carga_horaria = (
            sum(numero(fila[8]) * numero(fila[9]) for fila in filas) / total_dotacion
            if total_dotacion else numero(filas[0][8], 6)
        )
        cliente, campania = nombres[clave]
        antes.append(snapshot_proyeccion(registro))
        aplicar_calculo_proyeccion(registro, {
            'cliente': cliente, 'campania': campania,
            'horas_carga_manual': True,
            'horas_requeridas_manual': total_horas,
            'carga_semanal_plp': texto(filas[0][6]) or 'L a V',
            'carga_horaria_plp': carga_horaria or 6,
            'porcentaje_cumplimiento': cumplimiento,
            'tipo_plp': registro.tipo_plp,
        }, clave[2], fechas_feriadas_activas(YEAR))
        db.session.flush()
        despues.append(snapshot_proyeccion(registro))
    if despues:
        registrar_historial(
            'edicion', 'proyeccion_matriz', 'circuitos_excel',
            f'Circuitos de horas auxiliares restaurados: {len(despues)} proyecciones',
            antes={'proyecciones': antes}, despues={'proyecciones': despues},
            detalle='Las fórmulas de Matriz definen si el calendario modifica horas o sólo la dotación.',
        )
    return {'restauradas': len(despues), 'faltantes': len(faltantes)}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('archivo')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    app = create_app()
    with app.test_request_context('/restaurar-circuitos-horas-excel'):
        resultado = restaurar(args.archivo)
        db.session.commit() if args.apply else db.session.rollback()
        print('APLICADO' if args.apply else 'SIMULADO', resultado)
