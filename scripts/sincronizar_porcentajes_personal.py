"""Sincroniza porcentajes y penalidades de Personal desde el Excel fuente."""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from app import create_app, db  # noqa: E402
from app.models import (  # noqa: E402
    PersonalDistribucionHoras,
    ProyeccionMatriz,
    TarifacionCampania,
    VariableCampania,
)
from app.routes import registrar_historial  # noqa: E402
from app.routes import (  # noqa: E402
    aplicar_calculo_proyeccion,
    buscar_proyeccion_plp,
    fechas_feriadas_activas,
    normalizar_header,
    snapshot_proyeccion,
)


FILAS = {
    "Personal": (32, 51, "Personal", "Personal nocturnidad"),
    "Personal CX": (33, 52, "Personal CX", "Personal CX nocturnidad"),
    "Personal Soporte": (53, 54, "Personal Soporte", "Personal Soporte nocturnidad"),
    "Personal SMB": (34, 73, "Personal SMB", "Personal SMB nocturnidad"),
}


def numero(valor) -> float:
    return float(valor or 0)


def sincronizar(archivo: Path, year: int) -> None:
    libro = load_workbook(archivo, data_only=True, read_only=True)
    requeridas = {"Horas", "Precios", "Variable", "Tarifacion"}
    requeridas.add("Desglose HORAS PERSONAL")
    faltantes = requeridas.difference(libro.sheetnames)
    if faltantes:
        raise ValueError(f"Faltan hojas requeridas: {', '.join(sorted(faltantes))}")

    horas = libro["Horas"]
    precios = libro["Precios"]
    variables = libro["Variable"]
    tarifaciones = libro["Tarifacion"]
    app = create_app()
    with app.app_context(), app.test_request_context("/"):
        antes = []
        despues = []
        feriados = fechas_feriadas_activas(year)
        tipos = {normalizar_header(nombre): nombre for nombre in FILAS}
        desglose = libro["Desglose HORAS PERSONAL"]
        horas_por_campania = defaultdict(float)
        for tipo_origen, fecha, campania, valor in desglose.iter_rows(
            min_row=2, min_col=1, max_col=4, values_only=True
        ):
            if not hasattr(fecha, "year") or fecha.year != year or not campania:
                continue
            tipo = tipos.get(normalizar_header(tipo_origen))
            if not tipo or not isinstance(valor, (int, float)):
                continue
            mes = fecha.strftime("%Y-%m")
            horas_por_campania[(tipo, str(campania).strip(), mes)] += numero(valor)

        horas_base_servicio = defaultdict(float)
        for (tipo, _campania, mes), valor in horas_por_campania.items():
            horas_base_servicio[(tipo, mes)] += valor

        cumplimiento_servicio = {}
        for indice_mes, columna in enumerate(range(8, 20), start=1):
            mes = f"{year}-{indice_mes:02d}"
            for servicio, (fila_dia, fila_noche, _campania_dia, _campania_noche) in FILAS.items():
                horas_base = horas_base_servicio.get((servicio, mes), 0)
                horas_facturables = numero(horas.cell(fila_dia, columna).value) + numero(
                    horas.cell(fila_noche, columna).value
                )
                cumplimiento_servicio[(servicio, mes)] = (
                    horas_facturables / horas_base * 100 if horas_base else 0
                )

        claves_excel = set()
        for (tipo, campania, mes), valor in horas_por_campania.items():
            proyeccion = buscar_proyeccion_plp(tipo, campania, mes)
            if not proyeccion:
                proyeccion = ProyeccionMatriz(
                    cliente="Personal", campania=campania, year=year, mes=mes,
                    tipo_plp=tipo,
                )
                db.session.add(proyeccion)
                db.session.flush()
            claves_excel.add(proyeccion.id)
            antes.append(snapshot_proyeccion(proyeccion))
            aplicar_calculo_proyeccion(
                proyeccion,
                {
                    "cliente": "Personal",
                    "campania": campania,
                    "tipo_plp": tipo,
                    "horas_carga_manual": True,
                    "horas_requeridas_manual": numero(valor),
                    "carga_semanal_plp": proyeccion.carga_semanal or "L a V",
                    "carga_horaria_plp": proyeccion.carga_horaria or 6,
                    # El desglose contiene las horas base por campaña. La hoja
                    # Horas contiene el total facturable luego del cumplimiento.
                    # La misma relación servicio/mes se aplica a cada apertura.
                    "porcentaje_cumplimiento": cumplimiento_servicio[(tipo, mes)],
                    "jornadas": [],
                },
                mes,
                feriados,
            )
            despues.append(snapshot_proyeccion(proyeccion))

        # Una celda vacía en el desglose significa que esa campaña no aporta
        # horas ese mes. Se eliminan aperturas PLP viejas para que una carga
        # posterior no conserve datos de un archivo anterior.
        obsoletas = ProyeccionMatriz.query.filter(
            ProyeccionMatriz.year == year,
            ProyeccionMatriz.tipo_plp.isnot(None),
            ProyeccionMatriz.tipo_plp != "",
        ).all()
        for proyeccion in obsoletas:
            if proyeccion.id in claves_excel:
                continue
            antes.append(snapshot_proyeccion(proyeccion))
            db.session.delete(proyeccion)

        for indice_mes, columna in enumerate(range(8, 20), start=1):
            mes = f"{year}-{indice_mes:02d}"
            for servicio, (fila_dia, fila_noche, campania_dia, campania_noche) in FILAS.items():
                horas_dia = numero(horas.cell(fila_dia, columna).value)
                horas_noche = numero(horas.cell(fila_noche, columna).value)
                total_horas = horas_dia + horas_noche
                if total_horas <= 0:
                    continue

                distribucion = PersonalDistribucionHoras.query.filter_by(
                    servicio=servicio, mes=mes
                ).first()
                if not distribucion:
                    distribucion = PersonalDistribucionHoras(
                        servicio=servicio, year=year, mes=mes
                    )
                    db.session.add(distribucion)
                antes.append(distribucion.to_dict())
                distribucion.porcentaje_diurno = horas_dia / total_horas * 100
                despues.append(distribucion.to_dict())

                for fila, campania in ((fila_dia, campania_dia), (fila_noche, campania_noche)):
                    importe_horas = numero(horas.cell(fila, columna).value) * numero(
                        precios.cell(fila, columna).value
                    )
                    importe_variable = numero(variables.cell(fila, columna).value)
                    porcentaje = importe_variable / importe_horas * 100 if importe_horas else 0
                    registro_variable = VariableCampania.query.filter_by(
                        campania=campania, mes=mes
                    ).first()
                    if registro_variable:
                        antes.append(registro_variable.to_dict())
                        registro_variable.porcentaje = porcentaje
                        despues.append(registro_variable.to_dict())

                    penalidad = TarifacionCampania.query.filter_by(
                        campania=campania, concepto="Penalidad ADH", mes=mes
                    ).first()
                    if penalidad:
                        antes.append(penalidad.to_dict())
                        penalidad.site = "Personal"
                        penalidad.cliente = "Personal"
                        penalidad.monto = numero(tarifaciones.cell(fila, columna).value)
                        despues.append(penalidad.to_dict())

        registrar_historial(
            "edicion",
            "tarifaciones",
            str(year),
            f"Horas, porcentajes y penalidades de Personal {year} sincronizados desde {archivo.name}",
            antes={"registros": antes},
            despues={"registros": despues},
        )
        db.session.commit()
        print(f"Sincronizados {len(despues)} valores para {year} desde {archivo}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path, help="Ruta del Excel fuente")
    parser.add_argument("--year", type=int, default=2026)
    args = parser.parse_args()
    if not args.archivo.is_file():
        parser.error(f"No existe el archivo: {args.archivo}")
    sincronizar(args.archivo.resolve(), args.year)


if __name__ == "__main__":
    main()
