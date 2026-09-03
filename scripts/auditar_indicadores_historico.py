"""Concilia los indicadores de la app contra la hoja Historico de un Excel."""
import argparse
import os
import sys
from collections import defaultdict

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app
from app.models import HistoricoClienteMensual


CAMPOS = {
    "horas_dotacion_activa": 8,
    "horas_ausentismo": 9,
    "dotacion_promedio": 10,
    "bajas": 11,
    "horas_requeridas": 12,
    "horas_realizadas": 13,
    "pagadas": 14,
    "logueo": 15,
}


def numero(valor):
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def indicadores(datos):
    def cociente(numerador, denominador):
        return datos[numerador] / datos[denominador] if datos[denominador] else None

    return {
        "cumplimiento": cociente("horas_realizadas", "horas_requeridas"),
        "ausentismo": cociente("horas_ausentismo", "horas_dotacion_activa"),
        "rotacion": cociente("bajas", "dotacion_promedio"),
        "eficiencia": cociente("pagadas", "logueo"),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo")
    parser.add_argument("--year", type=int, default=2026)
    args = parser.parse_args()

    excel = defaultdict(lambda: defaultdict(float))
    hoja = load_workbook(args.archivo, data_only=True, read_only=True)["Historico"]
    for fila in hoja.iter_rows(min_row=3, max_col=16, values_only=True):
        fecha = fila[5]
        if not fecha or fecha.year != args.year:
            continue
        mes = fecha.strftime("%Y-%m")
        for campo, indice in CAMPOS.items():
            excel[mes][campo] += numero(fila[indice])

    app = create_app()
    with app.app_context():
        base = defaultdict(lambda: defaultdict(float))
        for registro in HistoricoClienteMensual.query.filter_by(year=args.year).all():
            datos = registro.base_dict()
            for campo in CAMPOS:
                if campo == "pagadas":
                    valor = registro.pagadas
                elif campo == "logueo":
                    valor = registro.logueo
                else:
                    valor = datos.get(campo)
                base[registro.mes][campo] += numero(valor)

    diferencias = 0
    print("mes       indicador        excel          base          diferencia")
    for mes in sorted(set(excel) | set(base)):
        excel_ind = indicadores(excel[mes])
        base_ind = indicadores(base[mes])
        for nombre in excel_ind:
            esperado, actual = excel_ind[nombre], base_ind[nombre]
            if esperado is None and actual is None:
                diferencia = 0.0
            elif esperado is None or actual is None:
                diferencia = float("inf")
            else:
                diferencia = actual - esperado
            if abs(diferencia) > 0.0000005:
                diferencias += 1
            print(f"{mes}  {nombre:14}  {esperado!s:12}  {actual!s:12}  {diferencia:+.10f}")
    print(f"DIFERENCIAS={diferencias}")
    raise SystemExit(1 if diferencias else 0)


if __name__ == "__main__":
    main()
