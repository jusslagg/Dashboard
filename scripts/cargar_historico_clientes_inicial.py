"""Carga el histórico de clientes desde un libro indicado explícitamente."""

import argparse
from collections import defaultdict
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import create_app, db
from app.models import DashboardOperativo, FacturacionAnio, HistoricoClienteMensual


def numero(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path, help="Libro Excel fuente")
    parser.add_argument("--sheet", default="Historico")
    parser.add_argument("--apply", action="store_true", help="Confirma la escritura en DATABASE_URL")
    args = parser.parse_args()
    if not args.archivo.is_file():
        parser.error(f"No existe el archivo: {args.archivo}")

    workbook = load_workbook(args.archivo, data_only=True, read_only=True)
    if args.sheet not in workbook.sheetnames:
        parser.error(f"No existe la hoja {args.sheet!r}")
    groups = defaultdict(lambda: {
        "empresa": "", "site": "", "industria": "", "horas_dotacion_activa": 0,
        "horas_ausentismo": 0, "dotacion_promedio": 0, "bajas": 0,
        "horas_requeridas": 0, "horas_realizadas": 0, "pagadas": 0,
        "logueo": 0, "dotacion_requerida": 0,
    })
    manual_present = defaultdict(lambda: [False, False])
    for row in workbook[args.sheet].iter_rows(min_row=3, max_col=21, values_only=True):
        fecha, cliente = row[5], str(row[3] or "").strip()
        if not fecha or not cliente:
            continue
        month = fecha.strftime("%Y-%m")
        data = groups[(month, cliente)]
        data["empresa"], data["site"], data["industria"] = (
            row[0] or data["empresa"], row[1] or data["site"], row[2] or data["industria"]
        )
        for key, index in (
            ("horas_dotacion_activa", 8), ("horas_ausentismo", 9),
            ("dotacion_promedio", 10), ("bajas", 11), ("horas_requeridas", 12),
            ("horas_realizadas", 13), ("pagadas", 14), ("logueo", 15),
            ("dotacion_requerida", 20),
        ):
            data[key] += numero(row[index])
        manual_present[(month, cliente)][0] |= row[14] not in (None, "")
        manual_present[(month, cliente)][1] |= row[15] not in (None, "")

    app = create_app()
    with app.app_context():
        for (month, cliente), data in groups.items():
            record = HistoricoClienteMensual.query.filter_by(mes=month, cliente=cliente).first()
            if not record:
                record = HistoricoClienteMensual(year=int(month[:4]), mes=month, cliente=cliente)
                db.session.add(record)
            record.datos_base = json.dumps(
                {key: value for key, value in data.items() if key not in ("pagadas", "logueo")},
                ensure_ascii=False,
            )
            record.pagadas = data["pagadas"] if manual_present[(month, cliente)][0] else None
            record.logueo = data["logueo"] if manual_present[(month, cliente)][1] else None

        invoices = defaultdict(lambda: [0, 0])
        for row in FacturacionAnio.query.all():
            key = (row.fecha.strftime("%Y-%m"), row.cliente)
            invoices[key][0] += numero(row.horas_objetivo)
            invoices[key][1] += numero(row.horas_facturadas)
        for (month, cliente), (target, actual) in invoices.items():
            record = HistoricoClienteMensual.query.filter_by(mes=month, cliente=cliente).first()
            if not record:
                record = HistoricoClienteMensual(
                    year=int(month[:4]), mes=month, cliente=cliente, pagadas=None, logueo=None,
                )
                db.session.add(record)
                base = {"_ocultar_sin_cambios": True}
            else:
                base = record.base_dict()
            base["_fact_snapshot_obj"] = target
            base["_fact_snapshot_real"] = actual
            record.datos_base = json.dumps(base, ensure_ascii=False)
        for row in DashboardOperativo.query.all():
            if HistoricoClienteMensual.query.filter_by(mes=row.mes, cliente=row.cliente).first():
                continue
            db.session.add(HistoricoClienteMensual(
                year=row.year, mes=row.mes, cliente=row.cliente,
                datos_base=json.dumps({"_ocultar_sin_cambios": True}, ensure_ascii=False),
            ))
        if args.apply:
            db.session.commit()
            print(f"APLICADO: histórico, {len(groups)} registro(s) por cliente y mes.")
        else:
            db.session.rollback()
            print(f"SIMULADO: histórico, {len(groups)} registro(s). Use --apply para guardar.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
