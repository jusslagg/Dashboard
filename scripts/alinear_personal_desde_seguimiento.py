"""Alinea jefaturas y segmentos de Personal con el libro de seguimiento."""

from argparse import ArgumentParser
from pathlib import Path
import sys

import openpyxl
from flask import session

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from app import create_app, db  # noqa: E402
from app.models import AsignacionComercial, Facturacion2026, Usuario, normalizar_jefe_site  # noqa: E402
from app.routes import normalizar_header, registrar_historial, segmento_personal_por_servicio, serializar_json  # noqa: E402


JEFATURAS_HOJAS = {
    'Antelo': 'Antelo, Micaela',
    'Del Soto': 'Del Soto, Alejandro',
    'Canalini': 'Canalini, Walter',
    'Chanampa': 'Chanampa, Pablo',
}
JEFATURAS_VALIDAS = set(JEFATURAS_HOJAS.values())
JEFATURA_POR_SEGMENTO = {
    'Ventas': 'Chanampa, Pablo',
    'Customer': 'Antelo, Micaela',
    'Customer Digital': 'Antelo, Micaela',
    'Retención': 'Del Soto, Alejandro',
    'Soporte': 'Canalini, Walter',
    'Personal Pay': 'Antelo, Micaela',
    'SMB': 'Canalini, Walter',
}


def clave(valor):
    return normalizar_header(valor).replace(' ', '')


def segmento_canonico(valor):
    normalizado = normalizar_header(valor)
    compacto = normalizado.replace(' ', '')
    if compacto.startswith('retenci'):
        return 'Retención'
    equivalencias = {
        'ventas': 'Ventas',
        'customer': 'Customer',
        'customerdigital': 'Customer Digital',
        'soporte': 'Soporte',
        'personalpay': 'Personal Pay',
        'smb': 'SMB',
    }
    return equivalencias.get(compacto)


def leer_reglas(ruta):
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    jefaturas = {}
    segmentos = {}

    for hoja, jefatura in JEFATURAS_HOJAS.items():
        if hoja not in libro.sheetnames:
            continue
        for fila in libro[hoja].iter_rows(values_only=True):
            servicio = str(fila[2] or '').strip() if len(fila) > 2 else ''
            indicador = str(fila[3] or '').strip() if len(fila) > 3 else ''
            if servicio and indicador == 'Horas Objetivo' and not servicio.upper().startswith('TOTAL'):
                jefaturas[clave(servicio)] = jefatura

    if 'aux' in libro.sheetnames:
        for fila in libro['aux'].iter_rows(min_row=2, values_only=True):
            servicio = str(fila[1] or '').strip() if len(fila) > 1 else ''
            jefatura = normalizar_jefe_site(fila[2]) if len(fila) > 2 else None
            segmento = segmento_canonico(fila[3] if len(fila) > 3 else None)
            if not servicio:
                continue
            if jefatura in JEFATURAS_VALIDAS:
                jefaturas.setdefault(clave(servicio), jefatura)
            if segmento:
                segmentos[clave(servicio)] = segmento

    if 'Base objetivo Masivo' in libro.sheetnames:
        for fila in libro['Base objetivo Masivo'].iter_rows(min_row=3, values_only=True):
            segmento = segmento_canonico(fila[0] if fila else None)
            servicio = str(fila[3] or '').strip() if len(fila) > 3 else ''
            if servicio and segmento:
                segmentos[clave(servicio)] = segmento

    return jefaturas, segmentos


def regla_para(registro, reglas):
    for nombre in (registro.campania, registro.subcampania):
        encontrada = reglas.get(clave(nombre))
        if encontrada:
            return encontrada
    return None


def ejecutar(ruta, aplicar=False):
    jefaturas, segmentos = leer_reglas(ruta)
    app = create_app()
    with app.app_context(), app.test_request_context('/'):
        usuario = Usuario.query.filter_by(activo=True).order_by(Usuario.id).first()
        if usuario:
            session['usuario_id'] = usuario.id

        asignaciones = [
            item for item in AsignacionComercial.query.order_by(AsignacionComercial.id).all()
            if normalizar_header(item.cliente).startswith('personal')
        ]
        antes = []
        cambiadas = []
        sin_segmento = []
        for asignacion in asignaciones:
            segmento_nuevo = (
                None if clave(asignacion.campania) == 'capacitaciones'
                else regla_para(asignacion, segmentos) or segmento_personal_por_servicio(
                    asignacion.campania, asignacion.subcampania, asignacion.tipo_negocio,
                )
            )
            jefe_nuevo = (
                regla_para(asignacion, jefaturas)
                or JEFATURA_POR_SEGMENTO.get(segmento_nuevo)
                or ('Canalini, Walter' if clave(asignacion.campania) == 'capacitaciones' else None)
                or normalizar_jefe_site(asignacion.jefe_site)
            )
            cambios = {}
            if jefe_nuevo and asignacion.jefe_site != jefe_nuevo:
                cambios['jefe_site'] = jefe_nuevo
            if (asignacion.segmento or None) != (segmento_nuevo or None):
                cambios['segmento'] = segmento_nuevo
            if cambios:
                antes.append(asignacion.to_dict())
                for campo, valor in cambios.items():
                    setattr(asignacion, campo, valor)
                cambiadas.append(asignacion)
            if not segmento_nuevo:
                sin_segmento.append(asignacion.campania)

        jefes_maestro = {}
        for asignacion in asignaciones:
            for servicio in (asignacion.campania, asignacion.subcampania):
                jefes_maestro[clave(servicio)] = asignacion.jefe_site

        facturas_cambiadas = 0
        for registro in Facturacion2026.query.filter(
            Facturacion2026.cliente.ilike('personal%')
        ).all():
            jefe_nuevo = regla_para(registro, jefaturas) or regla_para(registro, jefes_maestro)
            if jefe_nuevo and registro.jefe_site != jefe_nuevo:
                registro.jefe_site = jefe_nuevo
                facturas_cambiadas += 1

        if aplicar:
            db.session.flush()
            registrar_historial(
                'edicion', 'asignaciones_personal_excel', Path(ruta).name,
                'Jefaturas y segmentos de Personal alineados con el seguimiento 2026',
                antes={'asignaciones': antes},
                despues={'asignaciones': [item.to_dict() for item in cambiadas]},
                detalle=serializar_json({
                    'asignaciones_actualizadas': len(cambiadas),
                    'filas_facturacion_actualizadas': facturas_cambiadas,
                    'servicios_sin_segmento': sorted(set(sin_segmento)),
                }),
            )
            db.session.commit()
        else:
            db.session.rollback()

        print(f'Jefaturas leídas: {len(jefaturas)}')
        print(f'Segmentos leídos: {len(segmentos)}')
        print(f'Asignaciones a actualizar: {len(cambiadas)}')
        print(f'Filas de facturación a actualizar: {facturas_cambiadas}')
        print(f'Sin segmento: {sorted(set(sin_segmento))}')
        print('APLICADO' if aplicar else 'VISTA PREVIA')


if __name__ == '__main__':
    parser = ArgumentParser()
    parser.add_argument('archivo', type=Path)
    parser.add_argument('--aplicar', action='store_true')
    argumentos = parser.parse_args()
    ejecutar(argumentos.archivo, argumentos.aplicar)
