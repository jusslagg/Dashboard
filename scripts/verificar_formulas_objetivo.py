"""Compara G/S/T/U con los resultados guardados en el Excel de referencia."""
import os
from pathlib import Path
import sys
from math import isclose

os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from app.personal_objetivos import recalcular_objetivos, reglas_plantilla
from app.routes import celda_seguimiento_automatica, normalizar_header, periodo_valor_seguimiento

path = Path(sys.argv[1])
formulas = openpyxl.load_workbook(path, data_only=False)
values = openpyxl.load_workbook(path, data_only=True)

def sheet(name):
    return [[celda_seguimiento_automatica(
        values[name].cell(cell.row, cell.column).value,
        cell.value if cell.data_type == 'f' else None,
    ) for cell in row] for row in formulas[name]]

base = sheet('Base objetivo Masivo')
aux = sheet('Auxiliar Horas PLP masivo')
reglas = reglas_plantilla(base[2:], normalizar_header)
recalcular_objetivos(base, aux, reglas, normalizar_header, periodo_valor_seguimiento)
checked = 0
errors = []
for row_num, row in enumerate(base[2:], start=3):
    for column in (6, 18, 19, 20):
        original = values['Base objetivo Masivo'].cell(row_num, column + 1).value
        if isinstance(original, (int, float)):
            checked += 1
            if not isclose(row[column]['v'], original, rel_tol=1e-9, abs_tol=1e-6):
                errors.append((row_num, column + 1, row[3]['v'], row[4]['v'], original, row[column]['v']))
assert not errors, f'{len(errors)} diferencias: {errors[:20]}'
print(f'OK: {checked} resultados de G/S/T/U coinciden con el Excel de referencia.')
