"""Compara el endpoint de Indicadores con las formulas de la hoja Historico."""
import os
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app
from app.models import Usuario

ARCHIVO = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
    r"C:\Users\jegil\Downloads\2026 Facturacion (36).xlsx"
)
CAMPOS = {
    "cumplimiento": ("horas_realizadas", "horas_requeridas", 13, 12),
    "ausentismo": ("horas_ausentismo", "horas_dotacion_activa", 9, 8),
    "rotacion": ("bajas", "dotacion_promedio", 11, 10),
    "eficiencia": ("pagadas", "logueo", 14, 15),
}


def numero(valor):
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def ratios(filas, origen):
    salida = {}
    for indicador, (num, den, col_num, col_den) in CAMPOS.items():
        if origen == "excel":
            numerador = sum(numero(fila[col_num]) for fila in filas)
            denominador = sum(numero(fila[col_den]) for fila in filas)
        else:
            numerador = sum(numero(fila.get(num)) for fila in filas)
            denominador = sum(numero(fila.get(den)) for fila in filas)
        salida[indicador] = numerador / denominador if denominador else None
    return salida


def diferencia(esperado, actual):
    if esperado is None and actual is None:
        return 0.0
    if esperado is None or actual is None:
        return float("inf")
    return actual - esperado


def main():
    libro = load_workbook(ARCHIVO, data_only=True, read_only=True)
    hoja = libro["Historico"]
    hoja_indicadores = libro["Indicadores"]
    excel = defaultdict(list)
    for fila in hoja.iter_rows(min_row=3, max_col=16, values_only=True):
        fecha = fila[5]
        if fecha and 2019 <= fecha.year <= 2026:
            excel[fecha.year].append(fila)

    app = create_app()
    errores = 0
    with app.app_context():
        usuario = Usuario.query.filter_by(activo=True).first()
        cliente = app.test_client()
        with cliente.session_transaction() as sesion:
            sesion["usuario_id"] = usuario.id
        api_por_year = {}
        for year in range(2019, 2027):
            respuesta = cliente.get(f"/api/historico-clientes?year={year}").get_json()
            filas_api = respuesta["filas"]
            api_por_year[year] = filas_api
            periodos = (
                ("Anual", lambda f: True, lambda f: True),
                ("Ene-Jul", lambda f: f[5].month <= 7, lambda f: int(f["mes"][5:7]) <= 7),
            )
            for periodo, filtro_excel, filtro_api in periodos:
                esperado = ratios([f for f in excel[year] if filtro_excel(f)], "excel")
                actual = ratios([f for f in filas_api if filtro_api(f)], "api")
                for indicador in CAMPOS:
                    delta = diferencia(esperado[indicador], actual[indicador])
                    if abs(delta) > 0.0000005:
                        errores += 1
                    print(
                        f"{year} {periodo:7} {indicador:13} "
                        f"excel={esperado[indicador]!s:18} api={actual[indicador]!s:18} "
                        f"dif={delta:+.10f}"
                    )
        # La hoja visible reemplaza, en el año vigente, el último mes con el
        # acumulado Ene-último mes. Se valida la matriz tal como la ve el usuario.
        bloques = {"cumplimiento": 4, "ausentismo": 20, "rotacion": 36, "eficiencia": 52}
        for indicador, fila_inicial in bloques.items():
            for year in range(2019, 2027):
                columna = 2 + year - 2019
                for mes in range(1, 13):
                    esperado = hoja_indicadores.cell(fila_inicial + mes - 1, columna).value
                    esperado = numero(esperado) if isinstance(esperado, (int, float)) else None
                    filas_mes = [f for f in api_por_year[year] if int(f["mes"][5:7]) == mes]
                    actual = ratios(filas_mes, "api")[indicador]
                    if year == 2026:
                        campo_den = CAMPOS[indicador][1]
                        meses_con_dato = [int(f["mes"][5:7]) for f in api_por_year[year] if numero(f.get(campo_den))]
                        corte = max(meses_con_dato) if meses_con_dato else 0
                        if mes > corte:
                            actual = None
                        elif mes == corte and indicador in ("cumplimiento", "ausentismo"):
                            actual = ratios(
                                [f for f in api_por_year[year] if int(f["mes"][5:7]) <= corte], "api"
                            )[indicador]
                    delta = diferencia(esperado, actual)
                    if abs(delta) > 0.0000005:
                        errores += 1
                        print(f"MATRIZ {year}-{mes:02d} {indicador}: excel={esperado} api={actual} dif={delta}")
    print(f"DIFERENCIAS_CUADROS={errores}")
    raise SystemExit(1 if errores else 0)


if __name__ == "__main__":
    main()
