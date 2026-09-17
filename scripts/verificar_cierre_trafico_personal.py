"""Regresión del cierre de tráfico Personal, sin tocar la base local."""
import io
import json
import os
from pathlib import Path
import sys
import zlib
from datetime import datetime

os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
os.environ['APP_ENV'] = 'development'
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from app import create_app, db
from app.models import HistorialCambio, ProformaPersonal, SeguimientoPersonalImportacion, Usuario
from app.routes import ENCABEZADOS_PROFORMA_PERSONAL, clave_horas_auxiliar_personal


def cierre(periodo=(2026, 8), campania='Soporte', escala=1):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Soporte'
    hoja.cell(3, 2, 'Segmento')
    hoja.cell(3, 3, 'FECHA')
    hoja.cell(3, 7, 'Agentes productivos logueados')
    hoja.cell(3, 8, 'Planificación')
    hoja.cell(3, 17, 'Tipo de Horas')
    for fila, (tipo, brutas, planificadas) in enumerate([
        ('Diurnas', 180, 200), ('Nocturnas', 36, 40), ('Feriados', 18, 20),
    ], start=4):
        hoja.cell(fila, 2, campania)
        hoja.cell(fila, 3, datetime(periodo[0], periodo[1], fila - 3))
        hoja.cell(fila, 7, brutas * escala)
        hoja.cell(fila, 8, planificadas * escala)
        hoja.cell(fila, 17, tipo)
    hoja.cell(8, 13, 'Indice de Adhesión')
    hoja.cell(8, 15, 0.9)
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def proforma_definitiva():
    libro = Workbook()
    hoja = libro.active
    hoja.append(list(ENCABEZADOS_PROFORMA_PERSONAL))
    for tipo in ('Diurnas', 'Nocturnas', 'Feriados'):
        hoja.append(['FDV', 202608, 'SOPORTE', 'CAT', 'Soporte', '', tipo,
                     50, 1, 50, 0, 0, 0, 0, 0, 50])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


app = create_app()
app.config['TESTING'] = True
with app.app_context(), app.test_client() as client:
    usuario = Usuario(nombre='Prueba', email='cierre@example.test', rol='Admin',
                      puesto='Administrador', activo=True, debe_cambiar_password=False)
    usuario.set_password('Prueba12345')
    db.session.add(usuario)
    db.session.flush()
    with client.session_transaction() as sesion:
        sesion['usuario_id'] = usuario.id
        sesion['csrf_token'] = 'test-cierre'
    headers = {'X-CSRF-Token': 'test-cierre'}
    ajustes = {}
    for tipo in ('Diurnas', 'Nocturnas', 'Feriados'):
        ajustes[clave_horas_auxiliar_personal('202608', 'SOPORTE', 'Soporte', tipo)] = {
            'periodo': '202608', 'grupo': 'SOPORTE', 'campania': 'Soporte',
            'tipo_vh': tipo, 'proyectada': 50, 'definitiva': None,
            'brutas': None, 'adh': None,
        }
        db.session.add(ProformaPersonal(
            tipo_proforma='Soporte', estado_proforma='Proyectada', fdv='FDV', periodo='202608', negocio='SOPORTE',
            sitio_proveedor='CAT', segmento='Soporte', subsitio=None, tipo_hora=tipo,
            total_horas=50, precio=1, monto_fijo=50, porcentaje_bono_kpi_vs=0,
            monto_variable_kpi_vs=0, porcentaje_bono_ac=0, monto_variable_ac=0,
            monto_variable=0, total_proyeccion=50, bono_porcentaje_total=0,
        ))
    configuracion = SeguimientoPersonalImportacion(
        archivo='seguimiento.xlsx', contenido_xlsx=b'',
        contenido_hojas=zlib.compress(json.dumps({
            'version': 1, 'hojas': {}, 'ajustes_auxiliar_horas': ajustes,
            'cierres_trafico_importados': [],
        }).encode()),
    )
    db.session.add(configuracion)
    db.session.commit()

    respuesta_pendiente = client.post(
        '/api/seguimiento-personal/cierre-trafico/importar', headers=headers,
        data={'importacion_id': str(configuracion.id), 'archivo': (cierre(), 'cierre.xlsx')},
    )
    assert respuesta_pendiente.status_code == 200
    assert respuesta_pendiente.get_json()['aplicado'] is False
    respuesta = client.post(
        '/api/proforma-personal/importar', headers=headers,
        data={'tipo_proforma': 'Soporte', 'estado_proforma': 'Definitiva',
              'archivo': (proforma_definitiva(), 'proforma_definitiva.xlsx')},
    )
    assert respuesta.status_code == 200, respuesta.get_data(as_text=True)
    datos = respuesta.get_json()
    assert datos['periodos'] == ['202608'] and datos['estado_proforma'] == 'Definitiva'
    paquete = json.loads(zlib.decompress(configuracion.contenido_hojas))
    esperado = {'Diurnas': (100, 90), 'Nocturnas': (20, 18), 'Feriados': (10, 9)}
    for tipo, (planificadas, brutas) in esperado.items():
        valor = paquete['ajustes_auxiliar_horas'][
            clave_horas_auxiliar_personal('202608', 'SOPORTE', 'Soporte', tipo)
        ]
        assert valor['proyectada'] == 50
        assert valor['definitiva'] == planificadas
        assert valor['brutas'] == brutas
        assert valor['adh'] == 0.9
    assert len(paquete['cierres_trafico_importados']) == 1
    assert len(paquete['cierres_trafico_importados'][0]['detalle']) == 3
    control = client.get(
        f'/api/seguimiento-personal/cierres-trafico?importacion_id={configuracion.id}'
    )
    assert control.status_code == 200
    assert control.get_json()['cierres'][0]['detalle'][0]['hoja'] == 'Soporte'

    reemplazo = client.post(
        '/api/seguimiento-personal/cierre-trafico/importar', headers=headers,
        data={'importacion_id': str(configuracion.id),
              'archivo': (cierre(escala=2), 'cierre_corregido.xlsx')},
    )
    assert reemplazo.status_code == 200, reemplazo.get_data(as_text=True)
    paquete = json.loads(zlib.decompress(configuracion.contenido_hojas))
    assert len(paquete['cierres_trafico_importados']) == 1
    assert paquete['cierres_trafico_importados'][0]['archivo'] == 'cierre_corregido.xlsx'
    assert paquete['ajustes_auxiliar_horas'][
        clave_horas_auxiliar_personal('202608', 'SOPORTE', 'Soporte', 'Diurnas')
    ]['definitiva'] == 200

    historiales = HistorialCambio.query.filter_by(
        entidad='seguimiento_personal_cierre', accion='importacion',
    ).order_by(HistorialCambio.id.desc()).all()
    deshacer_reemplazo = client.post(f'/api/historial/{historiales[0].id}/deshacer', headers=headers)
    assert deshacer_reemplazo.status_code == 200
    paquete = json.loads(zlib.decompress(configuracion.contenido_hojas))
    assert paquete['ajustes_auxiliar_horas'][
        clave_horas_auxiliar_personal('202608', 'SOPORTE', 'Soporte', 'Diurnas')
    ]['definitiva'] == 100
    deshacer = client.post(f'/api/historial/{historiales[1].id}/deshacer', headers=headers)
    assert deshacer.status_code == 200, deshacer.get_data(as_text=True)
    paquete = json.loads(zlib.decompress(configuracion.contenido_hojas))
    assert paquete['cierres_trafico_importados'] == []
    assert all(valor['definitiva'] is None for valor in paquete['ajustes_auxiliar_horas'].values())
    pantalla = client.get('/carga-datos-personal')
    html = pantalla.get_data(as_text=True)
    assert pantalla.status_code == 200
    assert 'archivoCierreTraficoPersonal' in html
    assert 'importarCierreTraficoPersonal' in html
    assert 'tabCargaCierreTrafico' in html
    assert 'tbodyCierreTraficoControl' in html
    print('OK: carga definitiva, período, nombres, valores y deshacer verificados.')
