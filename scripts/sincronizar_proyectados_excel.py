"""Sincroniza las fuentes de Proyectados desde un libro operativo anual."""

import argparse
import os
import re
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app, db
from app.models import (
    NextGenDolar,
    NextGenProducto,
    ProyeccionMatriz,
    ProyeccionMatrizJornada,
    ProyeccionPrecio,
    SiteProyeccion,
    TarifacionCampania,
    VariableCampania,
    _canonizar_campania,
    _canonizar_cliente,
)
from app.routes import (aplicar_calculo_proyeccion, calcular_dias_objetivo, fechas_feriadas_activas,
                        normalizar_header, registrar_historial, site_original_next_gen, snapshot_proyeccion)


YEAR = int(os.getenv("APP_DEFAULT_YEAR", datetime.now().year))

# El libro publica estas aperturas con precio, pero su hoja Facturacion Horas
# las deja expresamente en cero antes del primer mes con dotacion propia.
SALIDAS_SIN_HERENCIA = {
    ("piano referidos", f"{YEAR}-01"),
    ("santander getnet supervisor exclusivo", f"{YEAR}-06"),
}


def mes(valor):
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m")
    if hasattr(valor, "year") and hasattr(valor, "month"):
        return f"{valor.year:04d}-{valor.month:02d}"
    return None


def numero(valor, default=0):
    if valor in (None, ""):
        return default
    return float(valor)


def texto(valor):
    return str(valor or "").strip()


def identidad_canonica(cliente, campania):
    cliente = _canonizar_cliente(texto(cliente))
    return cliente, _canonizar_campania(cliente, texto(campania))


def claves_horas_auxiliares(workbook):
    """Claves cuya K no es días × carga horaria × dotación."""
    if workbook is None:
        return set()
    claves = set()
    for row in workbook["Matriz"].iter_rows(min_row=3):
        cliente, campania = identidad_canonica(row[2].value, row[3].value)
        clave_mes = mes(row[5].value)
        if not cliente or not campania or not clave_mes or not clave_mes.startswith(str(YEAR)):
            continue
        numero_fila = row[0].row
        formula = str(row[10].value or '').upper().replace('$', '').replace(' ', '').lstrip('=+')
        referencias = [parte.lstrip('+') for parte in formula.split('*')]
        formula_jornada = sorted(referencias) == sorted(
            [f'H{numero_fila}', f'I{numero_fila}', f'J{numero_fila}']
        )
        if not formula_jornada:
            claves.add((normalizar_header(cliente), normalizar_header(campania), clave_mes))
    return claves


def jornadas_naturgy_auxiliares(workbook):
    """Obtiene las jornadas que alimentan los VLOOKUP de Naturgy en Matriz."""
    jornadas = defaultdict(list)
    if workbook is None or "Desglose Horas NATURGY" not in workbook.sheetnames:
        return jornadas
    for fila in workbook["Desglose Horas NATURGY"].iter_rows(min_row=2, values_only=True):
        clave_mes = mes(fila[5])
        campania = texto(fila[3])
        if not campania or not clave_mes or not clave_mes.startswith(str(YEAR)):
            continue
        jornadas[(normalizar_header(campania), clave_mes)].append({
            "dotacion_requerida": numero(fila[9]),
            "carga_semanal": texto(fila[6]) or "L a V",
            "carga_horaria": numero(fila[8]),
        })
    return jornadas


def jornadas_derivadas_naturgy(campania, clave_mes, filas, fuentes, es_auxiliar):
    if not es_auxiliar:
        return None
    campania_base = re.sub(r"\s+NOCTURNAS$", "", texto(campania), flags=re.IGNORECASE)
    jornadas_base = fuentes.get((normalizar_header(campania_base), clave_mes))
    if not jornadas_base:
        return None
    porcentajes = [numero(fila[15], 1) for fila in filas if fila[15] not in (None, "")]
    factor = max(porcentajes, default=1)
    return [
        {**jornada, "dotacion_requerida": jornada["dotacion_requerida"] * factor}
        for jornada in jornadas_base
    ]


def snapshot_simple(registro, tipo=None):
    if not registro:
        return None
    data = registro.to_dict()
    if tipo:
        data["tipo_registro"] = tipo
    return data


def guardar_historial(entidad, clave, resumen, antes, despues):
    claves = {"proyeccion_matriz": "proyecciones", "matriz_precios": "precios"}
    contenedor = claves.get(entidad, entidad)
    registrar_historial(
        "importacion",
        entidad,
        clave,
        resumen,
        antes={contenedor: antes},
        despues={contenedor: despues},
    )


def sincronizar_matriz(workbook, workbook_formulas=None):
    sheet = workbook["Matriz"]
    grupos = defaultdict(list)
    nombres = {}
    for row in sheet.iter_rows(min_row=3, values_only=True):
        cliente, campania = identidad_canonica(row[2], row[3])
        year, fecha = row[4], row[5]
        clave_mes = mes(fecha)
        # El período de la fila manda sobre la columna Año. El libro incluye,
        # por ejemplo, una jornada de mayo-2026 rotulada accidentalmente 2027
        # y sus propias fórmulas la incorporan al total de mayo.
        if not cliente or not campania or not clave_mes or not clave_mes.startswith(str(YEAR)):
            continue
        clave = (normalizar_header(cliente), normalizar_header(campania), clave_mes)
        grupos[clave].append(row)
        nombres.setdefault(clave, (cliente, campania, clave_mes))

    existentes = {
        (normalizar_header(r.cliente), normalizar_header(r.campania), r.mes): r
        for r in ProyeccionMatriz.query.filter_by(year=YEAR).all()
    }
    # La hoja Matriz no reemplaza las aperturas PLP mantenidas por la app.
    # Se actualizan las claves presentes en el libro y se preservan las demás.
    claves = sorted(set(grupos))
    horas_auxiliares = claves_horas_auxiliares(workbook_formulas)
    fuentes_naturgy = jornadas_naturgy_auxiliares(workbook)
    feriados = fechas_feriadas_activas(YEAR)
    antes, despues = [], []
    for clave in claves:
        registro = existentes.get(clave)
        filas = grupos.get(clave)
        cliente, campania, clave_mes = nombres[clave]
        if not registro:
            registro = ProyeccionMatriz.query.filter(
                db.func.lower(ProyeccionMatriz.cliente) == cliente.lower(),
                db.func.lower(ProyeccionMatriz.campania) == campania.lower(),
                ProyeccionMatriz.mes == clave_mes,
            ).first()
            if not registro:
                registro = ProyeccionMatriz(cliente=cliente, campania=campania, year=YEAR, mes=clave_mes)
                db.session.add(registro)
        antes.append(snapshot_proyeccion(registro))
        total_dotacion = sum(numero(f[9]) for f in filas)
        total_requeridas = sum(numero(f[10]) for f in filas)
        total_proyectadas = sum(numero(f[11]) for f in filas)
        registro.cliente, registro.campania, registro.year, registro.mes = cliente, campania, YEAR, clave_mes
        registro.dotacion_requerida = total_dotacion
        registro.carga_semanal = texto(filas[0][6]) or "L a V"
        registro.carga_horaria = sum(numero(f[8]) * numero(f[9]) for f in filas) / total_dotacion if total_dotacion else 0
        registro.dias_objetivo = int(max(numero(f[7]) for f in filas))
        registro.horas_requeridas = total_requeridas
        registro.porcentaje_cumplimiento = total_proyectadas / total_requeridas * 100 if total_requeridas else 100
        porcentajes_nocturnidad = [numero(f[15]) for f in filas if f[15] not in (None, "")]
        registro.tiene_nocturnidad = bool(porcentajes_nocturnidad and max(porcentajes_nocturnidad) > 0)
        registro.porcentaje_nocturnidad = max(porcentajes_nocturnidad, default=0) * 100
        # Personal conserva las horas del auxiliar. Naturgy conserva sus
        # jornadas base para responder también a cambios de feriados.
        jornadas_naturgy = jornadas_derivadas_naturgy(
            campania, clave_mes, filas, fuentes_naturgy, clave in horas_auxiliares
        )
        registro.horas_carga_manual = clave in horas_auxiliares and not jornadas_naturgy
        registro.dias_objetivo_manual = None
        registro.horas_requeridas_manual = total_requeridas if registro.horas_carga_manual else None
        registro.jornadas.clear()
        jornadas_origen = jornadas_naturgy or [
            {
                "dotacion_requerida": numero(fila[9]),
                "carga_semanal": texto(fila[6]) or "L a V",
                "carga_horaria": numero(fila[8]),
                "dias_objetivo": int(numero(fila[7])),
                "horas_requeridas": numero(fila[10]),
            }
            for fila in filas
        ]
        for jornada in jornadas_origen:
            registro.jornadas.append(ProyeccionMatrizJornada(**jornada))
        if registro.horas_carga_manual:
            aplicar_calculo_proyeccion(registro, {
                'cliente': cliente, 'campania': campania,
                'horas_carga_manual': True,
                'horas_requeridas_manual': total_requeridas,
                'carga_semanal_plp': registro.carga_semanal,
                'carga_horaria_plp': registro.carga_horaria,
                'porcentaje_cumplimiento': registro.porcentaje_cumplimiento,
            }, clave_mes, feriados)
        elif jornadas_naturgy or any(
            calcular_dias_objetivo(YEAR, int(clave_mes[-2:]), jornada.carga_semanal, feriados) != jornada.dias_objetivo
            for jornada in registro.jornadas
        ):
            aplicar_calculo_proyeccion(registro, {
                'cliente': cliente, 'campania': campania,
                'jornadas': [jornada.to_dict() for jornada in registro.jornadas],
                'porcentaje_cumplimiento': registro.porcentaje_cumplimiento,
            }, clave_mes, feriados)
        db.session.flush()
        despues.append(snapshot_proyeccion(registro))
    guardar_historial("proyeccion_matriz", str(YEAR), f"Actualización integral Matriz {YEAR}: {len(grupos)} registros", antes, despues)
    return len(grupos)


def datos_tabla_mensual(sheet, fila_inicio=5):
    headers = list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True)))
    columnas = [(i, mes(v)) for i, v in enumerate(headers) if mes(v)]
    for row in sheet.iter_rows(min_row=fila_inicio, values_only=True):
        cuenta, campania = identidad_canonica(row[2], row[4])
        site = texto(row[3])
        if not cuenta or not campania:
            continue
        yield cuenta, site, campania, row, columnas


def sincronizar_precios_y_fijos(workbook):
    deseados = {}
    for cliente, site, campania, row, columnas in datos_tabla_mensual(workbook["Precios"]):
        for index, clave_mes in columnas:
            if clave_mes.startswith(str(YEAR)) and row[index] not in (None, ""):
                deseados[(cliente, campania, clave_mes)] = {"site": site, "precio": numero(row[index]), "fijo": 0}

    precios_por_campania = defaultdict(list)
    precios_por_cliente = defaultdict(list)
    for clave in deseados:
        precios_por_cliente[normalizar_header(clave[0])].append(clave)
        precios_por_campania[normalizar_header(clave[1])].append(clave)
    fija = workbook["Suma fija"]
    encabezados = list(next(fija.iter_rows(min_row=2, max_row=2, values_only=True)))
    # El segundo bloque mensual es el importe total por dotación que utiliza la app.
    columnas_fijas = [(i, mes(v)) for i, v in enumerate(encabezados) if i >= 14 and mes(v)]
    for row in fija.iter_rows(min_row=3, values_only=True):
        nombre = normalizar_header(row[0])
        # La campaña exacta manda. Usar todas las filas del mismo cliente
        # duplicaba, por ejemplo, el fijo de BBVA Seguros en Autos.
        claves_destino = precios_por_campania.get(nombre, [])
        if not claves_destino and len(precios_por_cliente.get(nombre, [])) == 1:
            claves_destino = precios_por_cliente[nombre]
        for clave in claves_destino:
            indice = next((i for i, m in columnas_fijas if m == clave[2]), None)
            if indice is not None and row[indice] not in (None, ""):
                deseados[clave]["fijo"] = numero(row[indice])

    # Bice se factura enteramente como fijo: su precio horario es cero y el
    # libro no lo repite en Suma fija. Recuperamos sus doce importes de la hoja
    # calculada para no perder el valor al hacer una sincronización integral.
    facturacion_horas = workbook["Facturacion Horas"]
    for row in facturacion_horas.iter_rows(min_row=5, values_only=True):
        cuenta = normalizar_header(row[2])
        campania = normalizar_header(row[4])
        if cuenta != "bice" and campania != "bice":
            continue
        claves_bice = precios_por_campania.get("bice", []) or precios_por_cliente.get("bice", [])
        for clave in claves_bice:
            indice = int(clave[2][-2:]) + 6  # enero H=7 en índice base cero
            if indice < len(row) and row[indice] not in (None, ""):
                deseados[clave]["fijo"] = numero(row[indice])

    existentes = {(r.cliente, r.campania, r.mes): r for r in ProyeccionPrecio.query.filter_by(year=YEAR).all()}
    claves = sorted(set(existentes) | set(deseados))
    antes, despues = [], []
    for clave in claves:
        registro = existentes.get(clave)
        antes.append(snapshot_simple(registro))
        data = deseados.get(clave)
        if not data:
            db.session.delete(registro); despues.append(None); continue
        if not registro:
            registro = ProyeccionPrecio(cliente=clave[0], campania=clave[1], year=YEAR, mes=clave[2])
            db.session.add(registro)
        registro.site, registro.precio_base = data["site"], Decimal(str(data["precio"]))
        registro.alcance_porcentaje = (
            0 if (normalizar_header(clave[1]), clave[2]) in SALIDAS_SIN_HERENCIA else 100
        )
        registro.importe_fijo_mensual = Decimal(str(data["fijo"]))
        registro.recalcular(); db.session.flush(); despues.append(snapshot_simple(registro))
    guardar_historial("matriz_precios", str(YEAR), f"Actualización integral Precios y suma fija {YEAR}: {len(deseados)} registros", antes, despues)
    return len(deseados)


def sincronizar_identidades(workbook):
    """Alinea aperturas de Matriz con la Cuenta canónica definida en Precios."""
    canonicas = defaultdict(set)
    for cuenta, site, campania, _row, _columnas in datos_tabla_mensual(workbook["Precios"]):
        canonicas[normalizar_header(campania)].add((cuenta, site))

    existentes = {
        (normalizar_header(r.cliente), normalizar_header(r.campania)): r
        for r in SiteProyeccion.query.all()
    }
    antes, despues = [], []
    procesadas = set()
    proyecciones = ProyeccionMatriz.query.filter(
        ProyeccionMatriz.year == YEAR,
        db.or_(ProyeccionMatriz.tipo_plp.is_(None), ProyeccionMatriz.tipo_plp == ""),
    ).all()
    for proyeccion in proyecciones:
        destinos = canonicas.get(normalizar_header(proyeccion.campania), set())
        if len(destinos) != 1:
            continue
        cuenta, site = next(iter(destinos))
        if normalizar_header(proyeccion.cliente) == normalizar_header(cuenta):
            continue
        clave = (normalizar_header(proyeccion.cliente), normalizar_header(proyeccion.campania))
        if clave in procesadas:
            continue
        procesadas.add(clave)
        registro = existentes.get(clave)
        antes.append(snapshot_simple(registro))
        if not registro:
            registro = SiteProyeccion(cliente=proyeccion.cliente, campania=proyeccion.campania)
            db.session.add(registro)
            existentes[clave] = registro
        registro.site = site
        registro.cliente_destino = cuenta
        registro.campania_destino = cuenta
        db.session.flush()
        despues.append(snapshot_simple(registro))

    if despues:
        guardar_historial(
            "sites_proyeccion", str(YEAR),
            f"Alineación de identidades canónicas {YEAR}: {len(despues)} registro(s)",
            antes, despues,
        )
    return len(despues)


def sincronizar_variables(workbook):
    deseados = {}
    for cliente, site, campania, row, columnas in datos_tabla_mensual(workbook["Estimacion Variable"]):
        for index, clave_mes in columnas:
            if clave_mes.startswith(str(YEAR)) and row[index] not in (None, ""):
                valor = numero(row[index]); valor = valor * 100 if abs(valor) <= 1 else valor
                deseados[(cliente, campania, clave_mes)] = (site, valor)
    existentes = {(r.cliente, r.campania, r.mes): r for r in VariableCampania.query.filter_by(year=YEAR).all()}
    claves = sorted(set(existentes) | set(deseados)); antes, despues = [], []
    for clave in claves:
        registro = existentes.get(clave); antes.append(snapshot_simple(registro)); data = deseados.get(clave)
        if not data: db.session.delete(registro); despues.append(None); continue
        if not registro:
            registro = VariableCampania(cliente=clave[0], campania=clave[1], year=YEAR, mes=clave[2]); db.session.add(registro)
        registro.site, registro.porcentaje = data; db.session.flush(); despues.append(snapshot_simple(registro))
    guardar_historial("variables", str(YEAR), f"Actualización integral Variables {YEAR}: {len(deseados)} registros", antes, despues)
    return len(deseados)


def sincronizar_tarifacion(workbook):
    deseados = {}
    for cliente, site, campania, row, columnas in datos_tabla_mensual(workbook["Tarifacion"]):
        concepto = texto(row[5]) or "Tarifación"
        for index, clave_mes in columnas:
            if clave_mes.startswith(str(YEAR)) and row[index] not in (None, ""):
                deseados[(cliente, campania, concepto, clave_mes)] = (site, numero(row[index]))
    existentes = {(r.cliente, r.campania, r.concepto, r.mes): r for r in TarifacionCampania.query.filter_by(year=YEAR).all()}
    claves = sorted(set(existentes) | set(deseados)); antes, despues = [], []
    for clave in claves:
        registro = existentes.get(clave); antes.append(snapshot_simple(registro)); data = deseados.get(clave)
        if not data: db.session.delete(registro); despues.append(None); continue
        if not registro:
            registro = TarifacionCampania(cliente=clave[0], campania=clave[1], concepto=clave[2], year=YEAR, mes=clave[3]); db.session.add(registro)
        registro.site, registro.monto = data; db.session.flush(); despues.append(snapshot_simple(registro))
    guardar_historial("tarifaciones", str(YEAR), f"Actualización integral Tarifación {YEAR}: {len(deseados)} registros", antes, despues)
    return len(deseados)


def sincronizar_next_gen(workbook):
    sheet = workbook["Valores Next Gen"]
    headers = list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True)))
    columnas = [(i, mes(v)) for i, v in enumerate(headers) if mes(v)]
    # La proyección final (Unificado/RESUMEN) usa la serie mensual de la fila 1
    # de Valores Next Gen. Precios conserva una cotización auxiliar fija que no
    # representa abril en adelante.
    dolar_valores = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    dolar_por_mes = {
        clave_mes: dolar_valores[index]
        for index, clave_mes in columnas
        if index < len(dolar_valores)
    }
    precios = workbook["Precios"]
    headers_precios = list(next(precios.iter_rows(min_row=4, max_row=4, values_only=True)))
    dolar_precios = list(next(precios.iter_rows(min_row=1, max_row=1, values_only=True)))
    dolar_auxiliar_por_mes = {
        mes(header): dolar_precios[index]
        for index, header in enumerate(headers_precios)
        if mes(header) and index < len(dolar_precios)
    }
    # Solo se factura una combinación campaña/mes si la hoja de salida Next Gen
    # tiene un valor activo. Esto evita cargar cantidades maestras todavía no
    # incorporadas a la proyección final (Bi Bank desde marzo en este libro).
    salida = workbook["Next Gen"]
    activos = set()
    for row in salida.iter_rows(min_row=5, values_only=True):
        campania = texto(row[2])
        if not campania:
            continue
        for indice_mes, clave_mes in enumerate((f"{YEAR}-{m:02d}" for m in range(1, 13)), start=7):
            if indice_mes < len(row) and row[indice_mes] not in (None, "", 0):
                activos.add((normalizar_header(campania), clave_mes))
    productos_crudos = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        campania, producto = texto(row[3]), texto(row[4])
        # Las fórmulas de Next Gen hacen VLOOKUP sobre la columna D. Las filas
        # sin clave allí son auxiliares y no participan, aunque E tenga nombre.
        if not campania or not producto:
            continue
        cliente = campania
        for index, clave_mes in columnas:
            if (normalizar_header(campania), clave_mes) in activos and row[index] not in (None, ""):
                productos_crudos.setdefault((cliente, campania, producto, clave_mes), numero(row[index]))
    existentes_p = {(r.cliente, r.campania, r.producto, r.mes): r for r in NextGenProducto.query.filter_by(year=YEAR).all()}
    existentes_por_producto = {
        (normalizar_header(r.producto), r.mes): (r.cliente, r.campania, r.producto, r.mes)
        for r in existentes_p.values()
    }
    productos = {}
    for clave, valor in productos_crudos.items():
        clave_existente = existentes_por_producto.get((normalizar_header(clave[2]), clave[3]))
        # Replica la primera coincidencia de VLOOKUP cuando el producto aparece
        # duplicado (caso CONSUMER SANTANDER).
        productos.setdefault(clave_existente or clave, valor)
    importes_resumen = defaultdict(float)
    resumen = workbook["RESUMEN"]
    for row in resumen.iter_rows(min_row=5, values_only=True):
        if normalizar_header(row[3]) != "next gen" or not row[4]:
            continue
        for indice_mes, clave_mes in enumerate((f"{YEAR}-{m:02d}" for m in range(1, 13)), start=6):
            if indice_mes < len(row) and isinstance(row[indice_mes], (int, float)):
                importes_resumen[(normalizar_header(row[4]), clave_mes)] += numero(row[indice_mes])
    cantidades_por_campania = defaultdict(float)
    for (cliente, campania, _producto, clave_mes), cantidad in productos.items():
        cantidades_por_campania[(normalizar_header(campania or cliente), clave_mes)] += cantidad
    existentes_d = {r.mes: r for r in NextGenDolar.query.filter_by(year=YEAR).all()}
    antes, despues = [], []
    for clave_mes in [f"{YEAR}-{m:02d}" for m in range(1, 13)]:
        registro = existentes_d.get(clave_mes)
        antes.append(snapshot_simple(registro, "dolar")); valor = dolar_por_mes.get(clave_mes)
        if valor in (None, ""):
            if registro: db.session.delete(registro)
            despues.append(None); continue
        if not registro: registro = NextGenDolar(year=YEAR, mes=clave_mes); db.session.add(registro)
        registro.valor = numero(valor); db.session.flush(); despues.append(snapshot_simple(registro, "dolar"))
    for clave in sorted(set(existentes_p) | set(productos)):
        registro = existentes_p.get(clave); antes.append(snapshot_simple(registro, "producto")); valor = productos.get(clave)
        if valor is None:
            db.session.delete(registro); despues.append(None); continue
        if not registro:
            registro = NextGenProducto(cliente=clave[0], campania=clave[1], producto=clave[2], year=YEAR, mes=clave[3]); db.session.add(registro)
        registro.cantidad_usd = valor
        clave_resumen = (normalizar_header(clave[1] or clave[0]), clave[3])
        cantidad_total = cantidades_por_campania.get(clave_resumen, 0)
        importe_resumen = importes_resumen.get(clave_resumen)
        if cantidad_total and importe_resumen is not None:
            registro.cotizacion_aplicada = importe_resumen / cantidad_total
        else:
            registro.cotizacion_aplicada = numero(
                dolar_por_mes.get(clave[3])
                if normalizar_header(site_original_next_gen(clave[0], clave[1])) == "n/a"
                else dolar_auxiliar_por_mes.get(clave[3])
            )
        db.session.flush(); despues.append(snapshot_simple(registro, "producto"))
    guardar_historial("next_gen", str(YEAR), f"Actualización integral Next Gen {YEAR}: {len(productos) + 12} valores", antes, despues)
    return len(productos) + 12


def main():
    global YEAR, SALIDAS_SIN_HERENCIA
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo")
    parser.add_argument("--year", type=int, default=int(os.getenv("APP_DEFAULT_YEAR", datetime.now().year)))
    parser.add_argument("--apply", action="store_true")
    parser.add_argument(
        "--section",
        choices=("all", "matriz", "precios", "identidades", "variables", "tarifacion", "next_gen"),
        default="all",
    )
    args = parser.parse_args()
    YEAR = args.year
    SALIDAS_SIN_HERENCIA = {
        ("piano referidos", f"{YEAR}-01"),
        ("santander getnet supervisor exclusivo", f"{YEAR}-06"),
    }
    workbook = openpyxl.load_workbook(args.archivo, data_only=True, read_only=True)
    workbook_formulas = openpyxl.load_workbook(args.archivo, data_only=False, read_only=True)
    app = create_app()
    with app.test_request_context("/sincronizacion-proyectados"):
        tareas = {
            "matriz": lambda libro: sincronizar_matriz(libro, workbook_formulas),
            "precios": sincronizar_precios_y_fijos,
            "identidades": sincronizar_identidades,
            "variables": sincronizar_variables,
            "tarifacion": sincronizar_tarifacion,
            "next_gen": sincronizar_next_gen,
        }
        seleccionadas = tareas.items() if args.section == "all" else [(args.section, tareas[args.section])]
        resultados = {nombre: tarea(workbook) for nombre, tarea in seleccionadas}
        if args.apply:
            db.session.commit()
        else:
            db.session.rollback()
        print(("APLICADO" if args.apply else "SIMULADO"), resultados)


if __name__ == "__main__":
    main()
