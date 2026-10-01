"""Concilia la base anual contra la solapa del año del libro operativo."""

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import json
import os
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app, db
from app.models import Facturacion2026, HistorialCambio
from app.routes import normalizar_header, normalizar_tipo_vh


def texto(valor):
    return normalizar_header(valor)


def numero(valor):
    try:
        return round(float(valor or 0), 4)
    except (TypeError, ValueError):
        return 0.0


def periodo(valor):
    if isinstance(valor, (date, datetime)):
        return valor.strftime("%Y-%m")
    return str(valor or "")[:7]


def identidad_fuente(fila):
    return (texto(fila[4]), texto(fila[7]), texto(normalizar_tipo_vh(fila[8])))


def identidad_registro(registro):
    return (texto(registro.cliente), texto(registro.subcampania), texto(registro.tipo_jornada))


def clave_semilla_fuente(fila, nombre):
    return (periodo(fila[9]), texto(nombre), numero(fila[10]))


def clave_semilla_registro(registro):
    cliente = texto(registro.cliente)
    if cliente == "santander getnet" and "onboarding" in texto(registro.campania):
        cliente = "santander getnet onboarding"
    return (registro.mes, cliente, numero(registro.horas_objetivo))


def vector_fuente(fila):
    return tuple(numero(fila[indice]) for indice in (10, 11, 12, 13, 17, 18, 21, 26, 30))


def vector_registro(registro):
    return (
        numero(registro.horas_objetivo), numero(registro.horas_facturadas),
        numero(registro.valor_hora), numero(registro.facturado_horas),
        numero(registro.variable_productivo_calculo),
        numero(registro.facturado_bono + registro.penalizaciones_incumplimientos),
        numero(registro.total_dashboard), numero(registro.total_teorico),
        numero(registro.total_facturado_control),
    )


def distancia(registro, fila):
    total = 0.0
    for actual, deseado in zip(vector_registro(registro), vector_fuente(fila)):
        total += abs(actual - deseado) / max(abs(actual), abs(deseado), 1.0)
    return total


def construir_asociaciones(registros, filas):
    """Asocia filas aunque el libro nuevo haya cambiado horas o importes."""
    fuentes_semilla = defaultdict(list)
    for indice, fila in filas:
        for nombre in {fila[4], fila[6]}:
            fuentes_semilla[clave_semilla_fuente(fila, nombre)].append((indice, fila))

    usados = defaultdict(int)
    traducciones = defaultdict(Counter)
    for registro in registros:
        clave = clave_semilla_registro(registro)
        posicion = usados[clave]
        candidatos = fuentes_semilla.get(clave, ())
        if posicion >= len(candidatos):
            continue
        _, fila = candidatos[posicion]
        usados[clave] += 1
        # Las filas en cero se repiten entre muchos servicios y no identifican
        # un alias de forma segura. Los alias se aprenden solo de filas con base.
        if numero(fila[10]) == 0:
            continue
        traducciones[identidad_fuente(fila)][identidad_registro(registro)] += 1

    identidades_db = {identidad_registro(registro) for registro in registros}
    fuentes_por_grupo = defaultdict(list)
    for indice, fila in filas:
        identidad = identidad_fuente(fila)
        if identidad not in identidades_db and identidad in traducciones:
            identidad = traducciones[identidad].most_common(1)[0][0]
        fuentes_por_grupo[(periodo(fila[9]), *identidad)].append((indice, fila))

    registros_por_grupo = defaultdict(list)
    for registro in registros:
        registros_por_grupo[(registro.mes, *identidad_registro(registro))].append(registro)

    asociaciones, sin_fuente = [], []
    for clave, propios in registros_por_grupo.items():
        disponibles = list(fuentes_por_grupo.get(clave, ()))
        for registro in propios:
            if not disponibles:
                sin_fuente.append(registro)
                continue
            mejor = min(range(len(disponibles)), key=lambda i: distancia(registro, disponibles[i][1]))
            indice, fila = disponibles.pop(mejor)
            asociaciones.append((registro, indice, fila))
    return asociaciones, sin_fuente


def aplicar_fila(registro, fila):
    horas_objetivo = numero(fila[10])
    valor_hora = numero(fila[12])
    valor_hora_objetivo = numero(fila[22]) or valor_hora
    objetivo_horas = numero(fila[24])

    registro.horas_objetivo = horas_objetivo
    registro.horas_facturadas = numero(fila[11])
    registro.valor_hora = valor_hora
    registro.valor_hora_objetivo = valor_hora_objetivo
    registro.objetivo_separar_ajuste_vh = bool(
        horas_objetivo
        and abs(objetivo_horas - horas_objetivo * valor_hora_objetivo) > 0.02
        and abs(objetivo_horas - horas_objetivo * valor_hora) <= 0.02
    )
    registro.facturado_horas_manual = numero(fila[13])
    registro.tarifacion = numero(fila[14])
    registro.variable_productivo = numero(fila[17])
    ajuste = numero(fila[18])
    registro.bonos = max(ajuste, 0)
    registro.penalizaciones = abs(min(ajuste, 0))
    registro.netx_gen = numero(fila[19])
    registro.otros = numero(fila[20])
    registro.total_facturado_manual = numero(fila[21])
    registro.variable_objetivo = numero(fila[25])
    registro.control_objetivo_total = numero(fila[26])
    registro.control_facturado_horas = numero(fila[27])
    registro.control_variable_productivo = numero(fila[28])
    registro.control_penalizaciones_bonos = numero(fila[29])
    registro.control_total_facturado = numero(fila[30])


def main(archivo, year, aplicar=False):
    hoja = load_workbook(archivo, data_only=True, read_only=True)[str(year)]
    filas = [
        (indice, fila)
        for indice, fila in enumerate(hoja.iter_rows(min_row=3, max_col=31, values_only=True), start=3)
        if fila[4] and periodo(fila[9]).startswith(f"{year:04d}-")
    ]

    app = create_app()
    with app.app_context():
        meses_fuente = sorted({periodo(fila[9]) for _, fila in filas})
        registros = Facturacion2026.query.filter(
            Facturacion2026.mes.in_(meses_fuente),
            Facturacion2026.es_next_gen.is_(False),
        ).order_by(Facturacion2026.id).all()
        asociaciones, sin_fuente = construir_asociaciones(registros, filas)
        cambios = []
        for registro, fila_excel, fila in asociaciones:
            antes = registro.to_dict()
            aplicar_fila(registro, fila)
            despues = registro.to_dict()
            if antes != despues:
                cambios.append({"id": registro.id, "fila_excel": fila_excel, "antes": antes, "despues": despues})

        resultado = {
            "archivo": str(archivo),
            "sha256": hashlib.sha256(Path(archivo).read_bytes()).hexdigest(),
            "filas_db": len(registros), "filas_asociadas": len(asociaciones),
            "filas_actualizadas": len(cambios), "filas_sin_fuente": len(sin_fuente),
            "ids_sin_fuente": [registro.id for registro in sin_fuente],
        }
        if aplicar:
            historial = HistorialCambio(
                accion="reprocesamiento", entidad="facturacion_anio", entidad_id=str(year),
                resumen=f"Conciliación Facturación {year} desde solapa {year}: {len(cambios)} filas",
                detalle=f"Fuente: {Path(archivo).name}; SHA-256: {resultado['sha256']}",
                antes=json.dumps({"cambios": cambios}, ensure_ascii=False, default=str),
                despues=json.dumps(resultado, ensure_ascii=False, default=str),
            )
            db.session.add(historial)
            db.session.commit()
            resultado["history_id"] = historial.id
        else:
            db.session.rollback()
        print(json.dumps(resultado, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path)
    parser.add_argument("--year", type=int, default=int(os.getenv("APP_DEFAULT_YEAR", datetime.now().year)))
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    main(args.archivo, args.year, args.apply)
