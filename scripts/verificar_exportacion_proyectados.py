"""Pruebas de aperturas, agrupación PLP y tipos de celda de la exportación."""
from pathlib import Path
import sys
import unittest
from io import BytesIO

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook, load_workbook
from app.proyecciones_export import projection_rows, table, PERCENT


class ExportacionProyectadosTest(unittest.TestCase):
    def item(self, **changes):
        return dict(cliente='Personal', campania='Servicio', year=2026, mes='2026-01', mes_label='Ene-26',
                    dotacion_requerida=10, carga_semanal='L a V', carga_horaria=6, dias_objetivo=20,
                    horas_requeridas=1200, horas_proyectadas=1080, porcentaje_cumplimiento=90,
                    tiene_nocturnidad=True, porcentaje_nocturnidad=20, **changes)

    def test_todas_separa_nocturnidad_y_excluye_plp(self):
        rows = list(projection_rows([self.item(), self.item(tipo_plp='Personal')]))
        self.assertEqual(len(rows), 2)
        self.assertEqual([row[4] for row in rows], [8, 2])
        self.assertEqual(sum(row[9] for row in rows), 1080)
        self.assertEqual(rows[1][1], 'Servicio nocturnidad')
        self.assertEqual(rows[0][10], .9)

    def test_plp_agrupa_servicios_y_conserva_horas(self):
        rows = list(projection_rows([self.item(tipo_plp='Personal'), self.item(tipo_plp='Personal')], True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(sum(row[8] for row in rows), 2400)
        self.assertEqual(sum(row[9] for row in rows), 2160)
        self.assertEqual(rows[0][10], .9)

    def test_celdas_numericas_y_texto_literal(self):
        book = Workbook(); book.remove(book.active)
        table(book, 'Prueba', ['Nombre', 'Porcentaje'], [['=1+1', .2]], {2: PERCENT})
        output = BytesIO(); book.save(output); output.seek(0)
        sheet = load_workbook(output).active
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertEqual(sheet['B2'].value, .2)
        self.assertEqual(sheet['B2'].number_format, PERCENT)


if __name__ == '__main__':
    unittest.main()
