"""Verifica los nombres y datos de las hojas de ratios del Excel de Directorio."""
from pathlib import Path
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from app.directorio_export import crear_libro_directorio


class ExportacionRatiosTest(unittest.TestCase):
    def test_nombres_actuales_y_anteriores(self):
        for titles in (('Ratio Staff Operaciones 1', 'Ratio Staff Operaciones 2'),
                       ('Ratio Eli', 'Ratio Eli II')):
            with self.subTest(titles=titles):
                data = {'serie': [{'mes': '2026-08', 'requerido': 100, 'staff': 20,
                                   'operaciones': 100, 'total': 120, 'ratio': 0.2,
                                   'operaciones_importe': 1000, 'staff_importe': 200,
                                   'total_importe': 1200, 'ratio_importe': 0.2}]}
                book = load_workbook(crear_libro_directorio([(title, data) for title in titles], 2026))
                self.assertEqual(book.sheetnames, list(titles))
                for title in titles:
                    sheet = book[title]
                    self.assertEqual(sheet['A1'].value, title)
                    row = next(row for row in sheet.iter_rows() if row[0].value == '2026-08')
                    self.assertTrue(any(cell.value == 0.2 and '%' in cell.number_format for cell in row))
                    self.assertEqual(len(sheet._charts), 1)


if __name__ == '__main__':
    unittest.main()
