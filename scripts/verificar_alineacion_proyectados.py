"""Verifica Proyectados por cuenta, mes y concepto contra el libro operativo."""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from app import create_app  # noqa: E402
from app.models import ProyeccionMatriz  # noqa: E402
from app.routes import (  # noqa: E402
    construir_resumen_proyeccion_data,
    normalizar_header,
    normalizar_nombre_variable,
)


COMPONENTES = ("Facturacion Horas", "Variable", "Tarifacion", "Next Gen")


def moneda(valor: float) -> float:
    return float(Decimal(str(valor or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def clave_cuenta(valor: object) -> str:
    clave = normalizar_nombre_variable(valor)
    if clave.endswith(" nocturnidad"):
        clave = clave[: -len(" nocturnidad")]
    return clave


def datos_app(year: int):
    meses = [f"{year}-{numero:02d}" for numero in range(1, 13)]
    salida = {nombre: defaultdict(lambda: defaultdict(float)) for nombre in COMPONENTES}
    data = construir_resumen_proyeccion_data(year)
    for fila in data["filas"]:
        if fila.get("tipo") == "total_site":
            continue
        concepto = fila.get("concepto")
        componente = (
            "Facturacion Horas" if concepto in ("Horas", "Fijo mensual")
            else "Variable" if concepto == "Variable"
            else "Next Gen" if concepto == "Next Gen"
            else "Tarifacion"
        )
        clave = clave_cuenta(fila.get("campania"))
        for mes in meses:
            salida[componente][clave][mes] += float(fila.get("meses", {}).get(mes, 0) or 0)
    return salida, meses


def datos_excel(workbook, meses):
    salida = {nombre: defaultdict(lambda: defaultdict(float)) for nombre in COMPONENTES}
    for nombre in COMPONENTES:
        if nombre == "Next Gen":
            # TOTAL CAT se alimenta de RESUMEN/Unificado, cuya serie de dólar
            # puede diferir de la hoja auxiliar Next Gen.
            sheet = workbook["RESUMEN"]
            for row in sheet.iter_rows(min_row=5, values_only=True):
                if normalizar_header(row[3]) != "next gen" or not row[4]:
                    continue
                clave = clave_cuenta(row[4])
                for indice, mes_actual in enumerate(meses, 6):
                    valor = row[indice]
                    if isinstance(valor, (int, float)):
                        salida[nombre][clave][mes_actual] += float(valor)
            continue
        sheet = workbook[nombre]
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if not row[2]:
                continue
            clave = clave_cuenta(row[2])
            for indice, mes in enumerate(meses, 7):
                valor = row[indice]
                if isinstance(valor, (int, float)):
                    salida[nombre][clave][mes] += float(valor)
    return salida


def datos_personal_app(year: int):
    salida = defaultdict(float)
    filas = ProyeccionMatriz.query.filter(
        ProyeccionMatriz.year == year,
        ProyeccionMatriz.tipo_plp.isnot(None),
        ProyeccionMatriz.tipo_plp != "",
    ).all()
    for fila in filas:
        clave = (normalizar_header(fila.tipo_plp), normalizar_header(fila.campania), fila.mes)
        salida[clave] += float(fila.horas_requeridas or 0)
    return salida


def datos_personal_excel(workbook, year: int):
    salida = defaultdict(float)
    for tipo, fecha, campania, horas, *_ in workbook["Desglose HORAS PERSONAL"].iter_rows(
        min_row=2, values_only=True
    ):
        if getattr(fecha, "year", None) != year or not tipo or not campania:
            continue
        if not isinstance(horas, (int, float)):
            continue
        clave = (normalizar_header(tipo), normalizar_header(campania), fecha.strftime("%Y-%m"))
        salida[clave] += float(horas)
    return salida


def formula_bi_bank_ausente(workbook, componente, cuenta, mes, app, excel) -> bool:
    if componente != "Next Gen" or cuenta != "bi bank" or mes < "2026-03" or not app or excel:
        return False
    columna = int(mes[-2:]) + 7  # enero H=8
    return workbook["Next Gen"].cell(8, columna).value in (None, "")


def verificar(archivo: Path, year: int, estricto: bool) -> int:
    workbook = load_workbook(archivo, data_only=True, read_only=False)
    workbook_formulas = load_workbook(archivo, data_only=False, read_only=False)
    app = create_app()
    with app.app_context():
        actual, meses = datos_app(year)
        personal_actual = datos_personal_app(year)
    esperado = datos_excel(workbook, meses)
    personal_esperado = datos_personal_excel(workbook, year)

    inesperadas = []
    formulas_faltantes = []
    for componente in COMPONENTES:
        cuentas = set(actual[componente]) | set(esperado[componente])
        for cuenta in sorted(cuentas):
            for mes in meses:
                valor_app = moneda(actual[componente].get(cuenta, {}).get(mes, 0))
                valor_excel = moneda(esperado[componente].get(cuenta, {}).get(mes, 0))
                if valor_app == valor_excel:
                    continue
                diferencia = (componente, cuenta, mes, valor_app, valor_excel)
                if formula_bi_bank_ausente(
                    workbook_formulas, componente, cuenta, mes, valor_app, valor_excel
                ):
                    formulas_faltantes.append(diferencia)
                else:
                    inesperadas.append(diferencia)

    for componente in COMPONENTES:
        total = len((set(actual[componente]) | set(esperado[componente]))) * 12
        errores = sum(1 for item in inesperadas if item[0] == componente)
        faltantes = sum(1 for item in formulas_faltantes if item[0] == componente)
        print(f"{componente}: {total - errores - faltantes}/{total} coinciden; "
              f"errores={errores}; formulas_excel_faltantes={faltantes}")

    for componente, cuenta, mes, valor_app, valor_excel in inesperadas:
        print(f"ERROR {componente} | {cuenta} | {mes} | app={valor_app:.2f} | excel={valor_excel:.2f}")
    for componente, cuenta, mes, valor_app, valor_excel in formulas_faltantes:
        print(f"FORMULA_FALTANTE {componente} | {cuenta} | {mes} | "
              f"calculado={valor_app:.2f} | cache_excel={valor_excel:.2f}")

    personal_errores = []
    for clave in sorted(set(personal_actual) | set(personal_esperado)):
        valor_app = personal_actual.get(clave, 0)
        valor_excel = personal_esperado.get(clave, 0)
        if abs(valor_app - valor_excel) > 0.000001:
            personal_errores.append((*clave, valor_app, valor_excel))
    print(f"Personal campaña/mes: {len(set(personal_actual) | set(personal_esperado)) - len(personal_errores)}/"
          f"{len(set(personal_actual) | set(personal_esperado))} coinciden; errores={len(personal_errores)}")
    for servicio, campania, mes, valor_app, valor_excel in personal_errores:
        print(f"ERROR Personal | {servicio} | {campania} | {mes} | "
              f"app={valor_app:.6f} | excel={valor_excel:.6f}")

    return 1 if inesperadas or personal_errores or (estricto and formulas_faltantes) else 0


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path)
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--strict", action="store_true", help="Falla también por fórmulas ausentes en Excel")
    args = parser.parse_args()
    if not args.archivo.is_file():
        parser.error(f"No existe el archivo: {args.archivo}")
    raise SystemExit(verificar(args.archivo.resolve(), args.year, args.strict))


if __name__ == "__main__":
    main()
