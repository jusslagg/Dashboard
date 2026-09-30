"""Auditoría de solo lectura para las fórmulas y fuentes de Directorio."""
from math import isclose
import os
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app, db
from sqlalchemy import func

from app.models import DotacionClienteMensual, Facturacion2026, HistoricoClienteMensual, Usuario


def cerca(a, b, tolerancia=0.02):
    return isclose(float(a or 0), float(b or 0), abs_tol=tolerancia)


def verificar():
    app = create_app()
    year = app.config['DEFAULT_YEAR']
    future_year = year + 1
    errores = []
    urls = [
        f'/api/kpis?year={year}', f'/api/indicadores?year={year}', f'/api/dotaciones-clientes?year={year}',
        '/api/graficos-dotaciones', '/api/graficos-evolutivo-dotaciones',
        f'/api/variacion-anual?year={year}', f'/api/variacion-horas-clientes?year={year}',
        f'/api/dashboard-operativo?year={year}', f'/api/historico-clientes?year={year}',
        f'/api/cumplimiento-facturacion-clientes?year={year}',
        f'/api/comparativo-interanual?year={year}', '/api/ratio-eli',
        '/api/ratio-eli-ii', '/api/share',
    ]
    with app.test_client() as cliente:
        with cliente.session_transaction() as sesion:
            with app.app_context():
                usuario = Usuario.query.first()
                if not usuario:
                    raise RuntimeError('No existe un usuario para ejecutar la auditoría.')
                sesion['usuario_id'] = usuario.id

        respuestas = {}
        for url in urls:
            respuesta = cliente.get(url)
            if respuesta.status_code != 200:
                errores.append(f'{url}: HTTP {respuesta.status_code}')
                continue
            data = respuesta.get_json()
            if not data or not data.get('success'):
                errores.append(f'{url}: respuesta sin success')
                continue
            respuestas[url.split('?')[0]] = data

        reloj = respuestas.get('/api/indicadores', {}).get('reloj', {})
        with app.app_context():
            historico_reloj = HistoricoClienteMensual.query.filter_by(year=year).all()
            meses_reloj = sorted({item.mes for item in historico_reloj})
            corte_reloj = meses_reloj[-1] if meses_reloj else None
            filas_reloj = [item for item in historico_reloj if not corte_reloj or item.mes <= corte_reloj]
            requerido_reloj = sum(float(item.base_dict().get('horas_requeridas') or 0) for item in filas_reloj)
            realizado_reloj = sum(float(item.base_dict().get('horas_realizadas') or 0) for item in filas_reloj)
        if historico_reloj:
            if reloj.get('fuente') != 'historico_clientes_mensuales':
                errores.append('Reloj: no utiliza la base histórica de Directorio')
            if reloj.get('corte') != corte_reloj:
                errores.append('Reloj: el corte no coincide con el último cierre histórico')
            if not cerca(reloj.get('requerido'), requerido_reloj):
                errores.append('Reloj: horas requeridas no coinciden con Histórico clientes')
            if not cerca(reloj.get('realizado'), realizado_reloj):
                errores.append('Reloj: horas realizadas no coinciden con Histórico clientes')
            if meses_reloj:
                respuesta_corte = cliente.get(f'/api/indicadores?year={year}&mes={meses_reloj[0]}').get_json()['reloj']
                filas_primer_mes = [item for item in historico_reloj if item.mes <= meses_reloj[0]]
                requerido_primer_mes = sum(float(item.base_dict().get('horas_requeridas') or 0) for item in filas_primer_mes)
                if respuesta_corte.get('corte') != meses_reloj[0] or not cerca(respuesta_corte.get('requerido'), requerido_primer_mes):
                    errores.append('Reloj: el selector de corte mensual no recalcula el acumulado')
            resumen_horas = cliente.get(f'/api/resumen-horas?year={year}').get_json()
            total_resumen = resumen_horas.get('total', {})
            if not cerca(reloj.get('requerido'), total_resumen.get('horas_objetivo')):
                errores.append('Reloj: requerido distinto de TOTAL CAT en Resumen de horas')
            if not cerca(reloj.get('realizado'), total_resumen.get('horas_facturadas')):
                errores.append('Reloj: realizado distinto de TOTAL CAT en Resumen de horas')
            if not cerca(float(reloj.get('cumplimiento') or 0) * 100, total_resumen.get('porcentaje_cumplimiento_horas')):
                errores.append('Reloj: porcentaje distinto de % Alcance Hs de TOTAL CAT')

        for fila in respuestas.get('/api/ratio-eli', {}).get('serie', []):
            esperado = fila['staff'] / fila['requerido'] if fila['requerido'] else None
            if esperado is not None and not cerca(fila['ratio'], esperado, 1e-10):
                errores.append(f"Ratio Eli {fila['mes']}: ratio inconsistente")

        for fila in respuestas.get('/api/ratio-eli-ii', {}).get('serie', []):
            controles = (
                ('total', fila['operaciones'] + fila['staff']),
                ('ratio', fila['staff'] / fila['operaciones']),
                ('total_importe', fila['operaciones_importe'] + fila['staff_importe']),
                ('ratio_importe', fila['staff_importe'] / fila['operaciones_importe']),
            )
            for campo, esperado in controles:
                if not cerca(fila[campo], esperado, 1e-10):
                    errores.append(f"Ratio Eli II {fila['mes']}: {campo} inconsistente")

        for bloque in ('horas', 'dotaciones'):
            for fila in respuestas.get('/api/share', {}).get(bloque, []):
                grupos = ('Personal', 'Santander', 'Multicampania', 'USA')
                if not cerca(fila['Total'], sum(fila[g] for g in grupos)):
                    errores.append(f"Share {bloque} {fila['year']}: total inconsistente")
                if fila['Total'] and not cerca(sum(fila['shares'][g] for g in grupos), 1, 1e-9):
                    errores.append(f"Share {bloque} {fila['year']}: shares no suman 100%")
                if not cerca(fila['personal_santander'], fila['Personal'] + fila['Santander']):
                    errores.append(f"Share {bloque} {fila['year']}: Personal + Santander inconsistente")

        cumplimiento = respuestas.get('/api/cumplimiento-facturacion-clientes', {})
        filas_cumplimiento = cumplimiento.get('filas', [])
        anual_actual = next(
            (item for item in cumplimiento.get('resumen_anual', []) if item['year'] == cumplimiento.get('year')),
            None,
        )
        if filas_cumplimiento and anual_actual:
            controles = {
                'horas': sum(item['desvio_horas'] for item in filas_cumplimiento),
                'variable': sum(item['desvio_variable'] for item in filas_cumplimiento),
                'bonos_penalizaciones': sum(item['desvio_penalizaciones_bonos'] for item in filas_cumplimiento),
            }
            for campo, esperado in controles.items():
                if not cerca(anual_actual[campo], esperado):
                    errores.append(f'Cumplimiento facturación anual: {campo} inconsistente')
            if not cerca(anual_actual['total'], sum(controles.values())):
                errores.append('Cumplimiento facturación anual: total inconsistente')

        with app.app_context():
            anios_historico = {fila[0] for fila in db.session.query(HistoricoClienteMensual.year).distinct().all()}
            anios_dotacion = {
                int(fila[0]) for fila in db.session.query(
                    func.extract('year', DotacionClienteMensual.fecha)
                ).distinct().all() if fila[0]
            }
            anios_facturacion = {
                int(fila[0]) for fila in db.session.query(
                    func.extract('year', Facturacion2026.fecha)
                ).distinct().all() if fila[0]
            }
            anios_esperados = sorted(
                int(anio) for anio in anios_historico | anios_dotacion | anios_facturacion
                if anio and int(anio) >= 2020
            )
        if respuestas.get('/api/share', {}).get('anios') != anios_esperados:
            errores.append('Share: los años no se derivan de todos los datos disponibles')

        futuro = Facturacion2026(
            fecha=__import__('datetime').date(future_year, 8, 1), mes=f'{future_year}-08', cliente='Prueba dinamica',
            tipo_jornada='Diurna', horas_objetivo=100, horas_facturadas=90,
            valor_hora_objetivo=20, valor_hora=20, variable_objetivo=100,
            variable_productivo=80, bonos=10, penalizaciones=5, netx_gen=0, otros=0,
        )
        if not cerca(futuro.total_teorico, 2100):
            errores.append('Carga futura: Facturacion Objetivo no se calcula dinamicamente')
        if not cerca(futuro.total_facturado_control, 1885):
            errores.append('Carga futura: total de control no usa las formulas dinamicas')

        ruta_auditoria = os.getenv('DIRECTORIO_EXCEL_AUDIT', '').strip()
        archivo_excel = Path(ruta_auditoria) if ruta_auditoria else None
        if archivo_excel and archivo_excel.exists():
            from openpyxl import load_workbook
            hoja = load_workbook(archivo_excel, data_only=True, read_only=True)[str(year)]
            total_excel = objetivo_excel = 0.0
            for fila in hoja.iter_rows(min_row=3, max_col=31, values_only=True):
                if not hasattr(fila[9], 'year'):
                    continue
                total_excel += float(fila[21] or 0) if isinstance(fila[21], (int, float)) else 0
                objetivo_excel += float(fila[26] or 0) if isinstance(fila[26], (int, float)) else 0
            kpis = respuestas.get('/api/kpis', {}).get('kpis', {})
            if not cerca(kpis.get('total_facturado'), total_excel):
                errores.append(
                    f"KPI Total Facturado {year} no coincide con el SUBTOTAL del Excel: "
                    f"app={float(kpis.get('total_facturado') or 0):.2f}; excel={total_excel:.2f}"
                )
            if not cerca(kpis.get('total_teorico'), objetivo_excel):
                errores.append(
                    f"KPI Facturacion Objetivo {year} no coincide con el SUBTOTAL del Excel: "
                    f"app={float(kpis.get('total_teorico') or 0):.2f}; excel={objetivo_excel:.2f}"
                )

    if errores:
        raise AssertionError('\n'.join(errores))
    print(f'Auditoría Directorio OK: {len(urls)} APIs y fórmulas derivadas verificadas.')


if __name__ == '__main__':
    verificar()
