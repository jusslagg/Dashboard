"""Readable, editable Excel report for the Directorio screens."""

import io
import unicodedata
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NUMBER = '#,##0.00;[Red](#,##0.00);–'
INTEGER = '#,##0;[Red](#,##0);–'
PERCENT = '0.0%;[Red](0.0%);–'
MONEY = '$ #,##0.00;[Red]($ #,##0.00);–'
FORMATS = {'number': NUMBER, 'integer': INTEGER, 'percent': PERCENT, 'money': MONEY}
MONTHS = ('Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
          'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre')
LABELS = {
    'cliente': 'Cliente', 'campania': 'Campaña', 'fecha': 'Fecha', 'mes': 'Mes',
    'year': 'Año', 'gerente': 'Gerente', 'jefe_site': 'Jefe de site',
    'dotacion_requerida': 'Dotación requerida', 'dotacion_activa': 'Dotación activa',
    'horas_requeridas': 'Horas requeridas', 'horas_realizadas': 'Horas realizadas',
    'horas_objetivo': 'Horas objetivo', 'horas_cumplidas': 'Horas cumplidas',
    'cumplimiento_horas': 'Cumplimiento de horas', 'diferencia_horas': 'Diferencia de horas',
    'dotacion_promedio': 'Dotación promedio', 'rotacion': 'Rotación',
    'agentes_baja': 'Agentes baja', 'bajas_manual': 'Bajas manuales',
    'ausentismo_total': 'Ausentismo total', 'logueados': 'Logueados',
    'personal': 'Personal', 'requerido': 'Requerido', 'realizado': 'Realizado',
    'porcentaje_staff': '% staff', 'site': 'Site', 'industria': 'Industria',
}


def _col(key, label=None, kind=None):
    if kind is None:
        kind = ('percent' if key in {'cumplimiento', 'cumplimiento_horas', 'variacion',
                                    'rotacion', 'abs', 'porcentaje_staff', 'porcentaje_desvio',
                                    'over_requerido', 'ratio', 'ratio_importe'}
                or key.startswith(('share_', 'shares.', 'abs_')) else 'number')
    return key, label or LABELS.get(key) or key.replace('_', ' ').capitalize(), kind


def _table(title, rows, columns, chart=None):
    return {'title': title, 'rows': rows, 'columns': columns, 'chart': chart}


def _safe(value):
    if value is None or isinstance(value, (int, float, bool)):
        return value
    value = str(value)[:32767]
    return "'" + value if value.startswith(('=', '+', '-', '@')) else value


def indicadores_desde_historico(anterior, actual, year):
    """Use the same numerator and denominator as the four Indicadores charts."""
    fields = {
        'cumplimiento': ('horas_realizadas', 'horas_requeridas'),
        'ausentismo': ('horas_ausentismo', 'horas_dotacion_activa'),
        'rotacion': ('bajas', 'dotacion_promedio'),
        'eficiencia': ('pagadas', 'logueo'),
    }
    rows = []
    for current_year, source in ((year - 1, anterior), (year, actual)):
        by_month = defaultdict(list)
        for item in source.get('filas', []):
            by_month[str(item.get('mes') or '')[:7]].append(item)
        for month in range(1, 13):
            monthly = by_month[f'{current_year}-{month:02d}']
            row = {'year': current_year, 'mes': f'{current_year}-{month:02d}'}
            for key, (numerator, denominator) in fields.items():
                top = sum(float(item.get(numerator) or 0) for item in monthly)
                bottom = sum(float(item.get(denominator) or 0) for item in monthly)
                row[key] = top / bottom if bottom else None
            rows.append(row)
    return {'filas': rows}


def comparativo_clientes_desde_fuentes(anterior, actual, dot_anterior, dot_actual, year):
    current_rows = actual.get('filas', [])
    months = [int(str(row.get('mes', ''))[5:7]) for row in current_rows
              if float(row.get('horas_realizadas') or 0) and len(str(row.get('mes', ''))) >= 7]
    cutoff = max(months, default=1)

    def canon(value):
        original = value or ''
        normalized = ''.join(char for char in unicodedata.normalize('NFD', original)
                             if unicodedata.category(char) != 'Mn').lower()
        for prefix, name in (('personal', 'Personal'), ('santander', 'Santander'),
                             ('naturgy', 'Naturgy')):
            if normalized.startswith(prefix):
                return name
        return {'bbva seguros': 'BBVA', 'gire': 'Gire',
                'bsf carrefour': 'Bsf Carrefour', 'galicia seguros': 'Galicia'}.get(normalized, original)

    def hours(data):
        totals = defaultdict(float)
        for row in data.get('filas', []):
            if int(str(row.get('mes', '0000-00'))[5:7]) <= cutoff:
                totals[canon(row.get('cliente'))] += float(row.get('horas_realizadas') or 0)
        return totals

    def staffing(data, month):
        return {canon(row.get('cliente')): float(row.get('meses', {}).get(month) or 0)
                for row in data.get('clientes', [])}

    prev, curr = hours(anterior), hours(actual)
    prev_staff = staffing(dot_anterior, f'{year - 1}-12')
    curr_staff = staffing(dot_actual, f'{year}-{cutoff:02d}')
    rows = []
    for client in set(prev) | set(curr) | set(prev_staff) | set(curr_staff):
        old, new = prev.get(client, 0), curr.get(client, 0)
        rows.append({'cliente': client, 'anterior': old, 'actual': new,
                     'variacion': new / old - 1 if old else (1 if new else None),
                     'dotacion_anterior': prev_staff.get(client, 0),
                     'dotacion_actual': curr_staff.get(client, 0)})
    rows.sort(key=lambda row: (-row['actual'], row['cliente']))
    return {'corte': f'{year}-{cutoff:02d}', 'filas': rows}


def _spec(title, data, year):
    """Return an explanation, key facts, and typed tables for one screen."""
    if title == 'Proyectado vs Facturado':
        monthly = []
        for month in data.get('meses', []):
            items = [row for row in data.get('filas', []) if row['mes'] == month]
            monthly.append({'mes': month, **{key: sum(row.get(key, 0) for row in items)
                                            for key in ('teorico', 'proyectado', 'facturado')}})
        return ('Teórico: objetivo registrado. Proyectado: plan valorizado. Facturado: cierre del dashboard.',
                [('Año', year, 'integer')],
                [_table('Totales mensuales', monthly, [
                    _col('mes', 'Mes', 'text'), _col('teorico', 'Teórico', 'money'),
                    _col('proyectado', 'Proyectado', 'money'),
                    _col('facturado', 'Facturado', 'money')],
                    chart=('bar', 'mes', ('teorico', 'proyectado', 'facturado'))),
                 _table('Detalle por servicio', data.get('filas', []), [
                    _col('mes', 'Mes', 'text'), _col('site', 'Site', 'text'),
                    _col('cliente', 'Cliente', 'text'), _col('servicio', 'Servicio', 'text'),
                    _col('teorico', 'Teórico', 'money'),
                    _col('proyectado', 'Proyectado', 'money'),
                    _col('facturado', 'Facturado', 'money')])])
    if title == 'Reloj':
        item = data['reloj']
        return ('Horas facturadas frente al objetivo. 100% significa que se alcanzó la meta.',
                [('Año', year, 'integer'), ('Cumplimiento', item['cumplimiento'], 'percent'),
                 ('Diferencia de horas', item['diferencia'], 'number')],
                [_table('Horas del período', [item], [
                    _col('requerido', 'Horas requeridas'), _col('realizado', 'Horas realizadas'),
                    _col('diferencia', 'Diferencia'), _col('cumplimiento', 'Cumplimiento', 'percent')]),
                 _table('Comparación de horas', [
                    {'concepto': 'Requeridas', 'horas': item['requerido']},
                    {'concepto': 'Realizadas', 'horas': item['realizado']}],
                    [_col('concepto', 'Concepto', 'text'), _col('horas', 'Horas')],
                    chart=('bar', 'concepto', ('horas',)))])
    if title == 'Indicadores':
        return (f'Ratios mensuales de {year - 1} y {year}. Los porcentajes se calculan con los totales de cada mes.',
                [], [_table('Indicadores por mes', data['filas'], [
                    _col('year', 'Año', 'integer'), _col('mes', 'Mes', 'text'),
                    _col('cumplimiento', 'Cumplimiento de horas', 'percent'),
                    _col('ausentismo', 'Ausentismo', 'percent'),
                    _col('rotacion', 'Rotación', 'percent'),
                    _col('eficiencia', 'Horas pagadas / logueo', 'number')],
                    chart=('line', 'mes', ('cumplimiento',)))])
    if title == 'Dotaciones':
        months = ([data['diciembre_base']] if data.get('diciembre_base') else []) + data.get('meses', [])
        rows = []
        for item in data.get('clientes', []):
            row = {'cliente': item['cliente'], **item.get('meses', {}),
                   'promedio': item.get('promedio'), 'diferencia': item.get('diferencia')}
            if data.get('diciembre_base'):
                row[data['diciembre_base']] = data.get('base_por_cliente', {}).get(item['cliente'])
            rows.append(row)
        rows.append({'cliente': 'TOTAL', **data.get('totales', {})})
        columns = [_col('cliente', 'Cliente', 'text')]
        columns += [_col(month, month, 'number') for month in months]
        columns += [_col('promedio', 'Promedio'), _col('diferencia', 'Dif. vs diciembre')]
        return ('Dotación por cliente y mes. La última fila reúne el total mensual.',
                [('Último mes', data.get('ultimo_mes'), 'text'),
                 ('Base de comparación', data.get('diciembre_base'), 'text')],
                [_table('Dotación mensual', rows, columns)])
    if title == 'Gráficos dotaciones':
        return ('Serie mensual de dotación requerida y activa, con apertura por sede.', [],
                [_table('Datos mensuales', data.get('serie', []), [
                    _col('fecha', 'Mes', 'text'), _col('dotacion_requerida'),
                    _col('dotacion_activa'), _col('activa_sl', 'Activa SL'),
                    _col('activa_ba', 'Activa BA'), _col('diferencia', 'Diferencia'),
                    _col('porcentaje_staff', '% staff', 'percent')],
                    chart=('line', 'fecha', ('dotacion_requerida', 'dotacion_activa')))])
    if title == 'Evolutivo dotaciones':
        return ('Evolución mensual de la dotación requerida y del personal.', [],
                [_table('Serie mensual', data.get('serie', []), [
                    _col('fecha', 'Mes', 'text'), _col('dotacion_requerida'),
                    _col('personal')], chart=('line', 'fecha', ('dotacion_requerida', 'personal')))])
    if title == 'Variación interanual':
        rows = data.get('filas', []) + [{'nombre': 'TOTAL', **data.get('total', {})}]
        return (f'Comparación de horas entre {data.get("anterior")} y {year}. Real y proyección se distinguen en Origen.',
                [('Corte de datos reales', data.get('corte'), 'text'),
                 ('Variación total', data.get('total', {}).get('variacion'), 'percent'),
                 ('Variación acumulada', data.get('acumulado', {}).get('variacion'), 'percent')],
                [_table('Comparación mensual', rows, [
                    _col('nombre', 'Mes', 'text'), _col('anterior', f'Horas {year - 1}'),
                    _col('actual', f'Horas {year}'), _col('tipo', 'Origen', 'text'),
                    _col('diferencia', 'Diferencia'), _col('variacion', 'Variación', 'percent')],
                    chart=('bar', 'nombre', ('anterior', 'actual')))])
    if title == 'Variación horas clientes':
        rows = data.get('filas', []) + [{'cliente': 'TOTAL', **data.get('total', {})}]
        return (f'Horas por cliente: base {year - 1}, real cargado y proyección de {year}.',
                [('Corte real', data.get('corte'), 'text'),
                 ('Meses proyectados', ', '.join(data.get('meses_proyectados', [])), 'text')],
                [_table('Comparación por cliente', rows, [
                    _col('cliente', 'Cliente', 'text'), _col('anterior', f'Horas {year - 1}'),
                    _col('real', 'Real'), _col('proyectado', 'Proyectado'),
                    _col('mixto', 'Real + proyectado'), _col('diferencia', 'Diferencia'),
                    _col('variacion', 'Variación', 'percent'), _col('control', 'Control'),
                    _col('desvio_control', 'Desvío control')])])
    if title == 'Dashboard operativo':
        source = data.get('filas', [])
        rows = [{'id': item.get('id'), 'mes': item.get('mes'),
                 'cliente': item.get('cliente'), 'campania': item.get('campania'),
                 **item.get('datos', {})} for item in source]
        ordered = list(dict.fromkeys(k for k in data.get('columnas', [])
                                     if not k.startswith(('separador_', 'aux_'))))
        keys = list(dict.fromkeys(['id', 'mes', 'cliente', 'campania'] + ordered))
        aux_keys = [k for k in data.get('columnas', []) if k.startswith('aux_')]
        aux_rows = [{key: row.get(key) for key in ['id', 'mes', 'cliente', 'campania'] + aux_keys}
                    for row in rows]
        return ('Detalle operativo por campaña. Las columnas auxiliares se presentan en una tabla aparte.',
                [('Registros', len(rows), 'integer'), ('Año', year, 'integer')],
                [_table('Detalle operativo', rows, [_col(k, kind='integer' if k == 'id' else None) for k in keys]),
                 _table('Columnas auxiliares', aux_rows,
                        [_col(k, kind='integer' if k == 'id' else None)
                         for k in ['id', 'mes', 'cliente', 'campania'] + aux_keys])])
    if title == 'Histórico clientes':
        keys = ['empresa', 'site', 'industria', 'cliente', 'mes',
                'horas_dotacion_activa', 'horas_ausentismo', 'dotacion_promedio',
                'bajas', 'horas_requeridas', 'horas_realizadas', 'pagadas', 'logueo',
                'cumplimiento_horas', 'abs', 'rotacion', 'eficiencia', 'dotacion_requerida']
        return ('Base mensual por cliente que alimenta los indicadores del Directorio.',
                [('Registros', len(data.get('filas', [])), 'integer'), ('Año', year, 'integer')],
                [_table('Detalle mensual', data.get('filas', []),
                        [_col(k, kind='text' if k in keys[:5] else None) for k in keys])])
    if title == 'Cumplimiento horas':
        totals = defaultdict(lambda: {'objetivo': 0.0, 'realizado': 0.0})
        for item in data.get('filas', []):
            row = totals[item.get('cliente') or '']
            row['objetivo'] += float(item.get('horas_requeridas') or 0)
            row['realizado'] += float(item.get('horas_realizadas') or 0)
        summary = []
        for client, values in totals.items():
            target, done = values['objetivo'], values['realizado']
            summary.append({'cliente': client, 'objetivo': target, 'realizado': done,
                            'diferencia': done - target,
                            'cumplimiento': done / target if target else None})
        summary.sort(key=lambda row: -row['objetivo'])
        return ('Resumen anual por cliente y, debajo, detalle de cada mes cargado.',
                [('Clientes', len(summary), 'integer'), ('Año', year, 'integer')],
                [_table('Resumen por cliente', summary, [
                    _col('cliente', 'Cliente', 'text'), _col('objetivo', 'Horas objetivo'),
                    _col('realizado', 'Horas realizadas'), _col('diferencia', 'Diferencia'),
                    _col('cumplimiento', 'Cumplimiento', 'percent')]),
                 _table('Detalle mensual', data.get('filas', []), [
                    _col('cliente', 'Cliente', 'text'), _col('mes', 'Mes', 'text'),
                    _col('horas_requeridas'), _col('horas_realizadas')])])
    if title == 'Cumplimiento facturación':
        money_keys = ('horas', 'variable', 'penalizaciones_bonos', 'total')
        cols = [_col('cliente', 'Cliente', 'text')]
        for group in ('desvio', 'objetivo', 'alcanzado'):
            cols.extend(_col(f'{group}_{key}', f'{group.capitalize()} · {key.replace("_", " ")}', 'money')
                        for key in money_keys)
        cols.append(_col('porcentaje_desvio', '% desvío', 'percent'))
        annual = [_col('year', 'Año', 'integer')] + [
            _col(key, key.replace('_', ' ').capitalize(), 'money')
            for key in ('horas', 'variable', 'bonos_penalizaciones', 'total')]
        annual.append(_col('porcentaje_desvio', '% desvío', 'percent'))
        return ('Desvío de facturación por cliente. Objetivo y alcanzado se expresan en pesos.',
                [('Año', year, 'integer')],
                [_table('Control por cliente', data.get('filas', []), cols),
                 _table('Control anual', data.get('resumen_anual', []), annual)])
    if title == 'Comparativo interanual':
        quarterly = []
        for label, payload in ((str(year - 1), data.get('anterior_datos', {})),
                               (str(year), data.get('actual', {}))):
            quarterly.append({'year': label, 'periodo': 'Anual', **payload.get('anual', {})})
            quarterly.extend({'year': label, 'periodo': f'Trimestre {index}', **quarter}
                             for index, quarter in enumerate(payload.get('trimestres', []), 1))
        cols = [_col('year', 'Año', 'integer'), _col('periodo', 'Período', 'text'),
                _col('cumplimiento_facturacion', 'Cumplimiento facturación', 'percent'),
                _col('cumplimiento_horas', 'Cumplimiento horas', 'percent'),
                _col('desvio_facturacion', 'Desvío facturación', 'money')]
        return (f'Comparación de cumplimiento y desvío entre {year - 1} y {year}.', [],
                [_table('Control anual y trimestral', quarterly, cols),
                 _table('Control mensual', data.get('meses', []), [
                    _col('nombre', 'Mes', 'text'),
                    _col('cumplimiento_anterior', f'Cumplimiento {year - 1}', 'percent'),
                    _col('cumplimiento_actual', f'Cumplimiento {year}', 'percent'),
                    _col('desvio_anterior', f'Desvío {year - 1}', 'money'),
                    _col('desvio_actual', f'Desvío {year}', 'money'),
                    _col('cargado', 'Real cargado', 'text')],
                    chart=('line', 'nombre', ('cumplimiento_anterior', 'cumplimiento_actual')))])
    if title == 'Comparativo anual clientes':
        return (f'Horas realizadas acumuladas hasta {data.get("corte")}; dotación al cierre de cada período.',
                [('Corte', data.get('corte'), 'text')],
                [_table('Comparación por cliente', data.get('filas', []), [
                    _col('cliente', 'Cliente', 'text'), _col('anterior', f'Horas {year - 1}'),
                    _col('actual', f'Horas {year}'), _col('variacion', 'Variación', 'percent'),
                    _col('dotacion_anterior', f'Dotación dic {year - 1}'),
                    _col('dotacion_actual', f'Dotación {data.get("corte")}')])])
    if title in ('Ratio Staff Operaciones 1', 'Ratio Eli'):
        return ('Relación mensual entre staff y dotación requerida.', [],
                [_table('Serie mensual', data.get('serie', []), [
                    _col('mes', 'Mes', 'text'), _col('requerido', 'Requerido'),
                    _col('staff', 'Staff'), _col('ratio', 'Ratio', 'percent')],
                    chart=('line', 'mes', ('ratio',)))])
    if title in ('Ratio Staff Operaciones 2', 'Ratio Eli II'):
        return ('Operaciones y staff, en dotación e importe, con sus ratios mensuales.', [],
                [_table('Serie mensual', data.get('serie', []), [
                    _col('mes', 'Mes', 'text'), _col('operaciones', 'Operaciones'),
                    _col('staff', 'Staff'), _col('total', 'Total'),
                    _col('ratio', 'Ratio', 'percent'),
                    _col('operaciones_importe', 'Importe operaciones', 'money'),
                    _col('staff_importe', 'Importe staff', 'money'),
                    _col('total_importe', 'Importe total', 'money'),
                    _col('ratio_importe', 'Ratio importe', 'percent')],
                    chart=('line', 'mes', ('ratio', 'ratio_importe')))])
    if title == 'Share':
        columns = [_col('year', 'Año', 'integer')]
        columns += [_col(key, key, 'number') for key in ('Personal', 'Santander', 'Multicampania', 'USA', 'Total')]
        columns += [_col(f'shares.{key}', f'Share {key}', 'percent')
                    for key in ('Personal', 'Santander', 'Multicampania', 'USA')]
        columns += [_col('share_personal_santander', 'Share Personal + Santander', 'percent')]
        def flat(rows):
            return [{**row, **{f'shares.{key}': value for key, value in row.get('shares', {}).items()}}
                    for row in rows]
        return ('Participación anual por negocio. Horas y dotaciones se muestran por separado.', [],
                [_table('Share de horas', flat(data.get('horas', [])), columns),
                 _table('Share de dotaciones', flat(data.get('dotaciones', [])), columns)])
    raise ValueError(f'Pantalla sin definición de exportación: {title}')


def _write_sheet(workbook, title, data, year, index):
    description, facts, tables = _spec(title, data, year)
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.tabColor = '2563EB'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.orientation = 'landscape'
    sheet.merge_cells('A1:H2')
    heading = sheet['A1']
    heading.value = title
    heading.fill = PatternFill('solid', fgColor='0F172A')
    heading.font = Font(name='Aptos Display', size=18, bold=True, color='FFFFFF')
    heading.alignment = Alignment(vertical='center', indent=1)
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 15
    sheet.merge_cells('A3:H3')
    sheet['A3'] = description
    sheet['A3'].font = Font(name='Aptos', size=10, color='334155')
    sheet['A3'].alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[3].height = 32
    sheet['A4'] = f'Período seleccionado: {year}'
    sheet['A4'].font = Font(name='Aptos', size=10, color='64748B', italic=True)
    row = 6
    if facts:
        sheet.cell(row, 1, 'RESUMEN')
        sheet.cell(row, 1).font = Font(bold=True, color='1D4ED8')
        row += 1
        for label, value, kind in facts:
            sheet.cell(row, 1, label).font = Font(bold=True, color='334155')
            cell = sheet.cell(row, 2, _safe(value))
            cell.font = Font(bold=True, color='0F172A')
            cell.number_format = FORMATS.get(kind, 'General')
            row += 1
        row += 1
    first_header = None
    for table_index, table in enumerate(tables, 1):
        columns = table['columns']
        if not columns:
            continue
        sheet.cell(row, 1, table['title']).font = Font(size=12, bold=True, color='1E3A8A')
        row += 1
        header_row = row
        if first_header is None:
            first_header = header_row
        for col_number, (_, label, _) in enumerate(columns, 1):
            cell = sheet.cell(row, col_number, label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1E3A8A')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            cell.border = None
            sheet.column_dimensions[get_column_letter(col_number)].width = min(
                max(sheet.column_dimensions[get_column_letter(col_number)].width or 0,
                    len(label) + 3, 14), 31)
        sheet.row_dimensions[row].height = 30
        for item in table['rows']:
            row += 1
            for col_number, (key, _, kind) in enumerate(columns, 1):
                cell = sheet.cell(row, col_number, _safe(item.get(key)))
                cell.number_format = FORMATS.get(kind, 'General')
                cell.alignment = Alignment(vertical='center', horizontal='right' if kind != 'text' else 'left')
                if item.get('cliente') == 'TOTAL' or item.get('nombre') == 'TOTAL':
                    cell.font = Font(bold=True, color='0F172A')
                    cell.fill = PatternFill('solid', fgColor='DBEAFE')
        if table['rows']:
            ref = f'A{header_row}:{get_column_letter(len(columns))}{row}'
            excel_table = Table(displayName=f'Directorio{index}Tabla{table_index}', ref=ref)
            excel_table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(excel_table)
        if table.get('chart') and table['rows']:
            kind, category_key, value_keys = table['chart']
            chart = LineChart() if kind == 'line' else BarChart()
            chart.title = table['title']
            chart.height, chart.width = 8, 15
            chart.style = 13
            keys = [key for key, _, _ in columns]
            for key in value_keys:
                if key in keys:
                    col_number = keys.index(key) + 1
                    chart.add_data(Reference(sheet, min_col=col_number, min_row=header_row,
                                             max_row=row), titles_from_data=True)
            category_col = keys.index(category_key) + 1
            chart.set_categories(Reference(sheet, min_col=category_col,
                                           min_row=header_row + 1, max_row=row))
            position = f'H{header_row}' if len(columns) <= 6 and table_index == 1 else f'A{row + 3}'
            sheet.add_chart(chart, position)
            if position.startswith('A'):
                row += 19
        row += 3
    sheet.freeze_panes = (f'E{first_header + 1}' if title == 'Dashboard operativo'
                          else f'C{first_header + 1}') if first_header else 'A5'
    sheet.column_dimensions['A'].width = max(sheet.column_dimensions['A'].width or 0, 22)
    sheet.column_dimensions['B'].width = max(sheet.column_dimensions['B'].width or 0, 18)
    return sheet


def crear_libro_directorio(screens, year):
    """Create one usable worksheet per Directorio screen, with native Excel data."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, (title, data) in enumerate(screens, 1):
        _write_sheet(workbook, title, data, year, index)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
