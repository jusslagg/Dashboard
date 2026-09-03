# filepath: app/routes.py
from flask import Blueprint, Response, current_app, redirect, render_template, request, jsonify, send_file, session, url_for
from app import db, get_csrf_token
from app.models import AsignacionComercial, Campania, DashboardOperativo, DotacionClienteMensual, DotacionMensual, ExcepcionCalculo, Facturacion2026, FeriadoOperativo, GraficoDotacionMensual, HistorialCambio, HistoricoClienteMensual, JustificacionAjuste, NextGenDolar, NextGenProducto, PersonalDistribucionHoras, ProyeccionMatriz, ProyeccionMatrizJornada, ProyeccionPrecio, RatioEliMensual, RatioEliIIMensual, ROLES_USUARIO, PUESTOS_POR_ROL, permisos_perfil, SiteProyeccion, TarifacionCampania, Usuario, VariableCampania, redondear_moneda
from datetime import date, datetime, timedelta
import calendar
from sqlalchemy import func, or_
from html import escape
from html.parser import HTMLParser
from functools import wraps
from difflib import SequenceMatcher
import csv
import io
import json
import os
import re
import textwrap
import time
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

main_bp = Blueprint('main', __name__)
ADMIN_KEY = os.getenv('ADMIN_ACTION_KEY')
CLAVE_PRUEBA_ACCIONES = os.getenv('DEV_ACTION_KEY')
LOGIN_ATTEMPTS = {}
LOGIN_WINDOW_SECONDS = int(os.getenv('LOGIN_WINDOW_SECONDS', '900'))
LOGIN_MAX_ATTEMPTS = int(os.getenv('LOGIN_MAX_ATTEMPTS', '5'))
TIPO_VH_PERSONAL_COBRANZAS = 'Personal Cobranzas'
TIPOS_VH = [
    'Diurna',
    'Nocturna',
    'Feriado',
    'Capacitación',
    'Diurnas Feriado',
    'Nocturnas Feriado',
    'horas líder',
    'radio',
    TIPO_VH_PERSONAL_COBRANZAS,
]

COLUMNAS_IMPORTACION = [
    ('cliente', 'Cuenta'),
    ('subcampania', 'Sub campaña'),
    ('tipo_jornada', 'Tipo VH'),
    ('mes', 'Mes'),
    ('horas_objetivo', 'Objetivo Horas'),
    ('horas_facturadas', 'Horas Facturadas'),
    ('valor_hora', 'Valor Hora'),
    ('facturado_horas_manual', 'Facturado Horas'),
    ('tarifacion', 'Tarificacion'),
    ('unidad', 'Unidad'),
    ('porcentaje_variable', '% Variable'),
    ('variable_productivo', 'Variable productivo'),
    ('ajuste_bono_penalizacion', 'Penalizaciones bonos'),
    ('netx_gen', 'Facturacion Next Gen'),
    ('otros', 'Otros'),
    ('total_facturado_informado', 'Total Facturado'),
]

ALIAS_IMPORTACION = {
    'fecha': 'fecha',
    'fecha de carga': 'fecha',
    'mes': 'mes',
    'mes facturacion': 'mes',
    'mes facturación': 'mes',
    'cliente': 'cliente',
    'cuenta': 'cliente',
    'gerente': 'gerente',
    'jefe de site': 'jefe_site',
    'jefe_site': 'jefe_site',
    'campana': 'campania',
    'campaña': 'campania',
    'campania': 'campania',
    'sub campana': 'subcampania',
    'sub campaña': 'subcampania',
    'subcampania': 'subcampania',
    'tipo negocio': 'tipo_negocio',
    'tipo de negocio': 'tipo_negocio',
    'tipo_negocio': 'tipo_negocio',
    'negocio': 'tipo_negocio',
    'tipo de vh': 'tipo_jornada',
    'tipo vh': 'tipo_jornada',
    'tipo_jornada': 'tipo_jornada',
    'horas objetivo': 'horas_objetivo',
    'objetivo horas': 'horas_objetivo',
    'horas_objetivo': 'horas_objetivo',
    'horas facturadas': 'horas_facturadas',
    'horas_facturadas': 'horas_facturadas',
    'horas penalizadas': 'horas_penalizadas',
    'horas penalizacion adh': 'horas_penalizadas',
    'horas penalización adh': 'horas_penalizadas',
    'penalizacion adh horas': 'horas_penalizadas',
    'penalización adh horas': 'horas_penalizadas',
    'horas_penalizadas': 'horas_penalizadas',
    'valor hora objetivo': 'valor_hora_objetivo',
    'valor_hora_objetivo': 'valor_hora_objetivo',
    'valor hora facturado': 'valor_hora',
    'valor hora': 'valor_hora',
    'valor_hora': 'valor_hora',
    'facturado en horas': 'facturado_horas_manual',
    'facturado horas': 'facturado_horas_manual',
    'facturado_horas_manual': 'facturado_horas_manual',
    'tarifacion': 'tarifacion',
    'tarificacion': 'tarifacion',
    'tarifacion adicional': 'tarifacion',
    'tarificacion adicional': 'tarifacion',
    'importe fijo': 'importe_fijo',
    'importe fijo facturado': 'importe_fijo',
    'tarifa plana': 'importe_fijo',
    'facturacion fija': 'importe_fijo',
    'variable objetivo': 'variable_objetivo',
    'variable_objetivo': 'variable_objetivo',
    'bono objetivo': 'variable_objetivo',
    'objetivo facturacion bono': 'variable_objetivo',
    'objetivo facturación bono': 'variable_objetivo',
    'facturacion objetivo': 'control_objetivo_total',
    'facturación objetivo': 'control_objetivo_total',
    'variable productivo': 'variable_productivo',
    'variable_productivo': 'variable_productivo',
    'tarifación': 'tarifacion',
    'bonos': 'bonos',
    'bono': 'bonos',
    'facturado bono': 'variable_productivo',
    'penalizaciones': 'penalizaciones',
    'penalizacion': 'penalizaciones',
    'penalizaciones bonos': 'ajuste_bono_penalizacion',
    'penalizaciones/bonos': 'ajuste_bono_penalizacion',
    'penalizaciones por incumplimientos': 'penalizaciones',
    'penalizacion por incumplimientos': 'penalizaciones',
    'netx gen': 'netx_gen',
    'next gen': 'netx_gen',
    'facturacion next gen': 'netx_gen',
    'facturación next gen': 'netx_gen',
    'netx_gen': 'netx_gen',
    'otros': 'otros',
    'unidad': 'unidad',
    '% variable': 'porcentaje_variable',
    'porcentaje variable': 'porcentaje_variable',
    'total facturado': 'total_facturado_informado',
}

class TablaHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.current_row = None
        self.current_cell = None
        self.in_cell = False

    def handle_starttag(self, tag, attrs):
        if tag == 'tr':
            self.current_row = []
        elif tag in ('td', 'th') and self.current_row is not None:
            self.current_cell = []
            self.in_cell = True

    def handle_data(self, data):
        if self.in_cell and self.current_cell is not None:
            self.current_cell.append(data)

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.in_cell:
            self.current_row.append(''.join(self.current_cell).strip())
            self.current_cell = None
            self.in_cell = False
        elif tag == 'tr' and self.current_row is not None:
            if any(cell for cell in self.current_row):
                self.rows.append(self.current_row)
            self.current_row = None


def crear_xlsx(headers, rows):
    def columna_excel(indice):
        nombre = ''
        indice += 1
        while indice:
            indice, resto = divmod(indice - 1, 26)
            nombre = chr(65 + resto) + nombre
        return nombre

    def fila_xml(valores, numero):
        celdas = []
        for indice, valor in enumerate(valores):
            referencia = f'{columna_excel(indice)}{numero}'
            texto = xml_escape(str(valor if valor is not None else ''))
            celdas.append(f'<c r="{referencia}" t="inlineStr"><is><t>{texto}</t></is></c>')
        return f'<row r="{numero}">' + ''.join(celdas) + '</row>'

    filas_xml = [fila_xml(headers, 1)]
    filas_xml.extend(fila_xml(row, numero) for numero, row in enumerate(rows, start=2))
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData>' + ''.join(filas_xml) + '</sheetData></worksheet>'
    )

    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archivo:
        archivo.writestr('[Content_Types].xml', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '</Types>'
        ))
        archivo.writestr('_rels/.rels', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>'
        ))
        archivo.writestr('xl/workbook.xml', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Carga" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ))
        archivo.writestr('xl/_rels/workbook.xml.rels', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '</Relationships>'
        ))
        archivo.writestr('xl/worksheets/sheet1.xml', sheet_xml)
    output.seek(0)
    return output.getvalue()


def validar_clave(data):
    clave = str((data or {}).get('clave') or '')
    return requiere_edicion() or (ADMIN_KEY and clave == ADMIN_KEY)


def validar_confirmacion_accion(data):
    data = data or {}
    usuario = usuario_actual()
    password = str(data.get('password_confirmacion') or '')
    clave = str(data.get('clave') or '')
    if usuario and password and usuario.check_password(password):
        return True
    if current_app.debug and CLAVE_PRUEBA_ACCIONES and clave == CLAVE_PRUEBA_ACCIONES:
        return True
    return False


def serializar_json(valor):
    return json.dumps(valor, ensure_ascii=False, default=str)


def snapshot_modelo(modelo):
    if hasattr(modelo, 'to_dict'):
        return modelo.to_dict()
    return {}


def cambios_entre(antes, despues):
    cambios = {}
    for clave in sorted(set(antes) | set(despues)):
        if clave == 'justificaciones':
            continue
        valor_antes = antes.get(clave)
        valor_despues = despues.get(clave)
        if valor_antes != valor_despues:
            cambios[clave] = {'antes': valor_antes, 'despues': valor_despues}
    return cambios


def registrar_historial(accion, entidad, entidad_id, resumen, antes=None, despues=None, detalle=None):
    usuario = usuario_actual()
    item = HistorialCambio(
        usuario_id=usuario.id if usuario else None,
        usuario_nombre=usuario.nombre if usuario else 'Sistema',
        usuario_email=usuario.email if usuario else None,
        accion=accion,
        entidad=entidad,
        entidad_id=str(entidad_id) if entidad_id is not None else None,
        resumen=resumen,
        detalle=detalle,
        antes=serializar_json(antes) if antes is not None else None,
        despues=serializar_json(despues) if despues is not None else None,
    )
    db.session.add(item)
    return item


def usuario_actual():
    usuario_id = session.get('usuario_id')
    if not usuario_id:
        return None
    return Usuario.query.get(usuario_id)


def usuario_actual_dict():
    usuario = usuario_actual()
    return usuario.to_dict() if usuario else None


def requiere_login():
    usuario = usuario_actual()
    return usuario is not None and usuario.activo


def requiere_admin():
    usuario = usuario_actual()
    return usuario is not None and usuario.activo and usuario.puede_administrar_usuarios


def requiere_edicion():
    usuario = usuario_actual()
    return usuario is not None and usuario.activo and usuario.puede_editar


def requiere_carga():
    usuario = usuario_actual()
    return usuario is not None and usuario.activo and usuario.puede_cargar


def requiere_eliminacion():
    usuario = usuario_actual()
    return usuario is not None and usuario.activo and usuario.puede_eliminar


def respuesta_no_autorizado():
    if request.path.startswith('/api/'):
        return jsonify({'success': False, 'errores': ['No autorizado']}), 403
    if requiere_login():
        return redirect(url_for('main.index'))
    return redirect(url_for('main.login'))


def login_requerido(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        if not usuarios_registrados():
            return respuesta_no_autorizado()
        if not requiere_login():
            return respuesta_no_autorizado()
        usuario = usuario_actual()
        if usuario.debe_cambiar_password:
            if request.path.startswith('/api/'):
                return jsonify({'success': False, 'requiere_cambio_password': True, 'errores': ['Debe cambiar su contrasena temporal']}), 428
            return redirect(url_for('main.cambiar_password_obligatorio'))
        return func(*args, **kwargs)
    return wrapper


def permiso_requerido(verificador):
    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            if not usuarios_registrados():
                return respuesta_no_autorizado()
            if not verificador():
                return respuesta_no_autorizado()
            usuario = usuario_actual()
            if usuario.debe_cambiar_password:
                if request.path.startswith('/api/'):
                    return jsonify({'success': False, 'requiere_cambio_password': True, 'errores': ['Debe cambiar su contrasena temporal']}), 428
                return redirect(url_for('main.cambiar_password_obligatorio'))
            return func(*args, **kwargs)
        return wrapper
    return decorator


admin_requerido = permiso_requerido(requiere_admin)
carga_requerida = permiso_requerido(requiere_carga)
edicion_requerida = permiso_requerido(requiere_edicion)
eliminacion_requerida = permiso_requerido(requiere_eliminacion)


def edicion_si_mutacion_requerida(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        verificador = requiere_login if request.method in ('GET', 'HEAD', 'OPTIONS') else requiere_edicion
        if not verificador():
            return respuesta_no_autorizado()
        usuario = usuario_actual()
        if usuario.debe_cambiar_password:
            return respuesta_no_autorizado()
        return func(*args, **kwargs)
    return wrapper


RUTAS_DIRECTORIO = (
    '/reloj', '/indicadores', '/dotaciones', '/graficos-dotaciones',
    '/graficos-evolutivo-dotaciones', '/variacion-anual', '/variacion-horas-clientes',
    '/dashboard-operativo', '/historico-clientes', '/cumplimiento-horas-clientes',
    '/cumplimiento-facturacion-clientes', '/comparativo-interanual',
    '/comparativo-anual-clientes', '/ratio-eli', '/ratio-eli-ii', '/share',
    '/api/reloj', '/api/indicadores', '/api/dotaciones', '/api/graficos-dotaciones',
    '/api/graficos-evolutivo-dotaciones', '/api/variacion-anual',
    '/api/variacion-horas-clientes', '/api/dashboard-operativo',
    '/api/historico-clientes', '/api/cumplimiento-', '/api/comparativo-interanual',
    '/api/comparativo-anual-clientes', '/api/ratio-eli', '/api/ratio-eli-ii', '/api/share',
)
RUTAS_PROYECCIONES = (
    '/matriz-proyecciones', '/matriz-precios', '/suma-fija', '/resumen', '/variable',
    '/control-proyecciones', '/calendario-operativo', '/api/proyeccion',
    '/api/matriz-proyecciones', '/api/precios', '/api/suma-fija', '/api/variable',
    '/api/control-proyecciones', '/api/calendario-operativo', '/api/feriados',
)
RUTAS_PUBLICAS_AUTENTICADAS = (
    '/login', '/logout', '/cambiar-password', '/api/auth/', '/api/usuarios/setup',
    '/guia-usuario', '/api/sesion', '/favicon.ico',
)


def modulo_solicitado(path):
    if any(path.startswith(prefijo) for prefijo in RUTAS_DIRECTORIO):
        return 'directorio'
    if any(path.startswith(prefijo) for prefijo in RUTAS_PROYECCIONES):
        return 'proyecciones'
    if path.startswith(('/api/asignaciones', '/api/campanias')):
        return 'maestros_api'
    if path.startswith(('/usuarios', '/api/usuarios', '/catalogos', '/documentacion-tecnica')):
        return 'usuarios'
    if path.startswith('/historial') or path.startswith('/api/historial'):
        return 'historial'
    if path.startswith(RUTAS_PUBLICAS_AUTENTICADAS):
        return None
    return 'facturacion'


@main_bp.before_request
def proteger_modulos_por_perfil():
    usuario = usuario_actual()
    if not usuario or not usuario.activo or usuario.debe_cambiar_password:
        return None
    modulo = modulo_solicitado(request.path)
    autorizado = (
        modulo is None
        or (modulo == 'facturacion' and usuario.puede_ver_facturacion)
        or (modulo == 'directorio' and usuario.puede_ver_directorio)
        or (modulo == 'proyecciones' and usuario.puede_ver_proyecciones)
        or (modulo == 'maestros_api' and (usuario.puede_administrar_usuarios or usuario.puede_cargar))
        or (modulo in ('usuarios', 'historial') and usuario.puede_administrar_usuarios)
    )
    if autorizado:
        return None
    if request.path == '/':
        if usuario.puede_ver_directorio:
            return redirect(url_for('main.reloj'))
        if usuario.puede_ver_proyecciones:
            return redirect(url_for('main.matriz_proyecciones'))
        # Un perfil válido puede no tener módulos de negocio asignados todavía.
        # En ese caso conserva el acceso a Ayuda en lugar de recibir un 403 al
        # completar correctamente el inicio de sesión.
        return redirect(url_for('main.guia_usuario'))
    if request.path.startswith('/api/'):
        return jsonify({'success': False, 'errores': ['El perfil no tiene acceso a este módulo']}), 403
    return Response('Acceso no autorizado para el perfil asignado.', status=403)


def usuarios_registrados():
    return Usuario.query.count() > 0


def asegurar_administrador_inicial():
    """Promueve el primer usuario existente a administrador si aun no hay admin."""
    if not usuarios_registrados():
        return None
    if Usuario.query.filter(Usuario.rol.in_(('Admin', 'administrador'))).first():
        return None
    primer_usuario = Usuario.query.order_by(Usuario.creado_en.asc(), Usuario.id.asc()).first()
    if primer_usuario:
        rol_anterior = primer_usuario.rol
        primer_usuario.rol = 'Admin'
        registrar_historial(
            'edicion',
            'usuario',
            primer_usuario.id,
            f'Usuario inicial promovido a administrador: {primer_usuario.email}',
            antes={'rol': rol_anterior},
            despues={'rol': 'administrador'},
        )
        db.session.commit()
    return primer_usuario


def normalizar_email(email):
    return str(email or '').strip().lower()


def login_rate_key(email):
    forwarded = request.headers.get('X-Forwarded-For', '')
    ip = forwarded.split(',')[0].strip() or request.remote_addr or 'unknown'
    return f'{ip}:{normalizar_email(email)}'


def login_bloqueado(email):
    key = login_rate_key(email)
    now = time.time()
    attempts = [item for item in LOGIN_ATTEMPTS.get(key, []) if now - item < LOGIN_WINDOW_SECONDS]
    LOGIN_ATTEMPTS[key] = attempts
    return len(attempts) >= LOGIN_MAX_ATTEMPTS


def registrar_login_fallido(email):
    key = login_rate_key(email)
    now = time.time()
    attempts = [item for item in LOGIN_ATTEMPTS.get(key, []) if now - item < LOGIN_WINDOW_SECONDS]
    attempts.append(now)
    LOGIN_ATTEMPTS[key] = attempts


def limpiar_login_fallido(email):
    LOGIN_ATTEMPTS.pop(login_rate_key(email), None)


def validar_usuario_payload(data, require_password=False):
    errores = []
    nombre = str(data.get('nombre', '')).strip()
    email = normalizar_email(data.get('email'))
    password = str(data.get('password', '') or '')
    rol = str(data.get('rol', 'usuario')).strip()
    puesto = str(data.get('puesto', '') or '').strip()

    if not nombre:
        errores.append('El nombre es obligatorio')
    if not email or '@' not in email:
        errores.append('El email es obligatorio y debe ser valido')
    if require_password and len(password) < 12:
        errores.append('La contrasena debe tener al menos 12 caracteres')
    if require_password and password and (
        password.lower() == password
        or password.upper() == password
        or not any(caracter.isdigit() for caracter in password)
    ):
        errores.append('La contrasena debe combinar mayusculas, minusculas y numeros')
    if rol not in ROLES_USUARIO:
        errores.append('El rol no es valido')
    elif puesto and puesto not in PUESTOS_POR_ROL.get(rol, ()):
        errores.append('El puesto no corresponde al perfil seleccionado')
    elif rol != 'Full' and not puesto:
        errores.append('El puesto es obligatorio')
    return errores


MODULOS_PERMISO = {
    'facturacion_total', 'facturacion_asignada', 'rmo_santander', 'rmo_multi_co',
    'rmo_multi_sq', 'rmo_personal', 'directorio', 'proyecciones',
}


def normalizar_permisos_usuario(valor):
    if valor is None:
        return None
    if not isinstance(valor, dict):
        raise ValueError('Los permisos deben enviarse como un objeto')
    return {
        'visualizar': bool(valor.get('visualizar', True)),
        'cargar': bool(valor.get('cargar', False)),
        'editar': bool(valor.get('editar', False)),
        'eliminar': bool(valor.get('eliminar', False)),
        'administrar_perfiles': bool(valor.get('administrar_perfiles', False)),
        'modulos': sorted(set(valor.get('modulos') or ()) & MODULOS_PERMISO),
    }


def filtrar_valores_exactos(query, columna, valores):
    if isinstance(valores, list):
        return query.filter(columna.in_(valores))
    return query.filter(columna == valores)


def aplicar_filtros(
    query,
    year=None,
    mes=None,
    cliente=None,
    gerente=None,
    jefe_site=None,
    campania=None,
    subcampania=None,
    tipo_negocio=None,
    horas_objetivo_min=None,
    horas_objetivo_max=None,
    horas_facturadas_min=None,
    horas_facturadas_max=None,
):
    usuario = usuario_actual()
    if usuario and usuario.activo and not usuario.puede_acceder('facturacion_total'):
        # El gerente es el nivel jerárquico superior: al seleccionarlo se incluyen
        # automáticamente todos sus jefes, clientes y campañas de Datos Maestros.
        if usuario.gerente_asignado:
            query = query.filter(func.lower(Facturacion2026.gerente) == usuario.gerente_asignado.lower())
        elif usuario.jefe_site_asignado:
            query = query.filter(func.lower(Facturacion2026.jefe_site) == usuario.jefe_site_asignado.lower())
    if year:
        years = year if isinstance(year, list) else [year]
        condiciones = [Facturacion2026.mes.like(f'{valor}-%') for valor in years]
        query = query.filter(or_(*condiciones))
    if mes:
        if isinstance(mes, list):
            query = query.filter(Facturacion2026.mes.in_(mes))
        else:
            query = query.filter(Facturacion2026.mes == mes)
    if cliente:
        query = filtrar_valores_exactos(query, Facturacion2026.cliente, cliente)
    if gerente:
        query = filtrar_valores_exactos(query, Facturacion2026.gerente, gerente)
    if jefe_site:
        query = filtrar_valores_exactos(query, Facturacion2026.jefe_site, jefe_site)
    if campania:
        query = filtrar_valores_exactos(query, Facturacion2026.campania, campania)
    if subcampania:
        query = filtrar_valores_exactos(query, Facturacion2026.subcampania, subcampania)
    if tipo_negocio:
        query = filtrar_valores_exactos(query, Facturacion2026.tipo_negocio, tipo_negocio)
    if horas_objetivo_min is not None:
        query = query.filter(Facturacion2026.horas_objetivo >= horas_objetivo_min)
    if horas_objetivo_max is not None:
        query = query.filter(Facturacion2026.horas_objetivo <= horas_objetivo_max)
    if horas_facturadas_min is not None:
        query = query.filter(Facturacion2026.horas_facturadas >= horas_facturadas_min)
    if horas_facturadas_max is not None:
        query = query.filter(Facturacion2026.horas_facturadas <= horas_facturadas_max)
    return query


def numero_request(nombre):
    valor = request.args.get(nombre)
    if valor in (None, ''):
        return None
    try:
        return float(valor)
    except ValueError:
        return None


def mes_valido(valor):
    if not valor:
        return None
    texto = str(valor).strip()
    if texto.replace('.', '', 1).isdigit():
        return fecha_excel(float(texto)).strftime('%Y-%m')
    try:
        return datetime.strptime(texto, '%Y-%m').strftime('%Y-%m')
    except ValueError:
        pass
    for formato in ('%m/%Y', '%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(texto, formato).strftime('%Y-%m')
        except ValueError:
            continue
    return None


def normalizar_header(valor):
    texto = str(valor or '').strip().lower().replace('_', ' ')
    texto = ''.join(
        caracter for caracter in unicodedata.normalize('NFD', texto)
        if unicodedata.category(caracter) != 'Mn'
    )
    return ' '.join(texto.split())


def normalizar_site(valor):
    site = str(valor or '').strip()
    clave = normalizar_header(site).replace(' ', '')
    if 'multicuent' in clave or 'multicampan' in clave:
        return 'Gerencia Multicampaña'
    return site


def jefe_personal_por_servicio(*valores):
    texto = normalizar_header(' '.join(str(valor or '') for valor in valores))
    if 'soporte' in texto or 'smb' in texto:
        return 'Walter Canalini'
    if 'ventas' in texto:
        return 'Pablo Chanampa'
    if 'retencion' in texto:
        return 'Alejandro Del Soto'
    return 'Micaela Antelo'


def normalizar_tipo_vh(valor):
    texto = str(valor or '').strip()
    clave = normalizar_header(texto)
    equivalencias = {
        'diurna': 'Diurna', 'diurnas': 'Diurna', 'diurno': 'Diurna', 'horas diurnas': 'Diurna',
        'horas': 'Diurna',
        'nocturna': 'Nocturna', 'nocturnas': 'Nocturna', 'nocturno': 'Nocturna', 'horas nocturnas': 'Nocturna',
        'feriado': 'Feriado', 'feriados': 'Feriado',
        'horas feriado': 'Feriado',
        'capacitacion': 'Capacitación', 'capacitaciones': 'Capacitación',
        'diurna feriado': 'Diurnas Feriado', 'diurnas feriado': 'Diurnas Feriado',
        'diurna f': 'Diurnas Feriado', 'diurnas f': 'Diurnas Feriado',
        'nocturna feriado': 'Nocturnas Feriado', 'nocturnas feriado': 'Nocturnas Feriado',
        'nocturna f': 'Nocturnas Feriado', 'nocturnas f': 'Nocturnas Feriado',
        'horas lider': 'horas líder', 'hora lider': 'horas líder',
        'radio': 'radio',
        'personal cobranzas': TIPO_VH_PERSONAL_COBRANZAS,
    }
    if clave in equivalencias:
        return equivalencias[clave]
    return texto


def facturacion_por_anio(query, year):
    """Filtra por el período contable informado, no por la fecha de carga."""
    return query.filter(Facturacion2026.mes.like(f'{int(year):04d}-%'))


def anios_facturacion_disponibles():
    """Años presentes en el campo Mes de Facturación."""
    anios = set()
    for (mes,) in db.session.query(Facturacion2026.mes).distinct().all():
        texto = str(mes or '').strip()
        if len(texto) >= 4 and texto[:4].isdigit():
            anios.add(int(texto[:4]))
    return anios


def es_personal_cobranzas(data):
    return normalizar_tipo_vh(data.get('tipo_jornada')) == TIPO_VH_PERSONAL_COBRANZAS


def resolver_header(valor):
    normalizado = normalizar_header(valor)
    if normalizado in ALIAS_IMPORTACION:
        return ALIAS_IMPORTACION[normalizado]
    if 'camp' in normalizado and 'sub' in normalizado:
        return 'subcampania'
    if 'camp' in normalizado:
        return 'campania'
    if 'jefe' in normalizado and 'site' in normalizado:
        return 'jefe_site'
    return None


def parse_numero(valor):
    if valor in (None, ''):
        return 0
    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()
    if not texto:
        return 0

    negativo = False
    if texto.startswith('(') and texto.endswith(')'):
        negativo = True
        texto = texto[1:-1]

    texto = texto.replace('\xa0', '').replace(' ', '')
    for simbolo in ('$', '€', '£', 'ARS', 'USD'):
        texto = texto.replace(simbolo, '')
        texto = texto.replace(simbolo.lower(), '')
    texto = texto.replace("'", '')

    if texto.startswith('-'):
        negativo = True
        texto = texto[1:]
    elif texto.endswith('-'):
        negativo = True
        texto = texto[:-1]

    if not texto:
        return 0

    ultimo_punto = texto.rfind('.')
    ultima_coma = texto.rfind(',')
    if ultimo_punto != -1 and ultima_coma != -1:
        separador_decimal = '.' if ultimo_punto > ultima_coma else ','
        separador_miles = ',' if separador_decimal == '.' else '.'
        texto = texto.replace(separador_miles, '')
        texto = texto.replace(separador_decimal, '.')
    elif ultima_coma != -1:
        partes = texto.split(',')
        if len(partes) == 2 and len(partes[-1]) == 3 and len(partes[0]) <= 3:
            texto = ''.join(partes)
        elif len(partes[-1]) in (1, 2, 3):
            texto = ''.join(partes[:-1]).replace('.', '') + '.' + partes[-1]
        else:
            texto = ''.join(partes)
    elif ultimo_punto != -1:
        partes = texto.split('.')
        if len(partes) > 2 and len(partes[-1]) == 3:
            texto = ''.join(partes)
        elif len(partes) > 2:
            texto = ''.join(partes[:-1]) + '.' + partes[-1]
        elif len(partes) == 2 and len(partes[-1]) == 3 and len(partes[0]) <= 3:
            texto = ''.join(partes)

    numero = float(texto)
    if negativo:
        numero *= -1
    return numero


def fecha_excel(valor):
    return (datetime(1899, 12, 30) + timedelta(days=float(valor))).date()


def parse_fecha(valor):
    texto = str(valor or '').strip()
    if texto.replace('.', '', 1).isdigit():
        return fecha_excel(float(texto))
    for formato in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError('Fecha invalida. Use YYYY-MM-DD o DD/MM/YYYY')


def obtener_excepcion_calculo(cliente, campania):
    return ExcepcionCalculo.query.filter(
        func.lower(ExcepcionCalculo.cliente) == str(cliente or '').strip().lower(),
        func.lower(ExcepcionCalculo.campania) == str(campania or '').strip().lower(),
        ExcepcionCalculo.activa.is_(True),
    ).first()


def obtener_configuracion_campania(cliente, campania):
    return Campania.query.filter(
        func.lower(Campania.cliente) == str(cliente or '').strip().lower(),
        func.lower(Campania.nombre) == str(campania or '').strip().lower(),
    ).first()


def aplicar_configuracion_valor_hora(data, exigir_configuracion=True):
    campania = obtener_configuracion_campania(data.get('cliente'), data.get('campania'))
    if (not campania or campania.valor_hora_variable is None) and exigir_configuracion:
        raise ValueError('Debe indicar una sola vez si el valor hora de esta campaña es variable')
    if not campania or campania.valor_hora_variable is None:
        return campania
    if campania.valor_hora_variable is False:
        data['valor_hora'] = data.get('valor_hora_objetivo')
    return campania


def aplicar_excepcion_calculo(data, exigir_configuracion=True, preservar_facturado_manual=False):
    if data.get('es_next_gen'):
        data['facturado_horas_manual'] = None
        return None
    # Una importación ya conciliada es la fuente definitiva: no se recalculan
    # valores hora ni se aplican porcentajes configurados para la carga manual.
    if preservar_facturado_manual:
        if data.get('facturado_horas_manual') not in (None, '') and parse_numero(data.get('facturado_horas_manual')) < 0:
            raise ValueError('Facturado en horas no puede ser negativo')
        data['importe_fijo'] = None
        return obtener_excepcion_calculo(data.get('cliente'), data.get('campania'))
    if data.get('_excepcion_calculo_aplicada'):
        return obtener_excepcion_calculo(data.get('cliente'), data.get('campania'))
    aplicar_configuracion_valor_hora(data, exigir_configuracion=exigir_configuracion)
    excepcion = obtener_excepcion_calculo(data.get('cliente'), data.get('campania'))
    if not excepcion or excepcion.tipo_calculo not in ('facturado_horas_manual', 'facturado_manual_ajustes_vh'):
        data['facturado_horas_manual'] = None
        return None
    if not data.get('_excepcion_calculo_aplicada'):
        if excepcion.tipo_calculo == 'facturado_manual_ajustes_vh':
            vh_objetivo_base = parse_numero(data.get('valor_hora_objetivo') or data.get('valor_hora'))
            vh_alcanzado_base = parse_numero(data.get('valor_hora'))
            data['valor_hora_objetivo'] = vh_objetivo_base * (1 + (excepcion.ajuste_vh_objetivo_pct or 0) / 100)
            data['valor_hora'] = vh_alcanzado_base * (1 + (excepcion.ajuste_vh_alcanzado_pct or 0) / 100)
            data['variable_objetivo'] = parse_numero(data.get('horas_objetivo')) * (
                data['valor_hora_objetivo'] - data['valor_hora']
            )
            data['objetivo_separar_ajuste_vh'] = True
        data['_excepcion_calculo_aplicada'] = True
    horas = parse_numero(data.get('horas_facturadas'))
    facturado = parse_numero(data.get('facturado_horas_manual'))
    if horas <= 0:
        raise ValueError('Las horas facturadas deben ser mayores a 0 para calcular el valor hora')
    if facturado < 0:
        raise ValueError('Facturado en horas no puede ser negativo')
    if excepcion.tipo_calculo == 'facturado_horas_manual':
        data['valor_hora'] = facturado / horas
    data['importe_fijo'] = None
    return excepcion


def crear_registro_facturacion(data, exigir_configuracion_valor_hora=True, preservar_facturado_manual=False):
    fecha = parse_fecha(data['fecha'])
    mes = mes_valido(data.get('mes'))
    if not mes:
        raise ValueError('El mes de facturacion no es valido')

    aplicar_excepcion_calculo(
        data,
        exigir_configuracion=exigir_configuracion_valor_hora,
        preservar_facturado_manual=preservar_facturado_manual,
    )
    registro = Facturacion2026(
        fecha=fecha,
        mes=mes,
        cliente=data['cliente'].strip(),
        gerente=data.get('gerente', '').strip(),
        jefe_site=data.get('jefe_site', '').strip(),
        campania=data.get('campania', '').strip(),
        subcampania=data.get('subcampania', '').strip(),
        tipo_negocio=str(data.get('tipo_negocio') or '').strip() or None,
        es_next_gen=bool(data.get('es_next_gen')),
        objetivo_separar_ajuste_vh=bool(data.get('objetivo_separar_ajuste_vh')),
        tipo_jornada=normalizar_tipo_vh(data['tipo_jornada']),
        horas_objetivo=parse_numero(data.get('horas_objetivo')),
        horas_facturadas=parse_numero(data.get('horas_facturadas')),
        horas_penalizadas=parse_numero(data.get('horas_penalizadas')),
        valor_hora_objetivo=parse_numero(data.get('valor_hora_objetivo') or data.get('valor_hora')),
        valor_hora=parse_numero(data.get('valor_hora')),
        facturado_horas_manual=parse_numero(data.get('facturado_horas_manual')) if data.get('facturado_horas_manual') not in (None, '') else None,
        total_facturado_manual=parse_numero(data.get('total_facturado_manual')) if data.get('total_facturado_manual') not in (None, '') else None,
        control_facturado_horas=parse_numero(data.get('control_facturado_horas')) if data.get('control_facturado_horas') not in (None, '') else None,
        control_variable_productivo=parse_numero(data.get('control_variable_productivo')) if data.get('control_variable_productivo') not in (None, '') else None,
        control_penalizaciones_bonos=parse_numero(data.get('control_penalizaciones_bonos')) if data.get('control_penalizaciones_bonos') not in (None, '') else None,
        control_total_facturado=parse_numero(data.get('control_total_facturado')) if data.get('control_total_facturado') not in (None, '') else None,
        control_objetivo_total=parse_numero(data.get('control_objetivo_total')) if data.get('control_objetivo_total') not in (None, '') else None,
        tarifacion=parse_numero(data.get('tarifacion')) if data.get('tarifacion') not in (None, '') else None,
        importe_fijo=parse_numero(data.get('importe_fijo')) if data.get('importe_fijo') not in (None, '') else None,
        variable_objetivo=parse_numero(data.get('variable_objetivo')),
        variable_productivo=parse_numero(data.get('variable_productivo')),
        bonos=parse_numero(data.get('bonos')),
        penalizaciones=normalizar_importe_ajuste('penalizaciones', data.get('penalizaciones')),
        netx_gen=parse_numero(data.get('netx_gen')),
        otros=parse_numero(data.get('otros')),
    )
    db.session.add(registro)
    guardar_justificaciones(registro, data)
    asegurar_asignacion_desde_registro(registro)
    return registro


def actualizar_registro_facturacion(registro, data, exigir_configuracion_valor_hora=True, preservar_facturado_manual=False):
    fecha = parse_fecha(data['fecha'])
    mes = mes_valido(data.get('mes'))
    if not mes:
        raise ValueError('El mes de facturacion no es valido')

    aplicar_excepcion_calculo(
        data,
        exigir_configuracion=exigir_configuracion_valor_hora,
        preservar_facturado_manual=preservar_facturado_manual,
    )
    registro.fecha = fecha
    registro.mes = mes
    registro.cliente = data['cliente'].strip()
    registro.gerente = data.get('gerente', '').strip()
    registro.jefe_site = data.get('jefe_site', '').strip()
    registro.campania = data.get('campania', '').strip()
    registro.subcampania = data.get('subcampania', '').strip()
    registro.tipo_negocio = str(data.get('tipo_negocio') or '').strip() or None
    registro.es_next_gen = bool(data.get('es_next_gen'))
    registro.objetivo_separar_ajuste_vh = bool(data.get('objetivo_separar_ajuste_vh'))
    registro.tipo_jornada = normalizar_tipo_vh(data['tipo_jornada'])
    registro.horas_objetivo = parse_numero(data.get('horas_objetivo'))
    registro.horas_facturadas = parse_numero(data.get('horas_facturadas'))
    registro.horas_penalizadas = parse_numero(data.get('horas_penalizadas'))
    registro.valor_hora_objetivo = parse_numero(data.get('valor_hora_objetivo') or data.get('valor_hora'))
    registro.valor_hora = parse_numero(data.get('valor_hora'))
    registro.facturado_horas_manual = parse_numero(data.get('facturado_horas_manual')) if data.get('facturado_horas_manual') not in (None, '') else None
    registro.total_facturado_manual = parse_numero(data.get('total_facturado_manual')) if data.get('total_facturado_manual') not in (None, '') else None
    registro.control_facturado_horas = parse_numero(data.get('control_facturado_horas')) if data.get('control_facturado_horas') not in (None, '') else None
    registro.control_variable_productivo = parse_numero(data.get('control_variable_productivo')) if data.get('control_variable_productivo') not in (None, '') else None
    registro.control_penalizaciones_bonos = parse_numero(data.get('control_penalizaciones_bonos')) if data.get('control_penalizaciones_bonos') not in (None, '') else None
    registro.control_total_facturado = parse_numero(data.get('control_total_facturado')) if data.get('control_total_facturado') not in (None, '') else None
    registro.control_objetivo_total = parse_numero(data.get('control_objetivo_total')) if data.get('control_objetivo_total') not in (None, '') else None
    registro.tarifacion = parse_numero(data.get('tarifacion')) if data.get('tarifacion') not in (None, '') else None
    registro.importe_fijo = parse_numero(data.get('importe_fijo')) if data.get('importe_fijo') not in (None, '') else None
    registro.variable_objetivo = parse_numero(data.get('variable_objetivo'))
    registro.variable_productivo = parse_numero(data.get('variable_productivo'))
    registro.bonos = parse_numero(data.get('bonos'))
    registro.penalizaciones = normalizar_importe_ajuste('penalizaciones', data.get('penalizaciones'))
    registro.netx_gen = parse_numero(data.get('netx_gen'))
    registro.otros = parse_numero(data.get('otros'))
    guardar_justificaciones(registro, data)
    asegurar_asignacion_desde_registro(registro)
    return registro


def query_reemplazo_importacion(data):
    fecha = parse_fecha(data['fecha'])
    return Facturacion2026.query.filter(
        Facturacion2026.fecha == fecha,
        Facturacion2026.cliente == data['cliente'].strip(),
        Facturacion2026.gerente == data.get('gerente', '').strip(),
        Facturacion2026.jefe_site == data.get('jefe_site', '').strip(),
        Facturacion2026.campania == data.get('campania', '').strip(),
        Facturacion2026.subcampania == data.get('subcampania', '').strip(),
        Facturacion2026.tipo_jornada == normalizar_tipo_vh(data.get('tipo_jornada')),
        Facturacion2026.tipo_negocio == (str(data.get('tipo_negocio') or '').strip() or None),
    )


def clave_reemplazo_importacion(data):
    fecha = parse_fecha(data['fecha']).isoformat()
    return '|'.join([
        fecha,
        data['cliente'].strip(),
        data.get('gerente', '').strip(),
        data.get('jefe_site', '').strip(),
        data.get('campania', '').strip(),
        data.get('subcampania', '').strip(),
        normalizar_tipo_vh(data.get('tipo_jornada')),
        str(data.get('tipo_negocio') or '').strip(),
    ])


def detalle_conflicto_importacion(indice, item, existentes):
    return {
        'fila': indice,
        'fecha': parse_fecha(item['fecha']).isoformat(),
        'mes': mes_valido(item.get('mes')),
        'cliente': item['cliente'].strip(),
        'gerente': item.get('gerente', '').strip(),
        'jefe_site': item.get('jefe_site', '').strip(),
        'campania': item.get('campania', '').strip(),
        'subcampania': item.get('subcampania', '').strip(),
        'tipo_negocio': str(item.get('tipo_negocio') or '').strip(),
        'existentes': [registro.id for registro in existentes],
    }


def filas_desde_csv(contenido):
    muestra = contenido[:2048]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=';,|\t,')
    except csv.Error:
        dialecto = csv.excel
        dialecto.delimiter = ';'
    return list(csv.reader(io.StringIO(contenido), dialecto))


def decodificar_texto_importacion(contenido_bytes):
    for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            return contenido_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return contenido_bytes.decode('utf-8-sig', errors='replace')


def filas_desde_html(contenido):
    parser = TablaHTMLParser()
    parser.feed(contenido)
    return parser.rows


def indice_columna(celda_ref):
    letras = ''.join(caracter for caracter in celda_ref if caracter.isalpha())
    indice = 0
    for letra in letras:
        indice = indice * 26 + (ord(letra.upper()) - ord('A') + 1)
    return indice - 1


def filas_desde_xlsx(contenido_bytes):
    ns = {
        'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
        'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
    }
    with zipfile.ZipFile(io.BytesIO(contenido_bytes)) as archivo:
        shared_strings = []
        if 'xl/sharedStrings.xml' in archivo.namelist():
            root = ET.fromstring(archivo.read('xl/sharedStrings.xml'))
            for item in root.findall('main:si', ns):
                textos = [nodo.text or '' for nodo in item.findall('.//main:t', ns)]
                shared_strings.append(''.join(textos))

        sheet_path = 'xl/worksheets/sheet1.xml'
        if 'xl/workbook.xml' in archivo.namelist() and 'xl/_rels/workbook.xml.rels' in archivo.namelist():
            workbook = ET.fromstring(archivo.read('xl/workbook.xml'))
            rels = ET.fromstring(archivo.read('xl/_rels/workbook.xml.rels'))
            first_sheet = workbook.find('main:sheets/main:sheet', ns)
            if first_sheet is not None:
                rel_id = first_sheet.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                for rel in rels.findall('rel:Relationship', ns):
                    if rel.attrib.get('Id') == rel_id:
                        target = rel.attrib.get('Target', 'worksheets/sheet1.xml').replace('\\', '/').lstrip('/')
                        sheet_path = target if target.startswith('xl/') else f'xl/{target}'
                        break

        sheet = ET.fromstring(archivo.read(sheet_path))
        filas = []
        for row in sheet.findall('.//main:row', ns):
            valores = []
            for cell in row.findall('main:c', ns):
                ref = cell.attrib.get('r', '')
                indice = indice_columna(ref) if ref else len(valores)
                while len(valores) <= indice:
                    valores.append('')

                tipo = cell.attrib.get('t')
                valor = ''
                if tipo == 'inlineStr':
                    textos = [nodo.text or '' for nodo in cell.findall('.//main:t', ns)]
                    valor = ''.join(textos)
                else:
                    value_node = cell.find('main:v', ns)
                    if value_node is not None:
                        valor = value_node.text or ''
                        if tipo == 's':
                            valor = shared_strings[int(valor)] if valor.isdigit() and int(valor) < len(shared_strings) else ''
                valores[indice] = valor
            if any(str(valor).strip() for valor in valores):
                filas.append(valores)
        return filas


def datos_desde_filas(filas):
    if not filas:
        return []
    headers = []
    repeticiones = defaultdict(int)
    controles_repetidos = {
        'facturado_horas_manual': 'control_facturado_horas',
        'variable_productivo': 'control_variable_productivo',
        'total_facturado_informado': 'control_total_facturado',
    }
    for header in filas[0]:
        resuelto = resolver_header(header)
        repeticiones[resuelto] += 1
        if resuelto in controles_repetidos and repeticiones[resuelto] > 1:
            resuelto = controles_repetidos[resuelto]
        if normalizar_header(header) in ('penalizaciones por incumplimientos', 'penalizacion por incumplimientos'):
            resuelto = 'control_penalizaciones_bonos'
        headers.append(resuelto)
    datos = []
    for fila in filas[1:]:
        if not any(str(celda).strip() for celda in fila):
            continue
        item = {}
        for indice, valor in enumerate(fila):
            if indice < len(headers) and headers[indice]:
                if headers[indice] not in item:
                    item[headers[indice]] = valor.strip() if isinstance(valor, str) else valor
        datos.append(item)
    return datos


def resolver_asignacion_importacion(item):
    """Completa gerente y jefe desde Datos Maestros, sin confiar en el archivo."""
    cliente = str(item.get('cliente') or '').strip()
    campania = str(item.get('campania') or '').strip()
    subcampania = str(item.get('subcampania') or '').strip()
    tipo_negocio = str(item.get('tipo_negocio') or '').strip()

    cliente_clave = normalizar_header(cliente)
    subcampania_clave = normalizar_header(subcampania)
    campania_clave = normalizar_header(campania)
    tipo_clave = normalizar_header(tipo_negocio)
    compactar = lambda valor: normalizar_header(valor).replace(' ', '')
    cuenta_compacta = compactar(cliente)
    subcampania_partes = [parte.strip() for parte in re.split(r'\s*/\s*', subcampania) if parte.strip()]
    subclaves = {normalizar_header(subcampania), *(normalizar_header(p) for p in subcampania_partes)}
    subcompactas = {clave.replace(' ', '') for clave in subclaves}

    coincidencias_cuenta = []
    coincidencias_subcampania = []
    coincidencias_sub_exacta = []
    maestras_activas = AsignacionComercial.query.filter_by(activa=True).order_by(AsignacionComercial.id).all()
    for asignacion in maestras_activas:
        cliente_maestro = normalizar_header(asignacion.cliente)
        campania_maestra = normalizar_header(asignacion.campania)
        sub_maestra = normalizar_header(asignacion.subcampania)
        cuenta_coincide_servicio = (
            cliente_clave in (campania_maestra, sub_maestra)
            or cuenta_compacta in (campania_maestra.replace(' ', ''), sub_maestra.replace(' ', ''))
        )
        cuenta_coincide_cliente = (
            cliente_clave == cliente_maestro or cuenta_compacta == cliente_maestro.replace(' ', '')
        )
        sub_coincide = (
            campania_maestra in subclaves or sub_maestra in subclaves
            or campania_maestra.replace(' ', '') in subcompactas
            or sub_maestra.replace(' ', '') in subcompactas
        )
        if not (cuenta_coincide_servicio or sub_coincide or (cuenta_coincide_cliente and sub_coincide)):
            continue
        if campania_clave and campania_maestra != campania_clave:
            continue
        if tipo_clave and normalizar_header(asignacion.tipo_negocio) != tipo_clave:
            continue
        if sub_coincide:
            coincidencias_subcampania.append(asignacion)
        if sub_maestra == subcampania_clave or sub_maestra.replace(' ', '') == subcampania_clave.replace(' ', ''):
            coincidencias_sub_exacta.append(asignacion)
        if cuenta_coincide_servicio:
            coincidencias_cuenta.append(asignacion)
    coincidencias = coincidencias_sub_exacta or coincidencias_subcampania or coincidencias_cuenta
    mes_objetivo = mes_valido(item.get('mes'))
    if mes_objetivo and coincidencias:
        vigentes = [
            asignacion for asignacion in coincidencias
            if (not asignacion.vigencia_desde or asignacion.vigencia_desde <= mes_objetivo)
            and (not asignacion.vigencia_hasta or mes_objetivo <= asignacion.vigencia_hasta)
        ]
        # Los períodos explícitos son la fuente de verdad. La compatibilidad con
        # maestros antiguos sin fechas se conserva cuando todavía no hay ninguno.
        if vigentes:
            coincidencias = vigentes
    if not coincidencias_subcampania and len(coincidencias_cuenta) > 1:
        opciones = ', '.join(dict.fromkeys(a.subcampania for a in coincidencias_cuenta))
        return None, (
            f'no existe en Datos Maestros la Sub campaña {subcampania} para {cliente}. '
            f'Opciones disponibles: {opciones}'
        )
    if len(coincidencias) > 1:
        # Si la única diferencia es el jefe, se trata como una vigencia histórica.
        firmas_sin_jefe = {
            (a.cliente, a.gerente, a.campania, a.subcampania, a.tipo_negocio, bool(a.es_next_gen))
            for a in coincidencias
        }
        if len(firmas_sin_jefe) == 1:
            historicos = Facturacion2026.query.filter(
                func.lower(Facturacion2026.cliente) == cliente.lower(),
                func.lower(Facturacion2026.subcampania) == subcampania.lower(),
                Facturacion2026.jefe_site.in_([a.jefe_site for a in coincidencias]),
            ).order_by(Facturacion2026.mes).all()
            por_mes = {}
            for registro in historicos:
                por_mes.setdefault(registro.mes, registro.jefe_site)
            if mes_objetivo and por_mes:
                meses_anteriores = [mes for mes in por_mes if mes <= mes_objetivo]
                mes_referencia = max(meses_anteriores) if meses_anteriores else min(por_mes)
                jefe_referencia = por_mes[mes_referencia]
                candidatas = [a for a in coincidencias if a.jefe_site == jefe_referencia]
                if len(candidatas) == 1:
                    coincidencias = candidatas

    if len(coincidencias) > 1:
        # Altas duplicadas con exactamente la misma clasificación son equivalentes.
        firmas = {
            (a.cliente, a.gerente, a.jefe_site, a.campania, a.subcampania, a.tipo_negocio, bool(a.es_next_gen))
            for a in coincidencias
        }
        if len(firmas) == 1:
            coincidencias = [coincidencias[0]]
    if len(coincidencias) == 1:
        asignacion = coincidencias[0]
        item.update({
            'cliente': asignacion.cliente,
            'gerente': asignacion.gerente,
            'jefe_site': asignacion.jefe_site,
            'campania': asignacion.campania,
            'subcampania': asignacion.subcampania,
            'tipo_negocio': asignacion.tipo_negocio,
            'es_next_gen': bool(asignacion.es_next_gen),
        })
        return asignacion, None

    identidad = f'{cliente or "(sin cuenta)"} / {subcampania or "(sin subcampaña)"}'
    if campania:
        identidad += f' / {campania}'
    if tipo_negocio:
        identidad += f' / {tipo_negocio}'
    if not coincidencias:
        maestras_cliente = [
            a for a in AsignacionComercial.query.filter_by(activa=True).all()
            if normalizar_header(a.cliente) == cliente_clave
        ]
        sugeridas = sorted(
            maestras_cliente,
            key=lambda a: SequenceMatcher(None, subcampania_clave, normalizar_header(a.subcampania)).ratio(),
            reverse=True,
        )[:3]
        sugerencia = ''
        if sugeridas:
            opciones = ', '.join(dict.fromkeys(a.subcampania for a in sugeridas))
            sugerencia = f'. Opciones cercanas: {opciones}'
        return None, f'no existe en Datos Maestros: {identidad}{sugerencia}'
    alternativas = '; '.join(
        f'{a.gerente} / {a.jefe_site}' + (f' / {a.tipo_negocio}' if a.tipo_negocio else '')
        for a in coincidencias[:5]
    )
    return None, f'hay {len(coincidencias)} coincidencias en Datos Maestros para {identidad}: {alternativas}. Informe Tipo de negocio o revise el maestro'


def aplicar_asignacion_a_fila(item, asignacion):
    item.update({
        'cliente': asignacion.cliente,
        'gerente': asignacion.gerente,
        'jefe_site': asignacion.jefe_site,
        'campania': asignacion.campania,
        'subcampania': asignacion.subcampania,
        'tipo_negocio': asignacion.tipo_negocio,
        'es_next_gen': bool(asignacion.es_next_gen),
    })


def resolver_personal_importacion_directa(item):
    """Resuelve o crea el maestro de Personal sin detener la carga masiva."""
    cliente_archivo = str(item.get('cliente') or '').strip()
    subcampania_archivo = str(item.get('subcampania') or '').strip()
    cliente_clave = normalizar_header(cliente_archivo)
    subcampania_clave = normalizar_header(subcampania_archivo)
    maestras_activas = AsignacionComercial.query.filter_by(activa=True).order_by(AsignacionComercial.id).all()
    candidatas = [
        asignacion for asignacion in maestras_activas
        if normalizar_header(asignacion.cliente) == cliente_clave
    ]
    exactas = [
        asignacion for asignacion in candidatas
        if normalizar_header(asignacion.subcampania) == subcampania_clave
        or normalizar_header(asignacion.campania) == subcampania_clave
    ]
    if not exactas:
        # El archivo operativo puede traer el servicio como Cuenta (por ejemplo,
        # "Personal Cobranzas"), aunque el maestro lo tenga bajo Cliente Personal.
        exactas = [
            asignacion for asignacion in maestras_activas
            if 'personal' in normalizar_header(asignacion.cliente)
            and (
                normalizar_header(asignacion.campania) in (cliente_clave, subcampania_clave)
                or normalizar_header(asignacion.subcampania) in (cliente_clave, subcampania_clave)
            )
        ]
    if exactas:
        asignacion = exactas[0]
        jefe_correcto = jefe_personal_por_servicio(
            cliente_archivo,
            asignacion.campania,
            subcampania_archivo,
            asignacion.tipo_negocio,
        )
        if asignacion.jefe_site != jefe_correcto:
            asignacion.jefe_site = jefe_correcto
    else:
        if not candidatas and 'personal' in cliente_clave:
            candidatas = [
                asignacion for asignacion in maestras_activas
                if 'personal' in normalizar_header(asignacion.cliente)
            ]
        if not candidatas:
            return None, f'No existe una estructura base en Datos Maestros para {cliente_archivo}'
        base = max(
            candidatas,
            key=lambda asignacion: SequenceMatcher(
                None,
                subcampania_clave,
                normalizar_header(asignacion.subcampania),
            ).ratio(),
        )
        asignacion, _ = obtener_o_crear_asignacion({
            'cliente': cliente_archivo,
            'gerente': base.gerente,
            'jefe_site': jefe_personal_por_servicio(
                cliente_archivo,
                subcampania_archivo,
                base.tipo_negocio,
            ),
            'campania': subcampania_archivo,
            'subcampania': subcampania_archivo,
            'tipo_negocio': base.tipo_negocio,
            'es_next_gen': False,
        })
    aplicar_asignacion_a_fila(item, asignacion)
    return asignacion, None


def opciones_maestro_para_fila(item):
    cuenta = str(item.get('cliente') or '').strip()
    subcampania = str(item.get('subcampania') or '').strip()
    cuenta_clave = normalizar_header(cuenta).replace(' ', '')
    sub_clave = normalizar_header(subcampania).replace(' ', '')
    maestras = AsignacionComercial.query.filter_by(activa=True).all()
    puntuadas = []
    for asignacion in maestras:
        campos = [asignacion.cliente, asignacion.campania, asignacion.subcampania]
        claves = [normalizar_header(valor).replace(' ', '') for valor in campos]
        puntaje = max(
            max(SequenceMatcher(None, cuenta_clave, clave).ratio() for clave in claves),
            max(SequenceMatcher(None, sub_clave, clave).ratio() for clave in claves),
        )
        if puntaje >= 0.45:
            puntuadas.append((puntaje, asignacion))
    puntuadas.sort(key=lambda par: (-par[0], par[1].id))
    return [
        {
            'id': asignacion.id,
            'label': asignacion.label,
            'cliente': asignacion.cliente,
            'campania': asignacion.campania,
            'subcampania': asignacion.subcampania,
            'gerente': asignacion.gerente,
            'jefe_site': asignacion.jefe_site,
        }
        for _, asignacion in puntuadas[:6]
    ]


def preparar_fila_migracion_facturacion(item):
    """Traduce el formato operativo reducido al modelo de facturación."""
    identificacion_original = ' '.join(str(item.get(campo) or '') for campo in ('cliente', 'campania', 'subcampania'))
    item['_importacion_directa_personal'] = 'personal' in normalizar_header(identificacion_original)
    item['mes'] = mes_valido(item.get('mes')) or item.get('mes')
    if not item.get('fecha'):
        item['fecha'] = date.today().isoformat()
    item.setdefault('horas_penalizadas', 0)
    item.setdefault('importe_fijo', None)
    item.setdefault('variable_objetivo', 0)
    item.setdefault('otros', 0)
    item.setdefault('tarifacion', 0)
    item.setdefault('netx_gen', 0)
    item.setdefault('valor_hora_objetivo', item.get('valor_hora'))

    es_next_gen_sin_horas = (
        parse_numero(item.get('netx_gen')) > 0
        and parse_numero(item.get('horas_objetivo')) == 0
        and parse_numero(item.get('horas_facturadas')) == 0
        and parse_numero(item.get('valor_hora')) == 0
    )
    item['_es_next_gen_sin_horas'] = es_next_gen_sin_horas
    if es_next_gen_sin_horas:
        item['es_next_gen'] = True
        item['tipo_jornada'] = 'NextGen'

    ajuste = parse_numero(item.get('ajuste_bono_penalizacion'))
    item['bonos'] = ajuste if ajuste > 0 else 0
    item['penalizaciones'] = ajuste if ajuste < 0 else 0

    if item.get('variable_productivo') in (None, ''):
        unidad = parse_numero(item.get('unidad'))
        porcentaje = parse_numero(item.get('porcentaje_variable'))
        item['variable_productivo'] = unidad * porcentaje / 100
    else:
        item['variable_productivo'] = parse_numero(item.get('variable_productivo'))


def validar_total_fila_migracion(item):
    informado = item.get('total_facturado_informado')
    if informado in (None, ''):
        return None
    if item.get('_importacion_directa_personal'):
        item['total_facturado_manual'] = parse_numero(informado)
        return None
    horas_netas = max(parse_numero(item.get('horas_facturadas')) - parse_numero(item.get('horas_penalizadas')), 0)
    facturado_horas = (
        parse_numero(item.get('facturado_horas_manual'))
        if item.get('facturado_horas_manual') not in (None, '')
        else horas_netas * parse_numero(item.get('valor_hora'))
    )
    calculado = (
        facturado_horas
        + parse_numero(item.get('tarifacion'))
        + parse_numero(item.get('bonos'))
        + parse_numero(item.get('variable_productivo'))
        - abs(parse_numero(item.get('penalizaciones')))
        + parse_numero(item.get('netx_gen'))
        + parse_numero(item.get('otros'))
    )
    total_informado = parse_numero(informado)
    if abs(calculado - total_informado) > 1:
        return f'Total Facturado informado {total_informado:.2f} no coincide con el calculado {calculado:.2f}'
    item['total_facturado_manual'] = total_informado
    return None


def validar_variable_fila_migracion(item):
    if item.get('unidad') in (None, '') or item.get('porcentaje_variable') in (None, ''):
        return None
    esperado = parse_numero(item.get('unidad')) * parse_numero(item.get('porcentaje_variable')) / 100
    informado = parse_numero(item.get('variable_productivo'))
    if abs(esperado - informado) > 1:
        return f'Variable productivo {informado:.2f} no coincide con Unidad × % Variable ({esperado:.2f})'
    return None


def guardar_porcentaje_variable_importado(item):
    if item.get('porcentaje_variable') in (None, ''):
        return
    mes = mes_valido(item.get('mes'))
    porcentaje = parse_numero(item.get('porcentaje_variable'))
    registro = VariableCampania.query.filter_by(
        cliente=item['cliente'], campania=item['campania'], mes=mes
    ).first()
    if not registro:
        registro = VariableCampania(
            cliente=item['cliente'], campania=item['campania'], mes=mes,
            year=int(mes[:4]), site=item.get('gerente'), porcentaje=porcentaje,
        )
        db.session.add(registro)
    else:
        registro.site = item.get('gerente')
        registro.porcentaje = porcentaje


def valores_request(nombre):
    valores = request.args.getlist(nombre)
    if not valores:
        valor = request.args.get(nombre)
        valores = valor.split(',') if valor else []
    valores = [valor.strip() for valor in valores if valor and valor.strip()]
    if not valores:
        return None
    return valores if len(valores) > 1 else valores[0]


def filtros_request():
    return {
        'year': valores_request('year'),
        'mes': valores_request('mes'),
        'cliente': valores_request('cliente'),
        'gerente': valores_request('gerente'),
        'jefe_site': valores_request('jefe_site'),
        'campania': valores_request('campania'),
        'subcampania': valores_request('subcampania'),
        'tipo_negocio': valores_request('tipo_negocio'),
    }


def opciones_filtro(filtros, campo):
    filtros_base = dict(filtros)
    filtros_base[campo] = None
    registros = aplicar_filtros(Facturacion2026.query, **filtros_base).all()
    if campo == 'year':
        return sorted({registro.mes[:4] for registro in registros if registro.mes and len(registro.mes) >= 4})
    valores = {
        getattr(registro, campo)
        for registro in registros
        if getattr(registro, campo, None)
    }
    # Los gerentes se administran en Datos Maestros y deben poder elegirse antes
    # de que exista su primera facturación (por ejemplo, "Sin gerencia").
    usuario = usuario_actual()
    if campo == 'gerente' and (not usuario or usuario.puede_acceder('facturacion_total')):
        valores.update(
            gerente for (gerente,) in AsignacionComercial.query.with_entities(
                AsignacionComercial.gerente
            ).filter(AsignacionComercial.activa.is_(True)).distinct().all()
            if gerente
        )
    return sorted(valores)


def asignaciones_maestras_filtradas(filtros):
    """Aplica los filtros dimensionales a Datos Maestros (sin período)."""
    query = AsignacionComercial.query.filter(AsignacionComercial.activa.is_(True))
    usuario = usuario_actual()
    if usuario and not usuario.puede_acceder('facturacion_total'):
        if usuario.gerente_asignado:
            query = query.filter(func.lower(AsignacionComercial.gerente) == usuario.gerente_asignado.lower())
        elif usuario.jefe_site_asignado:
            query = query.filter(func.lower(AsignacionComercial.jefe_site) == usuario.jefe_site_asignado.lower())
    for clave, columna in (
        ('cliente', AsignacionComercial.cliente), ('gerente', AsignacionComercial.gerente),
        ('jefe_site', AsignacionComercial.jefe_site), ('campania', AsignacionComercial.campania),
        ('subcampania', AsignacionComercial.subcampania),
        ('tipo_negocio', AsignacionComercial.tipo_negocio),
    ):
        valor = filtros.get(clave)
        if valor:
            query = filtrar_valores_exactos(query, columna, valor)
    return query.all()


def filtros_comparativo_request():
    return {
        **filtros_request(),
        'horas_objetivo_min': numero_request('horas_objetivo_min'),
        'horas_objetivo_max': numero_request('horas_objetivo_max'),
        'horas_facturadas_min': numero_request('horas_facturadas_min'),
        'horas_facturadas_max': numero_request('horas_facturadas_max'),
    }


def registros_filtrados():
    filtros = filtros_request()
    return aplicar_filtros(Facturacion2026.query, **filtros).all()


def filtros_asignacion(asignacion):
    return {
        'cliente': asignacion.cliente,
        'gerente': asignacion.gerente,
        'jefe_site': asignacion.jefe_site,
        'campania': asignacion.campania,
        'subcampania': asignacion.subcampania,
        'tipo_negocio': asignacion.tipo_negocio,
    }


def query_por_asignacion(asignacion):
    return Facturacion2026.query.filter_by(**filtros_asignacion(asignacion))


def obtener_o_crear_asignacion(campos):
    asignacion = AsignacionComercial.query.filter_by(**campos).first()
    if asignacion:
        asignacion.activa = True
        if not asignacion.campania_id:
            asignacion.campania_catalogo = obtener_o_crear_campania(campos['cliente'], campos['campania'])
        return asignacion, False
    campania = obtener_o_crear_campania(campos['cliente'], campos['campania'])
    asignacion = AsignacionComercial(**campos, campania_catalogo=campania)
    db.session.add(asignacion)
    db.session.flush()
    return asignacion, True


def obtener_o_crear_campania(cliente, nombre):
    """Devuelve el ID canónico para la combinación cliente + campaña."""
    cliente = str(cliente or '').strip()
    nombre = str(nombre or '').strip()
    campania = Campania.query.filter_by(cliente=cliente, nombre=nombre).first()
    if campania:
        campania.activa = True
        return campania
    campania = Campania(cliente=cliente, nombre=nombre, activa=True)
    db.session.add(campania)
    db.session.flush()
    return campania


TIPOS_JUSTIFICACION = {
    'bonos': 'Bonos',
    'penalizaciones': 'Penalizaciones',
    'otros': 'Otros',
}


def normalizar_importe_ajuste(tipo, importe):
    valor = parse_numero(importe)
    if tipo == 'penalizaciones':
        return -abs(valor) if valor else 0
    return valor


def normalizar_justificaciones(data):
    items = data.get('justificaciones') or []
    salida = []
    for item in items:
        tipo = str(item.get('tipo', '')).strip()
        descripcion = str(item.get('descripcion', '')).strip()
        cantidad = parse_numero(item.get('cantidad')) if item.get('cantidad') not in (None, '') else 0
        precio = parse_numero(item.get('precio')) if item.get('precio') not in (None, '') else 0
        importe_base = parse_numero(item.get('importe')) if item.get('importe') not in (None, '') else cantidad * precio
        importe = normalizar_importe_ajuste(tipo, importe_base)
        if not tipo and not descripcion and cantidad == 0 and precio == 0 and importe == 0:
            continue
        salida.append({
            'tipo': tipo,
            'descripcion': descripcion,
            'cantidad': cantidad,
            'precio': precio,
            'importe': importe,
        })
    return salida


def validar_justificaciones(data):
    errores = []
    items = normalizar_justificaciones(data)
    if 'justificaciones' not in data and not items:
        return errores
    totales = {tipo: 0 for tipo in TIPOS_JUSTIFICACION}
    if not items:
        for tipo, label in TIPOS_JUSTIFICACION.items():
            if abs(normalizar_importe_ajuste(tipo, data.get(tipo))) > 0:
                errores.append(f'{label} requiere al menos una justificacion')
        return errores

    for indice, item in enumerate(items, start=1):
        if item['tipo'] not in TIPOS_JUSTIFICACION:
            errores.append(f'Justificacion {indice}: tipo invalido')
            continue
        if item['cantidad'] <= 0:
            errores.append(f'Justificacion {indice}: la cantidad debe ser mayor a 0')
        if item['precio'] <= 0:
            errores.append(f'Justificacion {indice}: el precio debe ser mayor a 0')
        importe_absoluto = abs(item['importe'])
        if importe_absoluto <= 0:
            errores.append(f'Justificacion {indice}: el importe debe ser mayor a 0')
        if abs(importe_absoluto - (item['cantidad'] * item['precio'])) > 0.01:
            errores.append(f'Justificacion {indice}: el importe debe coincidir con cantidad por precio')
        if not item['descripcion']:
            errores.append(f'Justificacion {indice}: la descripcion es obligatoria')
        totales[item['tipo']] += item['importe']

    for tipo, label in TIPOS_JUSTIFICACION.items():
        valor_campo = normalizar_importe_ajuste(tipo, data.get(tipo))
        if abs(valor_campo) > 0 and totales[tipo] == 0:
            errores.append(f'{label} requiere al menos una justificacion')
        if abs(totales[tipo] - valor_campo) > 0.01:
            errores.append(f'La suma de justificaciones de {label} debe coincidir con el importe cargado')

    return errores


def guardar_justificaciones(registro, data):
    registro.justificaciones.clear()
    for item in normalizar_justificaciones(data):
        registro.justificaciones.append(JustificacionAjuste(
            tipo=item['tipo'],
            cantidad=item['cantidad'],
            precio=item['precio'],
            importe=item['importe'],
            descripcion=item['descripcion'],
        ))


def resumen_registros(registros):
    total_real = sum(r.total_real for r in registros)
    total_teorico = sum(r.total_teorico for r in registros)
    facturado_horas = sum(r.facturado_horas for r in registros)
    objetivo_horas = sum(r.objetivo_facturacion_horas for r in registros)
    desvio_facturacion = facturado_horas - objetivo_horas
    desvio = total_real - total_teorico
    porcentaje = (total_real / total_teorico * 100) if total_teorico > 0 else 0
    porcentaje_facturacion = (facturado_horas / objetivo_horas * 100) if objetivo_horas > 0 else 0
    return {
        'horas_objetivo': round(sum(r.horas_objetivo for r in registros), 2),
        'horas_facturadas': round(sum(r.horas_facturadas for r in registros), 2),
        'horas_penalizadas': round(sum(r.horas_penalizadas or 0 for r in registros), 2),
        'bonos': round(sum(r.bonos or 0 for r in registros), 2),
        'variable_objetivo': round(sum(r.variable_objetivo or 0 for r in registros), 2),
        'variable_productivo': round(sum(r.variable_productivo_calculo for r in registros), 2),
        'penalizaciones': round(sum(r.penalizaciones_incumplimientos for r in registros), 2),
        'netx_gen': round(sum(r.netx_gen or 0 for r in registros), 2),
        'otros': round(sum(r.otros or 0 for r in registros), 2),
        'tarifacion': round(sum(r.tarifacion or 0 for r in registros), 2),
        'facturado_horas': round(facturado_horas, 2),
        'objetivo_facturacion_horas': round(objetivo_horas, 2),
        'desvio_facturacion': round(desvio_facturacion, 2),
        'porcentaje_cumplimiento_facturacion': round(porcentaje_facturacion, 2),
        'total_real': round(total_real, 2),
        'total_teorico': round(total_teorico, 2),
        'desvio': round(desvio, 2),
        'porcentaje_cumplimiento': round(porcentaje, 2)
    }


def resumen_dashboard(registros):
    registros_con_objetivo = [r for r in registros if not r.es_next_gen]
    resumen = resumen_registros(registros_con_objetivo)
    # Los conceptos independientes siguen informándose con la base completa.
    resumen['netx_gen'] = round(sum(r.netx_gen or 0 for r in registros), 2)
    resumen['otros'] = round(sum(r.otros or 0 for r in registros), 2)
    resumen['tarifacion'] = round(sum(r.tarifacion or 0 for r in registros), 2)
    total_facturado = sum(r.total_dashboard for r in registros)
    # Fórmula de la hoja RESUMEN: el alcance y el desvío comparan únicamente
    # Facturado Horas + Variable Productivo + Penalizaciones/Bonos contra el
    # objetivo. Next Gen, tarifación y otros integran Total Facturado, pero no
    # alteran este control comparable.
    total_comparable = sum(
        r.facturado_horas
        + r.variable_productivo_calculo
        + r.facturado_bono
        + r.penalizaciones_incumplimientos
        for r in registros_con_objetivo
    )
    total_teorico = sum(r.total_teorico for r in registros_con_objetivo)
    desvio = total_comparable - total_teorico
    porcentaje = (total_comparable / total_teorico * 100) if total_teorico > 0 else 0
    horas_objetivo = sum(r.horas_objetivo or 0 for r in registros)
    horas_facturadas = sum(r.horas_facturadas or 0 for r in registros)
    horas_penalizadas = sum(r.horas_penalizadas or 0 for r in registros)
    objetivo_bono = sum(r.objetivo_facturacion_bono for r in registros)
    facturado_bono = sum(r.facturado_bono for r in registros)
    horas_objetivo_tarifadas = sum(r.horas_objetivo or 0 for r in registros if (r.horas_objetivo or 0) > 0)
    horas_netas_tarifadas = sum(
        max((r.horas_facturadas or 0) - (r.horas_penalizadas or 0), 0)
        for r in registros
        if max((r.horas_facturadas or 0) - (r.horas_penalizadas or 0), 0) > 0
    )
    valor_hora_objetivo = (
        sum((r.valor_hora_objetivo_calculo or 0) * (r.horas_objetivo or 0) for r in registros) / horas_objetivo_tarifadas
        if horas_objetivo_tarifadas > 0 else 0
    )
    valor_hora_realizado = (
        sum((r.valor_hora_alcanzado or 0) * max((r.horas_facturadas or 0) - (r.horas_penalizadas or 0), 0) for r in registros) / horas_netas_tarifadas
        if horas_netas_tarifadas > 0 else 0
    )
    return {
        **resumen,
        'total_facturado': round(total_facturado, 2),
        'total_real': round(total_facturado, 2),
        'total_teorico': round(total_teorico, 2),
        'desvio': round(desvio, 2),
        'porcentaje_cumplimiento': round(porcentaje, 2),
        'porcentaje_cumplimiento_horas': round((horas_facturadas / horas_objetivo * 100) if horas_objetivo > 0 else 0, 2),
        'porcentaje_cumplimiento_horas_adh': round(((horas_facturadas - horas_penalizadas) / horas_objetivo * 100) if horas_objetivo > 0 else 0, 2),
        'porcentaje_valor_hora': round((valor_hora_realizado / valor_hora_objetivo * 100) if valor_hora_objetivo > 0 else 0, 2),
        'porcentaje_bono': round((facturado_bono / objetivo_bono * 100) if objetivo_bono > 0 else 0, 2),
        'porcentaje_cumplimiento_facturacion': round(porcentaje, 2),
    }


MESES_MATRIZ = [
    ('01', 'ene'),
    ('02', 'feb'),
    ('03', 'mar'),
    ('04', 'abr'),
    ('05', 'may'),
    ('06', 'jun'),
    ('07', 'jul'),
    ('08', 'ago'),
    ('09', 'sep'),
    ('10', 'oct'),
    ('11', 'nov'),
    ('12', 'dic'),
]


MESES_PROYECCION = [
    ('01', 'ene'),
    ('02', 'feb'),
    ('03', 'mar'),
    ('04', 'abr'),
    ('05', 'may'),
    ('06', 'jun'),
    ('07', 'jul'),
    ('08', 'ago'),
    ('09', 'sep'),
    ('10', 'oct'),
    ('11', 'nov'),
    ('12', 'dic'),
]


def opciones_meses_proyeccion(year):
    return [
        {
            'value': f'{year}-{numero}',
            'label': f'{label}-{str(year)[-2:]}',
        }
        for numero, label in MESES_PROYECCION
    ]


def etiqueta_mes_proyeccion(mes):
    try:
        year, numero = str(mes).split('-', 1)
    except ValueError:
        return str(mes)
    labels = dict(MESES_PROYECCION)
    return f"{labels.get(numero, numero)}-{year[-2:]}"


def meses_proyeccion_desde(mes):
    mes_normalizado = mes_valido(mes)
    if not mes_normalizado:
        return []
    year = int(mes_normalizado[:4])
    mes_inicio = int(mes_normalizado[-2:])
    return [
        f'{year}-{numero}'
        for numero, _ in MESES_PROYECCION
        if int(numero) >= mes_inicio
    ]


CARGAS_SEMANALES_PROYECCION = (
    'L a V',
    'L a V + S',
    'S + D + F',
    'S',
    'D',
    'F',
    'L a V +S+D+F',
)


def normalizar_carga_semanal(value):
    value = ' '.join(str(value or '').strip().split())
    alias = {
        'L a V +S': 'L a V + S',
        'S+D+F': 'S + D + F',
        'L a V + S+D+F': 'L a V +S+D+F',
        'L a V + S + D + F': 'L a V +S+D+F',
    }
    value = alias.get(value, value)
    return value if value in CARGAS_SEMANALES_PROYECCION else 'L a V'


def fechas_feriadas_activas(year):
    return {
        feriado.fecha
        for feriado in FeriadoOperativo.query.filter_by(year=year, activo=True).all()
        if feriado.fecha
    }


def calcular_dias_objetivo(year, mes_numero, carga_semanal, feriados=None):
    carga = normalizar_carga_semanal(carga_semanal)
    feriados = feriados if feriados is not None else fechas_feriadas_activas(year)
    _, ultimo_dia = calendar.monthrange(year, mes_numero)
    dias = 0
    for dia in range(1, ultimo_dia + 1):
        fecha = date(year, mes_numero, dia)
        weekday = fecha.weekday()
        es_feriado = fecha in feriados
        cuenta = False
        if carga == 'L a V':
            cuenta = weekday < 5 and not es_feriado
        elif carga == 'L a V + S':
            cuenta = weekday < 6 and not es_feriado
        elif carga == 'S + D + F':
            cuenta = weekday in (5, 6) or es_feriado
        elif carga == 'S':
            cuenta = weekday == 5 and not es_feriado
        elif carga == 'D':
            cuenta = weekday == 6 and not es_feriado
        elif carga == 'F':
            cuenta = es_feriado
        elif carga == 'L a V +S+D+F':
            cuenta = True
        if cuenta:
            dias += 1
    return dias


SERVICIOS_PERSONAL = ('Personal CX', 'Personal', 'Personal Soporte', 'Personal SMB')

SITES_NEXT_GEN_ORIGINALES = {
    normalizar_header('Johnson & Johnson'): 'Gerencia Multicampaña',
    normalizar_header('Mirgor'): 'Gerencia Multicampaña',
    normalizar_header('Mirgor BBTI'): 'Gerencia Multicampaña',
    normalizar_header('OpenPay'): 'Gerencia Multicampaña',
    normalizar_header('Supervielle Seguros'): 'Gerencia Multicampaña',
    normalizar_header('Yamaha'): 'Gerencia Multicampaña',
    normalizar_header('Bi Bank'): 'Gerencia Multicampaña',
    normalizar_header('Santander Consumer'): 'Mariano Quesada',
    normalizar_header('APM - TERMINAL 4'): 'N/A',
    normalizar_header('Banco Credicoop'): 'N/A',
    normalizar_header('Almundo'): 'N/A',
    normalizar_header('BAPRO'): 'N/A',
    normalizar_header('ANDREANI'): 'N/A',
}


def site_original_next_gen(cliente, campania, fallback='Sin site'):
    return (
        SITES_NEXT_GEN_ORIGINALES.get(normalizar_header(campania))
        or SITES_NEXT_GEN_ORIGINALES.get(normalizar_header(cliente))
        or fallback
    )


def buscar_proyeccion_plp(tipo_plp, campania, mes):
    """Encuentra una campaña PLP aunque Excel cambie acentos o traiga texto dañado."""
    candidatas = ProyeccionMatriz.query.filter_by(cliente='Personal', mes=mes).filter(
        ProyeccionMatriz.tipo_plp == tipo_plp,
    ).all()
    clave_buscada = ''.join(caracter for caracter in normalizar_header(campania) if caracter.isalnum())
    for candidata in candidatas:
        clave_candidata = ''.join(
            caracter for caracter in normalizar_header(candidata.campania) if caracter.isalnum()
        )
        if clave_candidata == clave_buscada:
            return candidata

    mejor = None
    mejor_ratio = 0
    for candidata in candidatas:
        if '\ufffd' not in str(candidata.campania or '') and '\ufffd' not in str(campania or ''):
            continue
        clave_candidata = ''.join(
            caracter for caracter in normalizar_header(candidata.campania) if caracter.isalnum()
        )
        ratio = SequenceMatcher(None, clave_buscada, clave_candidata).ratio()
        if ratio > mejor_ratio:
            mejor, mejor_ratio = candidata, ratio
    return mejor if mejor_ratio >= 0.88 else None


def servicio_personal_identificado(cliente, campania, tipo_plp=None):
    por_nombre = {normalizar_header(nombre): nombre for nombre in SERVICIOS_PERSONAL}
    return por_nombre.get(normalizar_header(tipo_plp)) or por_nombre.get(normalizar_header(campania)) or por_nombre.get(normalizar_header(cliente))


def distribucion_personal(cliente, campania, mes, tipo_plp=None):
    servicio = servicio_personal_identificado(cliente, campania, tipo_plp)
    if not servicio:
        return None
    exacta = PersonalDistribucionHoras.query.filter_by(servicio=servicio, mes=mes).first()
    if exacta:
        return exacta
    try:
        year, numero_mes = int(mes[:4]), mes[-2:]
    except (TypeError, ValueError):
        return None
    return PersonalDistribucionHoras.query.filter(
        PersonalDistribucionHoras.servicio == servicio,
        PersonalDistribucionHoras.year < year,
        PersonalDistribucionHoras.mes.like(f'%-{numero_mes}'),
    ).order_by(PersonalDistribucionHoras.year.desc()).first()


def mes_plp_importado(valor, year):
    directo = mes_valido(str(valor or '').strip())
    if directo:
        return directo
    texto = normalizar_header(valor).replace('_', '-').replace('/', '-').replace(' ', '-')
    partes = [parte for parte in texto.split('-') if parte]
    if len(partes) < 2:
        return None
    meses_por_nombre = {normalizar_header(label): numero for numero, label in MESES_PROYECCION}
    numero = meses_por_nombre.get(partes[0][:3])
    if not numero:
        return None
    try:
        year_valor = int(partes[-1])
    except ValueError:
        return None
    if year_valor < 100:
        year_valor += 2000
    return f'{year_valor}-{numero}' if year_valor == year else None


def aplicar_calculo_proyeccion(proyeccion, valores, mes, feriados=None):
    year = int(mes[:4])
    mes_numero = int(mes[-2:])
    horas_carga_manual = bool(valores.get('horas_carga_manual'))
    jornadas = [] if horas_carga_manual else (valores.get('jornadas') or [{
        'dotacion_requerida': valores.get('dotacion_requerida') or 0,
        'carga_semanal': valores.get('carga_semanal'),
        'carga_horaria': valores.get('carga_horaria') or 0,
    }])
    jornadas_calculadas = []
    dotacion_total = 0
    horas_total = 0
    dias_objetivo_total = 0
    carga_semanal_resumen = normalizar_carga_semanal(jornadas[0].get('carga_semanal') if jornadas else None)
    carga_horaria_resumen = jornadas[0].get('carga_horaria') if jornadas else 0
    dias_objetivo_manual = valores.get('dias_objetivo_manual')
    horas_requeridas_override = valores.get('horas_requeridas_override')

    if horas_carga_manual:
        horas_manual = valores.get('horas_requeridas_manual') or 0
        porcentaje_cumplimiento_manual = valores.get('porcentaje_cumplimiento') or 0
        horas_finales = horas_manual * porcentaje_cumplimiento_manual / 100
        carga_semanal_resumen = normalizar_carga_semanal(valores.get('carga_semanal_plp') or 'L a V')
        carga_horaria_resumen = valores.get('carga_horaria_plp') or 6
        dias_laborales = calcular_dias_objetivo(year, mes_numero, carga_semanal_resumen, feriados)
        dotacion_calculada = horas_finales / dias_laborales / carga_horaria_resumen if dias_laborales else 0
        dotacion_total = dotacion_calculada
        horas_total = horas_manual
        dias_objetivo_total = dias_laborales
        jornadas_calculadas.append({
            'dotacion_requerida': dotacion_calculada,
            'carga_semanal': carga_semanal_resumen,
            'carga_horaria': carga_horaria_resumen,
            'dias_objetivo': dias_laborales,
            'horas_requeridas': horas_manual,
        })

    for jornada in jornadas:
        dotacion = jornada.get('dotacion_requerida') or 0
        carga_semanal = normalizar_carga_semanal(jornada.get('carga_semanal'))
        carga_horaria = jornada.get('carga_horaria') or 0
        dias_objetivo = dias_objetivo_manual if dias_objetivo_manual is not None else calcular_dias_objetivo(year, mes_numero, carga_semanal, feriados)
        horas_requeridas = dotacion * dias_objetivo * carga_horaria
        dotacion_total += dotacion
        horas_total += horas_requeridas
        dias_objetivo_total = max(dias_objetivo_total, dias_objetivo)
        jornadas_calculadas.append({
            'dotacion_requerida': dotacion,
            'carga_semanal': carga_semanal,
            'carga_horaria': carga_horaria,
            'dias_objetivo': dias_objetivo,
            'horas_requeridas': horas_requeridas,
        })

    if not horas_carga_manual and horas_requeridas_override is not None:
        horas_total = horas_requeridas_override

    proyeccion.cliente = valores['cliente']
    proyeccion.campania = valores['campania']
    proyeccion.year = year
    proyeccion.mes = mes
    proyeccion.dotacion_requerida = dotacion_total
    proyeccion.carga_semanal = carga_semanal_resumen
    proyeccion.carga_horaria = carga_horaria_resumen or 0
    proyeccion.dias_objetivo = dias_objetivo_total
    proyeccion.horas_requeridas = horas_total
    proyeccion.porcentaje_cumplimiento = valores.get('porcentaje_cumplimiento') or 0
    proyeccion.tipo_plp = valores.get('tipo_plp') or None
    proyeccion.horas_carga_manual = horas_carga_manual
    proyeccion.dias_objetivo_manual = None if horas_carga_manual else dias_objetivo_manual
    proyeccion.horas_requeridas_manual = None if horas_carga_manual else horas_requeridas_override
    distribucion = distribucion_personal(valores['cliente'], valores['campania'], mes, proyeccion.tipo_plp)
    proyeccion.tiene_nocturnidad = bool(distribucion and distribucion.porcentaje_nocturno > 0)
    proyeccion.porcentaje_nocturnidad = distribucion.porcentaje_nocturno if distribucion else 0
    proyeccion.jornadas = [
        ProyeccionMatrizJornada(**jornada)
        for jornada in jornadas_calculadas
    ]
    return proyeccion


def recalcular_proyecciones_mes(year, mes_numero):
    mes = f'{year}-{mes_numero:02d}'
    feriados = fechas_feriadas_activas(year)
    proyecciones = ProyeccionMatriz.query.filter_by(mes=mes).all()
    for proyeccion in proyecciones:
        jornadas = [jornada.to_dict() for jornada in proyeccion.jornadas]
        if not jornadas:
            carga_horaria = proyeccion.carga_horaria or 0
            if carga_horaria <= 0 and (proyeccion.horas_requeridas or 0) > 0 and (proyeccion.dotacion_requerida or 0) > 0:
                dias_base = proyeccion.dias_objetivo or calcular_dias_objetivo(year, mes_numero, proyeccion.carga_semanal or 'L a V', feriados)
                if dias_base > 0:
                    carga_horaria = proyeccion.horas_requeridas / ((proyeccion.dotacion_requerida or 0) * dias_base)
            jornadas = [{
                'dotacion_requerida': proyeccion.dotacion_requerida,
                'carga_semanal': proyeccion.carga_semanal,
                'carga_horaria': carga_horaria,
            }]
        valores = {
            'cliente': proyeccion.cliente,
            'campania': proyeccion.campania,
            'jornadas': jornadas,
            'porcentaje_cumplimiento': proyeccion.porcentaje_cumplimiento,
            'tiene_nocturnidad': proyeccion.tiene_nocturnidad,
            'porcentaje_nocturnidad': proyeccion.porcentaje_nocturnidad,
            'tipo_plp': proyeccion.tipo_plp,
            'horas_carga_manual': proyeccion.horas_carga_manual,
            'horas_requeridas_manual': proyeccion.horas_requeridas,
            'carga_semanal_plp': proyeccion.carga_semanal,
            'carga_horaria_plp': proyeccion.carga_horaria,
            'dias_objetivo_manual': proyeccion.dias_objetivo_manual,
            'horas_requeridas_override': proyeccion.horas_requeridas_manual,
        }
        aplicar_calculo_proyeccion(proyeccion, valores, mes, feriados)
    return len(proyecciones)


def recalcular_proyecciones_plp_year(year):
    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    proyecciones = ProyeccionMatriz.query.filter(
        ProyeccionMatriz.mes.in_(meses),
        ProyeccionMatriz.tipo_plp.isnot(None),
        ProyeccionMatriz.tipo_plp != '',
    ).all()
    feriados = fechas_feriadas_activas(year)
    for proyeccion in proyecciones:
        aplicar_calculo_proyeccion(proyeccion, {
            'cliente': 'Personal',
            'campania': proyeccion.campania,
            'tipo_plp': proyeccion.tipo_plp,
            'horas_carga_manual': True,
            'horas_requeridas_manual': proyeccion.horas_requeridas,
            'carga_semanal_plp': proyeccion.carga_semanal,
            'carga_horaria_plp': proyeccion.carga_horaria,
            'porcentaje_cumplimiento': proyeccion.porcentaje_cumplimiento,
            'jornadas': [],
        }, proyeccion.mes, feriados)
    return len(proyecciones)


def snapshot_proyeccion(proyeccion):
    if not proyeccion:
        return None
    return proyeccion.to_dict()


def buscar_proyeccion_snapshot(snapshot):
    if not snapshot:
        return None
    proyeccion_id = snapshot.get('id')
    if proyeccion_id:
        encontrada = ProyeccionMatriz.query.get(proyeccion_id)
        if encontrada:
            return encontrada
    return ProyeccionMatriz.query.filter_by(
        cliente=snapshot.get('cliente'),
        campania=snapshot.get('campania'),
        mes=snapshot.get('mes'),
    ).first()


def restaurar_proyeccion_snapshot(proyeccion, snapshot):
    proyeccion.cliente = snapshot.get('cliente') or ''
    proyeccion.campania = snapshot.get('campania') or ''
    proyeccion.year = int(snapshot.get('year') or str(snapshot.get('mes', '0'))[:4] or 0)
    proyeccion.mes = snapshot.get('mes') or ''
    proyeccion.dotacion_requerida = snapshot.get('dotacion_requerida') or 0
    proyeccion.carga_semanal = normalizar_carga_semanal(snapshot.get('carga_semanal'))
    proyeccion.carga_horaria = snapshot.get('carga_horaria') or 0
    proyeccion.dias_objetivo = snapshot.get('dias_objetivo') or 0
    proyeccion.horas_requeridas = snapshot.get('horas_requeridas') or 0
    proyeccion.porcentaje_cumplimiento = snapshot.get('porcentaje_cumplimiento') or 0
    proyeccion.tiene_nocturnidad = bool(snapshot.get('tiene_nocturnidad'))
    proyeccion.porcentaje_nocturnidad = snapshot.get('porcentaje_nocturnidad') or 0
    proyeccion.tipo_plp = snapshot.get('tipo_plp') or None
    proyeccion.horas_carga_manual = bool(snapshot.get('horas_carga_manual'))
    proyeccion.dias_objetivo_manual = snapshot.get('dias_objetivo_manual')
    proyeccion.horas_requeridas_manual = snapshot.get('horas_requeridas_manual')
    proyeccion.jornadas = [
        ProyeccionMatrizJornada(
            dotacion_requerida=jornada.get('dotacion_requerida') or 0,
            carga_semanal=normalizar_carga_semanal(jornada.get('carga_semanal')),
            carga_horaria=jornada.get('carga_horaria') or 0,
            dias_objetivo=jornada.get('dias_objetivo') or 0,
            horas_requeridas=jornada.get('horas_requeridas') or 0,
        )
        for jornada in snapshot.get('jornadas', [])
    ]


def buscar_precio_snapshot(snapshot):
    if not snapshot:
        return None
    precio_id = snapshot.get('id')
    if precio_id:
        encontrada = ProyeccionPrecio.query.get(precio_id)
        if encontrada:
            return encontrada
    return ProyeccionPrecio.query.filter_by(
        cliente=snapshot.get('cliente'),
        campania=snapshot.get('campania'),
        mes=snapshot.get('mes'),
    ).first()


def restaurar_precio_snapshot(precio, snapshot):
    precio.site = snapshot.get('site') or ''
    precio.cliente = snapshot.get('cliente') or ''
    precio.campania = snapshot.get('campania') or ''
    precio.year = int(snapshot.get('year') or str(snapshot.get('mes', '0'))[:4] or 0)
    precio.mes = snapshot.get('mes') or ''
    precio.precio_base = redondear_moneda(snapshot.get('precio_base') or 0)
    precio.alcance_porcentaje = snapshot.get('alcance_porcentaje') or 0
    precio.importe_fijo_mensual = redondear_moneda(snapshot.get('importe_fijo_mensual') or 0)
    precio.recalcular()


def buscar_variable_snapshot(snapshot):
    if not snapshot:
        return None
    variable_id = snapshot.get('id')
    if variable_id:
        encontrada = VariableCampania.query.get(variable_id)
        if encontrada:
            return encontrada
    return VariableCampania.query.filter_by(
        cliente=snapshot.get('cliente'),
        campania=snapshot.get('campania'),
        mes=snapshot.get('mes'),
    ).first()


def restaurar_variable_snapshot(variable, snapshot):
    variable.site = snapshot.get('site') or ''
    variable.cliente = snapshot.get('cliente') or ''
    variable.campania = snapshot.get('campania') or ''
    variable.year = int(snapshot.get('year') or str(snapshot.get('mes', '0'))[:4] or 0)
    variable.mes = snapshot.get('mes') or ''
    variable.porcentaje = snapshot.get('porcentaje') or 0


def buscar_tarifacion_snapshot(snapshot):
    if not snapshot:
        return None
    if snapshot.get('id'):
        encontrada = db.session.get(TarifacionCampania, snapshot['id'])
        if encontrada:
            return encontrada
    return TarifacionCampania.query.filter_by(
        cliente=snapshot.get('cliente'), campania=snapshot.get('campania'),
        concepto=snapshot.get('concepto'), mes=snapshot.get('mes'),
    ).first()


def restaurar_tarifacion_snapshot(registro, snapshot):
    registro.site = snapshot.get('site') or ''
    registro.cliente = snapshot.get('cliente') or ''
    registro.campania = snapshot.get('campania') or ''
    registro.concepto = snapshot.get('concepto') or 'Tarifación'
    registro.year = int(snapshot.get('year') or str(snapshot.get('mes', '0'))[:4] or 0)
    registro.mes = snapshot.get('mes') or ''
    registro.monto = snapshot.get('monto') or 0


def buscar_next_gen_snapshot(snapshot):
    if not snapshot:
        return None
    modelo = NextGenDolar if snapshot.get('tipo_registro') == 'dolar' else NextGenProducto
    if snapshot.get('id'):
        encontrado = db.session.get(modelo, snapshot['id'])
        if encontrado:
            return encontrado
    if modelo is NextGenDolar:
        return NextGenDolar.query.filter_by(mes=snapshot.get('mes')).first()
    return NextGenProducto.query.filter_by(
        cliente=snapshot.get('cliente'), campania=snapshot.get('campania'),
        producto=snapshot.get('producto'), mes=snapshot.get('mes'),
    ).first()


def restaurar_next_gen_snapshot(registro, snapshot):
    registro.year = int(snapshot.get('year') or str(snapshot.get('mes', '0'))[:4] or 0)
    registro.mes = snapshot.get('mes') or ''
    if isinstance(registro, NextGenDolar):
        registro.valor = snapshot.get('valor') or 0
    else:
        registro.site = snapshot.get('site') or ''
        registro.cliente = snapshot.get('cliente') or ''
        registro.campania = snapshot.get('campania') or ''
        registro.producto = snapshot.get('producto') or ''
        registro.cantidad_usd = snapshot.get('cantidad_usd') or 0


def buscar_site_proyeccion_snapshot(snapshot):
    if not snapshot:
        return None
    if snapshot.get('id'):
        encontrado = db.session.get(SiteProyeccion, snapshot['id'])
        if encontrado:
            return encontrado
    return SiteProyeccion.query.filter_by(cliente=snapshot.get('cliente'), campania=snapshot.get('campania')).first()


def restaurar_site_proyeccion_snapshot(registro, snapshot):
    registro.cliente = snapshot.get('cliente') or ''
    registro.campania = snapshot.get('campania') or ''
    registro.site = snapshot.get('site') or ''
    registro.cliente_destino = snapshot.get('cliente_destino') or None
    registro.campania_destino = snapshot.get('campania_destino') or None


def lista_snapshots_historial(value):
    if not value:
        return []
    if isinstance(value, dict) and isinstance(value.get('proyecciones'), list):
        return value['proyecciones']
    if isinstance(value, dict) and isinstance(value.get('precios'), list):
        return value['precios']
    if isinstance(value, dict) and isinstance(value.get('variables'), list):
        return value['variables']
    if isinstance(value, dict) and isinstance(value.get('tarifaciones'), list):
        return value['tarifaciones']
    if isinstance(value, dict) and isinstance(value.get('next_gen'), list):
        return value['next_gen']
    if isinstance(value, dict) and isinstance(value.get('sites_proyeccion'), list):
        return value['sites_proyeccion']
    if isinstance(value, dict) and isinstance(value.get('filas'), list):
        return value['filas']
    if isinstance(value, dict) and value.get('tipo_registro') in ('dolar', 'producto'):
        return [value]
    if isinstance(value, dict) and {'cliente', 'campania', 'site'}.issubset(value.keys()) and 'mes' not in value:
        return [value]
    if isinstance(value, dict) and {'cliente', 'campania', 'mes'}.issubset(value.keys()):
        return [value]
    if isinstance(value, dict) and {'fecha', 'requerido', 'staff'}.issubset(value.keys()):
        return [value]
    if isinstance(value, dict) and {'fecha', 'operaciones', 'staff'}.issubset(value.keys()):
        return [value]
    if isinstance(value, list):
        return value
    return []


def historial_movimiento_deshace(historial_id):
    patron = f'Deshacer movimiento #{historial_id}:%'
    return HistorialCambio.query.filter(
        HistorialCambio.accion == 'deshacer',
        HistorialCambio.resumen.like(patron),
    ).first() is not None


def existe_proyeccion_snapshot(snapshot):
    if not snapshot:
        return False
    query = ProyeccionMatriz.query
    proyeccion_id = snapshot.get('id')
    if proyeccion_id and query.filter_by(id=proyeccion_id).first():
        return True
    return query.filter_by(
        cliente=snapshot.get('cliente'),
        campania=snapshot.get('campania'),
        mes=snapshot.get('mes'),
    ).first() is not None


def existe_precio_snapshot(snapshot):
    if not snapshot:
        return False
    query = ProyeccionPrecio.query
    precio_id = snapshot.get('id')
    if precio_id and query.filter_by(id=precio_id).first():
        return True
    return query.filter_by(
        cliente=snapshot.get('cliente'),
        campania=snapshot.get('campania'),
        mes=snapshot.get('mes'),
    ).first() is not None


def existe_variable_snapshot(snapshot):
    if not snapshot:
        return False
    query = VariableCampania.query
    variable_id = snapshot.get('id')
    if variable_id and query.filter_by(id=variable_id).first():
        return True
    return query.filter_by(
        cliente=snapshot.get('cliente'),
        campania=snapshot.get('campania'),
        mes=snapshot.get('mes'),
    ).first() is not None


def existe_tarifacion_snapshot(snapshot):
    return buscar_tarifacion_snapshot(snapshot) is not None if snapshot else False


def existe_next_gen_snapshot(snapshot):
    return buscar_next_gen_snapshot(snapshot) is not None if snapshot else False


def existe_site_proyeccion_snapshot(snapshot):
    return buscar_site_proyeccion_snapshot(snapshot) is not None if snapshot else False


def metricas_matriz(registros):
    registros_totales = list(registros)
    registros = [r for r in registros_totales if not r.es_next_gen]
    resumen = resumen_registros(registros)
    facturado_horas = sum(r.facturado_horas for r in registros)
    variable_real = sum(r.variable_productivo_calculo for r in registros)
    facturado_bono = sum(r.facturado_bono for r in registros)
    penalizaciones = sum(r.penalizaciones_incumplimientos for r in registros)
    tarifacion = sum(r.tarifacion or 0 for r in registros_totales)
    netx_gen = sum(r.netx_gen or 0 for r in registros_totales)
    otros = sum(r.otros or 0 for r in registros_totales)
    objetivo_horas = sum(r.objetivo_facturacion_horas for r in registros)
    objetivo_bono = sum(r.objetivo_facturacion_bono for r in registros)
    horas_netas_facturadas = max(resumen['horas_facturadas'] - resumen['horas_penalizadas'], 0)
    horas_objetivo_tarifadas = sum(r.horas_objetivo for r in registros if (r.horas_objetivo or 0) > 0)
    horas_netas_tarifadas = sum(
        max((r.horas_facturadas or 0) - (r.horas_penalizadas or 0), 0)
        for r in registros
        if max((r.horas_facturadas or 0) - (r.horas_penalizadas or 0), 0) > 0
    )
    valor_hora_objetivo = (
        sum((r.valor_hora_objetivo_calculo or 0) * (r.horas_objetivo or 0) for r in registros) / horas_objetivo_tarifadas
        if horas_objetivo_tarifadas > 0 else 0
    )
    valor_hora_realizado = (
        sum((r.valor_hora_alcanzado or 0) * max((r.horas_facturadas or 0) - (r.horas_penalizadas or 0), 0) for r in registros) / horas_netas_tarifadas
        if horas_netas_tarifadas > 0 else 0
    )
    horas_con_penalidad_adh = horas_netas_facturadas
    desvio_horas = horas_netas_facturadas - resumen['horas_objetivo']
    desvio_horas_monto = facturado_horas - objetivo_horas
    desvio_variable = variable_real - objetivo_bono
    neto_penalizaciones_bonos = facturado_bono + penalizaciones
    neto_otros_ajustes = 0
    total_objetivo = resumen['total_teorico']
    total_real = facturado_horas + variable_real + neto_penalizaciones_bonos + neto_otros_ajustes
    return {
        **resumen,
        'horas_netas_facturadas': round(horas_netas_facturadas, 2),
        'total_real': round(total_real, 2),
        'total_facturado': round(total_real, 2),
        'desvio': round(total_real - total_objetivo, 2),
        'porcentaje_cumplimiento': round(
            (total_real / total_objetivo * 100)
            if total_objetivo > 0 else 0,
            2
        ),
        'valor_hora_objetivo': round(valor_hora_objetivo, 2),
        'valor_hora_realizado': round(valor_hora_realizado, 2),
        'porcentaje_cumplimiento_horas': round(
            (resumen['horas_facturadas'] / resumen['horas_objetivo'] * 100)
            if resumen['horas_objetivo'] > 0 else 0,
            2
        ),
        'porcentaje_cumplimiento_logueo': round(
            (resumen['horas_facturadas'] / resumen['horas_objetivo'] * 100)
            if resumen['horas_objetivo'] > 0 else 0,
            2
        ),
        'porcentaje_cumplimiento_horas_adh': round(
            (horas_con_penalidad_adh / resumen['horas_objetivo'] * 100)
            if resumen['horas_objetivo'] > 0 else 0,
            2
        ),
        'porcentaje_valor_hora': round(
            (valor_hora_realizado / valor_hora_objetivo * 100)
            if valor_hora_objetivo > 0 else 0,
            2
        ),
        'porcentaje_bono': round(
            (facturado_bono / objetivo_bono * 100)
            if objetivo_bono > 0 else 0,
            2
        ),
        'desvio_horas': round(desvio_horas, 2),
        'desvio_horas_monto': round(desvio_horas_monto, 2),
        'horas_obj': round(objetivo_horas, 2),
        'variable_obj': round(objetivo_bono, 2),
        'total_obj': round(total_objetivo, 2),
        'horas_real': round(facturado_horas, 2),
        'variable_real': round(variable_real, 2),
        'bonos_real': round(facturado_bono, 2),
        'penalizaciones_real': round(penalizaciones, 2),
        'tarifacion_real': round(tarifacion, 2),
        'netx_gen_real': round(netx_gen, 2),
        'otros_real': round(otros, 2),
        'penalizaciones_bonos': round(neto_penalizaciones_bonos, 2),
        'otros_ajustes': round(neto_otros_ajustes, 2),
        'desvio_variable': round(desvio_variable, 2),
        'desvio_penalizaciones_bonos': round(neto_penalizaciones_bonos, 2),
        'desvio_otros_ajustes': round(neto_otros_ajustes, 2),
    }


def matriz_grupos(registros, campo, meses):
    grupos = {}
    for registro in registros:
        nombre = getattr(registro, campo) or 'Sin asignar'
        grupos.setdefault(nombre, []).append(registro)

    salida = []
    for nombre, registros_grupo in grupos.items():
        por_mes = {}
        for mes in meses:
            registros_mes = [r for r in registros_grupo if r.mes == mes]
            por_mes[mes] = metricas_matriz(registros_mes) if registros_mes else None
        salida.append({
            'nombre': nombre,
            'total': metricas_matriz(registros_grupo),
            'meses': por_mes,
        })
    salida.sort(key=lambda item: item['nombre'])
    return salida


def validar_payload_facturacion(data, exigir_configuracion_valor_hora=True, preservar_valores_importados=False):
    errores = []
    es_next_gen = bool(data.get('es_next_gen'))
    es_personal_directo = bool(data.get('_importacion_directa_personal'))
    try:
        aplicar_excepcion_calculo(
            data,
            exigir_configuracion=exigir_configuracion_valor_hora,
            preservar_facturado_manual=preservar_valores_importados,
        )
    except ValueError as error:
        errores.append(str(error))
    tipo_jornada = normalizar_tipo_vh(data.get('tipo_jornada'))
    sin_restriccion_horas = tipo_jornada == TIPO_VH_PERSONAL_COBRANZAS
    for campo, mensaje in [
        ('fecha', 'La fecha es obligatoria'),
        ('mes', 'El mes de facturacion es obligatorio'),
        ('cliente', 'El cliente es obligatorio'),
        ('gerente', 'El gerente es obligatorio'),
        ('jefe_site', 'El jefe de site es obligatorio'),
        ('campania', 'La campaña es obligatoria'),
        ('subcampania', 'La sub campaña es obligatoria'),
        ('tipo_jornada', 'El tipo de VH es obligatorio'),
    ]:
        if es_next_gen and campo not in ('fecha', 'mes', 'cliente'):
            continue
        if not str(data.get(campo, '')).strip():
            errores.append(mensaje)

    try:
        horas_objetivo = parse_numero(data.get('horas_objetivo'))
        horas_facturadas = parse_numero(data.get('horas_facturadas'))
        horas_penalizadas = parse_numero(data.get('horas_penalizadas'))
        valor_hora = parse_numero(data.get('valor_hora'))
        valor_hora_objetivo = parse_numero(data.get('valor_hora_objetivo', valor_hora))
        importe_fijo = parse_numero(data.get('importe_fijo')) if data.get('importe_fijo') not in (None, '') else None
        variable_objetivo = parse_numero(data.get('variable_objetivo'))
        variable_productivo = parse_numero(data.get('variable_productivo'))
    except ValueError:
        errores.append('Hay valores numéricos con formato inválido')
        horas_objetivo = horas_facturadas = horas_penalizadas = valor_hora = valor_hora_objetivo = 0
        importe_fijo = None
        variable_objetivo = 0
        variable_productivo = 0
    if horas_objetivo < 0 and not es_personal_directo:
        errores.append('Las horas objetivo no pueden ser negativas')
    if horas_facturadas < 0 and not es_personal_directo:
        errores.append('Las horas facturadas no pueden ser negativas')
    if horas_penalizadas < 0 and not es_personal_directo:
        errores.append('Las horas penalizadas no pueden ser negativas')
    if not es_next_gen and not es_personal_directo and not sin_restriccion_horas and horas_penalizadas > horas_facturadas:
        errores.append('Las horas penalizadas no pueden superar las horas facturadas')
    if not es_next_gen and not es_personal_directo and not sin_restriccion_horas and valor_hora <= 0:
        errores.append('El valor hora debe ser mayor a 0')
    if not es_next_gen and not es_personal_directo and not sin_restriccion_horas and valor_hora_objetivo <= 0:
        errores.append('El valor hora objetivo debe ser mayor a 0')
    if importe_fijo is not None and importe_fijo < 0:
        errores.append('El importe fijo facturado no puede ser negativo')
    if variable_objetivo < 0:
        errores.append('Variable Objetivo no puede ser negativo')
    if tipo_jornada and tipo_jornada not in TIPOS_VH and not es_next_gen and not es_personal_directo:
        errores.append(f'El tipo de VH "{data.get("tipo_jornada", "")}" no es válido')
    if data.get('mes') and not mes_valido(data.get('mes')):
        errores.append('El mes de facturacion no es valido')
    return errores


def asegurar_asignacion_desde_registro(registro):
    campos = {
        'cliente': registro.cliente,
        'gerente': registro.gerente,
        'jefe_site': registro.jefe_site,
        'campania': registro.campania,
        'subcampania': registro.subcampania,
        'tipo_negocio': registro.tipo_negocio,
        'es_next_gen': bool(registro.es_next_gen),
    }
    if not all(valor for campo, valor in campos.items() if campo != 'tipo_negocio'):
        return None
    existente = AsignacionComercial.query.filter_by(**campos).first()
    if existente:
        existente.activa = True
        if not existente.campania_id:
            existente.campania_catalogo = obtener_o_crear_campania(registro.cliente, registro.campania)
        return existente
    campania = obtener_o_crear_campania(registro.cliente, registro.campania)
    asignacion = AsignacionComercial(**campos, campania_catalogo=campania)
    db.session.add(asignacion)
    return asignacion


@main_bp.route('/')
@login_requerido
def index():
    """Dashboard principal"""
    return render_template('index.html')


@main_bp.route('/favicon.ico')
def favicon():
    return Response(status=204)


@main_bp.route('/cargar')
@carga_requerida
def cargar():
    """Vista de carga de datos"""
    return render_template('cargar.html')


@main_bp.route('/control')
@login_requerido
def control():
    """Vista de control de datos"""
    return render_template('control.html')


@main_bp.route('/resumen-horas')
@login_requerido
def resumen_horas():
    """Control mensual de dotación requerida y horas reales, como el cuadro Excel."""
    return render_template('resumen_horas.html')


@main_bp.route('/justificaciones')
@login_requerido
def justificaciones():
    """Vista de control de justificaciones de ajustes."""
    return render_template('justificaciones.html')


@main_bp.route('/comparativo')
@login_requerido
def comparativo():
    """Vista comparativa de horas objetivo contra horas facturadas"""
    return render_template('comparativo.html')


@main_bp.route('/matriz')
@login_requerido
def matriz():
    """Vista matricial mensual por gerencia y jefe de site."""
    return render_template('matriz.html')


@main_bp.route('/matriz-proyecciones')
@login_requerido
def matriz_proyecciones():
    """Vista de planificación mensual de dotación y horas proyectadas."""
    return render_template('matriz_proyecciones.html')


@main_bp.route('/matriz-precios')
@login_requerido
def matriz_precios():
    """Vista de planificación mensual de precios por campaña."""
    return render_template('matriz_precios.html')


@main_bp.route('/suma-fija')
@login_requerido
def suma_fija():
    """Importes mensuales por dotación que integran Facturación horas."""
    return render_template('suma_fija.html')


@main_bp.route('/resumen')
@login_requerido
def resumen():
    """Resumen anual de horas proyectadas por precio mensual."""
    return render_template('resumen.html')


@main_bp.route('/variable')
@login_requerido
def variable():
    """Vista de variables por campaña sobre facturación horas."""
    return render_template('variable.html')


@main_bp.route('/control-proyecciones')
@login_requerido
def control_proyecciones():
    """Vista de control mensual de horas y dotaciones proyectadas."""
    return render_template('control_proyecciones.html')


@main_bp.route('/indicadores')
@login_requerido
def indicadores():
    return render_template('indicadores.html')


@main_bp.route('/reloj')
@login_requerido
def reloj():
    """Reloj de cumplimiento de horas."""
    return render_template('reloj.html')


@main_bp.route('/dotaciones')
@login_requerido
def dotaciones():
    """Evolución y carga mensual de dotaciones."""
    return render_template('dotaciones.html')


@main_bp.route('/graficos-dotaciones')
@login_requerido
def graficos_dotaciones():
    return render_template('graficos_dotaciones.html')


@main_bp.route('/graficos-evolutivo-dotaciones')
@login_requerido
def graficos_evolutivo_dotaciones():
    return render_template('graficos_evolutivo_dotaciones.html')


@main_bp.route('/comparativo-anual-clientes')
@login_requerido
def comparativo_anual_clientes():
    anios = {fila[0] for fila in db.session.query(HistoricoClienteMensual.year).distinct().all()}
    anios.update(anios_facturacion_disponibles())
    anios.update(fila[0].year for fila in db.session.query(DotacionClienteMensual.fecha).distinct().all() if fila[0])
    anios = sorted(anio for anio in anios if anio >= 2024)
    return render_template('comparativo_anual_clientes.html', anios_comparativo=anios,
                           anio_comparativo=anios[-1] if anios else current_app.config['DEFAULT_YEAR'])


@main_bp.route('/ratio-eli')
@login_requerido
def ratio_eli():
    return render_template('ratio_eli.html')


@main_bp.route('/api/ratio-eli', methods=['GET'])
@login_requerido
def api_ratio_eli():
    ruta = os.path.join(current_app.root_path, 'data', 'ratio_eli_historico.json')
    with open(ruta, encoding='utf-8') as archivo:
        base = json.load(archivo)
    for mes, requerido, staff in base:
        fecha = datetime.strptime(mes, '%Y-%m').date()
        if not RatioEliMensual.query.filter_by(fecha=fecha).first():
            db.session.add(RatioEliMensual(fecha=fecha, requerido=requerido, staff=staff))
    db.session.commit()
    registros = RatioEliMensual.query.filter(RatioEliMensual.fecha >= date(2023, 1, 1)).order_by(RatioEliMensual.fecha).all()
    return jsonify({'success': True, 'serie': [x.to_dict() for x in registros]})


@main_bp.route('/api/ratio-eli', methods=['POST'])
@edicion_requerida
def api_guardar_ratio_eli():
    data=request.get_json(silent=True) or {}
    try:
        fecha=datetime.strptime(str(data.get('mes') or ''),'%Y-%m').date(); requerido=parse_numero(data.get('requerido')); staff=parse_numero(data.get('staff'))
        if requerido<=0 or staff<0: raise ValueError
    except (TypeError,ValueError): return jsonify({'success':False,'errores':['Datos inválidos']}),400
    r=RatioEliMensual.query.filter_by(fecha=fecha).first(); antes=r.to_dict() if r else None
    if not r: r=RatioEliMensual(fecha=fecha,requerido=requerido,staff=staff);db.session.add(r)
    else: r.requerido,r.staff=requerido,staff
    db.session.flush(); registrar_historial('edicion' if antes else 'creacion','ratio_eli_mensual',fecha.isoformat(),f'Ratio Eli: {fecha:%m/%Y}',antes=antes,despues=r.to_dict());db.session.commit()
    return jsonify({'success':True,'mensaje':'Mes guardado'})


@main_bp.route('/ratio-eli-ii')
@login_requerido
def ratio_eli_ii():
    return render_template('ratio_eli_ii.html')


@main_bp.route('/api/ratio-eli-ii', methods=['GET'])
@login_requerido
def api_ratio_eli_ii():
    ruta = os.path.join(current_app.root_path, 'data', 'ratio_eli_ii_historico.json')
    with open(ruta, encoding='utf-8') as archivo:
        base = json.load(archivo)
    for mes, operaciones, staff, operaciones_importe, staff_importe in base:
        fecha = datetime.strptime(mes, '%Y-%m').date()
        registro = RatioEliIIMensual.query.filter_by(fecha=fecha).first()
        if not registro:
            db.session.add(RatioEliIIMensual(fecha=fecha, operaciones=operaciones, staff=staff,
                                             operaciones_importe=operaciones_importe, staff_importe=staff_importe))
        elif not registro.operaciones_importe and not registro.staff_importe:
            registro.operaciones_importe, registro.staff_importe = operaciones_importe, staff_importe
    db.session.commit()
    registros = RatioEliIIMensual.query.order_by(RatioEliIIMensual.fecha).all()
    return jsonify({'success': True, 'serie': [x.to_dict() for x in registros]})


@main_bp.route('/api/ratio-eli-ii', methods=['POST'])
@edicion_requerida
def api_guardar_ratio_eli_ii():
    data = request.get_json(silent=True) or {}
    try:
        fecha = datetime.strptime(str(data.get('mes') or ''), '%Y-%m').date()
        operaciones = parse_numero(data.get('operaciones'))
        staff = parse_numero(data.get('staff'))
        operaciones_importe = parse_numero(data.get('operaciones_importe'))
        staff_importe = parse_numero(data.get('staff_importe'))
        if operaciones <= 0 or staff < 0 or operaciones_importe <= 0 or staff_importe < 0:
            raise ValueError
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['Datos invalidos']}), 400
    registro = RatioEliIIMensual.query.filter_by(fecha=fecha).first()
    antes = registro.to_dict() if registro else None
    if not registro:
        registro = RatioEliIIMensual(fecha=fecha, operaciones=operaciones, staff=staff,
                                     operaciones_importe=operaciones_importe, staff_importe=staff_importe)
        db.session.add(registro)
    else:
        registro.operaciones, registro.staff = operaciones, staff
        registro.operaciones_importe, registro.staff_importe = operaciones_importe, staff_importe
    db.session.flush()
    registrar_historial('edicion' if antes else 'creacion', 'ratio_eli_ii_mensual', fecha.isoformat(),
                        f'Ratio Eli II: {fecha:%m/%Y}', antes=antes, despues=registro.to_dict())
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Mes guardado'})


@main_bp.route('/share')
@login_requerido
def share():
    return render_template('share.html')


@main_bp.route('/api/share', methods=['GET'])
@login_requerido
def api_share():
    ruta_usa = os.path.join(current_app.root_path, 'data', 'share_usa_historico.json')
    with open(ruta_usa, encoding='utf-8') as archivo:
        base_usa = json.load(archivo)
        usa_historico = {int(year): float(valor) for year, valor in base_usa['horas'].items()}
        usa_dotaciones = {int(year): float(valor) for year, valor in base_usa['dotaciones'].items()}
    anios_historico = {x[0] for x in db.session.query(HistoricoClienteMensual.year).distinct().all()}
    anios_dotacion = {x[0] for x in db.session.query(func.extract('year', DotacionClienteMensual.fecha)).distinct().all()}
    anios_facturacion = anios_facturacion_disponibles()
    anios = sorted({
        int(x) for x in anios_historico | anios_dotacion | anios_facturacion
        if x and int(x) >= 2020
    })
    horas_por_anio = {}
    dotaciones_por_anio = {}

    for year in anios:
        por_clave = {}
        for registro in HistoricoClienteMensual.query.filter_by(year=year).all():
            base = registro.base_dict()
            por_clave[(registro.mes, registro.cliente)] = {
                'real': float(base.get('horas_realizadas') or 0),
                'snapshot': float(base.get('_fact_snapshot_real') or 0),
            }
        facturado = facturacion_por_anio(db.session.query(
            Facturacion2026.mes, Facturacion2026.cliente,
            func.sum(Facturacion2026.horas_facturadas),
        ), year).group_by(Facturacion2026.mes, Facturacion2026.cliente).all()
        for mes, cliente, real_actual in facturado:
            clave = (mes, cliente)
            item = por_clave.setdefault(clave, {'real': 0, 'snapshot': 0})
            item['real'] += float(real_actual or 0) - item['snapshot']
        clientes = {}
        for (_, cliente), item in por_clave.items():
            clientes[cliente] = clientes.get(cliente, 0) + item['real']
        def suma_clientes(predicado):
            return sum(valor for nombre, valor in clientes.items() if predicado(normalizar_header(nombre.strip("'\""))))
        if year <= 2022:
            personal = suma_clientes(lambda nombre: nombre in ('cablevision', 'telecom'))
        else:
            personal = suma_clientes(lambda nombre: nombre in ('telecom', 'personal'))
        santander = suma_clientes(lambda nombre: nombre == 'santander')
        usa = usa_historico.get(year, suma_clientes(lambda nombre: nombre in ('abbey usa', 'usa')))
        total = sum(clientes.values())
        multicampania = total - personal - santander - usa
        horas_por_anio[year] = {'Personal': personal, 'Santander': santander,
                                'Multicampania': multicampania, 'USA': usa, 'Total': total}

        meses = {}
        for fila in DotacionClienteMensual.query.filter(func.extract('year', DotacionClienteMensual.fecha) == year).all():
            mes = fila.fecha.strftime('%Y-%m')
            item = meses.setdefault(mes, {'Personal': 0, 'Santander': 0, 'Multicampania': 0, 'USA': 0})
            nombre = normalizar_header(fila.cliente)
            if nombre in ('personal', 'telecom', 'cablevision'):
                grupo = 'Personal'
            elif nombre == 'santander':
                grupo = 'Santander'
            elif nombre in ('abbey usa', 'usa'):
                grupo = 'USA'
            else:
                grupo = 'Multicampania'
            item[grupo] += float(fila.dotacion or 0)
        cantidad_meses = len(meses)
        promedios = {grupo: (sum(m[grupo] for m in meses.values()) / cantidad_meses if cantidad_meses else 0)
                     for grupo in ('Personal', 'Santander', 'Multicampania', 'USA')}
        promedios['Total'] = sum(promedios.values())
        if year in usa_dotaciones:
            promedios['USA'] = usa_dotaciones[year]
            promedios['Multicampania'] = promedios['Total'] - promedios['Personal'] - promedios['Santander'] - promedios['USA']
        dotaciones_por_anio[year] = promedios

    def con_share(base):
        salida = []
        for year in anios:
            valores = base.get(year, {})
            total = float(valores.get('Total') or 0)
            fila = {'year': year, **valores}
            fila['shares'] = {grupo: (float(valores.get(grupo) or 0) / total if total else 0)
                              for grupo in ('Personal', 'Santander', 'Multicampania', 'USA')}
            fila['personal_santander'] = float(valores.get('Personal') or 0) + float(valores.get('Santander') or 0)
            fila['share_personal_santander'] = fila['personal_santander'] / total if total else 0
            salida.append(fila)
        return salida

    return jsonify({'success': True, 'anios': anios, 'horas': con_share(horas_por_anio),
                    'dotaciones': con_share(dotaciones_por_anio)})


@main_bp.route('/variacion-anual')
@login_requerido
def variacion_anual():
    return render_template('variacion_anual.html')


@main_bp.route('/variacion-horas-clientes')
@login_requerido
def variacion_horas_clientes():
    return render_template('variacion_horas_clientes.html')


@main_bp.route('/dashboard-operativo')
@login_requerido
def dashboard_operativo():
    return render_template('dashboard_operativo.html')


@main_bp.route('/historico-clientes')
@login_requerido
def historico_clientes():
    return render_template('historico_clientes.html')


@main_bp.route('/cumplimiento-horas-clientes')
@login_requerido
def cumplimiento_horas_clientes():
    return render_template('cumplimiento_horas_clientes.html')


@main_bp.route('/cumplimiento-facturacion-clientes')
@login_requerido
def cumplimiento_facturacion_clientes():
    return render_template('cumplimiento_facturacion_clientes.html')


@main_bp.route('/comparativo-interanual')
@login_requerido
def comparativo_interanual():
    return render_template('comparativo_interanual.html')


@main_bp.route('/calendario-operativo')
@login_requerido
def calendario_operativo():
    """Configuración anual de feriados operativos."""
    return render_template('calendario_operativo.html')


@main_bp.route('/catalogos')
@edicion_requerida
def catalogos():
    """Vista de alta de datos maestros"""
    return render_template('catalogos.html')


@main_bp.route('/catalogos/crear', methods=['GET', 'POST'])
@edicion_requerida
def catalogos_crear():
    """Alta simple de datos maestros desde formulario HTML."""
    if request.method == 'GET':
        return redirect(url_for('main.catalogos'))
    respuesta = crear_asignacion_desde_request()
    status = respuesta[1] if isinstance(respuesta, tuple) and len(respuesta) > 1 else getattr(respuesta, 'status_code', 200)
    if status >= 400:
        return redirect(url_for('main.catalogos', estado='error'))
    return redirect(url_for('main.catalogos', estado='creado'))


@main_bp.route('/login')
def login():
    """Vista de acceso y primera configuracion."""
    asegurar_administrador_inicial()
    if requiere_login():
        if usuario_actual().debe_cambiar_password:
            return redirect(url_for('main.cambiar_password_obligatorio'))
        return redirect(url_for('main.index'))
    return render_template('login.html', requiere_setup=not usuarios_registrados())


@main_bp.route('/logout', methods=['POST'])
def logout():
    session.clear()
    return redirect(url_for('main.login'))


@main_bp.route('/cambiar-contrasena')
def cambiar_password_obligatorio():
    usuario = usuario_actual()
    if not usuario or not usuario.activo:
        return redirect(url_for('main.login'))
    if not usuario.debe_cambiar_password:
        return redirect(url_for('main.index'))
    return render_template('cambiar_password.html')


@main_bp.route('/api/auth/cambiar-contrasena', methods=['POST'])
def api_cambiar_password_obligatorio():
    usuario = usuario_actual()
    if not usuario or not usuario.activo:
        return jsonify({'success': False, 'errores': ['No autorizado']}), 403
    data = request.get_json(silent=True) or {}
    actual = str(data.get('password_actual') or '')
    nueva = str(data.get('password_nueva') or '')
    confirmacion = str(data.get('password_confirmacion') or '')
    errores = []
    if not usuario.check_password(actual):
        errores.append('La contrasena temporal no es correcta')
    if len(nueva) < 12:
        errores.append('La nueva contrasena debe tener al menos 12 caracteres')
    elif nueva.lower() == nueva or nueva.upper() == nueva or not any(c.isdigit() for c in nueva):
        errores.append('La nueva contrasena debe combinar mayusculas, minusculas y numeros')
    if nueva != confirmacion:
        errores.append('La confirmacion no coincide')
    if nueva and usuario.check_password(nueva):
        errores.append('La nueva contrasena debe ser diferente de la temporal')
    if errores:
        return jsonify({'success': False, 'errores': errores}), 400
    usuario.set_password(nueva)
    usuario.debe_cambiar_password = False
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Contrasena actualizada correctamente'})


@main_bp.route('/usuarios')
@admin_requerido
def usuarios():
    """Vista de administracion de usuarios."""
    plantillas = {}
    for rol in ROLES_USUARIO:
        plantillas[rol] = {}
        for puesto in PUESTOS_POR_ROL.get(rol, ('',)):
            permisos = permisos_perfil(rol, puesto)
            plantillas[rol][puesto] = {**permisos, 'modulos': sorted(permisos['modulos'])}
    return render_template('usuarios.html', roles=ROLES_USUARIO, puestos_por_rol=PUESTOS_POR_ROL,
                           plantillas_permisos=plantillas)


@main_bp.route('/historial')
@admin_requerido
def historial():
    """Vista de historial de modificaciones y eliminaciones."""
    return render_template('historial.html')


@main_bp.route('/guia-usuario')
@login_requerido
def guia_usuario():
    """Abre la guia funcional, independiente de la documentacion tecnica."""
    ruta = os.path.abspath(os.path.join(current_app.root_path, '..', 'docs', 'guia_usuario.html'))
    # Se devuelve como respuesta HTML normal para que after_request pueda
    # incorporar el nonce CSP al <style> y al <script> embebidos.
    contenido = Path(ruta).read_text(encoding='utf-8')
    return Response(contenido, content_type='text/html; charset=utf-8')


class _TextoDesdeHtml(HTMLParser):
    """Convierte la guia HTML en texto legible para su version PDF."""

    def __init__(self):
        super().__init__()
        self.partes = []
        self.omitir = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('style', 'script'):
            self.omitir += 1
        elif tag in ('h1', 'h2', 'h3', 'h4', 'p', 'li', 'tr', 'section'):
            self.partes.append('\n')

    def handle_endtag(self, tag):
        if tag in ('style', 'script') and self.omitir:
            self.omitir -= 1
        elif tag in ('h1', 'h2', 'h3', 'h4', 'p', 'li', 'tr'):
            self.partes.append('\n')

    def handle_data(self, data):
        if not self.omitir:
            self.partes.append(data)

    def texto(self):
        texto = ''.join(self.partes)
        texto = re.sub(r'[ \t]+', ' ', texto)
        texto = re.sub(r'\n\s*\n\s*\n+', '\n\n', texto)
        return texto.strip()


def _pdf_desde_texto(titulo, texto):
    """Genera un PDF simple y portable sin dependencias externas."""
    lineas = [titulo, '']
    for linea_original in texto.splitlines():
        linea = linea_original.strip()
        if not linea:
            lineas.append('')
            continue
        lineas.extend(textwrap.wrap(linea, width=92, break_long_words=False) or [''])

    lineas_por_pagina = 50
    paginas = [lineas[i:i + lineas_por_pagina] for i in range(0, len(lineas), lineas_por_pagina)]
    objetos = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'',  # Se completa cuando se conocen todas las paginas.
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
    ]
    referencias_paginas = []

    def escapar_pdf(linea):
        datos = linea.encode('cp1252', errors='replace')
        return datos.replace(b'\\', b'\\\\').replace(b'(', b'\\(').replace(b')', b'\\)')

    for pagina in paginas:
        comandos = [b'BT', b'/F1 10 Tf', b'48 800 Td', b'14 TL']
        for linea in pagina:
            comandos.append(b'(' + escapar_pdf(linea) + b') Tj')
            comandos.append(b'T*')
        comandos.append(b'ET')
        contenido = b'\n'.join(comandos)
        numero_pagina = len(objetos) + 1
        numero_contenido = numero_pagina + 1
        referencias_paginas.append(numero_pagina)
        objetos.append(
            f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] '
            f'/Resources << /Font << /F1 3 0 R >> >> /Contents {numero_contenido} 0 R >>'.encode()
        )
        objetos.append(
            f'<< /Length {len(contenido)} >>\nstream\n'.encode() + contenido + b'\nendstream'
        )

    hijos = ' '.join(f'{numero} 0 R' for numero in referencias_paginas)
    objetos[1] = f'<< /Type /Pages /Kids [{hijos}] /Count {len(referencias_paginas)} >>'.encode()

    salida = bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
    offsets = [0]
    for numero, objeto in enumerate(objetos, start=1):
        offsets.append(len(salida))
        salida.extend(f'{numero} 0 obj\n'.encode())
        salida.extend(objeto)
        salida.extend(b'\nendobj\n')
    inicio_xref = len(salida)
    salida.extend(f'xref\n0 {len(objetos) + 1}\n'.encode())
    salida.extend(b'0000000000 65535 f \n')
    for offset in offsets[1:]:
        salida.extend(f'{offset:010d} 00000 n \n'.encode())
    salida.extend(
        f'trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n'
        f'startxref\n{inicio_xref}\n%%EOF\n'.encode()
    )
    return bytes(salida)


@main_bp.route('/guia-usuario/descargar')
@login_requerido
def descargar_guia_usuario():
    """Descarga la guia maquetada en PDF para cualquier usuario autenticado."""
    ruta = os.path.abspath(os.path.join(current_app.root_path, '..', 'docs', 'guia_usuario.pdf'))
    return send_file(
        ruta,
        mimetype='application/pdf',
        as_attachment=True,
        download_name='Guia_de_usuario_Dashboard_Facturacion.pdf',
    )


@main_bp.route('/documentacion-tecnica/descargar')
@admin_requerido
def descargar_documentacion_tecnica():
    """Descarga la documentación técnica maquetada; solo para administradores."""
    ruta = os.path.abspath(os.path.join(current_app.root_path, '..', 'docs', 'documentacion_tecnica.pdf'))
    return send_file(
        ruta,
        mimetype='application/pdf',
        as_attachment=True,
        download_name='Documentacion_tecnica_Dashboard_Facturacion.pdf',
    )


@main_bp.route('/documentacion-tecnica/esquema/<motor>')
@admin_requerido
def descargar_esquema_base_datos(motor):
    """Descarga el DDL completo entregable al responsable de la base."""
    if str(motor or '').lower() == 'pdf':
        ruta = os.path.abspath(os.path.join(current_app.root_path, '..', 'docs', 'esquema_base_datos.pdf'))
        return send_file(
            ruta,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='Esquema_Base_Datos_Dashboard.pdf',
        )
    archivos = {
        'sqlite': 'esquema_base_datos_sqlite.sql',
        'postgresql': 'esquema_base_datos_postgresql.sql',
    }
    nombre = archivos.get(str(motor or '').lower())
    if not nombre:
        abort(404)
    ruta = os.path.abspath(os.path.join(current_app.root_path, '..', 'docs', nombre))
    return send_file(
        ruta,
        mimetype='application/sql',
        as_attachment=True,
        download_name=nombre,
    )


@main_bp.route('/documentacion-tecnica')
@admin_requerido
def documentacion_tecnica():
    """Abre la documentación técnica legible; solo para administradores."""
    ruta = os.path.abspath(os.path.join(current_app.root_path, '..', 'docs', 'documentacion_tecnica.html'))
    contenido = Path(ruta).read_text(encoding='utf-8')
    return Response(contenido, content_type='text/html; charset=utf-8')


# ========== ENDPOINTS API ==========

@main_bp.route('/api/auth/me', methods=['GET'])
def api_auth_me():
    asegurar_administrador_inicial()
    return jsonify({
        'success': True,
        'usuario': usuario_actual_dict(),
        'requiere_setup': not usuarios_registrados(),
        'csrf_token': get_csrf_token(),
    })


@main_bp.route('/api/auth/login', methods=['POST'])
def api_auth_login():
    data = request.get_json() or {}
    email = normalizar_email(data.get('email'))
    password = str(data.get('password', '') or '')

    if login_bloqueado(email):
        return jsonify({'success': False, 'errores': ['Demasiados intentos. Espere unos minutos e intente nuevamente']}), 429

    usuario = Usuario.query.filter_by(email=email).first()
    if not usuario or not usuario.activo or not usuario.check_password(password):
        registrar_login_fallido(email)
        return jsonify({'success': False, 'errores': ['Email o contrasena incorrectos']}), 401

    session.clear()
    session.permanent = True
    session['usuario_id'] = usuario.id
    csrf_token = get_csrf_token()
    limpiar_login_fallido(email)
    return jsonify({'success': True, 'usuario': usuario.to_dict(), 'csrf_token': csrf_token,
                    'requiere_cambio_password': usuario.debe_cambiar_password})


@main_bp.route('/api/auth/logout', methods=['POST'])
def api_auth_logout():
    session.clear()
    return jsonify({'success': True})


@main_bp.route('/api/usuarios/setup', methods=['POST'])
def api_usuarios_setup():
    """Crea el primer administrador cuando la base aun no tiene usuarios."""
    if usuarios_registrados():
        return jsonify({'success': False, 'errores': ['La configuracion inicial ya fue realizada']}), 403

    data = request.get_json() or {}
    data['rol'] = 'Admin'
    data['puesto'] = 'Administrador'
    errores = validar_usuario_payload(data, require_password=True)
    if errores:
        return jsonify({'success': False, 'errores': errores}), 400

    usuario = Usuario(
        nombre=data['nombre'].strip(),
        email=normalizar_email(data['email']),
        rol='Admin',
        puesto='Administrador',
        activo=True,
    )
    usuario.set_password(data['password'])
    db.session.add(usuario)
    db.session.commit()
    session.clear()
    session.permanent = True
    session['usuario_id'] = usuario.id
    csrf_token = get_csrf_token()
    return jsonify({'success': True, 'usuario': usuario.to_dict(), 'csrf_token': csrf_token})


@main_bp.route('/api/usuarios', methods=['GET'])
@admin_requerido
def api_usuarios():
    usuarios = Usuario.query.order_by(Usuario.creado_en.desc()).all()
    return jsonify({'success': True, 'usuarios': [usuario.to_dict() for usuario in usuarios]})


@main_bp.route('/api/usuarios', methods=['POST'])
@admin_requerido
def api_crear_usuario():
    data = request.get_json() or {}
    errores = validar_usuario_payload(data, require_password=True)
    email = normalizar_email(data.get('email'))
    if Usuario.query.filter_by(email=email).first():
        errores.append('Ya existe un usuario con ese email')
    if errores:
        return jsonify({'success': False, 'errores': errores}), 400

    usuario = Usuario(
        nombre=data['nombre'].strip(),
        email=email,
        rol=data.get('rol', 'RMO_OPS').strip(),
        puesto=str(data.get('puesto') or '').strip() or None,
        gerente_asignado=str(data.get('gerente_asignado') or '').strip() or None,
        jefe_site_asignado=str(data.get('jefe_site_asignado') or '').strip() or None,
        permisos_personalizados=serializar_json(normalizar_permisos_usuario(data.get('permisos'))) if data.get('permisos') is not None else None,
        activo=bool(data.get('activo', True)),
        debe_cambiar_password=True,
    )
    if usuario.gerente_asignado:
        usuario.jefe_site_asignado = None
    usuario.set_password(data['password'])
    db.session.add(usuario)
    db.session.flush()
    registrar_historial(
        'creacion',
        'usuario',
        usuario.id,
        f'Usuario creado: {usuario.email}',
        despues=usuario.to_dict(),
    )
    db.session.commit()
    return jsonify({'success': True, 'usuario': usuario.to_dict()})


@main_bp.route('/api/usuarios/<int:usuario_id>', methods=['PATCH'])
@admin_requerido
def api_actualizar_usuario(usuario_id):
    usuario = Usuario.query.get_or_404(usuario_id)
    antes = snapshot_modelo(usuario)
    data = request.get_json() or {}
    errores = []

    if 'nombre' in data:
        nombre = str(data.get('nombre', '')).strip()
        if not nombre:
            errores.append('El nombre es obligatorio')
        else:
            usuario.nombre = nombre
    if 'email' in data:
        email = normalizar_email(data.get('email'))
        existente = Usuario.query.filter(Usuario.email == email, Usuario.id != usuario.id).first()
        if not email or '@' not in email:
            errores.append('El email es obligatorio y debe ser valido')
        elif existente:
            errores.append('Ya existe un usuario con ese email')
        else:
            usuario.email = email
    if 'rol' in data:
        rol = str(data.get('rol', '')).strip()
        if rol not in ROLES_USUARIO:
            errores.append('El rol no es valido')
        else:
            usuario.rol = rol
    for campo in ('puesto', 'gerente_asignado', 'jefe_site_asignado'):
        if campo in data:
            setattr(usuario, campo, str(data.get(campo) or '').strip() or None)
    if usuario.gerente_asignado:
        usuario.jefe_site_asignado = None
    if 'permisos' in data:
        try:
            permisos = normalizar_permisos_usuario(data.get('permisos'))
            usuario.permisos_personalizados = serializar_json(permisos) if permisos is not None else None
        except ValueError as error:
            errores.append(str(error))
    if usuario.puesto and usuario.puesto not in PUESTOS_POR_ROL.get(usuario.rol, ()):
        errores.append('El puesto no corresponde al perfil seleccionado')
    if usuario.rol != 'Full' and not usuario.puesto:
        errores.append('El puesto es obligatorio')
    if 'activo' in data:
        usuario.activo = bool(data.get('activo'))
    if data.get('password'):
        password = str(data.get('password'))
        if len(password) < 12:
            errores.append('La contrasena debe tener al menos 12 caracteres')
        elif password.lower() == password or password.upper() == password or not any(caracter.isdigit() for caracter in password):
            errores.append('La contrasena debe combinar mayusculas, minusculas y numeros')
        else:
            usuario.set_password(password)
            usuario.debe_cambiar_password = True

    if errores:
        return jsonify({'success': False, 'errores': errores}), 400

    despues = snapshot_modelo(usuario)
    cambios = cambios_entre(antes, despues)
    if cambios:
        registrar_historial(
            'edicion',
            'usuario',
            usuario.id,
            f'Usuario actualizado: {usuario.email}',
            antes=antes,
            despues=despues,
            detalle=serializar_json(cambios),
        )
    db.session.commit()
    return jsonify({'success': True, 'usuario': usuario.to_dict()})


@main_bp.route('/api/usuarios/<int:usuario_id>', methods=['DELETE'])
@admin_requerido
def api_eliminar_usuario(usuario_id):
    actual = usuario_actual()
    usuario = Usuario.query.get_or_404(usuario_id)
    data = request.get_json(silent=True) or {}
    if not validar_confirmacion_accion(data):
        return jsonify({'success': False, 'errores': ['La contraseña de confirmación no es válida']}), 403
    if actual and usuario.id == actual.id:
        return jsonify({'success': False, 'errores': ['No podés eliminar tu propia cuenta']}), 400

    antes = usuario.to_dict()
    # La auditoría histórica conserva nombre y email, pero libera la FK para
    # permitir eliminar la cuenta sin perder quién realizó operaciones previas.
    HistorialCambio.query.filter_by(usuario_id=usuario.id).update(
        {'usuario_id': None}, synchronize_session=False,
    )
    registrar_historial(
        'eliminacion', 'usuario', usuario.id,
        f'Usuario eliminado: {usuario.email}',
        antes=antes,
        detalle='Cuenta eliminada definitivamente; su auditoría histórica fue conservada.',
    )
    db.session.delete(usuario)
    db.session.commit()
    return jsonify({'success': True, 'mensaje': f'Usuario eliminado: {antes["email"]}'})


@main_bp.route('/api/historial', methods=['GET'])
@admin_requerido
def api_historial():
    query = HistorialCambio.query
    accion = request.args.get('accion')
    entidad = request.args.get('entidad')
    if accion:
        query = query.filter(HistorialCambio.accion == accion)
    if entidad:
        query = query.filter(HistorialCambio.entidad == entidad)
    limite = min(numero_request('limite') or 100, 300)
    items = query.order_by(HistorialCambio.creado_en.desc()).limit(int(limite)).all()
    return jsonify({'success': True, 'historial': [item.to_dict() for item in items]})


@main_bp.route('/api/historial/<int:historial_id>/deshacer', methods=['POST'])
@admin_requerido
def api_deshacer_historial(historial_id):
    item = HistorialCambio.query.get_or_404(historial_id)
    resultado, status = deshacer_item_historial(item)
    return jsonify(resultado), status


def deshacer_item_historial(item):
    if item.accion == 'deshacer':
        return {'success': False, 'errores': ['Este movimiento ya es un deshacer']}, 400
    if historial_movimiento_deshace(item.id):
        return {'success': False, 'errores': ['Este movimiento ya fue deshecho']}, 400
    entidades_permitidas = ('proyeccion_matriz', 'matriz_precios', 'variables', 'tarifaciones', 'next_gen', 'sites_proyeccion', 'dotaciones_clientes_mensuales', 'graficos_dotaciones_mensuales', 'dashboard_operativo', 'historico_cliente', 'ratio_eli_mensual', 'ratio_eli_ii_mensual')
    if item.entidad not in entidades_permitidas:
        return {'success': False, 'errores': ['Este tipo de movimiento todavía no admite deshacer']}, 400

    antes = item._json(item.antes) or {}
    despues = item._json(item.despues) or {}
    antes_snapshots = lista_snapshots_historial(antes)
    despues_snapshots = lista_snapshots_historial(despues)

    if item.entidad == 'ratio_eli_ii_mensual':
        snapshots = [snapshot for snapshot in antes_snapshots + despues_snapshots if snapshot]
        fechas = {datetime.strptime(s['fecha'], '%Y-%m-%d').date() for s in snapshots if s.get('fecha')}
        for actual in RatioEliIIMensual.query.filter(RatioEliIIMensual.fecha.in_(fechas)).all():
            db.session.delete(actual)
        db.session.flush()
        for snapshot in antes_snapshots:
            db.session.add(RatioEliIIMensual(
                fecha=datetime.strptime(snapshot['fecha'], '%Y-%m-%d').date(),
                operaciones=parse_numero(snapshot.get('operaciones')), staff=parse_numero(snapshot.get('staff')),
                operaciones_importe=parse_numero(snapshot.get('operaciones_importe')),
                staff_importe=parse_numero(snapshot.get('staff_importe')),
            ))
        registrar_historial('deshacer', item.entidad, item.entidad_id,
                            f'Deshacer movimiento #{item.id}: {item.resumen}', antes=despues, despues=antes)
        db.session.commit()
        return {'success': True, 'mensaje': f'Movimiento #{item.id} deshecho'}, 200

    if item.entidad == 'ratio_eli_mensual':
        snapshots = [snapshot for snapshot in antes_snapshots + despues_snapshots if snapshot]
        fechas = {
            datetime.strptime(snapshot['fecha'], '%Y-%m-%d').date()
            for snapshot in snapshots if snapshot.get('fecha')
        }
        for actual in RatioEliMensual.query.filter(RatioEliMensual.fecha.in_(fechas)).all():
            db.session.delete(actual)
        db.session.flush()
        for snapshot in antes_snapshots:
            if snapshot:
                db.session.add(RatioEliMensual(
                    fecha=datetime.strptime(snapshot['fecha'], '%Y-%m-%d').date(),
                    requerido=parse_numero(snapshot.get('requerido')),
                    staff=parse_numero(snapshot.get('staff')),
                ))
        registrar_historial('deshacer', item.entidad, item.entidad_id,
                            f'Deshacer movimiento #{item.id}: {item.resumen}', antes=despues, despues=antes)
        db.session.commit()
        return {'success': True, 'mensaje': f'Movimiento #{item.id} deshecho'}, 200

    if item.entidad == 'graficos_dotaciones_mensuales':
        snapshots = antes_snapshots + despues_snapshots
        fechas = {datetime.strptime(snapshot['fecha'], '%Y-%m-%d').date() for snapshot in snapshots if snapshot and snapshot.get('fecha')}
        for actual in GraficoDotacionMensual.query.filter(GraficoDotacionMensual.fecha.in_(fechas)).all():
            db.session.delete(actual)
        db.session.flush()
        for snapshot in antes_snapshots:
            db.session.add(GraficoDotacionMensual(
                fecha=datetime.strptime(snapshot['fecha'], '%Y-%m-%d').date(),
                dotacion_requerida=parse_numero(snapshot.get('dotacion_requerida')),
                activa_sl=parse_numero(snapshot.get('activa_sl')),
                activa_ba=parse_numero(snapshot.get('activa_ba')),
            ))
        registrar_historial('deshacer', item.entidad, item.entidad_id,
                            f'Deshacer movimiento #{item.id}: {item.resumen}', antes=despues, despues=antes)
        db.session.commit()
        return {'success': True, 'mensaje': f'Movimiento #{item.id} deshecho'}, 200

    if item.entidad == 'dotaciones_clientes_mensuales':
        snapshots = antes_snapshots + despues_snapshots
        fechas = {
            datetime.strptime(snapshot['fecha'], '%Y-%m-%d').date()
            for snapshot in snapshots if snapshot and snapshot.get('fecha')
        }
        if not fechas:
            return {'success': False, 'errores': ['El movimiento no contiene un mes restaurable']}, 400
        actuales = DotacionClienteMensual.query.filter(DotacionClienteMensual.fecha.in_(fechas)).all()
        for actual in actuales:
            db.session.delete(actual)
        db.session.flush()
        for snapshot in antes_snapshots:
            db.session.add(DotacionClienteMensual(
                cliente=snapshot.get('cliente') or '',
                fecha=datetime.strptime(snapshot['fecha'], '%Y-%m-%d').date(),
                dotacion=parse_numero(snapshot.get('dotacion')),
            ))
        registrar_historial(
            'deshacer', item.entidad, item.entidad_id,
            f'Deshacer movimiento #{item.id}: {item.resumen}',
            antes=despues, despues=antes,
        )
        db.session.commit()
        return {'success': True, 'mensaje': f'Movimiento #{item.id} deshecho; se restauraron {len(antes_snapshots)} filas'}, 200

    if item.entidad == 'dashboard_operativo':
        snapshots = [snapshot for snapshot in antes_snapshots + despues_snapshots if snapshot]
        claves = {(snapshot.get('mes'), snapshot.get('cliente'), snapshot.get('campania')) for snapshot in snapshots}
        for mes, cliente, campania in claves:
            actual = DashboardOperativo.query.filter_by(mes=mes, cliente=cliente, campania=campania).first()
            if actual:
                db.session.delete(actual)
        db.session.flush()
        for snapshot in antes_snapshots:
            if snapshot:
                db.session.add(DashboardOperativo(
                    year=snapshot.get('year') or int(snapshot['mes'][:4]), mes=snapshot['mes'],
                    cliente=snapshot['cliente'], campania=snapshot['campania'],
                    datos=json.dumps(snapshot.get('datos') or {}, ensure_ascii=False),
                    bajas_manual=parse_numero(snapshot.get('bajas_manual') or 0),
                ))
        registrar_historial('deshacer', item.entidad, item.entidad_id,
                            f'Deshacer movimiento #{item.id}: {item.resumen}', antes=despues, despues=antes)
        db.session.commit()
        return {'success': True, 'mensaje': f'Movimiento #{item.id} deshecho'}, 200

    if item.entidad == 'historico_cliente':
        snapshot = antes_snapshots[0] if antes_snapshots else antes
        actual = HistoricoClienteMensual.query.get(item.entidad_id)
        if not actual:
            return {'success': False, 'errores': ['El registro histórico ya no existe']}, 404
        actual.pagadas = snapshot.get('pagadas')
        actual.logueo = snapshot.get('logueo')
        registrar_historial('deshacer', item.entidad, item.entidad_id,
                            f'Deshacer movimiento #{item.id}: {item.resumen}', antes=despues, despues=antes)
        db.session.commit()
        return {'success': True, 'mensaje': f'Movimiento #{item.id} deshecho'}, 200

    cantidad = max(len(antes_snapshots), len(despues_snapshots))

    for indice in range(cantidad):
        antes_snapshot = antes_snapshots[indice] if indice < len(antes_snapshots) else None
        despues_snapshot = despues_snapshots[indice] if indice < len(despues_snapshots) else None
        if item.entidad == 'matriz_precios':
            actual = buscar_precio_snapshot(despues_snapshot) or buscar_precio_snapshot(antes_snapshot)
        elif item.entidad == 'variables':
            actual = buscar_variable_snapshot(despues_snapshot) or buscar_variable_snapshot(antes_snapshot)
        elif item.entidad == 'tarifaciones':
            actual = buscar_tarifacion_snapshot(despues_snapshot) or buscar_tarifacion_snapshot(antes_snapshot)
        elif item.entidad == 'next_gen':
            actual = buscar_next_gen_snapshot(despues_snapshot) or buscar_next_gen_snapshot(antes_snapshot)
        elif item.entidad == 'sites_proyeccion':
            actual = buscar_site_proyeccion_snapshot(despues_snapshot) or buscar_site_proyeccion_snapshot(antes_snapshot)
        else:
            actual = buscar_proyeccion_snapshot(despues_snapshot) or buscar_proyeccion_snapshot(antes_snapshot)
        if antes_snapshot:
            if not actual:
                if item.entidad == 'matriz_precios':
                    actual = ProyeccionPrecio()
                elif item.entidad == 'variables':
                    actual = VariableCampania()
                elif item.entidad == 'tarifaciones':
                    actual = TarifacionCampania()
                elif item.entidad == 'next_gen':
                    snapshot_base = antes_snapshot or despues_snapshot or {}
                    actual = NextGenDolar() if snapshot_base.get('tipo_registro') == 'dolar' else NextGenProducto()
                elif item.entidad == 'sites_proyeccion':
                    actual = SiteProyeccion()
                else:
                    actual = ProyeccionMatriz()
                db.session.add(actual)
            if item.entidad == 'matriz_precios':
                restaurar_precio_snapshot(actual, antes_snapshot)
            elif item.entidad == 'variables':
                restaurar_variable_snapshot(actual, antes_snapshot)
            elif item.entidad == 'tarifaciones':
                restaurar_tarifacion_snapshot(actual, antes_snapshot)
            elif item.entidad == 'next_gen':
                restaurar_next_gen_snapshot(actual, antes_snapshot)
            elif item.entidad == 'sites_proyeccion':
                restaurar_site_proyeccion_snapshot(actual, antes_snapshot)
            else:
                restaurar_proyeccion_snapshot(actual, antes_snapshot)
        elif actual:
            db.session.delete(actual)

    registrar_historial(
        'deshacer',
        item.entidad,
        item.entidad_id,
        f'Deshacer movimiento #{item.id}: {item.resumen}',
        antes=despues,
        despues=antes,
    )
    db.session.flush()
    errores_validacion = []
    for indice in range(cantidad):
        antes_snapshot = antes_snapshots[indice] if indice < len(antes_snapshots) else None
        despues_snapshot = despues_snapshots[indice] if indice < len(despues_snapshots) else None
        if item.entidad == 'matriz_precios':
            existe_antes = existe_precio_snapshot(antes_snapshot)
            existe_despues = existe_precio_snapshot(despues_snapshot)
        elif item.entidad == 'variables':
            existe_antes = existe_variable_snapshot(antes_snapshot)
            existe_despues = existe_variable_snapshot(despues_snapshot)
        elif item.entidad == 'tarifaciones':
            existe_antes = existe_tarifacion_snapshot(antes_snapshot)
            existe_despues = existe_tarifacion_snapshot(despues_snapshot)
        elif item.entidad == 'next_gen':
            existe_antes = existe_next_gen_snapshot(antes_snapshot)
            existe_despues = existe_next_gen_snapshot(despues_snapshot)
        elif item.entidad == 'sites_proyeccion':
            existe_antes = existe_site_proyeccion_snapshot(antes_snapshot)
            existe_despues = existe_site_proyeccion_snapshot(despues_snapshot)
        else:
            existe_antes = existe_proyeccion_snapshot(antes_snapshot)
            existe_despues = existe_proyeccion_snapshot(despues_snapshot)
        if antes_snapshot and not existe_antes:
            errores_validacion.append(f'No se pudo restaurar {antes_snapshot.get("cliente")} / {antes_snapshot.get("campania")} / {antes_snapshot.get("mes")}')
        if not antes_snapshot and despues_snapshot and existe_despues:
            errores_validacion.append(f'No se pudo eliminar {despues_snapshot.get("cliente")} / {despues_snapshot.get("campania")} / {despues_snapshot.get("mes")}')
    if errores_validacion:
        db.session.rollback()
        return {'success': False, 'errores': errores_validacion}, 500
    db.session.commit()
    return {'success': True, 'mensaje': f'Movimiento #{item.id} deshecho'}, 200


@main_bp.route('/api/cargar', methods=['POST'])
@carga_requerida
def api_cargar():
    """Endpoint para cargar datos de facturación"""
    data = request.get_json() or {}
    data['cliente'] = data.get('cliente', '').strip()
    data['gerente'] = data.get('gerente', '').strip()
    data['jefe_site'] = data.get('jefe_site', '').strip()
    data['campania'] = data.get('campania', '').strip()
    data['subcampania'] = data.get('subcampania', '').strip()
    data['mes'] = data.get('mes', '').strip()
    data['es_next_gen'] = bool(data.get('es_next_gen'))
    if data['es_next_gen']:
        nombre_next_gen = data['cliente']
        data.update({
            'gerente': 'NextGen',
            'jefe_site': 'NextGen',
            'campania': nombre_next_gen,
            'subcampania': nombre_next_gen,
            'tipo_negocio': 'NextGen',
            'tipo_jornada': 'NextGen',
        })
    # Validaciones
    errores = []
    tipo_jornada = normalizar_tipo_vh(data.get('tipo_jornada'))
    sin_restriccion_horas = tipo_jornada == TIPO_VH_PERSONAL_COBRANZAS
    
    if not data.get('fecha'):
        errores.append('La fecha es obligatoria')
    if not data.get('mes'):
        errores.append('El mes de facturacion es obligatorio')
    elif not mes_valido(data.get('mes')):
        errores.append('El mes de facturacion no es valido')
    if not data.get('cliente'):
        errores.append('El cliente es obligatorio')
    if not data.get('tipo_jornada') and not data['es_next_gen']:
        errores.append('El tipo de VH es obligatorio')
    # Validar gerente
    if not data.get('gerente') and not data['es_next_gen']:
        errores.append('El gerente es obligatorio')
    if not data.get('jefe_site') and not data['es_next_gen']:
        errores.append('El jefe de site es obligatorio')
    if not data.get('campania') and not data['es_next_gen']:
        errores.append('La campaña es obligatoria')
    if not data.get('subcampania') and not data['es_next_gen']:
        errores.append('La sub campaña es obligatoria')
    if data.get('cliente') and data.get('campania'):
        try:
            aplicar_excepcion_calculo(data)
        except ValueError as error:
            errores.append(str(error))
    
    try:
        horas_objetivo = parse_numero(data.get('horas_objetivo'))
        horas_facturadas = parse_numero(data.get('horas_facturadas'))
        horas_penalizadas = parse_numero(data.get('horas_penalizadas'))
        valor_hora = parse_numero(data.get('valor_hora'))
        valor_hora_objetivo = parse_numero(data.get('valor_hora_objetivo', valor_hora))
        importe_fijo = parse_numero(data.get('importe_fijo')) if data.get('importe_fijo') not in (None, '') else None
        variable_objetivo = parse_numero(data.get('variable_objetivo'))
        variable_productivo = parse_numero(data.get('variable_productivo'))
    except ValueError:
        errores.append('Hay valores numéricos con formato inválido')
        horas_objetivo = horas_facturadas = horas_penalizadas = valor_hora = valor_hora_objetivo = 0
        importe_fijo = None
        variable_objetivo = 0
        variable_productivo = 0
    
    if horas_objetivo < 0:
        errores.append('Las horas objetivo no pueden ser negativas')
    if horas_facturadas < 0:
        errores.append('Las horas facturadas no pueden ser negativas')
    if horas_penalizadas < 0:
        errores.append('Las horas penalizadas no pueden ser negativas')
    if not sin_restriccion_horas and horas_penalizadas > horas_facturadas:
        errores.append('Las horas penalizadas no pueden superar las horas facturadas')
    if not data['es_next_gen'] and not sin_restriccion_horas and valor_hora <= 0:
        errores.append('El valor hora debe ser mayor a 0')
    if not data['es_next_gen'] and not sin_restriccion_horas and valor_hora_objetivo <= 0:
        errores.append('El valor hora objetivo debe ser mayor a 0')
    if importe_fijo is not None and importe_fijo < 0:
        errores.append('El importe fijo facturado no puede ser negativo')
    if variable_objetivo < 0:
        errores.append('Variable Objetivo no puede ser negativo')
    if tipo_jornada and tipo_jornada not in TIPOS_VH and not data['es_next_gen']:
        errores.append(f'El tipo de VH "{data.get("tipo_jornada", "")}" no es válido')
    
    errores.extend(validar_justificaciones(data))

    if errores:
        return jsonify({'success': False, 'errores': errores}), 400
    
    try:
        nuevo_registro = crear_registro_facturacion(data)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'mensaje': f'Datos cargados para {nuevo_registro.mes} con fecha de carga {nuevo_registro.fecha.strftime("%d/%m/%Y")}',
            'id': nuevo_registro.id,
            'mes': nuevo_registro.mes,
            'fecha': nuevo_registro.fecha.isoformat(),
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'errores': [str(e)]}), 500


@main_bp.route('/api/datos', methods=['GET'])
@login_requerido
def api_datos():
    """Endpoint para obtener datos filtrados por mes, cliente y gerente"""
    query = aplicar_filtros(Facturacion2026.query, **filtros_request())
    registros = query.order_by(Facturacion2026.fecha.desc()).all()

    return jsonify({
        'success': True,
        'data': [r.to_dict() for r in registros]
    })


@main_bp.route('/api/justificaciones', methods=['GET'])
@login_requerido
def api_justificaciones():
    """Listado de justificaciones de bonos, penalizaciones y otros."""
    registros = aplicar_filtros(
        Facturacion2026.query,
        **filtros_request()
    ).order_by(Facturacion2026.fecha.desc()).all()

    items = []
    for registro in registros:
        for justificacion in registro.justificaciones:
            item = justificacion.to_dict()
            item.update({
                'fecha': registro.fecha.isoformat() if registro.fecha else None,
                'mes': registro.mes,
                'cliente': registro.cliente,
                'gerente': registro.gerente,
                'jefe_site': registro.jefe_site,
                'campania': registro.campania,
                'subcampania': registro.subcampania,
                'tipo_jornada': registro.tipo_jornada,
            })
            items.append(item)

    return jsonify({'success': True, 'data': items})


@main_bp.route('/api/datos/<int:registro_id>', methods=['PUT'])
@edicion_requerida
def api_actualizar_dato(registro_id):
    data = request.get_json() or {}
    if not validar_confirmacion_accion(data):
        return jsonify({'success': False, 'errores': ['La confirmacion no es valida']}), 403

    registro = Facturacion2026.query.get_or_404(registro_id)
    antes = snapshot_modelo(registro)
    data['es_next_gen'] = bool(data.get('es_next_gen') or registro.es_next_gen)
    if data['es_next_gen']:
        nombre_next_gen = str(data.get('cliente') or registro.cliente).strip()
        data.update(gerente='NextGen', jefe_site='NextGen', campania=nombre_next_gen,
                    subcampania=nombre_next_gen, tipo_negocio='NextGen', tipo_jornada='NextGen')
    errores = validar_payload_facturacion(data, exigir_configuracion_valor_hora=False)
    if errores:
        return jsonify({'success': False, 'errores': errores}), 400

    try:
        fecha = datetime.strptime(data['fecha'], '%Y-%m-%d').date()
        registro.fecha = fecha
        registro.mes = mes_valido(data.get('mes'))
        registro.cliente = data['cliente'].strip()
        registro.gerente = data['gerente'].strip()
        registro.jefe_site = data['jefe_site'].strip()
        registro.campania = data['campania'].strip()
        registro.subcampania = data['subcampania'].strip()
        registro.tipo_negocio = str(data.get('tipo_negocio') or '').strip() or None
        registro.es_next_gen = bool(data.get('es_next_gen'))
        registro.tipo_jornada = normalizar_tipo_vh(data['tipo_jornada'])
        registro.horas_objetivo = parse_numero(data.get('horas_objetivo'))
        registro.horas_facturadas = parse_numero(data.get('horas_facturadas'))
        registro.horas_penalizadas = parse_numero(data.get('horas_penalizadas'))
        registro.valor_hora_objetivo = parse_numero(data.get('valor_hora_objetivo') or data.get('valor_hora'))
        registro.valor_hora = parse_numero(data.get('valor_hora'))
        registro.tarifacion = parse_numero(data.get('tarifacion')) if data.get('tarifacion') not in (None, '') else None
        registro.importe_fijo = parse_numero(data.get('importe_fijo')) if data.get('importe_fijo') not in (None, '') else None
        registro.variable_objetivo = parse_numero(data.get('variable_objetivo'))
        registro.variable_productivo = parse_numero(data.get('variable_productivo'))
        registro.bonos = float(data.get('bonos', 0) or 0)
        registro.penalizaciones = normalizar_importe_ajuste('penalizaciones', data.get('penalizaciones'))
        registro.netx_gen = float(data.get('netx_gen', 0) or 0)
        registro.otros = float(data.get('otros', 0) or 0)
        # Control edita los componentes del total. Los totales manuales traídos
        # por una importación anterior no deben congelar ni ocultar el cambio.
        registro.total_facturado_manual = None
        registro.control_total_facturado = None
        asegurar_asignacion_desde_registro(registro)
        despues = snapshot_modelo(registro)
        cambios = cambios_entre(antes, despues)
        if cambios:
            registrar_historial(
                'edicion',
                'facturacion',
                registro.id,
                f'Facturacion editada: {registro.cliente} / {registro.mes}',
                antes=antes,
                despues=despues,
                detalle=serializar_json(cambios),
            )
        db.session.commit()
        return jsonify({'success': True, 'mensaje': 'Registro actualizado', 'data': registro.to_dict()})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'errores': [str(e)]}), 500


@main_bp.route('/api/datos/<int:registro_id>', methods=['DELETE'])
@eliminacion_requerida
def api_eliminar_dato(registro_id):
    data = request.get_json() or {}
    if not validar_confirmacion_accion(data):
        return jsonify({'success': False, 'errores': ['La confirmacion no es valida']}), 403

    registro = Facturacion2026.query.get_or_404(registro_id)
    antes = snapshot_modelo(registro)
    registrar_historial(
        'eliminacion',
        'facturacion',
        registro.id,
        f'Facturacion eliminada: {registro.cliente} / {registro.mes}',
        antes=antes,
        detalle=f"Se elimino el registro de {registro.cliente} correspondiente a {registro.mes}.",
    )
    db.session.delete(registro)
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Registro eliminado'})


@main_bp.route('/api/resumen', methods=['GET'])
@login_requerido
def api_resumen():
    """Endpoint para resumen mensual"""
    filtros = filtros_request()
    
    query = db.session.query(
        Facturacion2026.mes,
        func.sum(Facturacion2026.horas_objetivo).label('horas_objetivo'),
        func.sum(Facturacion2026.horas_facturadas).label('horas_facturadas'),
        func.sum(Facturacion2026.bonos).label('bonos'),
        func.sum(Facturacion2026.netx_gen).label('netx_gen'),
        func.sum(Facturacion2026.otros).label('otros')
    )
    
    query = aplicar_filtros(query, **filtros)
    
    resultados = query.group_by(Facturacion2026.mes).order_by(Facturacion2026.mes).all()
    
    resumen = []
    for r in resultados:
        # Calcular totales usando la lógica del modelo
        registros_mes = aplicar_filtros(
            Facturacion2026.query,
            mes=r.mes,
            cliente=filtros['cliente'],
            gerente=filtros['gerente'],
            jefe_site=filtros['jefe_site'],
            campania=filtros['campania'],
            subcampania=filtros['subcampania'],
            tipo_negocio=filtros['tipo_negocio']
        ).all()
        
        resumen_mes = resumen_dashboard(registros_mes)
        
        resumen.append({
            'mes': r.mes,
            **resumen_mes
        })
    
    return jsonify({
        'success': True,
        'resumen': resumen
    })


@main_bp.route('/api/resumen-horas', methods=['GET'])
@login_requerido
def api_resumen_horas():
    """Resumen integral equivalente a la pestaña RESUMEN del libro operativo."""
    try:
        year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['El año no es válido']}), 400
    if not 2020 <= year <= 2100:
        return jsonify({'success': False, 'errores': ['El año debe estar entre 2020 y 2100']}), 400

    facturacion = facturacion_por_anio(Facturacion2026.query, year).all()
    facturacion_operativa = [registro for registro in facturacion if not registro.es_next_gen]
    facturacion_next_gen = [registro for registro in facturacion if (registro.netx_gen or 0) != 0]
    dotacion_por_mes = dict(
        db.session.query(
            func.strftime('%Y-%m', DotacionClienteMensual.fecha),
            func.sum(DotacionClienteMensual.dotacion),
        ).filter(
            func.extract('year', DotacionClienteMensual.fecha) == year,
        ).group_by(func.strftime('%Y-%m', DotacionClienteMensual.fecha)).all()
    )
    def metricas(registros_grupo, dotacion=0):
        detalle = metricas_matriz(registros_grupo)
        tablero = resumen_dashboard(registros_grupo)
        total_facturado = tablero['total_facturado']
        total_objetivo = tablero['total_teorico']
        horas_objetivo = detalle['horas_objetivo']
        horas_facturadas = detalle['horas_facturadas']
        valor_hora_promedio = total_facturado / horas_facturadas if horas_facturadas else 0
        vh_objetivo_resumen = total_objetivo / horas_objetivo if horas_objetivo else 0
        alcance_facturacion = (
            (detalle['horas_real'] + detalle['variable_real'] + detalle['penalizaciones_bonos']) / total_objetivo * 100
            if total_objetivo else 0
        )
        total_desvio = tablero['desvio']
        return {
            **detalle,
            'dotacion_requerida': round(float(dotacion or 0), 2),
            'total_facturado': total_facturado,
            'total_real': tablero['total_real'],
            'total_obj': total_objetivo,
            'total_teorico': total_objetivo,
            'desvio': tablero['desvio'],
            'porcentaje_cumplimiento': round(alcance_facturacion, 2),
            'valor_hora_promedio': round(valor_hora_promedio, 2),
            'vh_objetivo_resumen': round(vh_objetivo_resumen, 2),
            'porcentaje_alcance_vh': round((valor_hora_promedio / vh_objetivo_resumen * 100) if vh_objetivo_resumen else 0, 2),
            'total_desvio': round(total_desvio, 2),
            'porcentaje_desvio': round((total_desvio / total_objetivo * 100) if total_objetivo else 0, 2),
            'participacion_desvio_horas': round((detalle['desvio_horas_monto'] / total_desvio * 100) if total_desvio else 0, 2),
            'participacion_desvio_variable': round((detalle['desvio_variable'] / total_desvio * 100) if total_desvio else 0, 2),
            'participacion_desvio_penalizaciones': round((detalle['penalizaciones_bonos'] / total_desvio * 100) if total_desvio else 0, 2),
            'participacion_desvio_total': 100 if total_desvio else 0,
        }

    filas = []
    for mes in range(1, 13):
        clave = f'{year}-{mes:02d}'
        registros_mes = [registro for registro in facturacion if registro.mes == clave]
        filas.append({'nombre': clave, **metricas(registros_mes, dotacion_por_mes.get(clave))})
    dotaciones_con_dato = [fila['dotacion_requerida'] for fila in filas if fila['dotacion_requerida'] > 0]
    promedio_dotacion = sum(dotaciones_con_dato) / len(dotaciones_con_dato) if dotaciones_con_dato else 0

    def agrupar(clave):
        grupos = {}
        for registro in facturacion:
            nombre = clave(registro) or 'Sin asignar'
            grupos.setdefault(nombre, []).append(registro)
        return [{'nombre': nombre, **metricas(registros)} for nombre, registros in sorted(grupos.items(), key=lambda item: item[0].casefold())]

    trimestres = []
    for numero, meses in enumerate(((1, 2, 3), (4, 5, 6), (7, 8, 9), (10, 11, 12)), 1):
        claves = {f'{year}-{mes:02d}' for mes in meses}
        registros = [registro for registro in facturacion if registro.mes in claves]
        dotaciones = [float(dotacion_por_mes.get(clave) or 0) for clave in claves if float(dotacion_por_mes.get(clave) or 0) > 0]
        trimestres.append({'nombre': f'{numero}T', **metricas(registros, sum(dotaciones) / len(dotaciones) if dotaciones else 0)})

    def familia(registro):
        texto = f'{registro.cliente or ""} {registro.campania or ""}'.casefold()
        if 'personal' in texto:
            return 'Personal'
        if 'santander' in texto or 'getnet' in texto:
            return 'Santander'
        return 'Multicampaña'

    total = metricas(facturacion, promedio_dotacion)

    def participaciones(items):
        base = total['total_facturado']
        for item in items:
            item['participacion_operativa'] = round(item['total_facturado'] / base * 100, 2) if base else 0
        return items

    def agrupar_registros(registros_fuente, clave):
        grupos_salida = {}
        for registro in registros_fuente:
            nombre = clave(registro) or 'Sin asignar'
            grupos_salida.setdefault(nombre, []).append(registro)
        return [
            {'nombre': nombre, **metricas(registros)}
            for nombre, registros in sorted(grupos_salida.items(), key=lambda item: item[0].casefold())
        ]

    grupos = participaciones(agrupar_registros(facturacion_operativa, familia))
    gerencias = participaciones(agrupar_registros(facturacion_operativa, lambda registro: registro.gerente))
    detalle_jefaturas = []
    por_jefatura = {}
    for registro in facturacion_operativa:
        jefatura = registro.jefe_site or 'Sin asignar'
        clave_jefatura = normalizar_header(jefatura)
        grupo = por_jefatura.setdefault(clave_jefatura, {'nombre': jefatura, 'registros': []})
        grupo['registros'].append(registro)
    for grupo in sorted(por_jefatura.values(), key=lambda item: item['nombre'].casefold()):
        jefatura, registros = grupo['nombre'], grupo['registros']
        meses_jefatura = []
        for mes in range(1, 13):
            clave = f'{year}-{mes:02d}'
            meses_jefatura.append({'nombre': clave, **metricas([r for r in registros if r.mes == clave])})
        trimestres_jefatura = []
        for numero, meses in enumerate(((1, 2, 3), (4, 5, 6), (7, 8, 9), (10, 11, 12)), 1):
            claves = {f'{year}-{mes:02d}' for mes in meses}
            trimestres_jefatura.append({'nombre': f'{numero}T', **metricas([r for r in registros if r.mes in claves])})
        detalle_jefaturas.append({
            'nombre': jefatura,
            'total': metricas(registros),
            'meses': meses_jefatura,
            'trimestres': trimestres_jefatura,
        })
    jefaturas = participaciones([
        {'nombre': item['nombre'], **item['total']}
        for item in detalle_jefaturas
    ])
    next_gen = []
    for mes in range(1, 13):
        clave = f'{year}-{mes:02d}'
        registros_mes = [registro for registro in facturacion_next_gen if registro.mes == clave]
        next_gen.append({'nombre': clave, **metricas(registros_mes)})

    return jsonify({
        'success': True,
        'year': year,
        'filas': filas,
        'trimestres': trimestres,
        'grupos': grupos,
        'gerencias': gerencias,
        'jefaturas': jefaturas,
        'detalle_jefaturas': detalle_jefaturas,
        'next_gen': next_gen,
        'total': total,
    })


@main_bp.route('/api/kpis', methods=['GET'])
@login_requerido
def api_kpis():
    """Endpoint para obtener KPIs generales"""
    registros = registros_filtrados()
    
    if not registros:
        return jsonify({
            'success': True,
            'kpis': {
                'total_facturado': 0,
                'total_real': 0,
                'total_teorico': 0,
                'desvio': 0,
                'porcentaje_cumplimiento': 0,
                'porcentaje_cumplimiento_horas': 0,
                'porcentaje_cumplimiento_horas_adh': 0,
                'porcentaje_valor_hora': 0,
                'porcentaje_bono': 0,
                'porcentaje_cumplimiento_facturacion': 0,
                'horas_objetivo': 0,
                'horas_facturadas': 0,
                'bonos': 0,
                'variable_productivo': 0,
                'penalizaciones': 0,
                'tarifacion': 0
            }
        })

    kpis = resumen_dashboard(registros)
    
    return jsonify({
        'success': True,
        'kpis': kpis
    })


@main_bp.route('/api/grafico', methods=['GET'])
@login_requerido
def api_grafico():
    """Endpoint para datos del gráfico de evolución mensual"""
    filtros = filtros_request()

    query = db.session.query(
        Facturacion2026.mes,
        func.sum(Facturacion2026.horas_objetivo).label('horas_objetivo'),
        func.sum(Facturacion2026.horas_facturadas).label('horas_facturadas')
    )
    query = aplicar_filtros(query, **filtros)
    resultados = query.group_by(Facturacion2026.mes).order_by(Facturacion2026.mes).all()
    
    # Calcular totales reales y teóricos por mes
    datos = []
    for r in resultados:
        registros_mes = aplicar_filtros(
            Facturacion2026.query,
            mes=r.mes,
            cliente=filtros['cliente'],
            gerente=filtros['gerente'],
            jefe_site=filtros['jefe_site'],
            campania=filtros['campania'],
            subcampania=filtros['subcampania'],
            tipo_negocio=filtros['tipo_negocio']
        ).all()
        total_real = sum(reg.total_dashboard for reg in registros_mes)
        total_teorico = sum(reg.total_teorico for reg in registros_mes)
        
        datos.append({
            'mes': r.mes,
            'total_real': round(total_real, 2),
            'total_teorico': round(total_teorico, 2)
        })
    
    return jsonify({
        'success': True,
        'datos': datos
    })


@main_bp.route('/api/filtros', methods=['GET'])
@login_requerido
def api_filtros():
    """Opciones dinamicas disponibles para los filtros del dashboard."""
    filtros = filtros_request()
    return jsonify({
        'success': True,
        'filtros': {
            'years': opciones_filtro(filtros, 'year'),
            'meses': opciones_filtro(filtros, 'mes'),
            'clientes': opciones_filtro(filtros, 'cliente'),
            'gerentes': opciones_filtro(filtros, 'gerente'),
            'jefes_site': opciones_filtro(filtros, 'jefe_site'),
            'campanias': opciones_filtro(filtros, 'campania'),
            'subcampanias': opciones_filtro(filtros, 'subcampania'),
            'tipos_negocio': opciones_filtro(filtros, 'tipo_negocio'),
        }
    })


@main_bp.route('/api/por-cliente', methods=['GET'])
@login_requerido
def api_por_cliente():
    """Agrupacion ejecutiva por cliente."""
    registros = registros_filtrados()
    grupos = {}
    for registro in registros:
        grupos.setdefault(registro.cliente, []).append(registro)

    datos = []
    for cliente, registros_cliente in grupos.items():
        resumen = resumen_dashboard(registros_cliente)
        datos.append({
            'cliente': cliente,
            'gerente': registros_cliente[0].gerente,
            'jefe_site': registros_cliente[0].jefe_site,
            'registros': len(registros_cliente),
            **resumen
        })

    # Si la selección existe en Datos Maestros pero aún no fue facturada, se
    # muestra con cero para que el filtro no parezca vacío.
    clientes_presentes = {item['cliente'] for item in datos}
    for asignacion in asignaciones_maestras_filtradas(filtros_request()):
        if asignacion.cliente in clientes_presentes:
            continue
        datos.append({
            'cliente': asignacion.cliente,
            'gerente': asignacion.gerente,
            'jefe_site': asignacion.jefe_site,
            'registros': 0,
            'horas_objetivo': 0, 'horas_facturadas': 0,
            'total_facturado': 0, 'total_real': 0, 'total_teorico': 0,
            'desvio': 0, 'porcentaje_cumplimiento': 0,
        })
        clientes_presentes.add(asignacion.cliente)

    datos.sort(key=lambda item: item['total_real'], reverse=True)
    return jsonify({'success': True, 'clientes': datos})


@main_bp.route('/api/comparativo', methods=['GET'])
@login_requerido
def api_comparativo():
    """Comparativo de horas objetivo contra horas facturadas."""
    filtros = filtros_comparativo_request()
    registros = aplicar_filtros(Facturacion2026.query, **filtros).order_by(
        Facturacion2026.mes,
        Facturacion2026.cliente
    ).all()

    resumen = resumen_registros(registros)
    resumen['diferencia_horas'] = round(resumen['horas_facturadas'] - resumen['horas_objetivo'], 2)
    resumen['cumplimiento_horas'] = round(
        (resumen['horas_facturadas'] / resumen['horas_objetivo'] * 100)
        if resumen['horas_objetivo'] > 0 else 0,
        2
    )

    grupos_mes = {}
    grupos_cliente = {}
    for registro in registros:
        grupos_mes.setdefault(registro.mes, []).append(registro)
        clave_cliente = (
            registro.cliente,
            registro.gerente,
            registro.jefe_site,
            registro.campania,
            registro.subcampania,
        )
        grupos_cliente.setdefault(clave_cliente, []).append(registro)

    por_mes = []
    for mes, registros_mes in grupos_mes.items():
        horas_objetivo = sum(r.horas_objetivo for r in registros_mes)
        horas_facturadas = sum(r.horas_facturadas for r in registros_mes)
        por_mes.append({
            'mes': mes,
            'horas_objetivo': round(horas_objetivo, 2),
            'horas_facturadas': round(horas_facturadas, 2),
            'diferencia_horas': round(horas_facturadas - horas_objetivo, 2),
            'cumplimiento_horas': round((horas_facturadas / horas_objetivo * 100) if horas_objetivo > 0 else 0, 2),
        })

    por_cliente = []
    for (cliente, gerente, jefe_site, campania, subcampania), registros_cliente in grupos_cliente.items():
        horas_objetivo = sum(r.horas_objetivo for r in registros_cliente)
        horas_facturadas = sum(r.horas_facturadas for r in registros_cliente)
        por_cliente.append({
            'cliente': cliente,
            'gerente': gerente,
            'jefe_site': jefe_site,
            'campania': campania,
            'subcampania': subcampania,
            'registros': len(registros_cliente),
            'horas_objetivo': round(horas_objetivo, 2),
            'horas_facturadas': round(horas_facturadas, 2),
            'diferencia_horas': round(horas_facturadas - horas_objetivo, 2),
            'cumplimiento_horas': round((horas_facturadas / horas_objetivo * 100) if horas_objetivo > 0 else 0, 2),
        })

    por_mes.sort(key=lambda item: item['mes'])
    por_cliente.sort(key=lambda item: abs(item['diferencia_horas']), reverse=True)

    return jsonify({
        'success': True,
        'kpis': resumen,
        'por_mes': por_mes,
        'por_cliente': por_cliente,
        'detalle': [r.to_dict() for r in registros],
    })


@main_bp.route('/api/matriz', methods=['GET'])
@login_requerido
def api_matriz():
    """Matriz mensual de cumplimiento por gerencia y apertura por jefe de site."""
    year = request.args.get('year', str(current_app.config['DEFAULT_YEAR'])).strip()
    try:
        year_number = int(year)
    except ValueError:
        year_number = 0
    if not 2020 <= year_number <= 2100:
        return jsonify({'success': False, 'errores': ['El anio debe estar entre 2020 y 2100.']}), 400
    year = str(year_number)
    filtros = filtros_request()
    meses = [f'{year}-{numero}' for numero, _ in MESES_MATRIZ]
    columnas = [{'key': mes, 'label': f'{label}-{year[-2:]}'} for mes, (_, label) in zip(meses, MESES_MATRIZ)]

    base_year = Facturacion2026.query.filter(
        Facturacion2026.mes.in_(meses),
        Facturacion2026.es_next_gen.is_(False),
    )
    registros_year = aplicar_filtros(base_year, **filtros).all()
    registros_sin_filtros = aplicar_filtros(Facturacion2026.query).filter(
        Facturacion2026.mes.in_(meses),
        Facturacion2026.es_next_gen.is_(False),
    ).all()
    jefe_site = filtros.get('jefe_site')
    jefe_site_unico = jefe_site if isinstance(jefe_site, str) else ''

    return jsonify({
        'success': True,
        'year': year,
        'columnas': [{'key': 'total', 'label': year}, *columnas],
        'filtros': {
            'clientes': opciones_filtro({**filtros, 'mes': meses}, 'cliente'),
            'gerentes': opciones_filtro({**filtros, 'mes': meses}, 'gerente'),
            'jefes_site': opciones_filtro({**filtros, 'mes': meses}, 'jefe_site'),
            'campanias': opciones_filtro({**filtros, 'mes': meses}, 'campania'),
            'subcampanias': opciones_filtro({**filtros, 'mes': meses}, 'subcampania'),
            'tipos_negocio': opciones_filtro({**filtros, 'mes': meses}, 'tipo_negocio'),
        },
        'jefes_site': sorted({r.jefe_site for r in registros_sin_filtros if r.jefe_site}),
        'tipos_negocio': sorted({r.tipo_negocio for r in registros_sin_filtros if r.tipo_negocio}),
        'seleccion': filtros,
        'tipo_negocio': filtros.get('tipo_negocio') or '',
        'total_gerencia': matriz_grupos(registros_year, 'gerente', meses),
        'apertura_jefe_site': matriz_grupos(registros_year, 'jefe_site' if not jefe_site_unico else 'campania', meses),
        'jefe_site': jefe_site_unico,
    })


def payload_proyeccion(data):
    cliente = str(data.get('cliente') or '').strip()
    campania = str(data.get('campania') or '').strip()
    mes = str(data.get('mes') or '').strip()
    errores = []
    tipo_plp = str(data.get('tipo_plp') or '').strip()
    horas_carga_manual = bool(tipo_plp)
    if tipo_plp:
        cliente = 'Personal'
        campania = str(data.get('campania_plp') or campania).strip()
    horas_requeridas_manual = 0
    carga_semanal_plp = normalizar_carga_semanal(data.get('carga_semanal_plp') or 'L a V')
    try:
        carga_horaria_plp = parse_numero(data.get('carga_horaria_plp') if data.get('carga_horaria_plp') not in (None, '') else 6)
    except ValueError:
        carga_horaria_plp = 0
        errores.append('La jornada laboral PLP no es válida')
    if tipo_plp:
        if tipo_plp not in SERVICIOS_PERSONAL:
            errores.append('El tipo PLP no es válido')
        try:
            horas_requeridas_manual = parse_numero(data.get('horas_requeridas_manual'))
        except ValueError:
            errores.append('Las horas PLP no son válidas')
        if horas_requeridas_manual < 0:
            errores.append('Las horas PLP no pueden ser negativas')
        if carga_horaria_plp <= 0:
            errores.append('La jornada laboral PLP debe ser mayor a cero')

    try:
        year = int(data.get('year') or (mes.split('-')[0] if '-' in mes else datetime.utcnow().year))
    except (TypeError, ValueError):
        year = datetime.utcnow().year
        errores.append('El año no es valido')

    if not cliente:
        errores.append('El cliente es obligatorio')
    if not campania:
        errores.append('La campaña es obligatoria')
    if not mes_valido(mes):
        errores.append('El mes no es valido')
    else:
        mes = mes_valido(mes)
        year = int(mes[:4])

    try:
        porcentaje_cumplimiento = parse_numero(data.get('porcentaje_cumplimiento') if data.get('porcentaje_cumplimiento') not in (None, '') else 100)
    except ValueError:
        porcentaje_cumplimiento = 0
        errores.append('Hay valores numericos con formato invalido')

    dias_objetivo_manual = None
    horas_requeridas_override = None
    if not tipo_plp:
        try:
            if data.get('dias_objetivo_manual') not in (None, ''):
                dias_objetivo_manual = int(parse_numero(data.get('dias_objetivo_manual')))
            if data.get('horas_requeridas_override') not in (None, ''):
                horas_requeridas_override = parse_numero(data.get('horas_requeridas_override'))
        except (TypeError, ValueError):
            errores.append('Los días u horas manuales no tienen un formato válido')
        if dias_objetivo_manual is not None and dias_objetivo_manual <= 0:
            errores.append('Los días hábiles manuales deben ser mayores a cero')
        if horas_requeridas_override is not None and horas_requeridas_override < 0:
            errores.append('Las horas requeridas manuales no pueden ser negativas')

    es_personal = 'personal' in normalizar_header(f'{cliente} {campania}')
    tiene_nocturnidad = bool(data.get('tiene_nocturnidad')) and es_personal
    try:
        porcentaje_nocturnidad = parse_numero(data.get('porcentaje_nocturnidad') or 0)
    except ValueError:
        porcentaje_nocturnidad = 0
        errores.append('El porcentaje de nocturnidad no es valido')
    if tiene_nocturnidad and not (0 < porcentaje_nocturnidad < 100):
        errores.append('El porcentaje de nocturnidad debe ser mayor a 0 y menor a 100')
    if not tiene_nocturnidad:
        porcentaje_nocturnidad = 0

    jornadas = []
    raw_jornadas = data.get('jornadas') if isinstance(data.get('jornadas'), list) else []
    if not raw_jornadas:
        raw_jornadas = [{
            'dotacion_requerida': data.get('dotacion_requerida'),
            'carga_semanal': data.get('carga_semanal'),
            'carga_horaria': data.get('carga_horaria'),
        }]

    for index, item in enumerate(raw_jornadas, start=1):
        try:
            dotacion = parse_numero(item.get('dotacion_requerida'))
            carga_horaria = parse_numero(item.get('carga_horaria'))
        except (AttributeError, ValueError):
            errores.append(f'La jornada {index} tiene valores numericos invalidos')
            continue
        carga_semanal = normalizar_carga_semanal(item.get('carga_semanal'))
        if dotacion < 0:
            errores.append(f'La dotacion de la jornada {index} no puede ser negativa')
        if carga_horaria < 0:
            errores.append(f'La carga horaria de la jornada {index} no puede ser negativa')
        if dotacion > 0 or carga_horaria > 0:
            jornadas.append({
                'dotacion_requerida': dotacion,
                'carga_semanal': carga_semanal,
                'carga_horaria': carga_horaria,
            })

    if not jornadas and not horas_carga_manual:
        errores.append('Debe cargar al menos una jornada con dotacion y carga horaria')
    if porcentaje_cumplimiento < 0:
        errores.append('El porcentaje de cumplimiento no puede ser negativo')

    return {
        'errores': errores,
        'valores': {
            'cliente': cliente,
            'campania': campania,
            'year': year,
            'mes': mes,
            'jornadas': jornadas,
            'porcentaje_cumplimiento': porcentaje_cumplimiento,
            'tiene_nocturnidad': tiene_nocturnidad,
            'porcentaje_nocturnidad': porcentaje_nocturnidad,
            'tipo_plp': tipo_plp,
            'horas_carga_manual': horas_carga_manual,
            'horas_requeridas_manual': horas_requeridas_manual,
            'carga_semanal_plp': carga_semanal_plp,
            'carga_horaria_plp': carga_horaria_plp,
            'dias_objetivo_manual': dias_objetivo_manual,
            'horas_requeridas_override': horas_requeridas_override,
        }
    }


def payload_precio(data):
    cliente = str(data.get('cliente') or '').strip()
    campania = str(data.get('campania') or '').strip()
    mes = str(data.get('mes') or '').strip()
    errores = []

    try:
        year = int(data.get('year') or (mes.split('-')[0] if '-' in mes else datetime.utcnow().year))
    except (TypeError, ValueError):
        year = datetime.utcnow().year
        errores.append('El año no es valido')

    if not cliente:
        errores.append('El cliente es obligatorio')
    if not campania:
        errores.append('La campaña es obligatoria')
    if not mes_valido(mes):
        errores.append('El mes no es valido')
    else:
        mes = mes_valido(mes)
        year = int(mes[:4])

    try:
        precio_base = parse_numero(data.get('precio_base'))
        alcance_porcentaje = parse_numero(data.get('alcance_porcentaje') if data.get('alcance_porcentaje') not in (None, '') else 100)
        importe_fijo_mensual = parse_numero(data.get('importe_fijo_mensual') if data.get('importe_fijo_mensual') not in (None, '') else 0)
    except ValueError:
        precio_base = 0
        alcance_porcentaje = 0
        importe_fijo_mensual = 0
        errores.append('Hay valores numericos con formato invalido')

    if precio_base < 0:
        errores.append('El precio no puede ser negativo')
    if alcance_porcentaje < 0:
        errores.append('El alcance no puede ser negativo')
    if importe_fijo_mensual < 0:
        errores.append('El pago mensual fijo no puede ser negativo')

    return {
        'errores': errores,
        'valores': {
            'site': str(data.get('site') or '').strip(),
            'cliente': cliente,
            'campania': campania,
            'year': year,
            'mes': mes,
            'precio_base': precio_base,
            'alcance_porcentaje': alcance_porcentaje,
            'importe_fijo_mensual': importe_fijo_mensual,
        }
    }


def headers_template_precios(year):
    meses = [
        f'{label}-{str(year)[-2:]}'
        for _, label in MESES_PROYECCION
    ]
    fijos = [
        f'Fijo {label}-{str(year)[-2:]}'
        for _, label in MESES_PROYECCION
    ]
    return ['Site', 'Cliente'] + meses + fijos


def header_mes_precio(header):
    texto = normalizar_header(header)
    if texto.startswith('fijo '):
        return None
    partes = texto.split('-')
    if len(partes) != 2:
        return None
    labels = {label: numero for numero, label in MESES_PROYECCION}
    mes_numero = labels.get(partes[0])
    if not mes_numero:
        return None
    try:
        year = int(partes[1])
    except ValueError:
        return None
    year += 2000 if year < 100 else 0
    return f'{year}-{mes_numero}'


def header_mes_fijo_precio(header):
    texto = normalizar_header(header)
    if not texto.startswith('fijo '):
        return None
    return header_mes_precio(texto.replace('fijo ', '', 1))


def resolver_cliente_campania_precio(site, nombre):
    site_normalizado = normalizar_header(site)
    nombre_normalizado = normalizar_header(nombre)
    asignacion = AsignacionComercial.query.filter_by(activa=True).filter(
        func.lower(AsignacionComercial.gerente) == str(site or '').strip().lower(),
    ).filter(
        func.lower(AsignacionComercial.campania) == str(nombre or '').strip().lower(),
    ).first()
    if not asignacion:
        asignaciones = AsignacionComercial.query.filter_by(activa=True).all()
        for candidata in asignaciones:
            if site_normalizado and normalizar_header(candidata.gerente) != site_normalizado:
                continue
            if normalizar_header(candidata.campania) == nombre_normalizado or normalizar_header(candidata.cliente) == nombre_normalizado:
                asignacion = candidata
                break
    if asignacion:
        return asignacion.cliente, asignacion.campania
    return nombre, nombre


def sites_por_clave_precio(cliente, campania):
    asignacion = AsignacionComercial.query.filter_by(
        activa=True,
        cliente=cliente,
        campania=campania,
    ).first()
    if asignacion:
        return asignacion.gerente or asignacion.jefe_site or ''
    return ''


TIPOS_FERIADO = ('Inamovible', 'Trasladable', 'Puente', 'Manual')


def payload_feriado(data):
    errores = []
    nombre = str(data.get('nombre') or '').strip()
    tipo = str(data.get('tipo') or 'Manual').strip()
    activo = bool(data.get('activo', True))
    fecha_raw = str(data.get('fecha') or '').strip()

    if not nombre:
        errores.append('El nombre del feriado es obligatorio')
    if tipo not in TIPOS_FERIADO:
        errores.append('El tipo de feriado no es valido')
    try:
        fecha = datetime.strptime(fecha_raw, '%Y-%m-%d').date()
    except ValueError:
        fecha = None
        errores.append('La fecha del feriado no es valida')

    return {
        'errores': errores,
        'valores': {
            'nombre': nombre,
            'tipo': tipo,
            'activo': activo,
            'fecha': fecha,
            'year': fecha.year if fecha else None,
        }
    }


@main_bp.route('/api/calendario-operativo', methods=['GET'])
@login_requerido
def api_calendario_operativo():
    year = int(request.args.get('year') or datetime.utcnow().year)
    feriados = FeriadoOperativo.query.filter_by(year=year).order_by(FeriadoOperativo.fecha).all()
    return jsonify({
        'success': True,
        'year': year,
        'tipos': list(TIPOS_FERIADO),
        'feriados': [feriado.to_dict() for feriado in feriados],
    })


@main_bp.route('/api/calendario-operativo', methods=['POST'])
@edicion_requerida
def api_guardar_feriado():
    data = request.get_json(silent=True) or {}
    normalizado = payload_feriado(data)
    if normalizado['errores']:
        return jsonify({'success': False, 'errores': normalizado['errores']}), 400
    valores = normalizado['valores']
    feriado_id = data.get('id')

    if feriado_id:
        feriado = FeriadoOperativo.query.get_or_404(feriado_id)
        fecha_anterior = feriado.fecha
        duplicado = FeriadoOperativo.query.filter_by(fecha=valores['fecha']).filter(FeriadoOperativo.id != feriado.id).first()
        if duplicado:
            return jsonify({'success': False, 'errores': ['Ya existe un feriado cargado para esa fecha']}), 400
    else:
        feriado = FeriadoOperativo.query.filter_by(fecha=valores['fecha']).first()
        fecha_anterior = None
        if not feriado:
            feriado = FeriadoOperativo()
            db.session.add(feriado)

    for campo, valor in valores.items():
        setattr(feriado, campo, valor)
    db.session.flush()
    recalculadas = recalcular_proyecciones_mes(feriado.year, feriado.fecha.month)
    if fecha_anterior and fecha_anterior != feriado.fecha:
        recalculadas += recalcular_proyecciones_mes(fecha_anterior.year, fecha_anterior.month)
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'Feriado guardado. {recalculadas} proyeccion(es) recalculada(s).',
        'feriado': feriado.to_dict(),
    })


@main_bp.route('/api/calendario-operativo/<int:feriado_id>', methods=['DELETE'])
@eliminacion_requerida
def api_eliminar_feriado(feriado_id):
    feriado = FeriadoOperativo.query.get_or_404(feriado_id)
    year = feriado.year
    mes_numero = feriado.fecha.month
    db.session.delete(feriado)
    db.session.flush()
    recalculadas = recalcular_proyecciones_mes(year, mes_numero)
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'Feriado eliminado. {recalculadas} proyeccion(es) recalculada(s).',
    })


@main_bp.route('/api/matriz-precios', methods=['GET'])
@login_requerido
def api_matriz_precios():
    year = int(request.args.get('year') or datetime.utcnow().year)
    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    registros = ProyeccionPrecio.query.filter(ProyeccionPrecio.mes.in_(meses)).order_by(
        ProyeccionPrecio.cliente,
        ProyeccionPrecio.campania,
        ProyeccionPrecio.mes,
    ).all()
    asociaciones = AsignacionComercial.query.filter_by(activa=True).order_by(
        AsignacionComercial.cliente,
        AsignacionComercial.campania,
    ).all()
    pares = []
    vistos = set()
    sites_por_clave = {}
    for asignacion in asociaciones:
        clave = (asignacion.cliente, asignacion.campania)
        sites_por_clave.setdefault(clave, asignacion.gerente or asignacion.jefe_site or 'Sin site')
        if clave in vistos:
            continue
        vistos.add(clave)
        pares.append({
            'cliente': asignacion.cliente,
            'campania': asignacion.campania,
            'site': asignacion.gerente or asignacion.jefe_site or 'Sin site',
        })
    return jsonify({
        'success': True,
        'year': year,
        'meses': opciones_meses_proyeccion(year),
        'asociaciones': pares,
        'precios': [
            {
                **registro.to_dict(),
                'site': registro.site or sites_por_clave.get((registro.cliente, registro.campania), ''),
                'mes_label': etiqueta_mes_proyeccion(registro.mes),
            }
            for registro in registros
        ],
    })


@main_bp.route('/api/matriz-precios', methods=['POST'])
@edicion_requerida
def api_guardar_matriz_precio():
    data = request.get_json(silent=True) or {}
    normalizado = payload_precio(data)
    if normalizado['errores']:
        return jsonify({'success': False, 'errores': normalizado['errores']}), 400
    valores = normalizado['valores']
    precio_id = data.get('id')
    if precio_id:
        precio = ProyeccionPrecio.query.get_or_404(precio_id)
        if (
            precio.cliente != valores['cliente']
            or precio.campania != valores['campania']
            or precio.mes != valores['mes']
        ):
            return jsonify({'success': False, 'errores': ['La edición replica ajustes sobre la misma campaña y mes. Para cambiar cliente, campaña o mes, cree un precio nuevo.']}), 400

    guardados = []
    antes = []
    site_asignado = valores.get('site') or sites_por_clave_precio(valores['cliente'], valores['campania']) or ''
    meses_a_guardar = [valores['mes']] if valores.get('horas_carga_manual') else meses_proyeccion_desde(valores['mes'])
    for mes in meses_a_guardar:
        precio = ProyeccionPrecio.query.filter_by(
            cliente=valores['cliente'],
            campania=valores['campania'],
            mes=mes,
        ).first()
        antes.append(precio.to_dict() if precio else None)
        if not precio:
            precio = ProyeccionPrecio()
            db.session.add(precio)
        precio.site = site_asignado or precio.site or ''
        precio.cliente = valores['cliente']
        precio.campania = valores['campania']
        precio.year = int(mes[:4])
        precio.mes = mes
        precio.precio_base = valores['precio_base']
        precio.alcance_porcentaje = valores['alcance_porcentaje']
        precio.importe_fijo_mensual = valores['importe_fijo_mensual']
        precio.recalcular()
        guardados.append(precio)

    db.session.flush()
    despues = [precio.to_dict() for precio in guardados]
    registrar_historial(
        'edicion' if any(antes) else 'creacion',
        'matriz_precios',
        valores['mes'],
        f'Precios guardados: {valores["cliente"]} / {valores["campania"]} desde {valores["mes"]}',
        antes={'precios': antes},
        despues={'precios': despues},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'{len(guardados)} precio(s) guardado(s)',
        'precios': [
            {**precio.to_dict(), 'mes_label': etiqueta_mes_proyeccion(precio.mes)}
            for precio in guardados
        ],
    })


@main_bp.route('/api/matriz-precios/inflacion', methods=['POST'])
@edicion_requerida
def api_aplicar_inflacion_matriz_precios():
    data = request.get_json(silent=True) or {}
    mes = mes_valido(str(data.get('mes') or '').strip())
    errores = []
    if not mes:
        errores.append('El mes no es valido')

    try:
        indice_inflacion = parse_numero(data.get('indice_inflacion') if data.get('indice_inflacion') not in (None, '') else 0)
    except ValueError:
        indice_inflacion = 0
        errores.append('El indice de inflacion tiene un formato invalido')

    if indice_inflacion < -100:
        errores.append('El indice de inflacion no puede ser menor a -100%')

    if errores:
        return jsonify({'success': False, 'errores': errores}), 400

    meses_afectados = meses_proyeccion_desde(mes)
    precios = ProyeccionPrecio.query.filter(ProyeccionPrecio.mes.in_(meses_afectados)).order_by(
        ProyeccionPrecio.cliente,
        ProyeccionPrecio.campania,
        ProyeccionPrecio.mes,
    ).all()
    antes = [precio.to_dict() for precio in precios]
    factor_inflacion = 1 + (indice_inflacion / 100)

    for precio in precios:
        precio.precio_base = redondear_moneda(redondear_moneda(precio.precio_base) * factor_inflacion)
        precio.recalcular()

    db.session.flush()
    despues = [precio.to_dict() for precio in precios]
    registrar_historial(
        'edicion',
        'matriz_precios',
        mes,
        f'Inflacion aplicada: {indice_inflacion}% desde {mes}',
        antes={'precios': antes, 'indice_inflacion': indice_inflacion},
        despues={'precios': despues, 'indice_inflacion': indice_inflacion},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'Inflacion aplicada a {len(precios)} precio(s)',
        'precios': [
            {**precio.to_dict(), 'mes_label': etiqueta_mes_proyeccion(precio.mes)}
            for precio in precios
        ],
    })


@main_bp.route('/api/matriz-precios/<int:precio_id>', methods=['DELETE'])
@eliminacion_requerida
def api_eliminar_matriz_precio(precio_id):
    precio = ProyeccionPrecio.query.get_or_404(precio_id)
    antes = precio.to_dict()
    db.session.delete(precio)
    registrar_historial(
        'eliminacion',
        'matriz_precios',
        precio_id,
        f'Precio eliminado: {antes.get("cliente")} / {antes.get("campania")} / {antes.get("mes")}',
        antes={'precios': [antes]},
        despues={'precios': []},
    )
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Precio eliminado'})


@main_bp.route('/api/matriz-precios/template', methods=['GET'])
@login_requerido
def api_template_matriz_precios():
    year = int(request.args.get('year') or datetime.utcnow().year)
    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    registros = ProyeccionPrecio.query.filter(ProyeccionPrecio.mes.in_(meses)).all()
    asociaciones = AsignacionComercial.query.filter_by(activa=True).order_by(
        AsignacionComercial.gerente,
        AsignacionComercial.campania,
    ).all()

    mapa_precios = {
        (registro.cliente, registro.campania, registro.mes): registro.precio_base or 0
        for registro in registros
    }
    mapa_fijos = {
        (registro.cliente, registro.campania, registro.mes): registro.importe_fijo_mensual or ''
        for registro in registros
        if registro.importe_fijo_mensual
    }
    sites_precios = {
        (registro.cliente, registro.campania): registro.site or ''
        for registro in registros
        if registro.site
    }
    filas = []
    vistos = set()
    for asignacion in asociaciones:
        clave = (asignacion.cliente, asignacion.campania)
        if clave in vistos:
            continue
        vistos.add(clave)
        filas.append([
            sites_precios.get(clave) or asignacion.gerente or asignacion.jefe_site or '',
            asignacion.campania,
            *[
                mapa_precios.get((asignacion.cliente, asignacion.campania, mes), '')
                for mes in meses
            ],
            *[
                mapa_fijos.get((asignacion.cliente, asignacion.campania, mes), '')
                for mes in meses
            ],
        ])

    for registro in registros:
        clave = (registro.cliente, registro.campania)
        if clave in vistos:
            continue
        vistos.add(clave)
        filas.append([
            sites_precios.get(clave, ''),
            registro.campania or registro.cliente,
            *[
                mapa_precios.get((registro.cliente, registro.campania, mes), '')
                for mes in meses
            ],
            *[
                mapa_fijos.get((registro.cliente, registro.campania, mes), '')
                for mes in meses
            ],
        ])

    if not filas:
        filas.append(['Ej: Gerencia Multicampaña', 'Ej: Assurant', 13662.53, 13662.53, 13662.53, '', '', '', '', '', '', '', '', '', 250000, 250000, 250000, '', '', '', '', '', '', '', '', ''])

    contenido = crear_xlsx(headers_template_precios(year), filas)
    return Response(
        contenido,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="plantilla_precios_{year}.xlsx"'}
    )


@main_bp.route('/api/matriz-precios/importar', methods=['POST'])
@carga_requerida
def api_importar_matriz_precios():
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename:
        return jsonify({'success': False, 'errores': ['Seleccione un archivo para importar']}), 400

    nombre = archivo.filename.lower()
    contenido_bytes = archivo.read()
    try:
        if nombre.endswith('.xlsx'):
            filas = filas_desde_xlsx(contenido_bytes)
        elif nombre.endswith('.csv'):
            filas = filas_desde_csv(decodificar_texto_importacion(contenido_bytes))
        else:
            return jsonify({'success': False, 'errores': ['Use un archivo .xlsx o .csv con el template de precios']}), 400
    except Exception:
        return jsonify({'success': False, 'errores': ['No pude leer el archivo. Descargue nuevamente el template y vuelva a completarlo.']}), 400

    if len(filas) < 2:
        return jsonify({'success': False, 'errores': ['No se encontraron filas para importar']}), 400

    headers = [normalizar_header(celda) for celda in filas[0]]
    mapa = {header: index for index, header in enumerate(headers)}
    site_index = mapa.get('site')
    cliente_index = mapa.get('cliente')
    if site_index is None or cliente_index is None:
        return jsonify({'success': False, 'errores': ['El template debe tener las columnas Site y Cliente']}), 400

    columnas_mes = []
    columnas_fijo = []
    for index, header_original in enumerate(filas[0]):
        mes = header_mes_precio(header_original)
        if mes:
            columnas_mes.append((index, mes))
        mes_fijo = header_mes_fijo_precio(header_original)
        if mes_fijo:
            columnas_fijo.append((index, mes_fijo))
    if not columnas_mes:
        return jsonify({'success': False, 'errores': ['No encontre columnas de meses con formato ene-26, feb-26, etc.']}), 400

    errores = []
    guardados = []
    antes = []
    for numero_fila, fila in enumerate(filas[1:], start=2):
        if not any(str(celda or '').strip() for celda in fila):
            continue
        site = str(fila[site_index] if site_index < len(fila) else '').strip()
        nombre_cliente = str(fila[cliente_index] if cliente_index < len(fila) else '').strip()
        if normalizar_header(site).startswith('ej:') or normalizar_header(nombre_cliente).startswith('ej:'):
            continue
        if not nombre_cliente:
            errores.append(f'Fila {numero_fila}: Cliente es obligatorio')
            continue
        cliente, campania = resolver_cliente_campania_precio(site, nombre_cliente)
        meses_a_importar = sorted({mes for _, mes in columnas_mes} | {mes for _, mes in columnas_fijo})
        for mes in meses_a_importar:
            precio_columna = next((index for index, mes_columna in columnas_mes if mes_columna == mes), None)
            fijo_columna = next((index for index, mes_columna in columnas_fijo if mes_columna == mes), None)
            fijo_presente = fijo_columna is not None
            valor = fila[precio_columna] if precio_columna is not None and precio_columna < len(fila) else ''
            valor_fijo = fila[fijo_columna] if fijo_columna is not None and fijo_columna < len(fila) else ''
            precio_con_valor = valor not in (None, '') and str(valor).strip() != ''
            fijo_con_valor = fijo_presente and valor_fijo not in (None, '') and str(valor_fijo).strip() != ''
            if (
                not precio_con_valor
                and not fijo_con_valor
            ):
                continue
            try:
                precio_base = parse_numero(valor) if precio_con_valor else 0
                importe_fijo_mensual = parse_numero(valor_fijo) if fijo_con_valor else 0
            except ValueError:
                errores.append(f'Fila {numero_fila}, {mes}: precio o fijo invalido')
                continue
            registro = ProyeccionPrecio.query.filter_by(
                cliente=cliente,
                campania=campania,
                mes=mes,
            ).first()
            antes.append(registro.to_dict() if registro else None)
            if not registro:
                registro = ProyeccionPrecio()
                db.session.add(registro)
            registro.site = site
            registro.cliente = cliente
            registro.campania = campania
            registro.year = int(mes[:4])
            registro.mes = mes
            if precio_con_valor:
                registro.precio_base = precio_base
            elif not registro.id:
                registro.precio_base = 0
            if precio_con_valor or not registro.id:
                registro.alcance_porcentaje = 100
            if fijo_con_valor:
                registro.importe_fijo_mensual = importe_fijo_mensual
            elif not registro.id:
                registro.importe_fijo_mensual = 0
            registro.recalcular()
            guardados.append(registro)

    if errores:
        return jsonify({'success': False, 'errores': errores[:30]}), 400
    if not guardados:
        return jsonify({'success': False, 'errores': ['No hay precios validos para importar']}), 400

    db.session.flush()
    despues = [registro.to_dict() for registro in guardados]
    registrar_historial(
        'edicion' if any(antes) else 'creacion',
        'matriz_precios',
        ','.join(sorted({str(registro.year) for registro in guardados})),
        f'Importacion de precios: {len(guardados)} registro(s)',
        antes={'precios': antes},
        despues={'precios': despues},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'{len(guardados)} precio(s) importado(s)',
        'years': sorted({registro.year for registro in guardados}),
    })


def construir_resumen_proyeccion_data(year):
    meses_info = opciones_meses_proyeccion(year)
    meses = [item['value'] for item in meses_info]

    asociaciones = AsignacionComercial.query.filter_by(activa=True).all()
    sites_regularizados = {
        (normalizar_header(registro.cliente), normalizar_nombre_variable(registro.campania)): registro
        for registro in SiteProyeccion.query.all()
    }

    def regularizacion(cliente, campania):
        clave_campania = normalizar_nombre_variable(campania)
        if clave_campania.endswith(' nocturnidad'):
            clave_campania = clave_campania[:-len(' nocturnidad')]
        return sites_regularizados.get((normalizar_header(cliente), clave_campania))

    def site_regularizado(cliente, campania):
        registro = regularizacion(cliente, campania)
        return registro.site if registro else None
    sites_por_clave = {}
    sites_por_nombre = {}
    for asignacion in asociaciones:
        site_asignado = normalizar_site(asignacion.gerente or asignacion.jefe_site or 'Sin site')
        sites_por_clave.setdefault((asignacion.cliente, asignacion.campania), site_asignado)
        sites_por_nombre.setdefault(normalizar_header(asignacion.campania), site_asignado)
        sites_por_nombre.setdefault(normalizar_header(asignacion.cliente), site_asignado)

    # Personal se valoriza desde sus aperturas PLP, que distinguen servicio y
    # nocturnidad. La fila regular de la Matriz es el agregado de esas mismas
    # horas: conservar ambos circuitos duplica Facturación horas y Variable.
    proyecciones = ProyeccionMatriz.query.filter(
        ProyeccionMatriz.mes.in_(meses),
        db.or_(
            db.and_(ProyeccionMatriz.tipo_plp.isnot(None), ProyeccionMatriz.tipo_plp != ''),
            func.lower(func.trim(ProyeccionMatriz.cliente)) != 'personal',
        ),
    ).all()
    precios = ProyeccionPrecio.query.filter(ProyeccionPrecio.mes.in_(meses)).all()
    precios_por_clave = {}
    precios_por_nombre = {}
    precios_por_nombre_variable = {}
    for precio in precios:
        site_precio = (
            (precio.site or '').strip()
            or sites_por_clave.get((precio.cliente, precio.campania))
            or sites_por_nombre.get(normalizar_header(precio.campania))
            or sites_por_nombre.get(normalizar_header(precio.cliente))
        )
        precios_por_clave[(precio.cliente, precio.campania, precio.mes)] = precio
        precios_por_nombre.setdefault((normalizar_header(precio.campania), precio.mes), precio)
        precios_por_nombre.setdefault((normalizar_header(precio.cliente), precio.mes), precio)
        precios_por_nombre_variable.setdefault((normalizar_nombre_variable(precio.campania), precio.mes), precio)
        precios_por_nombre_variable.setdefault((normalizar_nombre_variable(precio.cliente), precio.mes), precio)

    filas = []
    total_general = {mes: 0 for mes in meses}
    totales_por_site = {}

    for proyeccion in proyecciones:
        clave = (proyeccion.cliente, proyeccion.campania)
        precio = (
            precios_por_clave.get((proyeccion.cliente, proyeccion.campania, proyeccion.mes))
            or precios_por_nombre.get((normalizar_header(proyeccion.campania), proyeccion.mes))
            or precios_por_nombre.get((normalizar_header(proyeccion.cliente), proyeccion.mes))
        )
        if not precio:
            continue
        site = 'Personal' if proyeccion.tipo_plp else (
            (precio.site or '').strip()
            or sites_por_clave.get(clave)
            or sites_por_nombre.get(normalizar_header(proyeccion.campania))
            or sites_por_nombre.get(normalizar_header(proyeccion.cliente))
            or site_original_next_gen(proyeccion.cliente, proyeccion.campania)
        )
        site = normalizar_site(site)
        if normalizar_header(site) == 'next gen':
            site = site_original_next_gen(proyeccion.cliente, proyeccion.campania)
        for nombre_campania, horas, _ in aperturas_nocturnidad(proyeccion):
            site_apertura = normalizar_site(site_regularizado(proyeccion.cliente, nombre_campania) or site)
            precio_apertura = precios_por_nombre_variable.get(
                (normalizar_nombre_variable(nombre_campania), proyeccion.mes),
                precio,
            )
            valor = horas * (precio_apertura.precio_final or 0)
            fila_clave = (site_apertura, proyeccion.cliente, nombre_campania, 'Horas')
            fila = next((item for item in filas if item['clave'] == fila_clave), None)
            if not fila:
                fila = {
                    'clave': fila_clave,
                    'site': site_apertura,
                    'cliente': proyeccion.cliente,
                    'campania': nombre_campania,
                    'campania_facturacion': nombre_campania if proyeccion.tipo_plp else proyeccion.cliente,
                    'concepto': 'Horas',
                    'meses': {mes: 0 for mes in meses},
                    'total': 0,
                }
                filas.append(fila)
            fila['meses'][proyeccion.mes] += valor
            fila['total'] += valor
            total_general[proyeccion.mes] += valor
            totales_por_site.setdefault(site_apertura, {mes: 0 for mes in meses})
            totales_por_site[site_apertura][proyeccion.mes] += valor

    # Si una apertura no tiene proyección propia en un mes, usa la proyección
    # base de su cliente. La relación apertura→cliente se aprende de los otros
    # meses cargados (caso Unicef Activos+Leads en enero).
    cliente_por_apertura = {
        normalizar_nombre_variable(item.campania): item.cliente
        for item in proyecciones
        if normalizar_nombre_variable(item.campania) != normalizar_nombre_variable(item.cliente)
        and not item.tipo_plp
    }
    proyecciones_exactas = {
        (normalizar_nombre_variable(item.campania), item.mes) for item in proyecciones
    }
    proyeccion_base = {
        (normalizar_header(item.cliente), item.mes): item
        for item in proyecciones
        if normalizar_nombre_variable(item.campania) == normalizar_nombre_variable(item.cliente)
        and not item.tipo_plp
    }
    for precio in precios:
        clave_apertura = normalizar_nombre_variable(precio.campania)
        cliente_padre = cliente_por_apertura.get(clave_apertura)
        if not cliente_padre or (clave_apertura, precio.mes) in proyecciones_exactas:
            continue
        base = proyeccion_base.get((normalizar_header(cliente_padre), precio.mes))
        if not base:
            continue
        site = normalizar_site(
            site_regularizado(cliente_padre, precio.campania)
            or (precio.site or '').strip()
            or sites_por_clave.get((cliente_padre, precio.campania))
            or sites_por_nombre.get(normalizar_header(cliente_padre))
            or 'Sin site'
        )
        valor = (base.horas_proyectadas or 0) * (precio.precio_final or 0)
        fila_clave = (site, cliente_padre, precio.campania, 'Horas')
        fila = next((item for item in filas if item['clave'] == fila_clave), None)
        if not fila:
            fila = {
                'clave': fila_clave, 'site': site, 'cliente': cliente_padre,
                'campania': precio.campania, 'campania_facturacion': cliente_padre,
                'concepto': 'Horas', 'meses': {mes: 0 for mes in meses}, 'total': 0,
            }
            filas.append(fila)
        fila['meses'][precio.mes] += valor
        fila['total'] += valor
        total_general[precio.mes] += valor
        totales_por_site.setdefault(site, {mes: 0 for mes in meses})
        totales_por_site[site][precio.mes] += valor

    for precio in precios:
        valor = precio.importe_fijo_mensual or 0
        if valor <= 0:
            continue
        clave = (precio.cliente, precio.campania)
        site = (
            (precio.site or '').strip()
            or sites_por_clave.get(clave)
            or sites_por_nombre.get(normalizar_header(precio.campania))
            or sites_por_nombre.get(normalizar_header(precio.cliente))
            or 'Sin site'
        )
        site = normalizar_site(site)
        site = normalizar_site(site_regularizado(precio.cliente, precio.campania) or site)
        fila_clave = (site, precio.cliente, precio.campania, 'Fijo mensual')
        fila = next((item for item in filas if item['clave'] == fila_clave), None)
        if not fila:
            fila = {
                'clave': fila_clave,
                'site': site,
                'cliente': precio.cliente,
                'campania': precio.campania,
                'concepto': 'Fijo mensual',
                'meses': {mes: 0 for mes in meses},
                'total': 0,
            }
            filas.append(fila)
        fila['meses'][precio.mes] += valor
        fila['total'] += valor
        total_general[precio.mes] += valor
        totales_por_site.setdefault(site, {mes: 0 for mes in meses})
        totales_por_site[site][precio.mes] += valor

    # Variable se incorpora como un concepto adicional de la misma campaña.
    # El cruce ignora mayúsculas, minúsculas, acentos y alias de nocturnidad.
    variables_year = VariableCampania.query.filter(VariableCampania.year <= year).all()
    filas_horas = [item for item in filas if item['concepto'] == 'Horas']
    # La planilla aplica el porcentaje sobre toda la base facturable de la
    # campaña. Esto incluye el fijo mensual cuando existe, no solamente horas.
    filas_base_variable = [item for item in filas if item['concepto'] in ('Horas', 'Fijo mensual')]
    for base in filas_base_variable:
        for mes in meses:
            campania_base_variable = normalizar_nombre_variable(base['campania'])
            if campania_base_variable.endswith(' nocturnidad'):
                campania_base_variable = campania_base_variable[:-len(' nocturnidad')]
            candidatas = [item for item in variables_year
                if normalizar_nombre_variable(item.campania) == normalizar_nombre_variable(base['campania'])
                and item.mes.endswith(f'-{mes[-2:]}')
                and (
                    normalizar_header(item.cliente) == normalizar_header(base['cliente'])
                    or normalizar_nombre_variable(item.cliente) == normalizar_nombre_variable(base['campania'])
                    or normalizar_nombre_variable(item.cliente) == campania_base_variable
                )]
            variable = max(candidatas, key=lambda item: item.year, default=None)
            if not variable:
                continue
            valor = base['meses'][mes] * ((variable.porcentaje or 0) / 100)
            fila_clave = (base['site'], base['cliente'], base['campania'], 'Variable')
            fila = next((item for item in filas if item['clave'] == fila_clave), None)
            if not fila:
                fila = {
                    'clave': fila_clave, 'site': base['site'], 'cliente': base['cliente'],
                    'campania': base['campania'], 'concepto': 'Variable',
                    'campania_facturacion': base.get('campania_facturacion'),
                    'meses': {item: 0 for item in meses}, 'total': 0,
                }
                filas.append(fila)
            fila['meses'][mes] += valor
            fila['total'] += valor
            total_general[mes] += valor
            totales_por_site.setdefault(base['site'], {item: 0 for item in meses})
            totales_por_site[base['site']][mes] += valor

    # Tarifación es un concepto monetario independiente que integra
    # Facturación horas y participa de sus totales con su propio signo.
    tarifaciones = TarifacionCampania.query.filter(
        TarifacionCampania.year == year,
        TarifacionCampania.mes.in_(meses),
    ).all()
    for tarifacion in tarifaciones:
        valor = tarifacion.monto or 0
        base_coincidente = next((item for item in filas_horas
            if normalizar_nombre_variable(item['campania']) == normalizar_nombre_variable(tarifacion.campania)
            and (
                normalizar_header(item['cliente']) == normalizar_header(tarifacion.cliente)
                or normalizar_nombre_variable(item['campania']) == normalizar_nombre_variable(tarifacion.cliente)
            )), None)
        site = normalizar_site(base_coincidente['site'] if base_coincidente else ((tarifacion.site or '').strip() or 'Sin site'))
        cliente = base_coincidente['cliente'] if base_coincidente else tarifacion.cliente
        campania = base_coincidente['campania'] if base_coincidente else tarifacion.campania
        concepto = tarifacion.concepto or 'Tarifación'
        fila_clave = (site, cliente, campania, concepto)
        fila = next((item for item in filas if item['clave'] == fila_clave), None)
        if not fila:
            fila = {
                'clave': fila_clave, 'site': site, 'cliente': cliente,
                'campania': campania, 'concepto': concepto,
                'campania_facturacion': base_coincidente.get('campania_facturacion') if base_coincidente else tarifacion.cliente,
                'meses': {mes: 0 for mes in meses}, 'total': 0,
            }
            filas.append(fila)
        fila['meses'][tarifacion.mes] += valor
        fila['total'] += valor
        total_general[tarifacion.mes] += valor
        totales_por_site.setdefault(site, {mes: 0 for mes in meses})
        totales_por_site[site][tarifacion.mes] += valor

    # Next Gen: cantidades mensuales en USD valorizadas con el dólar de cada mes.
    dolares_next_gen = {r.mes: r.valor or 0 for r in NextGenDolar.query.filter_by(year=year).all()}
    productos_next_gen = NextGenProducto.query.filter_by(year=year).all()
    for producto in productos_next_gen:
        cotizacion = producto.cotizacion_aplicada
        if cotizacion is None:
            cotizacion = dolares_next_gen.get(producto.mes, 0)
        # Igual que el libro: conservar precisión por producto y redondear una
        # vez consolidada la campaña/mes. Redondear aquí altera el TOTAL CAT.
        valor = (producto.cantidad_usd or 0) * cotizacion
        base_coincidente = next((item for item in filas_horas
            if normalizar_nombre_variable(item['campania']) == normalizar_nombre_variable(producto.campania)
            and (
                normalizar_header(item['cliente']) == normalizar_header(producto.cliente)
                or normalizar_nombre_variable(item['campania']) == normalizar_nombre_variable(producto.cliente)
            )), None)
        site_base = base_coincidente['site'] if base_coincidente else (producto.site or '').strip()
        if normalizar_header(site_base) in ('', 'next gen', 'sin site'):
            site_base = site_original_next_gen(producto.cliente, producto.campania)
        site_base = site_regularizado(producto.cliente, producto.campania) or site_base
        site = normalizar_site(site_base)
        cliente = base_coincidente['cliente'] if base_coincidente else producto.cliente
        campania = base_coincidente['campania'] if base_coincidente else producto.campania
        fila_clave = (site, cliente, campania, 'Next Gen')
        fila = next((item for item in filas if item['clave'] == fila_clave), None)
        if not fila:
            fila = {
                'clave': fila_clave, 'site': site, 'cliente': cliente,
                'campania': campania, 'concepto': 'Next Gen',
                'campania_facturacion': base_coincidente.get('campania_facturacion') if base_coincidente else producto.cliente,
                'meses': {mes: 0 for mes in meses}, 'total': 0,
            }
            filas.append(fila)
        if producto.mes in fila['meses']:
            fila['meses'][producto.mes] += valor
            fila['total'] += valor
            total_general[producto.mes] += valor
            totales_por_site.setdefault(site, {mes: 0 for mes in meses})
            totales_por_site[site][producto.mes] += valor

    filas_ordenadas = []
    for site in sorted({fila['site'] for fila in filas}):
        grupo = sorted(
            [fila for fila in filas if fila['site'] == site],
            key=lambda item: (item['campania'], item['concepto']),
        )
        filas_ordenadas.extend(grupo)
        total_site = totales_por_site.get(site, {mes: 0 for mes in meses})
        filas_ordenadas.append({
            'site': site,
            'cliente': '',
            'campania': site,
            'concepto': 'Total',
            'tipo': 'total_site',
            'meses': total_site,
            'total': sum(total_site.values()),
        })

    return {
        'success': True,
        'year': year,
        'meses': meses_info,
        'filas': [
            {
                'site': ((regularizacion(fila.get('cliente'), fila.get('campania')).site
                          if regularizacion(fila.get('cliente'), fila.get('campania')) else None) or fila['site']),
                'cliente': ((regularizacion(fila.get('cliente'), fila.get('campania')).cliente_destino
                             if regularizacion(fila.get('cliente'), fila.get('campania')) else None)
                            or fila.get('cliente') or ''),
                'campania': ((regularizacion(fila.get('cliente'), fila.get('campania')).campania_destino
                              if regularizacion(fila.get('cliente'), fila.get('campania')) else None)
                             or fila.get('campania_facturacion') or fila['campania']),
                'origenes': ([{'cliente': fila.get('cliente') or '', 'campania': fila.get('campania') or ''}]
                             if fila.get('tipo', 'detalle') == 'detalle' else []),
                'concepto': fila['concepto'],
                'tipo': fila.get('tipo', 'detalle'),
                # Se conserva precisión interna hasta consolidar la Cuenta.
                'meses': {mes: round(fila['meses'].get(mes, 0), 10) for mes in meses},
                'total': round(fila.get('total', 0), 10),
            }
            for fila in filas_ordenadas
        ],
        'total_general': {
            'meses': {mes: round(total_general.get(mes, 0), 10) for mes in meses},
            'total': round(sum(total_general.values()), 10),
        },
    }


@main_bp.route('/api/resumen-proyeccion', methods=['GET'])
@login_requerido
def api_resumen_proyeccion():
    year = int(request.args.get('year') or datetime.utcnow().year)
    data = construir_resumen_proyeccion_data(year)
    salida = consolidar_resumen_proyeccion(data)
    salida['filas_resumen'] = construir_filas_resumen_excel(data)
    return jsonify(salida)


CONCEPTOS_RESUMEN = ('Horas', 'Variable Productividad', 'Next Gen', 'Tarifacion + otros')


def construir_filas_resumen_excel(data):
    """Replica la topología de RESUMEN: concepto, campaña, total y meses."""
    meses = [item['value'] for item in data.get('meses', [])]
    orden_personal = {'personal cx': 0, 'personal soporte': 1, 'personal': 2, 'personal smb': 3}
    grupos = {}
    for fila in data.get('filas', []):
        if fila.get('tipo') == 'total_site':
            continue
        campania = str(fila.get('campania') or '').strip()
        clave_campania = normalizar_nombre_variable(campania)
        if clave_campania.endswith(' nocturnidad'):
            clave_campania = clave_campania[:-len(' nocturnidad')]
            campania = campania[:-len(' nocturnidad')].strip()
        concepto_origen = fila.get('concepto')
        if concepto_origen in ('Horas', 'Fijo mensual'):
            concepto = 'Horas'
        elif concepto_origen == 'Variable':
            concepto = 'Variable Productividad'
        elif concepto_origen == 'Next Gen':
            concepto = 'Next Gen'
        else:
            concepto = 'Tarifacion + otros'
        cliente_clave = 'personal' if clave_campania in orden_personal else normalizar_header(fila.get('cliente'))
        clave = (cliente_clave, clave_campania)
        grupo = grupos.setdefault(clave, {
            'site': fila.get('site') or '', 'cliente': fila.get('cliente') or '',
            'campania': campania, 'conceptos': {
                nombre: {mes: 0 for mes in meses} for nombre in CONCEPTOS_RESUMEN
            },
        })
        for mes in meses:
            grupo['conceptos'][concepto][mes] += numero_seguro(fila.get('meses', {}).get(mes))

    ordenados = sorted(grupos.values(), key=lambda item: (
        0 if normalizar_nombre_variable(item['campania']) in orden_personal else 1,
        orden_personal.get(normalizar_nombre_variable(item['campania']), 99),
        normalizar_header(item['site']), normalizar_header(item['campania']),
    ))
    salida = []
    total_personal = {mes: 0 for mes in meses}
    for grupo in ordenados:
        valores_total = {mes: 0 for mes in meses}
        for concepto in CONCEPTOS_RESUMEN:
            valores = grupo['conceptos'][concepto]
            for mes in meses:
                valores_total[mes] += valores[mes]
            salida.append({
                'site': grupo['site'], 'cliente': grupo['cliente'],
                'concepto': concepto, 'campania': grupo['campania'], 'tipo': 'concepto',
                'meses': {mes: redondear_moneda(valores[mes]) for mes in meses},
                'total': redondear_moneda(sum(valores.values())),
            })
        es_personal = normalizar_nombre_variable(grupo['campania']) in orden_personal
        if es_personal:
            for mes in meses:
                total_personal[mes] += valores_total[mes]
        salida.append({
            'site': grupo['site'], 'cliente': grupo['cliente'], 'concepto': 'Total',
            'campania': grupo['campania'], 'tipo': 'total_campania',
            'meses': {mes: redondear_moneda(valores_total[mes]) for mes in meses},
            'total': redondear_moneda(sum(valores_total.values())),
        })
        if es_personal and normalizar_nombre_variable(grupo['campania']) == 'personal smb':
            salida.append({
                'site': grupo['site'], 'cliente': 'Personal', 'concepto': '',
                'campania': 'TOTAL GRUPO PERSONAL', 'tipo': 'total_grupo',
                'meses': {mes: redondear_moneda(total_personal[mes]) for mes in meses},
                'total': redondear_moneda(sum(total_personal.values())),
            })
    total = data.get('total_general', {})
    salida.insert(0, {
        'site': '', 'cliente': '', 'concepto': '', 'campania': 'TOTAL CAT', 'tipo': 'total_cat',
        'meses': {mes: redondear_moneda(total.get('meses', {}).get(mes, 0)) for mes in meses},
        'total': redondear_moneda(total.get('total', 0)),
    })
    return salida


@main_bp.route('/api/resumen-proyeccion/exportar', methods=['GET'])
@login_requerido
def exportar_resumen_proyeccion():
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    year = int(request.args.get('year') or datetime.utcnow().year)
    data = construir_resumen_proyeccion_data(year)
    filas = construir_filas_resumen_excel(data)
    meses = data['meses']
    wb = Workbook(); ws = wb.active; ws.title = 'RESUMEN'
    nombres_meses = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
                     'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
    encabezados = ['Concepto', 'Campaña', str(year)] + nombres_meses
    azul = PatternFill('solid', fgColor='D9E5F6'); oscuro = PatternFill('solid', fgColor='40536D')
    borde = Border(bottom=Side(style='thin', color='B7C5D8'))
    fila_cat = next(fila for fila in filas if fila['tipo'] == 'total_cat')
    ws.append(['', fila_cat['campania'], fila_cat['total']] + [fila_cat['meses'][m['value']] for m in meses])
    ws.append(encabezados)
    for celda in ws[1]: celda.fill = oscuro; celda.font = Font(bold=True, color='FFFFFF')
    for celda in ws[2]:
        celda.font = Font(bold=True); celda.fill = azul; celda.alignment = Alignment(horizontal='center')
    for fila in (item for item in filas if item['tipo'] != 'total_cat'):
        ws.append([fila['concepto'], fila['campania'], fila['total']] + [fila['meses'][m['value']] for m in meses])
        numero_fila = ws.max_row
        if fila['tipo'] in ('total_campania', 'total_grupo'):
            for celda in ws[numero_fila]: celda.fill = azul; celda.font = Font(bold=True)
        for celda in ws[numero_fila]: celda.border = borde
        for columna in range(3, 16):
            ws.cell(numero_fila, columna).number_format = '$ #,##0;[Red]-$ #,##0;$ 0'
    for columna in range(3, 16): ws.cell(1, columna).number_format = '$ #,##0;[Red]-$ #,##0;$ 0'
    ws.freeze_panes = 'C3'; ws.auto_filter.ref = f'A2:{get_column_letter(ws.max_column)}2'
    ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 30
    for columna in range(3, 16): ws.column_dimensions[get_column_letter(columna)].width = 17
    salida = BytesIO(); wb.save(salida); salida.seek(0)
    return send_file(salida, as_attachment=True,
                     download_name=f'Facturacion_horas_RESUMEN_{year}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@main_bp.route('/api/sites-proyeccion', methods=['POST'])
@edicion_requerida
def api_guardar_site_proyeccion():
    data = request.get_json(silent=True) or {}
    site = str(data.get('site') or '').strip()
    cliente_destino = str(data.get('cliente_destino') or '').strip()
    campania_destino = str(data.get('campania_destino') or '').strip()
    origenes = data.get('origenes') or [{'cliente': data.get('cliente'), 'campania': data.get('campania')}]
    if not site or not cliente_destino or not campania_destino:
        return jsonify({'success': False, 'errores': ['Site, cliente y campaña son obligatorios']}), 400
    existentes = SiteProyeccion.query.all()
    antes, despues = [], []
    for origen in origenes:
        cliente = str(origen.get('cliente') or '').strip()
        campania = str(origen.get('campania') or '').strip()
        if not cliente or not campania:
            continue
        clave_cliente = normalizar_header(cliente)
        clave_campania = normalizar_nombre_variable(campania)
        if clave_campania.endswith(' nocturnidad'):
            clave_campania = clave_campania[:-len(' nocturnidad')]
            campania = campania[:-len(' nocturnidad')].strip()
        registro = next((item for item in existentes
                         if normalizar_header(item.cliente) == clave_cliente
                         and normalizar_nombre_variable(item.campania) == clave_campania), None)
        antes.append(registro.to_dict() if registro else None)
        if not registro:
            registro = SiteProyeccion(cliente=cliente, campania=campania, site=site)
            db.session.add(registro)
            existentes.append(registro)
        registro.site = site
        registro.cliente_destino = cliente_destino
        registro.campania_destino = campania_destino
        db.session.flush()
        despues.append(registro.to_dict())
    if not despues:
        return jsonify({'success': False, 'errores': ['No se recibió ningún nombre de origen válido']}), 400
    registrar_historial(
        'edicion' if any(antes) else 'creacion', 'sites_proyeccion', None,
        f'Facturación horas > Regularización: {cliente_destino} / {campania_destino} → {site}',
        detalle=f'{len(despues)} identificación(es) de origen',
        antes={'sites_proyeccion': antes}, despues={'sites_proyeccion': despues},
    )
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Site, cliente y campaña actualizados; los nombres coincidentes se consolidaron'})


def consolidar_resumen_proyeccion(data):
    """Devuelve una sola fila por campaña con Horas + Variable + Tarifación."""
    meses = [item['value'] for item in data.get('meses', [])]
    detalles = [fila for fila in data.get('filas', []) if fila.get('tipo') != 'total_site']
    grupos = {}
    # Horas primero para conservar su cliente y site como datos principales.
    for fila in sorted(detalles, key=lambda item: (
        0 if item.get('concepto') == 'Horas' else 1,
        1 if normalizar_nombre_variable(item.get('campania')).endswith(' nocturnidad') else 0,
    )):
        clave = normalizar_nombre_variable(fila.get('campania'))
        if clave.endswith(' nocturnidad'):
            clave = clave[:-len(' nocturnidad')]
        clave = (normalizar_header(fila.get('cliente')), clave)
        grupo = grupos.setdefault(clave, {
            'site': fila.get('site') or '', 'cliente': fila.get('cliente') or '',
            'campania': fila.get('campania') or '', 'concepto': 'Consolidado',
            'tipo': 'detalle', 'meses': {mes: 0 for mes in meses}, 'total': 0,
            'conceptos': set(), 'origenes': [],
        })
        grupo['conceptos'].add(fila.get('concepto') or '')
        for origen in fila.get('origenes', []):
            if origen not in grupo['origenes']:
                grupo['origenes'].append(origen)
        for mes in meses:
            grupo['meses'][mes] += numero_seguro(fila.get('meses', {}).get(mes))

    filas = []
    totales_site = {}
    for grupo in grupos.values():
        total_crudo = sum(grupo['meses'].values())
        grupo['meses'] = {mes: redondear_moneda(valor) for mes, valor in grupo['meses'].items()}
        grupo['total'] = redondear_moneda(total_crudo)
        grupo['conceptos'] = sorted(concepto for concepto in grupo['conceptos'] if concepto)
        filas.append(grupo)
        site = grupo['site'] or 'Sin site'
        totales_site.setdefault(site, {mes: 0 for mes in meses})
        for mes in meses:
            totales_site[site][mes] += grupo['meses'][mes]

    salida = []
    for site in sorted(totales_site):
        salida.extend(sorted([fila for fila in filas if (fila['site'] or 'Sin site') == site], key=lambda item: item['campania']))
        meses_site = {mes: round(totales_site[site][mes], 2) for mes in meses}
        salida.append({
            'site': site, 'cliente': '', 'campania': site, 'concepto': 'Total',
            'tipo': 'total_site', 'meses': meses_site, 'total': round(sum(meses_site.values()), 2),
        })
    return {**data, 'filas': salida}


def numero_seguro(valor):
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0


def site_por_proyeccion(cliente, campania, sites_por_clave, sites_por_nombre):
    return (
        sites_por_clave.get((cliente, campania))
        or sites_por_nombre.get(normalizar_header(campania))
        or sites_por_nombre.get(normalizar_header(cliente))
        or 'Sin site'
    )


def aperturas_nocturnidad(proyeccion):
    horas = proyeccion.horas_proyectadas or 0
    dotacion = proyeccion.dotacion_requerida or 0
    porcentaje = proyeccion.porcentaje_nocturnidad or 0
    nombre_base = proyeccion.tipo_plp or proyeccion.campania
    if not proyeccion.tiene_nocturnidad or porcentaje <= 0:
        return [(nombre_base, horas, dotacion)]
    horas_nocturnas = horas * porcentaje / 100
    dotacion_nocturna = dotacion * porcentaje / 100
    return [
        (nombre_base, horas - horas_nocturnas, dotacion - dotacion_nocturna),
        (f'{nombre_base} nocturnidad', horas_nocturnas, dotacion_nocturna),
    ]


def construir_control_proyecciones_data(year):
    meses_info = opciones_meses_proyeccion(year)
    meses = [item['value'] for item in meses_info]
    asociaciones = AsignacionComercial.query.filter_by(activa=True).all()
    sites_por_clave = {}
    sites_por_nombre = {}
    for asignacion in asociaciones:
        site = asignacion.gerente or asignacion.jefe_site or 'Sin site'
        sites_por_clave.setdefault((asignacion.cliente, asignacion.campania), site)
        sites_por_nombre.setdefault(normalizar_header(asignacion.campania), site)
        sites_por_nombre.setdefault(normalizar_header(asignacion.cliente), site)

    # El control replica la pestaña Horas del Excel. Las filas PLP son un
    # desglose auxiliar de Personal y ya están contenidas en la matriz base.
    proyecciones = ProyeccionMatriz.query.filter(
        ProyeccionMatriz.mes.in_(meses),
        db.or_(ProyeccionMatriz.tipo_plp.is_(None), ProyeccionMatriz.tipo_plp == ''),
    ).order_by(
        ProyeccionMatriz.cliente,
        ProyeccionMatriz.campania,
        ProyeccionMatriz.mes,
    ).all()
    filas_por_clave = {}
    total_horas_crudo = {mes: 0 for mes in meses}
    total_dotaciones_crudo = {mes: 0 for mes in meses}
    for proyeccion in proyecciones:
        site = 'Personal' if proyeccion.tipo_plp else site_por_proyeccion(proyeccion.cliente, proyeccion.campania, sites_por_clave, sites_por_nombre)
        for nombre_campania, horas, dotacion in aperturas_nocturnidad(proyeccion):
            clave = (site, proyeccion.cliente, nombre_campania)
            fila = filas_por_clave.setdefault(clave, {
                'site': site,
                'cliente': proyeccion.cliente,
                'campania': nombre_campania,
                'horas': {mes: 0 for mes in meses},
                'dotaciones': {mes: 0 for mes in meses},
                'total_horas': 0,
                'total_dotaciones': 0,
            })
            fila['horas'][proyeccion.mes] += horas
            fila['dotaciones'][proyeccion.mes] += dotacion
            total_horas_crudo[proyeccion.mes] += horas
            total_dotaciones_crudo[proyeccion.mes] += dotacion

    filas = []
    total_horas = {mes: 0 for mes in meses}
    total_dotaciones = {mes: 0 for mes in meses}
    for clave in sorted(filas_por_clave):
        fila = filas_por_clave[clave]
        fila['horas'] = {mes: round(fila['horas'][mes], 2) for mes in meses}
        fila['dotaciones'] = {mes: round(fila['dotaciones'][mes], 2) for mes in meses}
        fila['total_horas'] = round(sum(fila['horas'].values()), 2)
        fila['total_dotaciones'] = round(sum(fila['dotaciones'].values()), 2)
        filas.append(fila)

    total_horas = total_horas_crudo
    total_dotaciones = total_dotaciones_crudo

    return {
        'filas': filas,
        'total_horas': {
            'meses': {mes: round(total_horas[mes], 2) for mes in meses},
            'total': round(sum(total_horas.values()), 2),
        },
        'total_dotaciones': {
            'meses': {mes: round(total_dotaciones[mes], 2) for mes in meses},
            'total': round(sum(total_dotaciones.values()), 2),
        },
    }


def normalizar_nombre_variable(valor):
    nombre = normalizar_header(valor)
    for sufijo in (' nocturnos', ' nocturno', ' nocturna'):
        if nombre.endswith(sufijo):
            nombre = nombre[:-len(sufijo)] + ' nocturnidad'
    return nombre


def construir_variable_data(year):
    resumen_data = construir_resumen_proyeccion_data(year)
    meses_info = resumen_data['meses']
    meses = [item['value'] for item in meses_info]
    base_por_clave = {}

    for fila in resumen_data['filas']:
        if fila.get('tipo') == 'total_site':
            continue
        if fila.get('concepto') not in ('Horas', 'Fijo mensual'):
            continue
        clave = (fila.get('site') or '', fila.get('cliente') or '', fila.get('campania') or '')
        base = base_por_clave.setdefault(clave, {
            'site': fila.get('site') or '',
            'cliente': fila.get('cliente') or '',
            'campania': fila.get('campania') or '',
            'meses': {mes: 0 for mes in meses},
            'total_base': 0,
        })
        for mes in meses:
            base['meses'][mes] += numero_seguro(fila.get('meses', {}).get(mes))
        base['total_base'] += numero_seguro(fila.get('total'))

    variables = VariableCampania.query.filter(VariableCampania.year <= year).all()
    claves_base_normalizadas = {
        (normalizar_header(item['cliente']), normalizar_nombre_variable(item['campania']))
        for item in base_por_clave.values()
    }
    for variable in sorted(variables, key=lambda item: item.year, reverse=True):
        clave_normalizada = (normalizar_header(variable.cliente), normalizar_nombre_variable(variable.campania))
        if clave_normalizada in claves_base_normalizadas:
            continue
        base_por_clave[(variable.site or '', variable.cliente, variable.campania)] = {
            'site': variable.site or '',
            'cliente': variable.cliente,
            'campania': variable.campania,
            'meses': {mes: 0 for mes in meses},
            'total_base': 0,
        }
        claves_base_normalizadas.add(clave_normalizada)
    variables_por_clave = {
        (normalizar_header(variable.cliente), normalizar_nombre_variable(variable.campania), variable.mes): variable
        for variable in variables
    }

    filas = []
    total_general = {mes: 0 for mes in meses}
    for clave in sorted(base_por_clave):
        base = base_por_clave[clave]
        porcentajes = {}
        importes = {}
        total = 0
        for mes in meses:
            clave_variable = (normalizar_header(base['cliente']), normalizar_nombre_variable(base['campania']), mes)
            variable = variables_por_clave.get(clave_variable)
            heredado = False
            if not variable:
                numero_mes = mes[-2:]
                candidatas = [
                    item for item in variables
                    if normalizar_header(item.cliente) == normalizar_header(base['cliente'])
                    and normalizar_nombre_variable(item.campania) == normalizar_nombre_variable(base['campania'])
                    and item.year < year and item.mes.endswith(f'-{numero_mes}')
                ]
                variable = max(candidatas, key=lambda item: item.year, default=None)
                heredado = bool(variable)
            porcentaje = variable.porcentaje if variable else 0
            importe = redondear_moneda(base['meses'][mes] * (porcentaje / 100))
            porcentajes[mes] = round(porcentaje or 0, 2)
            importes[mes] = importe
            total += importe
            total_general[mes] += importe
        filas.append({
            'site': base['site'],
            'cliente': base['cliente'],
            'campania': base['campania'],
            'base_meses': {mes: round(base['meses'][mes], 2) for mes in meses},
            'porcentajes': porcentajes,
            'porcentajes_heredados': {
                mes: not bool(variables_por_clave.get((normalizar_header(base['cliente']), normalizar_nombre_variable(base['campania']), mes))) and bool(porcentajes[mes])
                for mes in meses
            },
            'meses': importes,
            'total_base': round(base['total_base'], 2),
            'total': round(total, 2),
        })

    return {
        'success': True,
        'year': year,
        'meses': meses_info,
        'campanias': [
            {
                'site': fila['site'],
                'cliente': fila['cliente'],
                'campania': fila['campania'],
                'label': f"{fila['site']} / {fila['cliente']} / {fila['campania']}",
            }
            for fila in filas
        ],
        'filas': filas,
        'total_general': {
            'meses': {mes: round(total_general.get(mes, 0), 2) for mes in meses},
            'total': round(sum(total_general.values()), 2),
        },
    }


def construir_tarifacion_data(year):
    resumen = construir_resumen_proyeccion_data(year)
    meses_info = resumen['meses']
    meses = [item['value'] for item in meses_info]
    bases = {}
    for fila in resumen['filas']:
        if fila.get('tipo') == 'total_site':
            continue
        if fila.get('concepto') not in ('Horas', 'Fijo mensual'):
            continue
        clave = (normalizar_header(fila.get('cliente')), normalizar_nombre_variable(fila.get('campania')))
        base = bases.setdefault(clave, {
            'site': fila.get('site') or '', 'cliente': fila.get('cliente') or '',
            'campania': fila.get('campania') or '', 'meses': {mes: 0 for mes in meses},
        })
        for mes in meses:
            base['meses'][mes] += numero_seguro(fila.get('meses', {}).get(mes))
    registros = TarifacionCampania.query.filter_by(year=year).all()
    for registro in registros:
        clave = (normalizar_header(registro.cliente), normalizar_nombre_variable(registro.campania))
        bases.setdefault(clave, {
            'site': registro.site or '', 'cliente': registro.cliente,
            'campania': registro.campania, 'meses': {mes: 0 for mes in meses},
        })
    filas = []
    total_ajustado = {mes: 0 for mes in meses}
    for clave, base in sorted(bases.items()):
        propios = [r for r in registros if (normalizar_header(r.cliente), normalizar_nombre_variable(r.campania)) == clave]
        conceptos = sorted({r.concepto for r in propios}) or ['Tarifación']
        for concepto in conceptos:
            montos = {mes: 0 for mes in meses}
            for registro in propios:
                if registro.concepto == concepto and registro.mes in montos:
                    montos[registro.mes] += registro.monto or 0
            ajustados = {mes: redondear_moneda(base['meses'][mes] + montos[mes]) for mes in meses}
            for mes in meses:
                total_ajustado[mes] += ajustados[mes]
            filas.append({
                'site': base['site'], 'cliente': base['cliente'], 'campania': base['campania'],
                'concepto': concepto,
                'facturacion_horas': {mes: round(base['meses'][mes], 2) for mes in meses},
                'montos': {mes: round(montos[mes], 2) for mes in meses},
                'ajustados': ajustados,
                'total_monto': round(sum(montos.values()), 2),
                'total_ajustado': round(sum(ajustados.values()), 2),
            })
    return {'success': True, 'year': year, 'meses': meses_info, 'filas': filas,
            'total_ajustado': {'meses': total_ajustado, 'total': round(sum(total_ajustado.values()), 2)}}


@main_bp.route('/api/tarifaciones', methods=['GET'])
@login_requerido
def api_tarifaciones():
    year = int(request.args.get('year') or datetime.utcnow().year)
    return jsonify(construir_tarifacion_data(year))


@main_bp.route('/api/tarifaciones', methods=['POST'])
@edicion_requerida
def api_guardar_tarifacion():
    data = request.get_json(silent=True) or {}
    mes = mes_valido(data.get('mes'))
    cliente = str(data.get('cliente') or '').strip()
    campania = str(data.get('campania') or '').strip()
    concepto = str(data.get('concepto') or 'Tarifación').strip() or 'Tarifación'
    if not mes or not cliente or not campania:
        return jsonify({'success': False, 'errores': ['Cliente, campaña y mes son obligatorios']}), 400
    try:
        monto = parse_numero(data.get('monto'))
    except ValueError:
        return jsonify({'success': False, 'errores': ['El monto no es válido']}), 400
    registro = TarifacionCampania.query.filter_by(
        cliente=cliente, campania=campania, concepto=concepto, mes=mes,
    ).first()
    antes = registro.to_dict() if registro else None
    if not registro:
        registro = TarifacionCampania(cliente=cliente, campania=campania, concepto=concepto, mes=mes)
        db.session.add(registro)
    registro.site = str(data.get('site') or '').strip()
    registro.year = int(mes[:4])
    registro.monto = monto
    db.session.flush()
    registrar_historial(
        'edicion' if antes else 'creacion', 'tarifaciones', mes,
        f'Tarifación guardada: {cliente} / {campania} / {concepto} / {mes}',
        antes=antes, despues=registro.to_dict(),
    )
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Tarifación guardada y aplicada en Facturación horas'})


@main_bp.route('/api/tarifaciones/edicion-masiva', methods=['POST'])
@edicion_requerida
def api_edicion_masiva_tarifaciones():
    """Guarda una grilla de tarifaciones en una sola transacción."""
    data = request.get_json(silent=True) or {}
    entradas = data.get('cambios') if isinstance(data.get('cambios'), list) else []
    if not entradas:
        return jsonify({'success': False, 'errores': ['No hay cambios para guardar']}), 400
    if len(entradas) > 1000:
        return jsonify({'success': False, 'errores': ['La edición admite hasta 1000 importes por operación']}), 400

    errores = []
    cambios = []
    claves = set()
    for numero, entrada in enumerate(entradas, start=1):
        cliente = str(entrada.get('cliente') or '').strip()
        campania = str(entrada.get('campania') or '').strip()
        concepto = str(entrada.get('concepto') or 'Tarifación').strip() or 'Tarifación'
        site = str(entrada.get('site') or '').strip()
        mes = mes_valido(entrada.get('mes'))
        try:
            monto = parse_numero(entrada.get('monto'))
        except (TypeError, ValueError):
            monto = None
        clave = (cliente, campania, concepto, mes)
        if not cliente or not campania or not mes:
            errores.append(f'Celda {numero}: cliente, campaña y mes son obligatorios')
        elif monto is None:
            errores.append(f'Celda {numero}: el monto no es válido')
        elif clave in claves:
            errores.append(f'Celda {numero}: el importe está repetido')
        else:
            claves.add(clave)
            cambios.append((site, cliente, campania, concepto, mes, monto))
    if errores:
        return jsonify({'success': False, 'errores': errores[:30]}), 400

    antes = []
    despues = []
    for site, cliente, campania, concepto, mes, monto in cambios:
        registro = TarifacionCampania.query.filter_by(
            cliente=cliente, campania=campania, concepto=concepto, mes=mes,
        ).first()
        snapshot = registro.to_dict() if registro else None
        antes.append(snapshot)
        if monto == 0:
            if registro:
                db.session.delete(registro)
            despues.append(None)
            continue
        if not registro:
            registro = TarifacionCampania(
                cliente=cliente, campania=campania, concepto=concepto, mes=mes,
            )
            db.session.add(registro)
        registro.site = site
        registro.year = int(mes[:4])
        registro.monto = monto
        db.session.flush()
        despues.append(registro.to_dict())

    registrar_historial(
        'edicion', 'tarifaciones', str(data.get('year') or ''),
        f'Edición en grilla de tarifaciones: {len(cambios)} importe(s)',
        antes={'tarifaciones': antes}, despues={'tarifaciones': despues},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'{len(cambios)} importe(s) guardado(s) y aplicado(s) en Facturación horas',
    })


@main_bp.route('/api/tarifaciones/template', methods=['GET'])
@login_requerido
def api_template_tarifaciones():
    year = int(request.args.get('year') or datetime.utcnow().year)
    data = construir_tarifacion_data(year)
    headers = ['Cuenta', 'Site', 'Cliente', 'Concepto'] + [mes['label'] for mes in data['meses']]
    rows = [[fila['cliente'], fila['site'], fila['campania'], fila['concepto'],
             *[fila['montos'][mes['value']] for mes in data['meses']]] for fila in data['filas']]
    return Response(crear_xlsx(headers, rows),
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="tarifaciones_{year}.xlsx"'})


@main_bp.route('/api/tarifaciones/importar', methods=['POST'])
@carga_requerida
def api_importar_tarifaciones():
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename:
        return jsonify({'success': False, 'errores': ['Seleccione el archivo de tarifación']}), 400
    year = int(request.form.get('year') or datetime.utcnow().year)
    try:
        contenido = archivo.read()
        filas = filas_desde_xlsx(contenido) if archivo.filename.lower().endswith('.xlsx') else filas_desde_csv(decodificar_texto_importacion(contenido))
    except Exception:
        return jsonify({'success': False, 'errores': ['No se pudo leer el archivo de tarifación']}), 400
    if len(filas) < 2:
        return jsonify({'success': False, 'errores': ['El archivo no contiene datos']}), 400
    headers = [normalizar_header(x) for x in filas[0]]
    idx_cuenta = next((i for i, h in enumerate(headers) if h == 'cuenta'), None)
    idx_site = next((i for i, h in enumerate(headers) if h == 'site'), None)
    idx_cliente = next((i for i, h in enumerate(headers) if h in ('cliente', 'campana', 'campania')), None)
    idx_concepto = next((i for i, h in enumerate(headers) if h == 'concepto'), None)
    columnas = {i: mes_plp_importado(v, year) for i, v in enumerate(filas[0])}
    columnas = {i: mes for i, mes in columnas.items() if mes}
    if idx_cliente is None or not columnas:
        return jsonify({'success': False, 'errores': ['Se esperan Cliente y columnas de enero a diciembre']}), 400
    guardados = []
    antes = []
    for fila in filas[1:]:
        celda = lambda i: fila[i] if i is not None and i < len(fila) else ''
        campania = str(celda(idx_cliente) or '').strip()
        if not campania:
            continue
        cliente = str(celda(idx_cuenta) or campania).strip()
        site = str(celda(idx_site) or '').strip()
        concepto = str(celda(idx_concepto) or 'Tarifación').strip() or 'Tarifación'
        for indice, mes in columnas.items():
            valor = celda(indice)
            if valor in (None, ''):
                continue
            try:
                monto = parse_numero(valor)
            except ValueError:
                continue
            registro = TarifacionCampania.query.filter_by(cliente=cliente, campania=campania, concepto=concepto, mes=mes).first()
            antes.append(registro.to_dict() if registro else None)
            if not registro:
                registro = TarifacionCampania(cliente=cliente, campania=campania, concepto=concepto, mes=mes)
                db.session.add(registro)
            registro.site, registro.year, registro.monto = site, year, monto
            guardados.append(registro)
    if not guardados:
        return jsonify({'success': False, 'errores': ['No se encontraron importes para guardar']}), 400
    db.session.flush()
    registrar_historial(
        'importacion', 'tarifaciones', str(year),
        f'Importación de tarifaciones {year}: {len(guardados)} importe(s)',
        antes={'tarifaciones': antes},
        despues={'tarifaciones': [item.to_dict() for item in guardados]},
    )
    db.session.commit()
    return jsonify({'success': True, 'mensaje': f'{len(guardados)} importe(s) de tarifación importado(s) para {year}'})


@main_bp.route('/api/variables', methods=['GET'])
@login_requerido
def api_variables():
    year = int(request.args.get('year') or datetime.utcnow().year)
    return jsonify(construir_variable_data(year))


def construir_next_gen_data(year):
    meses_info = opciones_meses_proyeccion(year)
    meses = [item['value'] for item in meses_info]
    dolares = {mes: 0 for mes in meses}
    for registro in NextGenDolar.query.filter_by(year=year).all():
        if registro.mes in dolares:
            dolares[registro.mes] = round(registro.valor or 0, 4)
    registros = NextGenProducto.query.filter_by(year=year).all()
    grupos = {}
    for registro in registros:
        site_registro = registro.site or site_original_next_gen(registro.cliente, registro.campania)
        clave = (site_registro, registro.cliente, registro.campania, registro.producto)
        grupo = grupos.setdefault(clave, {
            'site': site_registro, 'cliente': registro.cliente,
            'campania': registro.campania, 'producto': registro.producto,
            'usd': {mes: 0 for mes in meses}, 'pesos': {mes: 0 for mes in meses},
        })
        if registro.mes in grupo['usd']:
            grupo['usd'][registro.mes] += registro.cantidad_usd or 0
    filas = []
    total_pesos = {mes: 0 for mes in meses}
    total_usd = {mes: 0 for mes in meses}
    for grupo in grupos.values():
        for mes in meses:
            grupo['pesos'][mes] = redondear_moneda(grupo['usd'][mes] * dolares[mes])
            total_pesos[mes] += grupo['pesos'][mes]
            total_usd[mes] += grupo['usd'][mes]
        grupo['total_usd'] = round(sum(grupo['usd'].values()), 2)
        grupo['total_pesos'] = round(sum(grupo['pesos'].values()), 2)
        filas.append(grupo)
    filas.sort(key=lambda item: (normalizar_header(item['cliente']), normalizar_header(item['campania']), normalizar_header(item['producto'])))
    return {
        'success': True, 'year': year, 'meses': meses_info, 'dolares': dolares,
        'filas': filas,
        'total_usd': {'meses': {mes: round(valor, 2) for mes, valor in total_usd.items()}, 'total': round(sum(total_usd.values()), 2)},
        'total_pesos': {'meses': {mes: round(valor, 2) for mes, valor in total_pesos.items()}, 'total': round(sum(total_pesos.values()), 2)},
    }


@main_bp.route('/api/next-gen', methods=['GET', 'POST'])
@edicion_si_mutacion_requerida
def api_next_gen():
    if request.method == 'GET':
        year = int(request.args.get('year') or datetime.utcnow().year)
        return jsonify(construir_next_gen_data(year))
    data = request.get_json(silent=True) or {}
    mes = mes_valido(data.get('mes'))
    if not mes:
        return jsonify({'success': False, 'errores': ['El mes no es válido']}), 400
    if data.get('tipo') == 'dolar':
        try:
            valor = parse_numero(data.get('valor'))
        except ValueError:
            return jsonify({'success': False, 'errores': ['El valor del dólar no es válido']}), 400
        registro = NextGenDolar.query.filter_by(mes=mes).first()
        antes = registro.to_dict() if registro else None
        if not registro:
            registro = NextGenDolar(year=int(mes[:4]), mes=mes)
            db.session.add(registro)
        registro.valor = valor
        db.session.flush()
        despues = {**registro.to_dict(), 'tipo_registro': 'dolar'}
        if antes:
            antes = {**antes, 'tipo_registro': 'dolar'}
        registrar_historial('edicion' if antes else 'creacion', 'next_gen', mes,
                            f'Next Gen > Dólar actualizado: {mes} = {valor}', antes=antes, despues=despues)
    else:
        cliente = str(data.get('cliente') or '').strip()
        campania = str(data.get('campania') or '').strip()
        producto = str(data.get('producto') or '').strip()
        if not cliente or not campania or not producto:
            return jsonify({'success': False, 'errores': ['Cliente, campaña y producto son obligatorios']}), 400
        try:
            cantidad = parse_numero(data.get('cantidad_usd'))
        except ValueError:
            return jsonify({'success': False, 'errores': ['La cantidad en USD no es válida']}), 400
        registro = NextGenProducto.query.filter_by(cliente=cliente, campania=campania, producto=producto, mes=mes).first()
        antes = registro.to_dict() if registro else None
        if not registro:
            registro = NextGenProducto(cliente=cliente, campania=campania, producto=producto, mes=mes, year=int(mes[:4]))
            db.session.add(registro)
        registro.site = str(data.get('site') or '').strip()
        registro.cantidad_usd = cantidad
        db.session.flush()
        despues = {**registro.to_dict(), 'tipo_registro': 'producto'}
        if antes:
            antes = {**antes, 'tipo_registro': 'producto'}
        registrar_historial('edicion' if antes else 'creacion', 'next_gen', mes,
                            f'Next Gen > Producto: {cliente} / {campania} / {producto} / {mes}', antes=antes, despues=despues)
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Next Gen guardado y aplicado en Facturación horas'})


@main_bp.route('/api/next-gen/template', methods=['GET'])
@login_requerido
def api_template_next_gen():
    year = int(request.args.get('year') or datetime.utcnow().year)
    data = construir_next_gen_data(year)
    headers = ['CLIENTE'] + [mes['label'] for mes in data['meses']]
    rows = [[fila['producto'], *[fila['usd'][mes['value']] for mes in data['meses']]] for fila in data['filas']]
    if not rows:
        rows = [['APM - TERMINAL 4', *([0] * 12)]]
    return Response(crear_xlsx(headers, rows),
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="next_gen_{year}.xlsx"'})


@main_bp.route('/api/next-gen/importar', methods=['POST'])
@carga_requerida
def api_importar_next_gen():
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename:
        return jsonify({'success': False, 'errores': ['Seleccione el archivo de Next Gen']}), 400
    year = int(request.form.get('year') or datetime.utcnow().year)
    try:
        contenido = archivo.read()
        if archivo.filename.lower().endswith('.xlsx'):
            filas = filas_desde_xlsx(contenido)
        elif archivo.filename.lower().endswith('.csv'):
            filas = filas_desde_csv(decodificar_texto_importacion(contenido))
        else:
            return jsonify({'success': False, 'errores': ['Use un archivo .xlsx o .csv']}), 400
    except Exception:
        return jsonify({'success': False, 'errores': ['No se pudo leer el archivo de Next Gen']}), 400
    if len(filas) < 2:
        return jsonify({'success': False, 'errores': ['El archivo no contiene datos']}), 400
    headers = [normalizar_header(x) for x in filas[0]]
    idx_cuenta = next((i for i, h in enumerate(headers) if h == 'cuenta'), None)
    idx_site = next((i for i, h in enumerate(headers) if h == 'site'), None)
    idx_cliente_simple = next((i for i, h in enumerate(headers) if h == 'cliente'), None)
    idx_campania = next((i for i, h in enumerate(headers) if h in ('campana', 'campania')), None)
    idx_producto = next((i for i, h in enumerate(headers) if h == 'producto'), None)
    if idx_producto is None:
        idx_producto = idx_cliente_simple
    columnas = {i: mes_plp_importado(v, year) for i, v in enumerate(filas[0])}
    columnas = {i: mes for i, mes in columnas.items() if mes}
    if idx_producto is None or not columnas:
        return jsonify({'success': False, 'errores': ['Se esperan CLIENTE y los doce meses']}), 400
    aliases = {normalizar_header(clave): valor for clave, valor in {
        'APM - TERMINAL 4': 'APM - TERMINAL 4', 'BAPRO': 'BAPRO',
        'BCO. CREDICOOP': 'Banco Credicoop', 'LIFE': 'LIFE',
        'CONSUMER SANTANDER': 'Santander Consumer', 'J&J': 'Johnson & Johnson',
        'LIBRA': 'LIBRA', 'LEIVA JOYAS': 'Leiva Joyas', 'M & W': 'M & W',
        'MAR SALVAJE': 'MAR SALVAJE', 'MIRGOR - SC': 'Mirgor',
        'MIRGOR - TRADE IN': 'Mirgor BBTI', 'OPEN PAY': 'OpenPay',
        'SPV SEG.': 'Supervielle Seguros', 'SHIP NOW': 'SHIP NOW',
        'YAMAHA': 'Yamaha', 'DASC - ADVAL': 'DASC - ADVAL',
        'BSF CARREFOUR': 'Carrefour', 'BIBANK': 'Bi Bank', 'ALMUNDO': 'Almundo',
        'ANDREANI': 'ANDREANI',
    }.items()}
    campanias_base = construir_variable_data(year).get('campanias', [])
    antes, guardados = [], []
    for fila in filas[1:]:
        celda = lambda i: fila[i] if i is not None and i < len(fila) else ''
        producto_nombre = str(celda(idx_producto) or '').strip()
        if not producto_nombre:
            continue
        es_dolar = normalizar_header(producto_nombre) in ('valor dolar', 'dolar')
        existente = next((r for r in NextGenProducto.query.filter_by(year=year).all()
                          if normalizar_header(r.producto) == normalizar_header(producto_nombre)), None)
        cuenta_archivo = str(celda(idx_cuenta) or '').strip() if idx_cuenta is not None else ''
        campania_archivo = str(celda(idx_campania) or '').strip() if idx_campania is not None else ''
        canonica = aliases.get(normalizar_header(producto_nombre), campania_archivo or producto_nombre)
        base = next((item for item in campanias_base
                     if normalizar_nombre_variable(item['campania']) == normalizar_nombre_variable(canonica)), None)
        cliente = existente.cliente if existente else (cuenta_archivo or (base['cliente'] if base else canonica))
        campania = existente.campania if existente else (campania_archivo or (base['campania'] if base else canonica))
        site = existente.site if existente else (str(celda(idx_site) or '').strip() if idx_site is not None else (base['site'] if base else ''))
        for indice, mes in columnas.items():
            valor = celda(indice)
            if valor in (None, ''):
                continue
            try:
                numero = parse_numero(valor)
            except ValueError:
                continue
            if es_dolar:
                registro = NextGenDolar.query.filter_by(mes=mes).first()
                previo = registro.to_dict() if registro else None
                if not registro:
                    registro = NextGenDolar(year=year, mes=mes)
                    db.session.add(registro)
                registro.valor = numero
                antes.append({**previo, 'tipo_registro': 'dolar'} if previo else None)
                guardados.append(registro)
            else:
                registro = NextGenProducto.query.filter_by(cliente=cliente, campania=campania, producto=producto_nombre, mes=mes).first()
                previo = registro.to_dict() if registro else None
                if not registro:
                    registro = NextGenProducto(cliente=cliente, campania=campania, producto=producto_nombre, year=year, mes=mes)
                    db.session.add(registro)
                registro.site = site
                registro.cantidad_usd = numero
                antes.append({**previo, 'tipo_registro': 'producto'} if previo else None)
                guardados.append(registro)
    if not guardados:
        return jsonify({'success': False, 'errores': ['No se encontraron valores para importar']}), 400
    db.session.flush()
    despues = [{**item.to_dict(), 'tipo_registro': 'dolar' if isinstance(item, NextGenDolar) else 'producto'} for item in guardados]
    registrar_historial('importacion', 'next_gen', str(year),
                        f'Next Gen > Importación {year}: {len(guardados)} valor(es)',
                        antes={'next_gen': antes}, despues={'next_gen': despues})
    db.session.commit()
    return jsonify({'success': True, 'mensaje': f'{len(guardados)} valor(es) de Next Gen importado(s) para {year}'})


@main_bp.route('/api/variables/template', methods=['GET'])
@login_requerido
def api_template_variables():
    year = int(request.args.get('year') or datetime.utcnow().year)
    data = construir_variable_data(year)
    meses = data['meses']
    headers = ['Cuenta', 'Site', 'Cliente'] + [mes['label'] for mes in meses]
    rows = [
        [
            fila['cliente'],
            fila['site'],
            fila['campania'],
            *[fila.get('porcentajes', {}).get(mes['value'], 0) for mes in meses],
        ]
        for fila in data['filas']
    ]
    contenido = crear_xlsx(headers, rows)
    return Response(
        contenido,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="variables_{year}.xlsx"'},
    )


@main_bp.route('/api/variables/importar', methods=['POST'])
@carga_requerida
def api_importar_variables():
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename:
        return jsonify({'success': False, 'errores': ['Seleccione el archivo de variables']}), 400
    try:
        year = int(request.form.get('year') or datetime.utcnow().year)
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['El año no es válido']}), 400
    try:
        contenido = archivo.read()
        if archivo.filename.lower().endswith('.xlsx'):
            filas_archivo = filas_desde_xlsx(contenido)
        elif archivo.filename.lower().endswith('.csv'):
            filas_archivo = filas_desde_csv(decodificar_texto_importacion(contenido))
        else:
            return jsonify({'success': False, 'errores': ['Use un archivo .xlsx o .csv']}), 400
    except Exception:
        return jsonify({'success': False, 'errores': ['No se pudo leer el archivo de variables']}), 400
    if len(filas_archivo) < 2:
        return jsonify({'success': False, 'errores': ['El archivo no contiene datos']}), 400

    headers = [normalizar_header(valor) for valor in filas_archivo[0]]
    indice_cuenta = next((i for i, h in enumerate(headers) if h in ('cuenta', 'cliente')), None)
    indice_site = next((i for i, h in enumerate(headers) if h == 'site'), None)
    indice_nombre = next((i for i, h in enumerate(headers) if h in ('campana', 'campania')), None)
    if indice_nombre is None:
        indices_cliente = [i for i, h in enumerate(headers) if h == 'cliente']
        indice_nombre = indices_cliente[-1] if indices_cliente else None
    columnas_meses = {}
    for indice, header in enumerate(filas_archivo[0]):
        mes = mes_plp_importado(header, year)
        if mes:
            columnas_meses[indice] = mes
    if indice_nombre is None or not columnas_meses:
        return jsonify({'success': False, 'errores': ['Se esperan Cuenta/Cliente, Campaña o Cliente identificador y columnas mensuales']}), 400

    base = construir_variable_data(year)
    campanias_base = base['campanias']
    antes = []
    guardadas = []
    omitidas = []
    for numero_fila, fila in enumerate(filas_archivo[1:], start=2):
        if not any(str(valor or '').strip() for valor in fila):
            continue
        celda = lambda indice: fila[indice] if indice is not None and indice < len(fila) else ''
        nombre = str(celda(indice_nombre) or '').strip()
        cuenta = str(celda(indice_cuenta) or '').strip()
        if not nombre:
            omitidas.append(f'Fila {numero_fila}: falta el identificador de campaña')
            continue
        nombre_normalizado = normalizar_nombre_variable(nombre)
        candidatas = [item for item in campanias_base if normalizar_nombre_variable(item['campania']) == nombre_normalizado]
        if cuenta:
            por_cuenta = [item for item in candidatas if normalizar_header(item['cliente']) == normalizar_header(cuenta)]
            candidatas = por_cuenta or candidatas
        if len(candidatas) == 1:
            campania = candidatas[0]
        else:
            campania = {
                'site': str(celda(indice_site) or '').strip(),
                'cliente': cuenta or nombre,
                'campania': nombre,
            }
        for indice, mes in columnas_meses.items():
            valor = celda(indice)
            if valor in (None, ''):
                continue
            try:
                porcentaje = parse_numero(valor)
                if 0 < abs(porcentaje) <= 1:
                    porcentaje *= 100
            except ValueError:
                omitidas.append(f'Fila {numero_fila}, {mes}: porcentaje inválido')
                continue
            variable = VariableCampania.query.filter_by(
                cliente=campania['cliente'], campania=campania['campania'], mes=mes,
            ).first()
            antes.append(variable.to_dict() if variable else None)
            if not variable:
                variable = VariableCampania(cliente=campania['cliente'], campania=campania['campania'], mes=mes)
                db.session.add(variable)
            variable.site = campania['site']
            variable.year = year
            variable.porcentaje = porcentaje
            guardadas.append(variable)
    if not guardadas:
        return jsonify({'success': False, 'errores': omitidas[:30] or ['No hubo campañas coincidentes con Facturación horas']}), 400
    db.session.flush()
    registrar_historial(
        'importacion', 'variables', str(year),
        f'Importación de variables {year}: {len(guardadas)} porcentaje(s)',
        antes={'variables': antes},
        despues={'variables': [item.to_dict() for item in guardadas]},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'{len(guardadas)} porcentaje(s) importado(s) para {year}',
        'omitidas': omitidas[:30],
    })


@main_bp.route('/api/control-proyecciones', methods=['GET'])
@login_requerido
def api_control_proyecciones():
    year = int(request.args.get('year') or datetime.utcnow().year)
    meses_info = opciones_meses_proyeccion(year)
    data = construir_control_proyecciones_data(year)
    data.update({
        'success': True,
        'year': year,
        'meses': meses_info,
    })
    return jsonify(data)


def cargar_base_directorio(nombre):
    ruta = os.path.join(current_app.root_path, 'data', nombre)
    with open(ruta, encoding='utf-8') as archivo:
        return json.load(archivo)


def asegurar_graficos_dotaciones_iniciales():
    if GraficoDotacionMensual.query.first():
        return
    for year_texto, series in cargar_base_directorio('dotaciones_resumen_historico.json')['graficos'].items():
        year = int(year_texto)
        for indice, requerida in enumerate(series['requerida'], 1):
            db.session.add(GraficoDotacionMensual(
                fecha=date(year, indice, 1), dotacion_requerida=requerida,
                activa_sl=series['sl'][indice - 1], activa_ba=series['ba'][indice - 1],
            ))
    db.session.commit()


def asegurar_dotaciones_iniciales():
    if DotacionMensual.query.first():
        return
    for mes, requerida, personal in cargar_base_directorio('dotaciones_resumen_historico.json')['indicadores']:
        db.session.add(DotacionMensual(
            fecha=datetime.strptime(f'{mes}-01', '%Y-%m-%d').date(),
            dotacion_requerida=requerida,
            personal=personal,
        ))
    db.session.commit()


def asegurar_dotaciones_clientes_iniciales():
    ruta = os.path.join(current_app.root_path, 'dotaciones_clientes_iniciales.json')
    if not os.path.exists(ruta):
        return
    with open(ruta, encoding='utf-8') as archivo:
        registros = json.load(archivo)
    if DotacionClienteMensual.query.count() >= len(registros):
        return
    for cliente, mes, valor in registros:
        fecha = datetime.strptime(f'{mes}-01', '%Y-%m-%d').date()
        if not DotacionClienteMensual.query.filter_by(cliente=cliente, fecha=fecha).first():
            db.session.add(DotacionClienteMensual(cliente=cliente, fecha=fecha, dotacion=valor))
    db.session.commit()


def construir_dotaciones_clientes_data(year=None):
    asegurar_dotaciones_clientes_iniciales()
    registros = DotacionClienteMensual.query.order_by(DotacionClienteMensual.fecha, DotacionClienteMensual.cliente).all()
    anios = sorted({item.fecha.year for item in registros})
    year = year if year in anios else (anios[-1] if anios else datetime.utcnow().year)
    fechas_todas = sorted({item.fecha for item in registros})
    fechas = [fecha for fecha in fechas_todas if fecha.year == year]
    meses = [fecha.strftime('%Y-%m') for fecha in fechas]
    ultimo = fechas[-1] if fechas else None
    diciembre = next((fecha for fecha in reversed(fechas_todas) if fecha.month == 12 and fecha < ultimo), None) if ultimo else None
    por_cliente = {}
    for item in registros:
        por_cliente.setdefault(item.cliente, {})[item.fecha.strftime('%Y-%m')] = float(item.dotacion or 0)
    clientes = []
    mes_ultimo = ultimo.strftime('%Y-%m') if ultimo else None
    mes_diciembre = diciembre.strftime('%Y-%m') if diciembre else None
    for cliente, valores in sorted(por_cliente.items(), key=lambda item: (item[0] != 'Personal', item[0].lower())):
        valores_anio = [valor for mes, valor in valores.items() if mes.startswith(str(year))]
        clientes.append({
            'cliente': cliente,
            'meses': {mes: valor for mes, valor in valores.items() if mes.startswith(str(year))},
            'diferencia': round(valores.get(mes_ultimo, 0) - valores.get(mes_diciembre, 0), 2),
            'promedio': round(sum(valores_anio) / len(valores_anio), 2) if valores_anio else 0,
        })
    totales = {mes: round(sum(por_cliente[cliente].get(mes, 0) for cliente in por_cliente), 2) for mes in meses}
    if mes_diciembre:
        totales[mes_diciembre] = round(sum(valores.get(mes_diciembre, 0) for valores in por_cliente.values()), 2)
    personal_valores = por_cliente.get('Personal', {})
    personal = {mes: valor for mes, valor in personal_valores.items() if mes.startswith(str(year)) or mes == mes_diciembre}
    return {
        'year': year,
        'anios': anios,
        'meses': meses,
        'clientes': clientes,
        'totales': totales,
        'personal': personal,
        'ultimo_mes': mes_ultimo,
        'diciembre_base': mes_diciembre,
    }


@main_bp.route('/api/indicadores', methods=['GET'])
@login_requerido
def api_indicadores():
    asegurar_dotaciones_iniciales()
    registros = DotacionMensual.query.order_by(DotacionMensual.fecha).all()
    anios_dotacion = sorted({item.fecha.year for item in registros})
    year = int(request.args.get('year') or (anios_dotacion[-1] if anios_dotacion else datetime.utcnow().year))
    actuales = [item for item in registros if item.fecha.year == year]
    ultimo = actuales[-1] if actuales else None
    diciembre = next(
        (item for item in reversed(registros) if item.fecha.month == 12 and item.fecha.year < year),
        None,
    )

    def promedio(campo):
        valores = [float(getattr(item, campo) or 0) for item in actuales if float(getattr(item, campo) or 0) > 0]
        return round(sum(valores) / len(valores), 2) if valores else 0

    facturacion = facturacion_por_anio(Facturacion2026.query, year).all()
    requerido = round(sum(float(item.horas_objetivo or 0) for item in facturacion), 2)
    realizado = round(sum(float(item.horas_facturadas or 0) for item in facturacion), 2)
    cumplimiento = realizado / requerido if requerido else 0
    return jsonify({
        'success': True,
        'reloj': {'year': year, 'requerido': requerido, 'realizado': realizado,
                  'diferencia': round(realizado - requerido, 2), 'cumplimiento': round(cumplimiento, 6)},
        'dotaciones': {
            'serie': [item.to_dict() for item in registros],
            'ultimo': ultimo.to_dict() if ultimo else None,
            'diciembre_base': diciembre.to_dict() if diciembre else None,
            'diferencia_requerida': round(float(ultimo.dotacion_requerida) - float(diciembre.dotacion_requerida), 2) if ultimo and diciembre else 0,
            'diferencia_personal': round(float(ultimo.personal) - float(diciembre.personal), 2) if ultimo and diciembre else 0,
            'promedio_requerida': promedio('dotacion_requerida'),
            'promedio_personal': promedio('personal'),
        },
    })


@main_bp.route('/api/dotaciones', methods=['POST'])
@edicion_requerida
def api_guardar_dotacion():
    data = request.get_json(silent=True) or {}
    try:
        fecha = datetime.strptime(str(data.get('mes') or ''), '%Y-%m').date().replace(day=1)
        requerida = parse_numero(data.get('dotacion_requerida'))
        personal = parse_numero(data.get('personal'))
        if requerida <= 0 or personal <= 0:
            raise ValueError
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['Mes, dotación requerida y personal deben ser valores válidos mayores que cero']}), 400
    registro = DotacionMensual.query.filter_by(fecha=fecha).first()
    antes = registro.to_dict() if registro else None
    if not registro:
        registro = DotacionMensual(fecha=fecha)
        db.session.add(registro)
    registro.dotacion_requerida = requerida
    registro.personal = personal
    db.session.flush()
    registrar_historial('edicion' if antes else 'creacion', 'dotaciones_mensuales', fecha.isoformat(),
                        f'Dotaciones guardadas: {fecha:%m/%Y}', antes=antes, despues=registro.to_dict())
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Mes guardado', 'registro': registro.to_dict()})


@main_bp.route('/api/dotaciones-clientes', methods=['GET'])
@login_requerido
def api_dotaciones_clientes():
    try:
        year = int(request.args.get('year')) if request.args.get('year') else None
    except ValueError:
        year = None
    return jsonify({'success': True, **construir_dotaciones_clientes_data(year)})


@main_bp.route('/api/graficos-evolutivo-dotaciones', methods=['GET'])
@login_requerido
def api_graficos_evolutivo_dotaciones():
    """Replica las fórmulas de la hoja: total por mes y fila Personal."""
    asegurar_dotaciones_clientes_iniciales()
    registros = DotacionClienteMensual.query.order_by(DotacionClienteMensual.fecha).all()
    por_mes = {}
    for registro in registros:
        if registro.fecha < date(2023, 1, 1):
            continue
        clave = registro.fecha.strftime('%Y-%m')
        item = por_mes.setdefault(clave, {
            'fecha': registro.fecha.isoformat(), 'dotacion_requerida': 0.0, 'personal': 0.0,
        })
        valor = float(registro.dotacion or 0)
        item['dotacion_requerida'] += valor
        if normalizar_header(registro.cliente) == 'personal':
            item['personal'] += valor
    serie = []
    for clave in sorted(por_mes):
        item = por_mes[clave]
        item['dotacion_requerida'] = round(item['dotacion_requerida'], 2)
        item['personal'] = round(item['personal'], 2)
        serie.append(item)
    return jsonify({'success': True, 'serie': serie, 'origen': 'dotaciones_clientes_mensuales'})


@main_bp.route('/api/dotaciones-clientes', methods=['POST'])
@edicion_requerida
def api_guardar_dotaciones_clientes():
    data = request.get_json(silent=True) or {}
    try:
        fecha = datetime.strptime(str(data.get('mes') or ''), '%Y-%m').date().replace(day=1)
    except ValueError:
        return jsonify({'success': False, 'errores': ['El mes no es válido']}), 400
    filas = data.get('clientes') or []
    errores = []
    normalizadas = []
    vistos = set()
    for index, fila in enumerate(filas, 1):
        cliente = str(fila.get('cliente') or '').strip()
        try:
            valor = parse_numero(fila.get('dotacion'))
        except ValueError:
            errores.append(f'La dotación de la fila {index} no es válida')
            continue
        clave = normalizar_header(cliente)
        if not cliente:
            errores.append(f'El cliente de la fila {index} es obligatorio')
        elif clave in vistos:
            errores.append(f'El cliente {cliente} está repetido')
        elif valor < 0:
            errores.append(f'La dotación de {cliente} no puede ser negativa')
        else:
            vistos.add(clave)
            normalizadas.append((cliente, valor))
    if errores or not normalizadas:
        return jsonify({'success': False, 'errores': errores or ['Debe cargar al menos un cliente']}), 400
    antes = [item.to_dict() for item in DotacionClienteMensual.query.filter_by(fecha=fecha).all()]
    for cliente, valor in normalizadas:
        registro = DotacionClienteMensual.query.filter_by(cliente=cliente, fecha=fecha).first()
        if not registro:
            registro = DotacionClienteMensual(cliente=cliente, fecha=fecha)
            db.session.add(registro)
        registro.dotacion = valor
    db.session.flush()
    despues = [item.to_dict() for item in DotacionClienteMensual.query.filter_by(fecha=fecha).all()]
    registrar_historial('edicion' if antes else 'creacion', 'dotaciones_clientes_mensuales', fecha.isoformat(),
                        f'Dotaciones por cliente guardadas: {fecha:%m/%Y}', antes={'filas': antes}, despues={'filas': despues})
    db.session.commit()
    return jsonify({'success': True, 'mensaje': f'{len(normalizadas)} clientes guardados para {fecha:%m/%Y}'})


@main_bp.route('/api/graficos-dotaciones', methods=['GET'])
@login_requerido
def api_graficos_dotaciones():
    asegurar_graficos_dotaciones_iniciales()
    registros = GraficoDotacionMensual.query.order_by(GraficoDotacionMensual.fecha).all()
    return jsonify({'success': True, 'serie': [item.to_dict() for item in registros]})


@main_bp.route('/api/graficos-dotaciones', methods=['POST'])
@edicion_requerida
def api_guardar_grafico_dotacion():
    data = request.get_json(silent=True) or {}
    try:
        fecha = datetime.strptime(str(data.get('mes') or ''), '%Y-%m').date().replace(day=1)
        requerida = parse_numero(data.get('dotacion_requerida'))
        activa_sl = parse_numero(data.get('activa_sl'))
        activa_ba = parse_numero(data.get('activa_ba'))
        if min(requerida, activa_sl, activa_ba) < 0 or requerida == 0:
            raise ValueError
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['Mes y dotaciones deben ser valores válidos; la requerida debe ser mayor que cero']}), 400
    registro = GraficoDotacionMensual.query.filter_by(fecha=fecha).first()
    antes = registro.to_dict() if registro else None
    if not registro:
        registro = GraficoDotacionMensual(fecha=fecha)
        db.session.add(registro)
    registro.dotacion_requerida = requerida
    registro.activa_sl = activa_sl
    registro.activa_ba = activa_ba
    db.session.flush()
    despues = registro.to_dict()
    registrar_historial('edicion' if antes else 'creacion', 'graficos_dotaciones_mensuales', fecha.isoformat(),
                        f'Gráfico de dotaciones guardado: {fecha:%m/%Y}',
                        antes={'filas': [antes] if antes else []}, despues={'filas': [despues]})
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Mes guardado', 'registro': despues})


def cliente_variacion_horas(nombre, campania=None):
    """Agrupa las aperturas operativas como la hoja Variación horas ctes."""
    normalizado = normalizar_header(nombre)
    campania_normalizada = normalizar_header(campania)
    if normalizado == 'supervielle' and 'seguro' in campania_normalizada:
        return 'Supervielle Seguros'
    if normalizado.startswith('personal'):
        return 'Personal'
    if normalizado.startswith('santander'):
        return 'Santander'
    if normalizado.startswith('naturgy'):
        return 'Naturgy'
    if normalizado.startswith('bbva'):
        return 'BBVA'
    if normalizado.startswith('mirgor'):
        return 'Mirgor'
    if normalizado.startswith('unicef'):
        return 'Unicef'
    if normalizado in ('river', 'river plate'):
        return 'River Plate'
    if normalizado in ('carrefour', 'bsf carrefour'):
        return 'Bsf Carrefour'
    if normalizado in ('galicia', 'galicia seguros'):
        return 'Galicia'
    nombres = {
        'metrogas': 'Metrogas', 'omint': 'Omint', 'supervielle seguros': 'Supervielle Seguros',
        'assurant': 'Assurant', 'qualia': 'Qualia', 'leiva joyas': 'Leiva Joyas',
        'banco piano': 'Banco Piano', 'openpay': 'Openpay', 'gire': 'Gire',
        'supervielle': 'Supervielle', 'spazios': 'Spazios', 'farmapay': 'Farmapay',
        'bice': 'Bice', 'yamaha': 'Yamaha', 'bibank': 'BiBank', 'bi bank': 'BiBank',
        'johnson & johnson': 'Johnson & Johnson', 'cachamai': 'Cachamai',
        'banco comafi': 'Banco Comafi', 'bna': 'BNA', 'autotekne': 'Autotekne',
        'invertir online': 'Invertir Online', 'shipnow': 'Shipnow',
    }
    return nombres.get(normalizado, str(nombre or '').strip())


def base_horas_clientes(anio):
    filas = facturacion_por_anio(Facturacion2026.query.with_entities(
        Facturacion2026.cliente, func.sum(Facturacion2026.horas_facturadas)
    ), anio).group_by(Facturacion2026.cliente).all()
    if filas:
        resultado = {}
        for cliente, horas in filas:
            clave = cliente_variacion_horas(cliente)
            resultado[clave] = resultado.get(clave, 0) + float(horas or 0)
        return resultado
    resultado = {}
    for registro in HistoricoClienteMensual.query.filter_by(year=anio).all():
        clave = cliente_variacion_horas(registro.cliente)
        resultado[clave] = resultado.get(clave, 0) + float(registro.base_dict().get('horas_realizadas') or 0)
    return resultado


@main_bp.route('/api/variacion-horas-clientes', methods=['GET'])
@login_requerido
def api_variacion_horas_clientes():
    year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    anterior = year - 1
    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    reales_por_mes = dict(facturacion_por_anio(Facturacion2026.query.with_entities(
        Facturacion2026.mes, func.sum(Facturacion2026.horas_facturadas)
    ), year).group_by(Facturacion2026.mes).all())
    corte = max((mes for mes in meses if mes in reales_por_mes), default=None)

    reales = {}
    filas_reales = facturacion_por_anio(Facturacion2026.query.with_entities(
        Facturacion2026.cliente, func.sum(Facturacion2026.horas_facturadas)
    ), year).group_by(Facturacion2026.cliente).all()
    for cliente, horas in filas_reales:
        clave = cliente_variacion_horas(cliente)
        reales[clave] = reales.get(clave, 0) + float(horas or 0)

    meses_proyectados = [mes for mes in meses if not corte or mes > corte]
    proyectadas = {}
    filas_proyectadas = ProyeccionMatriz.query.with_entities(
        ProyeccionMatriz.cliente, ProyeccionMatriz.campania,
        func.sum(ProyeccionMatriz.horas_requeridas * ProyeccionMatriz.porcentaje_cumplimiento / 100),
    ).filter(
        ProyeccionMatriz.mes.in_(meses_proyectados),
        or_(ProyeccionMatriz.tipo_plp.is_(None), ProyeccionMatriz.tipo_plp == ''),
    ).group_by(ProyeccionMatriz.cliente, ProyeccionMatriz.campania).all()
    for cliente, campania, horas in filas_proyectadas:
        clave = cliente_variacion_horas(cliente, campania)
        proyectadas[clave] = proyectadas.get(clave, 0) + float(horas or 0)

    base = base_horas_clientes(anterior)
    orden = list(base)
    for cliente in list(reales) + list(proyectadas):
        if cliente and cliente not in orden:
            orden.append(cliente)
    filas = []
    for cliente in orden:
        hs_anterior = float(base.get(cliente) or 0)
        real = float(reales.get(cliente) or 0)
        proyeccion = float(proyectadas.get(cliente) or 0)
        mixto = real + proyeccion
        # La hoja de Excel omite las cuentas Next Gen sin horas. Se conservan
        # los clientes históricos que caen a cero, pero no filas 0 / 0 / 0.
        if abs(hs_anterior) <= .01 and abs(real) <= .01 and abs(proyeccion) <= .01:
            continue
        diferencia = mixto - hs_anterior
        control = float(proyectadas.get(cliente) or 0)
        filas.append({
            'cliente': cliente, 'anterior': round(hs_anterior, 2), 'real': round(real, 2),
            'proyectado': round(proyeccion, 2), 'mixto': round(mixto, 2),
            'diferencia': round(diferencia, 2),
            'variacion': round(diferencia / hs_anterior, 6) if hs_anterior else None,
            'control': round(control, 2), 'desvio_control': round(control - proyeccion, 2),
        })

    def totales(campo):
        return round(sum(float(fila.get(campo) or 0) for fila in filas), 2)
    total = {campo: totales(campo) for campo in ('anterior', 'real', 'proyectado', 'mixto', 'diferencia', 'control', 'desvio_control')}
    total['variacion'] = round(total['diferencia'] / total['anterior'], 6) if total['anterior'] else None
    control_real_fuente = round(sum(float(valor or 0) for valor in reales_por_mes.values()), 2)
    control_proyectado_fuente = round(sum(float(valor or 0) for _, _, valor in filas_proyectadas), 2)
    controles = {
        'real': {'tabla': total['real'], 'fuente': control_real_fuente, 'diferencia': round(total['real'] - control_real_fuente, 2)},
        'proyectado': {'tabla': total['proyectado'], 'fuente': control_proyectado_fuente, 'diferencia': round(total['proyectado'] - control_proyectado_fuente, 2)},
        'filas_descuadradas': sum(1 for fila in filas if abs(fila['desvio_control']) > .01),
    }
    return jsonify({'success': True, 'year': year, 'anterior': anterior, 'corte': corte,
                    'meses_proyectados': meses_proyectados, 'filas': filas, 'total': total, 'controles': controles})


DASHBOARD_COLUMNAS = [
    'gerente', 'jefe_site', 'conversion_dias', 'cliente_normalizado', 'mes', 'cliente', 'campania',
    'horas_objetivo', 'horas_cumplidas', 'diferencia_horas', 'cumplimiento_horas',
    'abs_sin_vacaciones', 'abs_total', 'ausentismo_total', 'abs_varias', 'ausencias_varias',
    'agentes_e_mas', 'agentes_lcon', 'agentes_ce', 'agentes_off', 'abs_tecnologico',
    'falta_insumos', 'sin_luz', 'sin_internet', 'abs_vacaciones', 'con_vacaciones',
    'logueados', 'agentes_ff', 'facturables_capa', 'facturables_total', 'facturables_os',
    'facturables_tt', 'dotacion_promedio', 'agentes_requeridos', 'agentes_requeridos_ss',
    'agentes_activos', 'over_requerido', 'activos_menos_requeridos', 'rotacion', 'agentes_baja',
    'separador_ap', 'separador_aq', 'aux_activos_horas', 'aux_ausencias_horas',
    'aux_insumos_horas', 'aux_luz_horas', 'aux_internet_horas', 'separador_aw',
    'separador_ax', 'aux_dotacion_promedio', 'aux_agentes_baja', 'separador_ba',
    'aux_agentes_activos', 'aux_ausencias',
]


def valor_dashboard_json(valor):
    if valor is None:
        return None
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    if isinstance(valor, (int, float, str, bool)):
        return valor
    return str(valor)


@main_bp.route('/api/dashboard-operativo/plantilla', methods=['GET'])
@login_requerido
def descargar_plantilla_dashboard_operativo():
    """Genera la plantilla con la misma posición de columnas que la hoja Dashboard."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    encabezados = [
        '', 'Gte', 'Gte', 'Conversion Diras', '', 'Mes', 'Cliente', 'Campaña',
        'Horas objetivo', 'Hs cumplidas a facturar', 'Dif Hs', 'Cumplimiento Hs',
        '%ABS Sin Vacaciones', '%ABS Total', 'Ausentismo Total', '%ABS Varias',
        'Ausencias Varias', 'Agentes E+', 'Agentes LCON', 'Agentes CE', 'Agentes OFF',
        '%ABS Tecnológico', 'Agentes falta de insumos', 'Agentes Sin Luz',
        'Agentes Sin Internet', '%ABS Vacaciones', 'Agentes con Vacaciones',
        'Agentes Logueados', 'Agentes FF', 'Agentes Facturables en Capa',
        'Total Agentes Facturables', 'Agentes Facturables OS', 'Agentes Facturables TT',
        'Dotación Promedio', 'Agentes Requeridos', 'Agentes Requeridos SS',
        'Agentes Activos', '%Over Sobre Requerido', 'Activos - Requeridos SS',
        '%Rotación', 'Agentes Baja', '', '', 'Agentes Activos', 'Ausencias Varias',
        'Agentes falta de insumos', 'Agentes Sin Luz', 'Agentes Sin Internet', '', '',
        'Dotación Promedio', 'Agentes Baja', '', 'Agentes Activos', 'Ausencias',
    ]
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Dashboard'
    hoja['B1'] = 'PLANTILLA DE IMPORTACIÓN - DASHBOARD OPERATIVO'
    hoja.merge_cells(start_row=1, start_column=2, end_row=1, end_column=len(encabezados))
    hoja['B1'].fill = PatternFill('solid', fgColor='1F4E78')
    hoja['B1'].font = Font(color='FFFFFF', bold=True, size=12)
    hoja['B1'].alignment = Alignment(horizontal='center')
    borde = Border(bottom=Side(style='thin', color='FFFFFF'))
    for columna, encabezado in enumerate(encabezados, 1):
        celda = hoja.cell(row=2, column=columna, value=encabezado)
        celda.fill = PatternFill('solid', fgColor='203864')
        celda.font = Font(color='FFFFFF', bold=True)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celda.border = borde
        hoja.column_dimensions[get_column_letter(columna)].width = 18
    hoja.column_dimensions['A'].width = 3
    hoja.column_dimensions['F'].width = 13
    hoja.column_dimensions['G'].width = 28
    hoja.column_dimensions['H'].width = 32
    hoja.freeze_panes = 'I3'
    hoja.auto_filter.ref = f'B2:{get_column_letter(len(encabezados))}2'
    hoja.row_dimensions[1].height = 24
    hoja.row_dimensions[2].height = 42
    hoja.sheet_view.showGridLines = False
    for fila in range(3, 1003):
        hoja.cell(fila, 44, f'=SUMIF($I$2:$AO$2,AR$2,$I{fila}:$AO{fila})*6')
        hoja.cell(fila, 45, f'=SUMIF($I$2:$AO$2,AS$2,$I{fila}:$AO{fila})*6')
        hoja.cell(fila, 46, f'=SUMIF($I$2:$AO$2,AT$2,$I{fila}:$AO{fila})*6')
        hoja.cell(fila, 47, f'=SUMIF($I$2:$AO$2,AU$2,$I{fila}:$AO{fila})*6')
        hoja.cell(fila, 48, f'=SUMIF($I$2:$AO$2,AV$2,$I{fila}:$AO{fila})*6')
        hoja.cell(fila, 51, f'=SUMIF($I$2:$AO$2,AY$2,$I{fila}:$AO{fila})')
        hoja.cell(fila, 52, f'=SUMIF($I$2:$AO$2,AZ$2,$I{fila}:$AO{fila})')
        hoja.cell(fila, 54, f'=SUMIF($I$2:$AO$2,BB$2,$I{fila}:$AO{fila})')
        hoja.cell(fila, 55, f'=SUM(AS{fila}:AV{fila})/6')
    libro.calculation.fullCalcOnLoad = True
    libro.calculation.forceFullCalc = True
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True,
                     download_name='Plantilla Dashboard operativo.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@main_bp.route('/api/dashboard-operativo', methods=['GET'])
@login_requerido
def api_dashboard_operativo():
    year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    mes = str(request.args.get('mes') or '').strip()
    query = DashboardOperativo.query.filter_by(year=year)
    if mes:
        query = query.filter_by(mes=mes)
    filas = query.order_by(DashboardOperativo.mes, DashboardOperativo.cliente, DashboardOperativo.campania).all()
    meses = sorted({fila.mes for fila in DashboardOperativo.query.filter_by(year=year).all()})
    datos = []
    for fila in filas:
        item = fila.to_dict()
        item['datos']['agentes_baja'] = fila.bajas_manual or 0
        def numero_aux(clave):
            try:
                return float(item['datos'].get(clave) or 0)
            except (TypeError, ValueError):
                return 0.0
        item['datos'].update({
            'aux_activos_horas': numero_aux('agentes_activos') * 6,
            'aux_ausencias_horas': numero_aux('ausencias_varias') * 6,
            'aux_insumos_horas': numero_aux('falta_insumos') * 6,
            'aux_luz_horas': numero_aux('sin_luz') * 6,
            'aux_internet_horas': numero_aux('sin_internet') * 6,
            'aux_dotacion_promedio': item['datos'].get('dotacion_promedio'),
            'aux_agentes_baja': fila.bajas_manual or 0,
            'aux_agentes_activos': item['datos'].get('agentes_activos'),
            'aux_ausencias': sum(numero_aux(clave) for clave in
                                 ('ausencias_varias', 'falta_insumos', 'sin_luz', 'sin_internet')),
        })
        datos.append(item)
    return jsonify({'success': True, 'year': year, 'meses': meses, 'columnas': DASHBOARD_COLUMNAS, 'filas': datos})


def importar_historico_desde_libro(libro):
    """Importa literalmente A:U de Historico, agrupado como sus SUMIFS."""
    nombre = next((n for n in libro.sheetnames if normalizar_header(n) == 'historico'), None)
    if not nombre:
        return 0
    grupos = {}
    campos = (
        ('horas_dotacion_activa', 8), ('horas_ausentismo', 9),
        ('dotacion_promedio', 10), ('bajas', 11), ('horas_requeridas', 12),
        ('horas_realizadas', 13), ('pagadas', 14), ('logueo', 15),
        ('dotacion_requerida', 20),
    )
    for row in libro[nombre].iter_rows(min_row=3, max_col=21, values_only=True):
        fecha, cliente = row[5], str(row[3] or '').strip()
        if not isinstance(fecha, (datetime, date)) or not cliente:
            continue
        mes = fecha.strftime('%Y-%m')
        clave = (mes, cliente)
        item = grupos.setdefault(clave, {
            'year': fecha.year, 'empresa': '', 'site': '', 'industria': '',
            **{campo: 0.0 for campo, _ in campos},
            '_pagadas_presente': False, '_logueo_presente': False,
        })
        item['empresa'] = row[0] or item['empresa']
        item['site'] = row[1] or item['site']
        item['industria'] = row[2] or item['industria']
        for campo, indice in campos:
            item[campo] += parse_numero(row[indice]) if row[indice] not in (None, '') else 0
        item['_pagadas_presente'] |= row[14] not in (None, '')
        item['_logueo_presente'] |= row[15] not in (None, '')

    # Cada hoja es una foto completa de los meses que contiene. Se reemplazan
    # sólo esos meses para que nombres corregidos no queden sumados dos veces.
    meses = {mes for mes, _cliente in grupos}
    if meses:
        HistoricoClienteMensual.query.filter(HistoricoClienteMensual.mes.in_(meses)).delete(
            synchronize_session=False
        )
    for (mes, cliente), item in grupos.items():
        base = {campo: item[campo] for campo, _ in campos
                if campo not in ('pagadas', 'logueo')}
        base.update({campo: item[campo] for campo in ('empresa', 'site', 'industria')})
        db.session.add(HistoricoClienteMensual(
            year=item['year'], mes=mes, cliente=cliente,
            datos_base=json.dumps(base, ensure_ascii=False),
            pagadas=item['pagadas'] if item['_pagadas_presente'] else None,
            logueo=item['logueo'] if item['_logueo_presente'] else None,
        ))
    return len(grupos)


@main_bp.route('/api/dashboard-operativo/importar', methods=['POST'])
@carga_requerida
def api_importar_dashboard_operativo():
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename.lower().endswith('.xlsx'):
        return jsonify({'success': False, 'errores': ['Seleccione un archivo Excel .xlsx']}), 400
    contenido = archivo.read()
    if len(contenido) > 25 * 1024 * 1024:
        return jsonify({'success': False, 'errores': ['El archivo supera el máximo de 25 MB']}), 400
    try:
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        nombre_hoja = next((nombre for nombre in libro.sheetnames if normalizar_header(nombre) == 'dashboard'), libro.sheetnames[0])
        hoja = libro[nombre_hoja]
    except Exception:
        return jsonify({'success': False, 'errores': ['No se pudo leer el archivo Excel']}), 400

    importadas, antes, despues, errores = 0, [], [], []
    for numero_fila, row in enumerate(hoja.iter_rows(min_row=3, values_only=True), 3):
        valores = list(row)
        columnas_necesarias = len(DASHBOARD_COLUMNAS) + 1  # A vacía + B:BC
        if len(valores) < columnas_necesarias:
            valores += [None] * (columnas_necesarias - len(valores))
        fecha, cliente, campania = valores[5], str(valores[6] or '').strip(), str(valores[7] or '').strip()
        if not isinstance(fecha, (datetime, date)) or not cliente or not campania:
            continue
        clave_mes = fecha.strftime('%Y-%m')
        year = fecha.year
        registro = DashboardOperativo.query.filter_by(mes=clave_mes, cliente=cliente, campania=campania).first()
        antes.append(registro.to_dict() if registro else None)
        if not registro:
            registro = DashboardOperativo(year=year, mes=clave_mes, cliente=cliente, campania=campania)
            registro.bajas_manual = parse_numero(valores[40]) if valores[40] not in (None, '') else 0
            db.session.add(registro)
        datos = {clave: valor_dashboard_json(valores[indice + 1]) for indice, clave in enumerate(DASHBOARD_COLUMNAS)}
        datos['mes'] = clave_mes
        datos['cliente'] = cliente
        datos['campania'] = campania
        datos.pop('agentes_baja', None)
        registro.year, registro.mes, registro.cliente, registro.campania = year, clave_mes, cliente, campania
        registro.datos = json.dumps(datos, ensure_ascii=False)
        db.session.flush()
        despues.append(registro.to_dict())
        importadas += 1
    if not importadas:
        db.session.rollback()
        return jsonify({'success': False, 'errores': errores or ['No se encontraron filas válidas en la hoja Dashboard']}), 400
    historicas = importar_historico_desde_libro(libro)
    registrar_historial('importacion', 'dashboard_operativo', archivo.filename,
                        f'Dashboard operativo importado: {importadas} fila(s); Histórico: {historicas} fila(s)',
                        antes={'filas': antes}, despues={'filas': despues})
    db.session.commit()
    mensaje = f'{importadas} fila(s) de Dashboard importada(s).'
    if historicas:
        mensaje += f' Histórico actualizado exactamente con {historicas} fila(s).'
    else:
        mensaje += ' El archivo no contenía una hoja Histórico.'
    return jsonify({'success': True, 'mensaje': mensaje})


@main_bp.route('/api/dashboard-operativo/<int:registro_id>/bajas', methods=['POST'])
@edicion_requerida
def api_guardar_bajas_dashboard(registro_id):
    registro = DashboardOperativo.query.get_or_404(registro_id)
    data = request.get_json(silent=True) or {}
    try:
        bajas = parse_numero(data.get('bajas'))
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['La cantidad de bajas no es válida']}), 400
    if bajas < 0:
        return jsonify({'success': False, 'errores': ['Las bajas no pueden ser negativas']}), 400
    antes = registro.to_dict()
    registro.bajas_manual = bajas
    db.session.flush()
    registrar_historial('edicion', 'dashboard_operativo', registro.id,
                        f'Bajas actualizadas: {registro.cliente} / {registro.campania} / {registro.mes}',
                        antes={'filas': [antes]}, despues={'filas': [registro.to_dict()]})
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Bajas guardadas', 'bajas': bajas})


@main_bp.route('/api/historico-clientes', methods=['GET'])
@login_requerido
def api_historico_clientes():
    year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    mes_filtro = str(request.args.get('mes') or '').strip()
    # Histórico es la única fuente de los indicadores. Los valores importados
    # no se sustituyen durante la lectura por Dashboard ni por Facturación.
    consolidadas = {}
    campos_suma = (
        'horas_dotacion_activa', 'horas_ausentismo', 'dotacion_promedio',
        'bajas', 'horas_requeridas', 'horas_realizadas', 'dotacion_requerida',
    )
    for registro in HistoricoClienteMensual.query.filter_by(year=year).all():
        if mes_filtro and registro.mes != mes_filtro:
            continue
        base = registro.base_dict()
        fila = {
            'id': registro.id, 'mes': registro.mes, 'year': year,
            'cliente': registro.cliente, 'empresa': base.get('empresa'),
            'site': base.get('site'), 'industria': base.get('industria'),
            **{campo: float(base.get(campo) or 0) for campo in campos_suma},
            'pagadas': float(registro.pagadas) if registro.pagadas is not None else None,
            'logueo': float(registro.logueo) if registro.logueo is not None else None,
            'fuentes': ['historico'],
            'fuentes_campos': {campo: 'historico' for campo in (
                'horas', 'operativos', 'dotacion_requerida', 'pagadas_logueo')},
        }
        clave = (registro.mes, normalizar_header(registro.cliente))
        destino = consolidadas.get(clave)
        if destino is None:
            consolidadas[clave] = dict(fila)
            continue
        for campo in campos_suma:
            destino[campo] = float(destino.get(campo) or 0) + float(fila.get(campo) or 0)
        for campo in ('pagadas', 'logueo'):
            valores = [valor for valor in (destino.get(campo), fila.get(campo)) if valor is not None]
            destino[campo] = sum(float(valor) for valor in valores) if valores else None
        if destino.get('id') is None and fila.get('id') is not None:
            destino['id'] = fila['id']
        for campo in ('empresa', 'site', 'industria'):
            destino[campo] = destino.get(campo) or fila.get(campo)
        destino['fuentes'] = list(dict.fromkeys((destino.get('fuentes') or []) + (fila.get('fuentes') or [])))
        for campo, fuente in (fila.get('fuentes_campos') or {}).items():
            actual = destino.setdefault('fuentes_campos', {}).get(campo)
            if actual and fuente and fuente not in actual.split('+'):
                destino['fuentes_campos'][campo] = f'{actual}+{fuente}'
            elif not actual:
                destino['fuentes_campos'][campo] = fuente
    filas = list(consolidadas.values())
    for fila in filas:
        req, real = float(fila.get('horas_requeridas') or 0), float(fila.get('horas_realizadas') or 0)
        activa, ausencia = float(fila.get('horas_dotacion_activa') or 0), float(fila.get('horas_ausentismo') or 0)
        dot, bajas = float(fila.get('dotacion_promedio') or 0), float(fila.get('bajas') or 0)
        pagadas, logueo = fila.get('pagadas'), fila.get('logueo')
        fila.update({
            'cumplimiento_horas': real / req if req else None,
            'abs': ausencia / activa if activa else None,
            'rotacion': bajas / dot if dot else None,
            'eficiencia': float(pagadas) / float(logueo) if pagadas is not None and logueo else None,
        })
    filas.sort(key=lambda fila: (fila['mes'], normalizar_header(fila['cliente'])))
    meses = sorted({fila['mes'] for fila in filas})
    return jsonify({'success': True, 'year': year, 'meses': meses, 'filas': filas})


def cliente_cumplimiento_facturacion(nombre):
    clave = normalizar_header(nombre)
    if clave.startswith('personal'):
        return 'Personal'
    if clave.startswith('santander'):
        return 'Santander'
    if clave.startswith('naturgy'):
        return 'Naturgy'
    equivalencias = {
        'bbva seguros': 'BBVA', 'bice': 'Bice', 'bsf carrefour': 'Bsf Carrefour',
        'cachamai': 'Cachamai', 'gire': 'GIRE', 'galicia seguros': 'Galicia',
        'river plate': 'River Plate', 'unicef': 'Unicef',
    }
    return equivalencias.get(clave, nombre)


@main_bp.route('/api/cumplimiento-horas-clientes', methods=['GET'])
@login_requerido
def api_cumplimiento_horas_clientes():
    """Cumplimiento de horas exactamente desde la base histórica importada."""
    year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    grupos = {}

    def nuevo(mes, cliente):
        return {'mes': mes, 'year': year, 'cliente': cliente,
                'objetivo': 0.0, 'realizado': 0.0}

    for registro in HistoricoClienteMensual.query.filter_by(year=year).all():
        cliente = cliente_cumplimiento_facturacion(registro.cliente)
        item = grupos.setdefault((registro.mes, cliente), nuevo(registro.mes, cliente))
        datos = registro.base_dict()
        item['objetivo'] += float(datos.get('horas_requeridas') or 0)
        item['realizado'] += float(datos.get('horas_realizadas') or 0)

    filas = []
    for item in grupos.values():
        objetivo, realizado = item['objetivo'], item['realizado']
        if objetivo or realizado:
            filas.append({'mes': item['mes'], 'year': year, 'cliente': item['cliente'],
                          'horas_requeridas': round(objetivo, 6),
                          'horas_realizadas': round(realizado, 6)})
    filas.sort(key=lambda fila: (fila['mes'], fila['cliente'].lower()))
    return jsonify({'success': True, 'year': year, 'filas': filas})


@main_bp.route('/api/cumplimiento-facturacion-clientes', methods=['GET'])
@login_requerido
def api_cumplimiento_facturacion_clientes():
    year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    filas = {}
    def obtener(cliente):
        return filas.setdefault(cliente, {
            'cliente': cliente, 'objetivo_horas': 0, 'objetivo_variable': 0,
            'objetivo_penalizaciones_bonos': 0, 'alcanzado_horas': 0,
            'alcanzado_variable': 0, 'alcanzado_penalizaciones_bonos': 0,
        })

    registros = facturacion_por_anio(Facturacion2026.query, year).all()
    for registro in registros:
        if registro.es_next_gen:
            continue
        item = obtener(cliente_cumplimiento_facturacion(registro.cliente))
        item['objetivo_horas'] += float(registro.objetivo_facturacion_horas or 0)
        item['objetivo_variable'] += float(registro.variable_objetivo or 0)
        item['alcanzado_horas'] += float(registro.facturado_horas_control or 0)
        item['alcanzado_variable'] += float(registro.variable_productivo_control or 0)
        item['alcanzado_penalizaciones_bonos'] += float(registro.penalizaciones_bonos_control or 0)

    salida = []
    for item in filas.values():
        item['objetivo_total'] = item['objetivo_horas'] + item['objetivo_variable'] + item['objetivo_penalizaciones_bonos']
        item['alcanzado_total'] = item['alcanzado_horas'] + item['alcanzado_variable'] + item['alcanzado_penalizaciones_bonos']
        item['desvio_horas'] = item['alcanzado_horas'] - item['objetivo_horas']
        item['desvio_variable'] = item['alcanzado_variable'] - item['objetivo_variable']
        item['desvio_penalizaciones_bonos'] = item['alcanzado_penalizaciones_bonos'] - item['objetivo_penalizaciones_bonos']
        item['desvio_total'] = item['alcanzado_total'] - item['objetivo_total']
        item['porcentaje_desvio'] = item['desvio_total'] / item['objetivo_total'] if item['objetivo_total'] else None
        if item['objetivo_total'] or item['alcanzado_total']:
            salida.append({clave: round(valor, 6) if isinstance(valor, float) else valor for clave, valor in item.items()})
    orden_excel = [
        'Personal', 'Santander', 'BBVA', 'Qualia', 'Assurant', 'GIRE', 'Autotekne',
        'Invertir Online', 'Cachamai', 'Bice', 'Farmapay', 'Yamaha', 'Johnson & Johnson',
        'Banco Comafi', 'BiBank', 'Openpay', 'Banco Piano', 'Metrogas', 'Spazios',
        'Supervielle Seguros', 'Shipnow', 'Supervielle', 'Leiva Joyas', 'Naturgy',
        'Bsf Carrefour', 'Mirgor', 'Omint', 'River Plate', 'Unicef', 'Galicia', 'BNA',
    ]
    indice_orden = {cliente: indice for indice, cliente in enumerate(orden_excel)}
    salida.sort(key=lambda item: (indice_orden.get(item['cliente'], 999), item['cliente'].lower()))
    ruta_historico = os.path.join(current_app.root_path, 'data', 'cumplimiento_facturacion_historico.json')
    with open(ruta_historico, encoding='utf-8') as archivo:
        historico = json.load(archivo)
    resumen_anual = []
    for anio, valores in historico.items():
        horas = float(valores.get('horas') or 0)
        variable = float(valores.get('variable') or 0)
        bonos = float(valores.get('bonos_penalizaciones') or 0)
        resumen_anual.append({
            'year': int(anio), 'horas': horas, 'variable': variable,
            'bonos_penalizaciones': bonos, 'total': horas + variable + bonos,
            'porcentaje_desvio': valores.get('porcentaje_desvio'),
        })
    objetivo_actual = sum(float(item['objetivo_total'] or 0) for item in salida)
    horas_actual = sum(float(item['desvio_horas'] or 0) for item in salida)
    variable_actual = sum(float(item['desvio_variable'] or 0) for item in salida)
    bonos_actual = sum(float(item['desvio_penalizaciones_bonos'] or 0) for item in salida)
    total_actual = horas_actual + variable_actual + bonos_actual
    if salida:
        resumen_anual = [item for item in resumen_anual if item['year'] != year]
        resumen_anual.append({
            'year': year, 'horas': horas_actual, 'variable': variable_actual,
            'bonos_penalizaciones': bonos_actual, 'total': total_actual,
            'porcentaje_desvio': total_actual / objetivo_actual if objetivo_actual else None,
        })
    anios_facturacion = anios_facturacion_disponibles()
    anios_disponibles = sorted({int(anio) for anio in historico} | anios_facturacion)
    resumen_anual.sort(key=lambda item: item['year'], reverse=True)
    return jsonify({'success': True, 'year': year, 'anios': anios_disponibles,
                    'filas': salida, 'resumen_anual': resumen_anual})


@main_bp.route('/api/comparativo-interanual', methods=['GET'])
@login_requerido
def api_comparativo_interanual():
    year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    anterior = year - 1
    ruta_historico = os.path.join(current_app.root_path, 'data', 'comparativo_interanual_historico.json')
    with open(ruta_historico, encoding='utf-8') as archivo:
        historico = json.load(archivo)
    base = historico.get(str(anterior), {})
    cumplimiento_anterior = base.get('cumplimiento_facturacion', [None] * 12)
    desvio_anterior = base.get('desvio_facturacion', [None] * 12)

    registros = facturacion_por_anio(Facturacion2026.query, year).all()
    acumulados = {mes: {'obj_fact': 0.0, 'real_fact': 0.0, 'obj_hs': 0.0, 'real_hs': 0.0, 'cantidad': 0} for mes in range(1, 13)}
    for registro in registros:
        if registro.es_next_gen:
            continue
        mes = int(registro.mes[5:7])
        item = acumulados[mes]
        item['obj_fact'] += float(registro.objetivo_facturacion_horas or 0) + float(registro.variable_objetivo or 0)
        item['real_fact'] += float(registro.total_facturado_control or 0)
        item['obj_hs'] += float(registro.horas_objetivo or 0)
        item['real_hs'] += float(registro.horas_facturadas or 0)
        item['cantidad'] += 1

    meses = []
    for numero in range(1, 13):
        actual = acumulados[numero]
        cargado = actual['cantidad'] > 0 and actual['obj_fact'] != 0
        desvio = actual['real_fact'] - actual['obj_fact'] if cargado else None
        cumplimiento = actual['real_fact'] / actual['obj_fact'] if cargado else None
        meses.append({
            'numero': numero, 'nombre': calendar.month_abbr[numero].lower(),
            'cumplimiento_anterior': cumplimiento_anterior[numero - 1] if numero <= len(cumplimiento_anterior) else None,
            'desvio_anterior': desvio_anterior[numero - 1] if numero <= len(desvio_anterior) else None,
            'cumplimiento_actual': cumplimiento, 'desvio_actual': desvio, 'cargado': cargado,
        })

    def resumen(indices):
        items = [acumulados[i] for i in indices if acumulados[i]['cantidad'] and acumulados[i]['obj_fact']]
        obj_fact = sum(x['obj_fact'] for x in items)
        real_fact = sum(x['real_fact'] for x in items)
        obj_hs = sum(x['obj_hs'] for x in items)
        real_hs = sum(x['real_hs'] for x in items)
        return {
            'cumplimiento_horas': real_hs / obj_hs if obj_hs else None,
            'cumplimiento_facturacion': real_fact / obj_fact if obj_fact else None,
            'desvio_facturacion': real_fact - obj_fact if obj_fact else None,
        }

    trimestres_actual = [resumen(range(inicio, inicio + 3)) for inicio in (1, 4, 7, 10)]
    anual_actual = resumen(range(1, 13))
    ant_hs = base.get('cumplimiento_horas_trimestral', [None] * 4)
    ant_fact = base.get('cumplimiento_facturacion_trimestral', [None] * 4)
    ant_desv = base.get('desvio_facturacion_trimestral', [None] * 4)
    trimestres_anterior = [{'cumplimiento_horas': ant_hs[i], 'cumplimiento_facturacion': ant_fact[i], 'desvio_facturacion': ant_desv[i]} for i in range(4)]
    def anual_desde_trimestres(items):
        validos = [x for x in items if x['cumplimiento_facturacion'] is not None]
        return {
            'cumplimiento_horas': sum(x['cumplimiento_horas'] for x in validos) / len(validos) if validos else None,
            'cumplimiento_facturacion': sum(x['cumplimiento_facturacion'] for x in validos) / len(validos) if validos else None,
            'desvio_facturacion': sum(x['desvio_facturacion'] for x in validos) if validos else None,
        }
    anual_anterior = {
        'cumplimiento_horas': base.get('cumplimiento_horas_anual'),
        'cumplimiento_facturacion': base.get('cumplimiento_facturacion_anual'),
        'desvio_facturacion': base.get('desvio_facturacion_anual'),
    }
    if not any(valor is not None for valor in anual_anterior.values()):
        anual_anterior = anual_desde_trimestres(trimestres_anterior)
    return jsonify({'success': True, 'year': year, 'anterior': anterior, 'meses': meses,
                    'actual': {'anual': anual_actual, 'trimestres': trimestres_actual},
                    'anterior_datos': {'anual': anual_anterior, 'trimestres': trimestres_anterior}})


@main_bp.route('/api/historico-clientes/manual', methods=['POST'])
@edicion_requerida
def api_guardar_manual_historico():
    data = request.get_json(silent=True) or {}
    mes, cliente = str(data.get('mes') or ''), str(data.get('cliente') or '').strip()
    if not re.match(r'^\d{4}-\d{2}$', mes) or not cliente:
        return jsonify({'success': False, 'errores': ['Mes y cliente son obligatorios']}), 400
    registro = HistoricoClienteMensual.query.filter_by(mes=mes, cliente=cliente).first()
    if not registro:
        registro = HistoricoClienteMensual(year=int(mes[:4]), mes=mes, cliente=cliente, datos_base='{}')
        db.session.add(registro)
    antes = registro.to_dict()
    registro.pagadas = parse_numero(data.get('pagadas')) if data.get('pagadas') not in (None, '') else None
    registro.logueo = parse_numero(data.get('logueo')) if data.get('logueo') not in (None, '') else None
    db.session.flush()
    registrar_historial('edicion', 'historico_cliente', registro.id,
                        f'Pagadas y logueo: {cliente} / {mes}', antes=antes, despues=registro.to_dict())
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Pagadas y logueo guardados'})


@main_bp.route('/api/variacion-anual', methods=['GET'])
@login_requerido
def api_variacion_anual():
    year = int(request.args.get('year') or current_app.config['DEFAULT_YEAR'])
    anterior = year - 1
    meses = opciones_meses_proyeccion(year)
    claves = [item['value'] for item in meses]
    nombres = [item['label'] for item in meses]

    def horas_reales(anio):
        filas = facturacion_por_anio(Facturacion2026.query.with_entities(
            Facturacion2026.mes, func.sum(Facturacion2026.horas_facturadas)
        ), anio).group_by(Facturacion2026.mes).all()
        resultado = {mes: round(float(total or 0), 2) for mes, total in filas}
        if resultado:
            return resultado
        historico = {}
        for registro in HistoricoClienteMensual.query.filter_by(year=anio).all():
            historico[registro.mes] = historico.get(registro.mes, 0) + float(registro.base_dict().get('horas_realizadas') or 0)
        return {mes: round(total, 2) for mes, total in historico.items()}

    reales_actual = horas_reales(year)
    reales_anterior = horas_reales(anterior)
    # Las aperturas PLP desglosan las mismas horas de Personal para valorizarlas;
    # no son horas adicionales y por eso no deben duplicar el total interanual.
    proyectadas = dict(ProyeccionMatriz.query.with_entities(
        ProyeccionMatriz.mes,
        func.sum(ProyeccionMatriz.horas_requeridas * ProyeccionMatriz.porcentaje_cumplimiento / 100),
    ).filter(
        ProyeccionMatriz.mes.in_(claves),
        db.or_(ProyeccionMatriz.tipo_plp.is_(None), ProyeccionMatriz.tipo_plp == ''),
    ).group_by(ProyeccionMatriz.mes).all())
    meses_reales = [mes for mes in claves if mes in reales_actual]
    corte = max(meses_reales) if meses_reales else None
    filas = []
    for indice, (mes, nombre) in enumerate(zip(claves, nombres), 1):
        clave_anterior = f'{anterior}-{indice:02d}'
        es_real = bool(corte and mes <= corte)
        actual = reales_actual.get(mes, 0) if es_real else float(proyectadas.get(mes) or 0)
        base = float(reales_anterior.get(clave_anterior) or 0)
        filas.append({'mes': mes, 'nombre': nombre, 'anterior': round(base, 2), 'actual': round(actual, 2),
                      'tipo': 'real' if es_real else 'proyectado', 'diferencia': round(actual - base, 2),
                      'variacion': round((actual - base) / base, 6) if base else 0})
    indice_corte = len(meses_reales)
    total_anterior = sum(item['anterior'] for item in filas)
    total_actual = sum(item['actual'] for item in filas)
    acum_anterior = sum(item['anterior'] for item in filas[:indice_corte])
    acum_actual = sum(item['actual'] for item in filas[:indice_corte])
    return jsonify({'success': True, 'year': year, 'anterior': anterior, 'corte': corte, 'filas': filas,
                    'total': {'anterior': round(total_anterior, 2), 'actual': round(total_actual, 2),
                              'diferencia': round(total_actual-total_anterior, 2),
                              'variacion': round((total_actual-total_anterior)/total_anterior, 6) if total_anterior else 0},
                    'acumulado': {'meses': indice_corte, 'anterior': round(acum_anterior, 2), 'actual': round(acum_actual, 2),
                                  'diferencia': round(acum_actual-acum_anterior, 2),
                                  'variacion': round((acum_actual-acum_anterior)/acum_anterior, 6) if acum_anterior else 0}})


@main_bp.route('/api/variables', methods=['POST'])
@edicion_requerida
def api_guardar_variable():
    data = request.get_json(silent=True) or {}
    mes = mes_valido(str(data.get('mes') or '').strip())
    errores = []
    if not mes:
        errores.append('El mes no es valido')

    site = str(data.get('site') or '').strip()
    cliente = str(data.get('cliente') or '').strip()
    campania = str(data.get('campania') or '').strip()
    if not cliente:
        errores.append('El cliente es obligatorio')
    if not campania:
        errores.append('La campaña es obligatoria')

    try:
        porcentaje = parse_numero(data.get('porcentaje') if data.get('porcentaje') not in (None, '') else 0)
    except ValueError:
        porcentaje = 0
        errores.append('El porcentaje tiene un formato invalido')

    if errores:
        return jsonify({'success': False, 'errores': errores}), 400

    year = int(mes[:4])
    base_data = construir_variable_data(year)
    campania_valida = next((
        item for item in base_data['campanias']
        if item['site'] == site and item['cliente'] == cliente and item['campania'] == campania
    ), None)
    if not campania_valida:
        return jsonify({'success': False, 'errores': ['La campaña no existe en Facturación horas para el año seleccionado']}), 400

    variable = VariableCampania.query.filter_by(cliente=cliente, campania=campania, mes=mes).first()
    antes = variable.to_dict() if variable else None
    if not variable:
        variable = VariableCampania()
        db.session.add(variable)
    variable.site = site
    variable.cliente = cliente
    variable.campania = campania
    variable.year = year
    variable.mes = mes
    variable.porcentaje = porcentaje
    db.session.flush()
    despues = variable.to_dict()
    registrar_historial(
        'edicion' if antes else 'creacion',
        'variables',
        mes,
        f'Variable guardada: {cliente} / {campania} / {mes}',
        antes=antes,
        despues=despues,
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': 'Variable guardada',
        'variable': variable.to_dict(),
    })


@main_bp.route('/api/matriz-proyecciones/nuevo-anio', methods=['POST'])
@edicion_requerida
def api_matriz_proyecciones_nuevo_anio():
    """Crea el año siguiente desde la fotografía mensual más recientemente editada."""
    data = request.get_json(silent=True) or {}
    try:
        ultimo_year = int(data.get('year_base'))
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['Seleccioná un año base válido']}), 400
    if not 2020 <= ultimo_year <= 2100:
        return jsonify({'success': False, 'errores': ['El año base está fuera del rango permitido']}), 400
    if not ProyeccionMatriz.query.filter_by(year=ultimo_year).first():
        return jsonify({'success': False, 'errores': ['No hay proyecciones cargadas para usar como base']}), 400

    candidatas = ProyeccionMatriz.query.filter_by(year=ultimo_year).all()
    if not candidatas:
        return jsonify({'success': False, 'errores': ['El último año no contiene proyecciones']}), 400
    ahora = datetime.now()
    limite_calendario = f'{ahora.year}-{ahora.month:02d}'
    meses_disponibles = sorted({fila.mes for fila in candidatas})
    meses_hasta_hoy = [mes for mes in meses_disponibles if mes <= limite_calendario]
    # Si el año vigente contiene meses futuros, no se usan como fotografía:
    # agosto toma agosto y, cuando septiembre esté cargado, toma septiembre.
    mes_base = max(meses_hasta_hoy) if meses_hasta_hoy else max(meses_disponibles)
    base = [fila for fila in candidatas if fila.mes == mes_base]
    year_destino = int(ultimo_year) + 1
    if ProyeccionMatriz.query.filter_by(year=year_destino).first():
        return jsonify({
            'success': False,
            'errores': [f'El año {year_destino} ya tiene proyecciones. No se sobrescribió ningún dato.'],
        }), 409

    feriados_destino = fechas_feriadas_activas(year_destino)
    creadas = []
    for numero_mes in range(1, 13):
        mes_destino = f'{year_destino}-{numero_mes:02d}'
        for origen in base:
            jornadas = [
                {
                    'dotacion_requerida': jornada.dotacion_requerida,
                    'carga_semanal': jornada.carga_semanal,
                    'carga_horaria': jornada.carga_horaria,
                }
                for jornada in origen.jornadas
            ]
            nueva = ProyeccionMatriz()
            aplicar_calculo_proyeccion(nueva, {
                'cliente': origen.cliente,
                'campania': origen.campania,
                'jornadas': jornadas,
                'dotacion_requerida': origen.dotacion_requerida,
                'carga_semanal': origen.carga_semanal,
                'carga_horaria': origen.carga_horaria,
                'porcentaje_cumplimiento': origen.porcentaje_cumplimiento,
                'tipo_plp': origen.tipo_plp,
                'horas_carga_manual': origen.horas_carga_manual,
                'horas_requeridas_manual': origen.horas_requeridas if origen.horas_carga_manual else None,
                'carga_semanal_plp': origen.carga_semanal,
                'carga_horaria_plp': origen.carga_horaria,
            }, mes_destino, feriados_destino)
            db.session.add(nueva)
            creadas.append(nueva)

    def filas_mes(modelo):
        return modelo.query.filter_by(year=ultimo_year, mes=mes_base).all()

    precios_base = filas_mes(ProyeccionPrecio)
    variables_base = filas_mes(VariableCampania)
    distribuciones_base = filas_mes(PersonalDistribucionHoras)
    for numero_mes in range(1, 13):
        mes_destino = f'{year_destino}-{numero_mes:02d}'
        for origen in precios_base:
            db.session.add(ProyeccionPrecio(
                site=origen.site, cliente=origen.cliente, campania=origen.campania,
                year=year_destino, mes=mes_destino, precio_base=origen.precio_base,
                alcance_porcentaje=origen.alcance_porcentaje, precio_final=origen.precio_final,
                importe_fijo_mensual=origen.importe_fijo_mensual,
            ))
        for origen in variables_base:
            db.session.add(VariableCampania(
                site=origen.site, cliente=origen.cliente, campania=origen.campania,
                year=year_destino, mes=mes_destino, porcentaje=origen.porcentaje,
            ))
        for origen in distribuciones_base:
            db.session.add(PersonalDistribucionHoras(
                servicio=origen.servicio, year=year_destino,
                mes=mes_destino, porcentaje_diurno=origen.porcentaje_diurno,
            ))

    db.session.flush()
    resumen = {
        'year_base': ultimo_year,
        'mes_base': mes_base,
        'year_destino': year_destino,
        'campanias': len(base),
        'proyecciones': len(creadas),
        'precios': len(precios_base) * 12,
        'variables': len(variables_base) * 12,
        'distribuciones_personal': len(distribuciones_base) * 12,
        'feriados': 0,
    }
    registrar_historial(
        'creacion', 'matriz_proyecciones', str(year_destino),
        f'Nuevo año {year_destino} generado desde {mes_base}',
        despues=resumen,
        detalle='Fotografía del último mes actualizado repetida en los 12 meses; feriados no copiados.',
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'Año {year_destino} creado desde la fotografía de {mes_base}. Los feriados quedaron vacíos.',
        **resumen,
    })


@main_bp.route('/api/matriz-proyecciones', methods=['GET'])
@login_requerido
def api_matriz_proyecciones():
    year = int(request.args.get('year') or datetime.utcnow().year)
    if recalcular_proyecciones_plp_year(year):
        db.session.commit()
    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    registros = ProyeccionMatriz.query.filter(ProyeccionMatriz.mes.in_(meses)).order_by(
        ProyeccionMatriz.cliente,
        ProyeccionMatriz.campania,
        ProyeccionMatriz.mes,
    ).all()
    asociaciones = AsignacionComercial.query.filter_by(activa=True).order_by(
        AsignacionComercial.cliente,
        AsignacionComercial.campania,
    ).all()
    pares = []
    vistos = set()
    for asignacion in asociaciones:
        clave = (asignacion.cliente, asignacion.campania)
        if clave in vistos:
            continue
        vistos.add(clave)
        pares.append({'cliente': asignacion.cliente, 'campania': asignacion.campania})
    return jsonify({
        'success': True,
        'year': year,
        'meses': opciones_meses_proyeccion(year),
        'asociaciones': pares,
        'feriados': [feriado.to_dict() for feriado in FeriadoOperativo.query.filter_by(year=year, activo=True).all()],
        'cargas_semanales': list(CARGAS_SEMANALES_PROYECCION),
        'proyecciones': [
            {**registro.to_dict(), 'mes_label': etiqueta_mes_proyeccion(registro.mes)}
            for registro in registros
        ],
    })


@main_bp.route('/api/suma-fija', methods=['GET'])
@login_requerido
def api_suma_fija():
    year = int(request.args.get('year') or datetime.utcnow().year)
    meses_info = opciones_meses_proyeccion(year)
    meses = [item['value'] for item in meses_info]
    registros = ProyeccionPrecio.query.filter(ProyeccionPrecio.mes.in_(meses)).order_by(
        ProyeccionPrecio.cliente, ProyeccionPrecio.campania, ProyeccionPrecio.mes,
    ).all()
    filas = {}
    for registro in registros:
        clave = (registro.cliente, registro.campania)
        fila = filas.setdefault(clave, {
            'site': registro.site or '',
            'cliente': registro.cliente,
            'campania': registro.campania,
            'meses': {mes: {'id': None, 'monto': 0} for mes in meses},
        })
        fila['meses'][registro.mes] = {
            'id': registro.id,
            'monto': round(registro.importe_fijo_mensual or 0, 2),
        }
    return jsonify({
        'success': True, 'year': year, 'meses': meses_info,
        'filas': list(filas.values()),
    })


@main_bp.route('/api/suma-fija', methods=['POST'])
@edicion_requerida
def api_guardar_suma_fija():
    data = request.get_json(silent=True) or {}
    precio_id = data.get('id')
    monto = parse_numero(data.get('monto'))
    if monto < 0:
        return jsonify({'success': False, 'errores': ['La suma fija no puede ser negativa']}), 400
    precio = ProyeccionPrecio.query.get_or_404(precio_id)
    antes = precio.to_dict()
    precio.importe_fijo_mensual = monto
    db.session.flush()
    registrar_historial(
        'edicion', 'matriz_precios', precio.id,
        f'Suma fija > {precio.cliente} / {precio.campania} / {precio.mes}: {monto}',
        detalle='Importe mensual por dotación aplicado en Facturación horas',
        antes={'precios': [antes]}, despues={'precios': [precio.to_dict()]},
    )
    db.session.commit()
    return jsonify({
        'success': True, 'mensaje': 'Suma fija actualizada',
        'registro': {'id': precio.id, 'monto': round(precio.importe_fijo_mensual or 0, 2)},
    })


@main_bp.route('/api/personal-distribucion', methods=['GET', 'POST'])
@edicion_si_mutacion_requerida
def api_personal_distribucion():
    if request.method == 'GET':
        year = int(request.args.get('year') or datetime.utcnow().year)
        filas = PersonalDistribucionHoras.query.filter_by(year=year).order_by(
            PersonalDistribucionHoras.mes,
            PersonalDistribucionHoras.servicio,
        ).all()
        por_clave = {(fila.servicio, fila.mes): fila.to_dict() for fila in filas}
        distribuciones = []
        for mes_info in opciones_meses_proyeccion(year):
            mes = mes_info['value']
            for servicio in SERVICIOS_PERSONAL:
                actual = por_clave.get((servicio, mes))
                if actual:
                    actual['heredado'] = False
                    distribuciones.append(actual)
                    continue
                anterior = distribucion_personal('', '', mes, servicio)
                porcentaje_diurno = anterior.porcentaje_diurno if anterior else 100
                distribuciones.append({
                    'id': None,
                    'servicio': servicio,
                    'year': year,
                    'mes': mes,
                    'porcentaje_diurno': porcentaje_diurno,
                    'porcentaje_nocturno': max(0, 100 - porcentaje_diurno),
                    'heredado': bool(anterior),
                })
        return jsonify({
            'success': True,
            'year': year,
            'servicios': list(SERVICIOS_PERSONAL),
            'meses': opciones_meses_proyeccion(year),
            'distribuciones': distribuciones,
        })

    data = request.get_json(silent=True) or {}
    try:
        year = int(data.get('year') or datetime.utcnow().year)
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['El año no es válido']}), 400
    entradas = data.get('distribuciones') if isinstance(data.get('distribuciones'), list) else []
    errores = []
    normalizadas = []
    for entrada in entradas:
        servicio = str(entrada.get('servicio') or '').strip()
        mes = mes_valido(entrada.get('mes'))
        try:
            porcentaje_diurno = parse_numero(entrada.get('porcentaje_diurno'))
        except (AttributeError, ValueError):
            porcentaje_diurno = -1
        if servicio not in SERVICIOS_PERSONAL:
            errores.append(f'Servicio de Personal no válido: {servicio or "vacío"}')
        elif not mes or int(mes[:4]) != year:
            errores.append(f'Mes no válido para {servicio}')
        elif not 0 <= porcentaje_diurno <= 100:
            errores.append(f'El porcentaje diurno de {servicio} debe estar entre 0 y 100')
        else:
            normalizadas.append((servicio, mes, porcentaje_diurno))
    if errores or not normalizadas:
        return jsonify({'success': False, 'errores': errores or ['No hay porcentajes para guardar']}), 400

    antes = []
    guardadas = []
    for servicio, mes, porcentaje_diurno in normalizadas:
        fila = PersonalDistribucionHoras.query.filter_by(servicio=servicio, mes=mes).first()
        antes.append(fila.to_dict() if fila else None)
        if not fila:
            fila = PersonalDistribucionHoras(servicio=servicio, year=year, mes=mes)
            db.session.add(fila)
        fila.porcentaje_diurno = porcentaje_diurno
        guardadas.append(fila)

    db.session.flush()
    proyecciones = ProyeccionMatriz.query.filter(ProyeccionMatriz.mes.in_([mes for _, mes, _ in normalizadas])).all()
    feriados = fechas_feriadas_activas(year)
    for proyeccion in proyecciones:
        if not servicio_personal_identificado(proyeccion.cliente, proyeccion.campania, proyeccion.tipo_plp):
            continue
        aplicar_calculo_proyeccion(proyeccion, {
            'cliente': proyeccion.cliente,
            'campania': proyeccion.campania,
            'jornadas': [jornada.to_dict() for jornada in proyeccion.jornadas],
            'porcentaje_cumplimiento': proyeccion.porcentaje_cumplimiento,
            'tipo_plp': proyeccion.tipo_plp,
            'horas_carga_manual': proyeccion.horas_carga_manual,
            'horas_requeridas_manual': proyeccion.horas_requeridas,
            'carga_semanal_plp': proyeccion.carga_semanal,
            'carga_horaria_plp': proyeccion.carga_horaria,
        }, proyeccion.mes, feriados)

    registrar_historial(
        'edicion',
        'personal_distribucion_horas',
        str(year),
        f'Distribución diurna/nocturna de Personal actualizada para {year}',
        antes={'distribuciones': antes},
        despues={'distribuciones': [fila.to_dict() for fila in guardadas]},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': 'Porcentajes de Personal guardados y proyecciones recalculadas',
        'distribuciones': [fila.to_dict() for fila in guardadas],
    })


@main_bp.route('/api/proyecciones-plp/edicion-masiva', methods=['POST'])
@edicion_requerida
def api_edicion_masiva_proyecciones_plp():
    """Actualiza varias filas PLP en una sola transacción, como una planilla."""
    data = request.get_json(silent=True) or {}
    try:
        year = int(data.get('year'))
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['El año no es válido']}), 400
    if not 2020 <= year <= 2100:
        return jsonify({'success': False, 'errores': ['El año debe estar entre 2020 y 2100']}), 400

    entradas = data.get('filas') if isinstance(data.get('filas'), list) else []
    if not entradas:
        return jsonify({'success': False, 'errores': ['No hay filas PLP para guardar']}), 400
    if len(entradas) > 500:
        return jsonify({'success': False, 'errores': ['La edición masiva admite hasta 500 filas por operación']}), 400

    errores = []
    normalizadas = []
    ids_recibidos = set()
    ids_a_eliminar = set()
    claves_recibidas = set()
    for numero, entrada in enumerate(entradas, start=1):
        try:
            registro_id = int(entrada['id']) if entrada.get('id') not in (None, '') else None
        except (TypeError, ValueError):
            registro_id = -1
        eliminar = entrada.get('eliminar') is True
        if eliminar:
            if not registro_id or registro_id == -1 or registro_id in ids_recibidos:
                errores.append(f'Fila {numero}: no se puede eliminar una fila sin identificador válido')
            else:
                ids_recibidos.add(registro_id)
                ids_a_eliminar.add(registro_id)
                normalizadas.append((registro_id, True, None, None, None, None, None, None, None))
            continue
        tipo = str(entrada.get('tipo_plp') or '').strip()
        mes = mes_valido(entrada.get('mes'))
        campania = str(entrada.get('campania') or '').strip()
        try:
            horas = parse_numero(entrada.get('horas'))
            carga_horaria = parse_numero(entrada.get('carga_horaria'))
            cumplimiento = parse_numero(entrada.get('porcentaje_cumplimiento'))
        except (TypeError, ValueError):
            horas, carga_horaria, cumplimiento = -1, -1, -1
        carga_semanal = normalizar_carga_semanal(entrada.get('carga_semanal') or 'L a V')
        clave = (tipo, mes, normalizar_header(campania))

        if registro_id == -1 or (registro_id and registro_id in ids_recibidos):
            errores.append(f'Fila {numero}: identificador inválido o repetido')
        elif tipo not in SERVICIOS_PERSONAL:
            errores.append(f'Fila {numero}: clasificación PLP no válida')
        elif not mes or int(mes[:4]) != year:
            errores.append(f'Fila {numero}: el mes no corresponde a {year}')
        elif not campania:
            errores.append(f'Fila {numero}: falta el nombre de campaña')
        elif len(campania) > 100:
            errores.append(f'Fila {numero}: la campaña supera los 100 caracteres')
        elif horas < 0:
            errores.append(f'Fila {numero}: las horas no pueden ser negativas')
        elif carga_horaria <= 0:
            errores.append(f'Fila {numero}: las horas de jornada deben ser mayores a cero')
        elif not 0 <= cumplimiento <= 1000:
            errores.append(f'Fila {numero}: el porcentaje de cumplimiento debe estar entre 0 y 1000')
        elif clave in claves_recibidas:
            errores.append(f'Fila {numero}: la campaña está repetida para la misma clasificación y mes')
        else:
            if registro_id:
                ids_recibidos.add(registro_id)
            claves_recibidas.add(clave)
            normalizadas.append((registro_id, False, tipo, mes, campania, horas, carga_semanal, carga_horaria, cumplimiento))
    if errores:
        return jsonify({'success': False, 'errores': errores[:30]}), 400

    registros_por_id = {}
    if ids_recibidos:
        registros_por_id = {
            registro.id: registro
            for registro in ProyeccionMatriz.query.filter(ProyeccionMatriz.id.in_(ids_recibidos)).all()
        }
        faltantes = sorted(ids_recibidos - set(registros_por_id))
        if faltantes:
            return jsonify({'success': False, 'errores': [f'No existen las proyecciones: {faltantes}']}), 404
        if any(not registro.tipo_plp or registro.year != year for registro in registros_por_id.values()):
            return jsonify({'success': False, 'errores': ['Una fila no pertenece al año o al módulo PLP seleccionado']}), 400

    antes = []
    guardadas = []
    eliminadas = []
    feriados = fechas_feriadas_activas(year)
    for registro_id in ids_a_eliminar:
        proyeccion = registros_por_id[registro_id]
        antes.append(snapshot_proyeccion(proyeccion))
        eliminadas.append(proyeccion.id)
        db.session.delete(proyeccion)
    if eliminadas:
        db.session.flush()
    for registro_id, eliminar, tipo, mes, campania, horas, carga_semanal, carga_horaria, cumplimiento in normalizadas:
        if eliminar:
            continue
        proyeccion = registros_por_id.get(registro_id) if registro_id else None
        duplicada = buscar_proyeccion_plp(tipo, campania, mes)
        if duplicada and duplicada.id != registro_id and duplicada.id not in ids_a_eliminar:
            # La grilla funciona como una planilla: la clave
            # clasificación + campaña + mes identifica el registro. Si el
            # usuario vuelve a cargar esa clave, se actualiza la fila existente
            # en lugar de rechazarla como duplicada.
            if proyeccion:
                antes.append(snapshot_proyeccion(proyeccion))
                eliminadas.append(proyeccion.id)
                db.session.delete(proyeccion)
                db.session.flush()
            proyeccion = duplicada
        antes.append(snapshot_proyeccion(proyeccion))
        if not proyeccion:
            proyeccion = ProyeccionMatriz()
            db.session.add(proyeccion)
        aplicar_calculo_proyeccion(proyeccion, {
            'cliente': 'Personal',
            'campania': campania,
            'tipo_plp': tipo,
            'horas_carga_manual': True,
            'horas_requeridas_manual': horas,
            'carga_semanal_plp': carga_semanal,
            'carga_horaria_plp': carga_horaria,
            'porcentaje_cumplimiento': cumplimiento,
            'jornadas': [],
        }, mes, feriados)
        guardadas.append(proyeccion)

    db.session.flush()
    registrar_historial(
        'edicion', 'proyeccion_matriz', str(year),
        f'Edición rápida PLP {year}: {len(guardadas)} guardada(s), {len(eliminadas)} eliminada(s)',
        antes={'proyecciones': antes},
        despues={'proyecciones': [snapshot_proyeccion(item) for item in guardadas], 'ids_eliminados': eliminadas},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'PLP actualizado: {len(guardadas)} guardada(s) y {len(eliminadas)} eliminada(s)',
        'proyecciones': [item.to_dict() for item in guardadas],
        'eliminadas': eliminadas,
    })


@main_bp.route('/api/proyecciones-plp/template', methods=['GET'])
@login_requerido
def api_template_proyecciones_plp():
    year = int(request.args.get('year') or datetime.utcnow().year)
    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    registros = ProyeccionMatriz.query.filter(
        ProyeccionMatriz.mes.in_(meses),
        ProyeccionMatriz.tipo_plp.isnot(None),
        ProyeccionMatriz.tipo_plp != '',
    ).order_by(ProyeccionMatriz.mes, ProyeccionMatriz.tipo_plp, ProyeccionMatriz.campania).all()
    rows = [[registro.tipo_plp, etiqueta_mes_proyeccion(registro.mes), registro.campania, round(registro.horas_requeridas or 0, 2), registro.carga_semanal or 'L a V', registro.carga_horaria or 6, registro.porcentaje_cumplimiento or 0] for registro in registros]
    if not rows:
        rows = [['Personal', f'ene-{str(year)[-2:]}', 'MÓVIL TELEFÓNICO', '', 'L a V', 6, 100]]
    contenido = crear_xlsx(['Clasificación PLP', 'MES', 'CAMPAÑA', 'HORAS', 'Carga semanal', 'Horas jornada', '% cumplimiento'], rows)
    return Response(
        contenido,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="proyecciones_plp_{year}.xlsx"'},
    )


@main_bp.route('/api/proyecciones-plp/importar', methods=['POST'])
@carga_requerida
def api_importar_proyecciones_plp():
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename:
        return jsonify({'success': False, 'errores': ['Seleccione el archivo anual de PLP']}), 400
    try:
        year = int(request.form.get('year') or datetime.utcnow().year)
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['El año no es válido']}), 400
    nombre = archivo.filename.lower()
    contenido = archivo.read()
    try:
        if nombre.endswith('.xlsx'):
            filas = filas_desde_xlsx(contenido)
        elif nombre.endswith('.csv'):
            filas = filas_desde_csv(decodificar_texto_importacion(contenido))
        else:
            return jsonify({'success': False, 'errores': ['Use un archivo .xlsx o .csv']}), 400
    except Exception:
        return jsonify({'success': False, 'errores': ['No se pudo leer el archivo PLP']}), 400
    if len(filas) < 2:
        return jsonify({'success': False, 'errores': ['El archivo no contiene filas para importar']}), 400

    headers = [normalizar_header(celda) for celda in filas[0]]
    indices_campania = [indice for indice, header in enumerate(headers) if header in ('campana', 'campania')]
    indice_tipo_explicito = next((i for i, h in enumerate(headers) if h in ('clasificacion plp', 'tipo plp', 'plp')), None)
    indice_tipo = indice_tipo_explicito if indice_tipo_explicito is not None else (indices_campania[0] if indices_campania else None)
    indice_nombre = indices_campania[-1] if indices_campania and indices_campania[-1] != indice_tipo else next((i for i, h in enumerate(headers) if h in ('nombre campana', 'nombre campania', 'servicio')), None)
    indice_mes = next((i for i, h in enumerate(headers) if h == 'mes'), None)
    indice_horas = next((i for i, h in enumerate(headers) if h in ('horas', 'hs')), None)
    indice_carga = next((i for i, h in enumerate(headers) if h == 'carga semanal'), None)
    indice_jornada = next((i for i, h in enumerate(headers) if h in ('horas jornada', 'jornada laboral')), None)
    indice_cumplimiento = next((i for i, h in enumerate(headers) if 'cumplimiento' in h), None)
    if None in (indice_tipo, indice_nombre, indice_mes, indice_horas):
        return jsonify({'success': False, 'errores': ['Se esperan las columnas Clasificación PLP, MES, CAMPAÑA y HORAS']}), 400

    por_tipo = {normalizar_header(tipo): tipo for tipo in SERVICIOS_PERSONAL}
    errores = []
    entradas = []
    for numero_fila, fila in enumerate(filas[1:], start=2):
        if not any(str(celda or '').strip() for celda in fila):
            continue
        celda = lambda indice: fila[indice] if indice < len(fila) else ''
        tipo = por_tipo.get(normalizar_header(celda(indice_tipo)))
        mes = mes_plp_importado(celda(indice_mes), year)
        campania = str(celda(indice_nombre) or '').strip()
        try:
            horas = parse_numero(celda(indice_horas))
        except ValueError:
            horas = -1
        carga_semanal = normalizar_carga_semanal(celda(indice_carga) if indice_carga is not None else 'L a V')
        try:
            carga_horaria = parse_numero(celda(indice_jornada)) if indice_jornada is not None and celda(indice_jornada) not in (None, '') else 6
        except ValueError:
            carga_horaria = 0
        try:
            cumplimiento = parse_numero(celda(indice_cumplimiento)) if indice_cumplimiento is not None and celda(indice_cumplimiento) not in (None, '') else 100
            # Excel almacena una celda con formato 98% como 0.98.
            if indice_cumplimiento is not None and 0 < cumplimiento <= 1:
                cumplimiento *= 100
        except ValueError:
            cumplimiento = -1
        if not tipo:
            errores.append(f'Fila {numero_fila}: clasificación PLP no válida')
        elif not mes:
            errores.append(f'Fila {numero_fila}: el mes no pertenece a {year}')
        elif not campania:
            errores.append(f'Fila {numero_fila}: falta el nombre de campaña')
        elif horas < 0:
            errores.append(f'Fila {numero_fila}: horas no válidas')
        elif carga_horaria <= 0:
            errores.append(f'Fila {numero_fila}: jornada laboral no válida')
        elif cumplimiento < 0:
            errores.append(f'Fila {numero_fila}: porcentaje de cumplimiento no válido')
        else:
            entradas.append((tipo, mes, campania, horas, carga_semanal, carga_horaria, cumplimiento))
    if errores:
        return jsonify({'success': False, 'errores': errores[:30]}), 400
    if not entradas:
        return jsonify({'success': False, 'errores': ['No se encontraron filas PLP válidas']}), 400

    antes = []
    guardadas = []
    actualizadas = 0
    creadas = 0
    feriados = fechas_feriadas_activas(year)
    for tipo, mes, campania, horas, carga_semanal, carga_horaria, cumplimiento in entradas:
        proyeccion = buscar_proyeccion_plp(tipo, campania, mes)
        antes.append(snapshot_proyeccion(proyeccion))
        if proyeccion:
            actualizadas += 1
        else:
            proyeccion = ProyeccionMatriz()
            db.session.add(proyeccion)
            creadas += 1
        aplicar_calculo_proyeccion(proyeccion, {
            'cliente': 'Personal',
            'campania': campania,
            'tipo_plp': tipo,
            'horas_carga_manual': True,
            'horas_requeridas_manual': horas,
            'carga_semanal_plp': carga_semanal,
            'carga_horaria_plp': carga_horaria,
            'porcentaje_cumplimiento': cumplimiento,
            'jornadas': [],
        }, mes, feriados)
        guardadas.append(proyeccion)
    db.session.flush()
    registrar_historial(
        'importacion', 'proyeccion_matriz', str(year),
        f'Importación anual PLP {year}: {len(guardadas)} registro(s)',
        antes={'proyecciones': antes},
        despues={'proyecciones': [snapshot_proyeccion(item) for item in guardadas]},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'PLP {year} actualizado: {actualizadas} reemplazadas y {creadas} nuevas',
        'actualizadas': actualizadas,
        'creadas': creadas,
    })


@main_bp.route('/api/matriz-proyecciones', methods=['POST'])
@edicion_requerida
def api_guardar_matriz_proyeccion():
    data = request.get_json(silent=True) or {}
    normalizado = payload_proyeccion(data)
    if normalizado['errores']:
        return jsonify({'success': False, 'errores': normalizado['errores']}), 400
    valores = normalizado['valores']
    proyeccion_id = data.get('id')
    if proyeccion_id:
        proyeccion = ProyeccionMatriz.query.get_or_404(proyeccion_id)
        if (
            proyeccion.cliente != valores['cliente']
            or proyeccion.campania != valores['campania']
            or proyeccion.mes != valores['mes']
        ):
            return jsonify({'success': False, 'errores': ['La edición replica ajustes sobre la misma campaña y mes. Para cambiar cliente, campaña o mes, cree una proyección nueva.']}), 400

    proyecciones_guardadas = []
    feriados = fechas_feriadas_activas(valores['year'])
    antes_snapshots = []

    meses_a_guardar = [valores['mes']] if valores.get('horas_carga_manual') else meses_proyeccion_desde(valores['mes'])
    for mes in meses_a_guardar:
        if valores.get('tipo_plp'):
            proyeccion = buscar_proyeccion_plp(valores['tipo_plp'], valores['campania'], mes)
        else:
            proyeccion = ProyeccionMatriz.query.filter_by(
                cliente=valores['cliente'],
                campania=valores['campania'],
                mes=mes,
            ).first()
        antes_snapshots.append(snapshot_proyeccion(proyeccion))
        if not proyeccion:
            proyeccion = ProyeccionMatriz()
            db.session.add(proyeccion)

        aplicar_calculo_proyeccion(proyeccion, valores, mes, feriados)
        proyecciones_guardadas.append(proyeccion)

    db.session.flush()
    despues_snapshots = [snapshot_proyeccion(proyeccion) for proyeccion in proyecciones_guardadas]
    registrar_historial(
        'edicion' if any(antes_snapshots) else 'creacion',
        'proyeccion_matriz',
        valores['mes'],
        f'Proyecciones guardadas: {valores["cliente"]} / {valores["campania"]} desde {valores["mes"]}',
        antes={'proyecciones': antes_snapshots},
        despues={'proyecciones': despues_snapshots},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'{len(proyecciones_guardadas)} proyeccion(es) guardada(s)',
        'proyecciones': [
            {**proyeccion.to_dict(), 'mes_label': etiqueta_mes_proyeccion(proyeccion.mes)}
            for proyeccion in proyecciones_guardadas
        ],
    })


@main_bp.route('/api/matriz-proyecciones/<int:proyeccion_id>', methods=['DELETE'])
@eliminacion_requerida
def api_eliminar_matriz_proyeccion(proyeccion_id):
    proyeccion = ProyeccionMatriz.query.get_or_404(proyeccion_id)
    antes = snapshot_proyeccion(proyeccion)
    db.session.delete(proyeccion)
    registrar_historial(
        'eliminacion',
        'proyeccion_matriz',
        proyeccion_id,
        f'Proyeccion eliminada: {antes.get("cliente")} / {antes.get("campania")} / {antes.get("mes")}',
        antes={'proyecciones': [antes]},
        despues={'proyecciones': []},
    )
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Proyección eliminada'})


@main_bp.route('/api/matriz-proyecciones/eliminar-todo', methods=['POST'])
@eliminacion_requerida
def api_eliminar_todas_matriz_proyecciones():
    data = request.get_json(silent=True) or {}
    try:
        year = int(data.get('year') or datetime.utcnow().year)
    except (TypeError, ValueError):
        return jsonify({'success': False, 'errores': ['El año no es valido']}), 400

    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    proyecciones = ProyeccionMatriz.query.filter(ProyeccionMatriz.mes.in_(meses)).order_by(
        ProyeccionMatriz.cliente,
        ProyeccionMatriz.campania,
        ProyeccionMatriz.mes,
    ).all()
    if not proyecciones:
        return jsonify({'success': False, 'errores': [f'No hay proyecciones cargadas para {year}']}), 404

    antes = [snapshot_proyeccion(proyeccion) for proyeccion in proyecciones]
    for proyeccion in proyecciones:
        db.session.delete(proyeccion)
    registrar_historial(
        'eliminacion',
        'proyeccion_matriz',
        year,
        f'Eliminacion total de proyecciones {year}: {len(proyecciones)} registro(s)',
        antes={'proyecciones': antes},
        despues={'proyecciones': []},
    )
    db.session.flush()
    pendientes = [snapshot for snapshot in antes if existe_proyeccion_snapshot(snapshot)]
    if pendientes:
        db.session.rollback()
        return jsonify({'success': False, 'errores': ['No se pudieron eliminar todas las proyecciones']}), 500
    db.session.commit()
    return jsonify({'success': True, 'mensaje': f'{len(proyecciones)} proyeccion(es) eliminada(s)'})


@main_bp.route('/api/matriz-proyecciones/deshacer-ultimo', methods=['POST'])
@edicion_requerida
def api_deshacer_ultima_proyeccion():
    usuario = usuario_actual()
    query = HistorialCambio.query.filter(
        HistorialCambio.entidad == 'proyeccion_matriz',
        HistorialCambio.accion != 'deshacer',
    )
    if usuario and not usuario.es_administrador:
        query = query.filter(HistorialCambio.usuario_id == usuario.id)
    item = None
    for candidato in query.order_by(HistorialCambio.creado_en.desc(), HistorialCambio.id.desc()).limit(50).all():
        if not historial_movimiento_deshace(candidato.id):
            item = candidato
            break
    if not item:
        return jsonify({'success': False, 'errores': ['No hay movimientos de proyecciones para deshacer']}), 404
    resultado, status = deshacer_item_historial(item)
    return jsonify(resultado), status


@main_bp.route('/api/matriz-proyecciones/template', methods=['GET'])
@login_requerido
def api_template_matriz_proyecciones():
    year = int(request.args.get('year') or datetime.utcnow().year)
    meses = [item['value'] for item in opciones_meses_proyeccion(year)]
    registros = ProyeccionMatriz.query.filter(ProyeccionMatriz.mes.in_(meses)).order_by(
        ProyeccionMatriz.cliente,
        ProyeccionMatriz.campania,
        ProyeccionMatriz.mes,
    ).all()

    headers = [
        'Cliente', 'Campaña', 'Año', 'Mes',
        'Dotación requerida', 'Carga semanal', 'Hs diarias', 'Q dias objetivo',
        'Hs requeridas jornada', 'Dotacion total', 'Hs requeridas total',
        '% cumplimiento', 'Horas proyectadas total',
    ]
    rows = [[
        'Ej: BNA',
        'Ej: BNA Mora',
        year,
        f'{year}-01',
        'Ej: 9',
        'Ej: L a V',
        'Ej: 6',
        'Calculado por sistema',
        'Calculado por sistema',
        'Calculado por sistema',
        'Calculado por sistema',
        'Ej: 100',
        'Calculado por sistema',
    ]]
    for registro in registros:
        jornadas = registro.jornadas or [ProyeccionMatrizJornada(
            dotacion_requerida=registro.dotacion_requerida,
            carga_semanal=registro.carga_semanal,
            carga_horaria=registro.carga_horaria,
            dias_objetivo=registro.dias_objetivo,
            horas_requeridas=registro.horas_requeridas,
        )]
        for jornada in jornadas:
            rows.append([
                registro.cliente,
                registro.campania,
                registro.year,
                registro.mes,
                jornada.dotacion_requerida or 0,
                jornada.carga_semanal or 'L a V',
                jornada.carga_horaria or 0,
                jornada.dias_objetivo or 0,
                round(jornada.horas_requeridas or 0, 2),
                registro.dotacion_requerida or 0,
                round(registro.horas_requeridas or 0, 2),
                registro.porcentaje_cumplimiento or 0,
                round(registro.horas_proyectadas or 0, 2),
            ])

    contenido = crear_xlsx(headers, rows)
    return Response(
        contenido,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="plantilla_proyecciones_{year}.xlsx"'}
    )


@main_bp.route('/api/matriz-proyecciones/importar', methods=['POST'])
@carga_requerida
def api_importar_matriz_proyecciones():
    archivo = request.files.get('archivo')
    if not archivo or not archivo.filename:
        return jsonify({'success': False, 'errores': ['Seleccione un archivo para importar']}), 400

    nombre = archivo.filename.lower()
    contenido_bytes = archivo.read()
    try:
        if nombre.endswith('.xlsx'):
            filas = filas_desde_xlsx(contenido_bytes)
        elif nombre.endswith('.csv'):
            filas = filas_desde_csv(decodificar_texto_importacion(contenido_bytes))
        else:
            return jsonify({'success': False, 'errores': ['Use un archivo .xlsx o .csv generado desde el template']}), 400
    except Exception:
        return jsonify({'success': False, 'errores': ['No pude leer el archivo. Descargue nuevamente el template y vuelva a completarlo.']}), 400

    if len(filas) < 2:
        return jsonify({'success': False, 'errores': ['No se encontraron filas para importar']}), 400

    headers = [normalizar_header(celda) for celda in filas[0]]
    mapa = {header: index for index, header in enumerate(headers)}
    requeridas = {
        'cliente': 'cliente',
        'campana': 'campania',
        'campaña': 'campania',
        'ano': 'year',
        'año': 'year',
        'mes': 'mes',
        'personas': 'dotacion_requerida',
        'dotacion requerida': 'dotacion_requerida',
        'carga semanal': 'carga_semanal',
        'hs diarias': 'carga_horaria',
        '% cumplimiento': 'porcentaje_cumplimiento',
        'cumplimiento': 'porcentaje_cumplimiento',
        'porcentaje cumplimiento': 'porcentaje_cumplimiento',
        'hs requeridas total': 'horas_requeridas_total',
        'horas requeridas total': 'horas_requeridas_total',
        'horas proyectadas total': 'horas_proyectadas_total',
        'horas proyectadas': 'horas_proyectadas_total',
    }
    columnas = {}
    for header, destino in requeridas.items():
        if header in mapa:
            columnas[destino] = mapa[header]
    faltantes = [campo for campo in ('cliente', 'campania', 'year', 'mes', 'dotacion_requerida', 'carga_semanal', 'carga_horaria') if campo not in columnas]
    if faltantes:
        return jsonify({'success': False, 'errores': [f'Faltan columnas requeridas: {", ".join(faltantes)}']}), 400

    grupos = {}
    errores = []
    for indice, fila in enumerate(filas[1:], start=2):
        if not any(str(celda or '').strip() for celda in fila):
            continue

        def celda(campo):
            posicion = columnas.get(campo)
            return fila[posicion] if posicion is not None and posicion < len(fila) else ''

        cliente = str(celda('cliente') or '').strip()
        if normalizar_header(cliente).startswith('ej:'):
            continue
        campania = str(celda('campania') or '').strip()
        mes = mes_valido(celda('mes'))
        try:
            year = int(float(str(celda('year')).strip()))
            dotacion = parse_numero(celda('dotacion_requerida'))
            carga_horaria = parse_numero(celda('carga_horaria'))
            porcentaje = parse_numero(celda('porcentaje_cumplimiento')) if 'porcentaje_cumplimiento' in columnas and str(celda('porcentaje_cumplimiento')).strip() else 100
            # Excel entrega una celda porcentual como decimal: 98% llega como 0.98.
            # Si el archivo ya trae 98 o "98%", parse_numero devuelve 98 y se conserva.
            if 0 < abs(porcentaje) <= 1:
                porcentaje *= 100
            horas_requeridas_total = parse_numero(celda('horas_requeridas_total')) if 'horas_requeridas_total' in columnas and str(celda('horas_requeridas_total')).strip() else None
            horas_proyectadas_total = parse_numero(celda('horas_proyectadas_total')) if 'horas_proyectadas_total' in columnas and str(celda('horas_proyectadas_total')).strip() else None
        except ValueError:
            errores.append(f'Fila {indice}: hay valores numericos invalidos')
            continue
        if not cliente or not campania or not mes:
            errores.append(f'Fila {indice}: cliente, campaña y mes son obligatorios')
            continue
        carga_semanal = normalizar_carga_semanal(celda('carga_semanal'))
        cumplimiento_explicito = (
            'porcentaje_cumplimiento' in columnas
            and str(celda('porcentaje_cumplimiento')).strip() != ''
        )
        clave = (cliente, campania, mes)
        grupos.setdefault(clave, {
            'cliente': cliente,
            'campania': campania,
            'year': year,
            'mes': mes,
            'porcentaje_cumplimiento': porcentaje,
            'cumplimiento_explicito': cumplimiento_explicito,
            'horas_requeridas_total': horas_requeridas_total,
            'horas_proyectadas_total': horas_proyectadas_total,
            'jornadas': [],
        })
        if horas_requeridas_total is not None:
            grupos[clave]['horas_requeridas_total'] = horas_requeridas_total
        if horas_proyectadas_total is not None:
            grupos[clave]['horas_proyectadas_total'] = horas_proyectadas_total
        if cumplimiento_explicito:
            grupos[clave]['porcentaje_cumplimiento'] = porcentaje
            grupos[clave]['cumplimiento_explicito'] = True
        grupos[clave]['jornadas'].append({
            'dotacion_requerida': dotacion,
            'carga_semanal': carga_semanal,
            'carga_horaria': carga_horaria,
        })

    if errores:
        return jsonify({'success': False, 'errores': errores[:20]}), 400
    if not grupos:
        return jsonify({'success': False, 'errores': ['No hay filas validas para importar']}), 400

    guardadas = []
    antes_snapshots = []
    years = set()
    for valores in grupos.values():
        normalizado = payload_proyeccion(valores)
        if normalizado['errores']:
            errores.extend(normalizado['errores'])
            continue
        datos = normalizado['valores']
        feriados = fechas_feriadas_activas(datos['year'])
        proyeccion = ProyeccionMatriz.query.filter_by(
            cliente=datos['cliente'],
            campania=datos['campania'],
            mes=datos['mes'],
        ).first()
        antes_snapshots.append(snapshot_proyeccion(proyeccion))
        if not proyeccion:
            proyeccion = ProyeccionMatriz()
            db.session.add(proyeccion)
        aplicar_calculo_proyeccion(proyeccion, datos, datos['mes'], feriados)
        horas_requeridas_total = valores.get('horas_requeridas_total')
        horas_proyectadas_total = valores.get('horas_proyectadas_total')
        if horas_requeridas_total is not None and horas_requeridas_total >= 0:
            proyeccion.horas_requeridas = horas_requeridas_total
            if (
                horas_proyectadas_total is not None
                and horas_requeridas_total > 0
                and not valores.get('cumplimiento_explicito')
            ):
                proyeccion.porcentaje_cumplimiento = (horas_proyectadas_total / horas_requeridas_total) * 100
        elif horas_proyectadas_total is not None and horas_proyectadas_total >= 0:
            proyeccion.horas_requeridas = horas_proyectadas_total
            proyeccion.porcentaje_cumplimiento = 100
        guardadas.append(proyeccion)
        years.add(datos['year'])

    if errores:
        return jsonify({'success': False, 'errores': errores[:20]}), 400
    db.session.flush()
    despues_snapshots = [snapshot_proyeccion(proyeccion) for proyeccion in guardadas]
    registrar_historial(
        'edicion' if any(antes_snapshots) else 'creacion',
        'proyeccion_matriz',
        ','.join(str(year) for year in sorted(years)),
        f'Importacion de proyecciones: {len(guardadas)} registro(s)',
        antes={'proyecciones': antes_snapshots},
        despues={'proyecciones': despues_snapshots},
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': f'{len(guardadas)} proyeccion(es) importada(s)',
        'years': sorted(years),
    })


@main_bp.route('/api/alertas', methods=['GET'])
@login_requerido
def api_alertas():
    """Alertas automaticas segun los datos filtrados."""
    registros = registros_filtrados()
    alertas = []

    if not registros:
        asignaciones = asignaciones_maestras_filtradas(filtros_request())
        if asignaciones:
            nombres = ', '.join(sorted({a.cliente for a in asignaciones})[:5])
            alertas.append({
                'tipo': 'info',
                'titulo': 'Asignación activa sin facturación',
                'detalle': f'{nombres}: existe en Datos Maestros pero todavía no tiene facturación para el período seleccionado.'
            })
            return jsonify({'success': True, 'alertas': alertas})
        alertas.append({
            'tipo': 'info',
            'titulo': 'Sin datos para el filtro',
            'detalle': 'No hay registros que coincidan con la seleccion actual.'
        })
        return jsonify({'success': True, 'alertas': alertas})

    resumen = resumen_registros(registros)
    if resumen['porcentaje_cumplimiento'] < 95:
        alertas.append({
            'tipo': 'danger',
            'titulo': 'Cumplimiento bajo',
            'detalle': f"El cumplimiento general es {resumen['porcentaje_cumplimiento']}%."
        })
    elif resumen['porcentaje_cumplimiento'] < 100:
        alertas.append({
            'tipo': 'warning',
            'titulo': 'Cumplimiento debajo del objetivo',
            'detalle': f"El cumplimiento general esta en {resumen['porcentaje_cumplimiento']}%."
        })

    if resumen['desvio'] < 0:
        alertas.append({
            'tipo': 'danger',
            'titulo': 'Desvio negativo',
            'detalle': f"El desvio acumulado es {round(resumen['desvio'], 2)}."
        })

    exceso_horas = [r for r in registros if r.horas_facturadas > r.horas_objetivo]
    if exceso_horas:
        alertas.append({
            'tipo': 'warning',
            'titulo': 'Horas facturadas sobre objetivo',
            'detalle': f"{len(exceso_horas)} registro(s) superan las horas objetivo."
        })

    grupos = {}
    for registro in registros:
        grupos.setdefault(registro.cliente, []).append(registro)
    clientes_en_riesgo = []
    for cliente, registros_cliente in grupos.items():
        resumen_cliente = resumen_registros(registros_cliente)
        if resumen_cliente['porcentaje_cumplimiento'] < 95:
            clientes_en_riesgo.append({'cliente': cliente, **resumen_cliente})
    if clientes_en_riesgo:
        nombres = ', '.join(item['cliente'] for item in clientes_en_riesgo[:3])
        alertas.append({
            'tipo': 'danger',
            'titulo': 'Clientes en riesgo',
            'detalle': f"{nombres} por debajo de 95% de cumplimiento."
        })

    if not alertas:
        alertas.append({
            'tipo': 'success',
            'titulo': 'Operacion dentro de objetivo',
            'detalle': 'No se detectaron desvios criticos en la seleccion actual.'
        })

    return jsonify({'success': True, 'alertas': alertas})


@main_bp.route('/api/exportar_excel', methods=['GET'])
@login_requerido
def api_exportar_excel():
    """Exporta los registros filtrados en un archivo compatible con Excel."""
    registros = aplicar_filtros(Facturacion2026.query, **filtros_request()).order_by(
        Facturacion2026.fecha.desc()
    ).all()

    headers = [
        'Fecha', 'Mes', 'Cliente', 'Gerente', 'Jefe de Site', 'Campaña', 'Sub campaña', 'Tipo de negocio', 'Tipo de VH', 'Horas objetivo',
        'Horas facturadas', 'Horas Penalizacion ADH', 'Valor hora objetivo', 'Valor hora alcanzado', 'Facturado en horas manual', 'Tarificacion', 'Importe fijo facturado', 'Variable Objetivo', 'Variable Productivo',
        '% cumplimiento horas',
        'Objetivo facturacion horas', 'Objetivo facturacion bono', 'Facturacion objetivo',
        'Facturado horas', 'Facturado bono', 'Variable Productivo', 'Penalizaciones por incumplimientos',
        'Next Gen', 'Otros', 'Total facturado', 'Desvio', '% cumplimiento'
    ]
    rows = []
    for r in registros:
        rows.append([
            r.fecha.isoformat(), r.mes, r.cliente, r.gerente or '', r.jefe_site or '',
            r.campania or '', r.subcampania or '', r.tipo_negocio or '', r.tipo_jornada,
            r.horas_objetivo, r.horas_facturadas, r.horas_penalizadas or 0,
            r.valor_hora_objetivo if r.valor_hora_objetivo else r.valor_hora,
            r.valor_hora_alcanzado,
            r.facturado_horas_manual if r.facturado_horas_manual is not None else '',
            r.tarifacion or 0,
            r.importe_fijo if r.importe_fijo is not None else '',
            r.variable_objetivo or 0,
            r.variable_productivo or 0,
            round(r.porcentaje_cumplimiento_horas, 2),
            round(r.objetivo_facturacion_horas, 2), round(r.objetivo_facturacion_bono, 2),
            round(r.facturacion_objetivo, 2), round(r.facturado_horas, 2),
            round(r.facturado_bono, 2), round(r.variable_productivo_calculo, 2), round(r.penalizaciones_incumplimientos, 2),
            r.netx_gen or 0, r.otros or 0,
            round(r.total_dashboard, 2),
            round(r.total_dashboard - r.total_teorico, 2),
            round((r.total_dashboard / r.total_teorico * 100) if r.total_teorico > 0 else 0, 2)
        ])

    table_rows = ['<tr>' + ''.join(f'<th>{escape(h)}</th>' for h in headers) + '</tr>']
    for row in rows:
        table_rows.append('<tr>' + ''.join(f'<td>{escape(str(value))}</td>' for value in row) + '</tr>')

    html = (
        '<html><head><meta charset="utf-8"></head><body>'
        '<table border="1">'
        + ''.join(table_rows) +
        '</table></body></html>'
    )
    filename = f"facturacion_{datetime.now().strftime('%Y%m%d_%H%M')}.xls"
    return Response(
        html,
        content_type='application/vnd.ms-excel; charset=utf-8',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )


@main_bp.route('/api/template_carga', methods=['GET'])
@carga_requerida
def api_template_carga():
    """Descarga una plantilla xlsx para carga masiva."""
    headers = [label for _, label in COLUMNAS_IMPORTACION]
    ayuda_por_campo = {
        'cliente': 'Debe existir en Datos Maestros',
        'subcampania': 'Debe hacer match único con Cuenta',
        'tipo_jornada': 'Diurna, Nocturna, Feriado, Capacitación, Diurnas Feriado, Nocturnas Feriado, horas líder o radio',
        'mes': 'YYYY-MM, MM/YYYY o fecha Excel',
        'horas_objetivo': 'Número',
        'horas_facturadas': 'Número',
        'valor_hora': 'Número',
        'facturado_horas_manual': 'Opcional; si está vacío se calcula Horas × Valor Hora',
        'unidad': 'Base para calcular Variable productivo',
        'porcentaje_variable': 'Porcentaje; se guarda para la campaña y el mes',
        'variable_productivo': 'Opcional; si está vacío se calcula Unidad × % Variable',
        'ajuste_bono_penalizacion': 'Positivo = bono; negativo = penalización',
        'total_facturado_informado': 'Control: debe coincidir con el total calculado',
    }
    ayuda = [ayuda_por_campo.get(campo, 'Opcional') for campo, _ in COLUMNAS_IMPORTACION]
    contenido = crear_xlsx(headers, [ayuda])
    return Response(
        contenido,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={
            'Content-Disposition': 'attachment; filename="plantilla_carga_facturacion_tipo_negocio.xlsx"',
            'Cache-Control': 'no-store, no-cache, must-revalidate, max-age=0',
            'Pragma': 'no-cache',
        }
    )


@main_bp.route('/api/importar_datos', methods=['POST'])
@carga_requerida
def api_importar_datos():
    """Importa registros desde la plantilla Excel/CSV."""
    archivo = request.files.get('archivo')
    if not archivo:
        return jsonify({'success': False, 'errores': ['Debe adjuntar un archivo']}), 400

    try:
        contenido_bytes = archivo.read()
        nombre = (archivo.filename or '').lower()
        if nombre.endswith('.xlsx'):
            filas = filas_desde_xlsx(contenido_bytes)
        else:
            try:
                contenido = contenido_bytes.decode('utf-8-sig')
            except UnicodeDecodeError:
                contenido = contenido_bytes.decode('cp1252')
            if nombre.endswith('.xls') and '<table' not in contenido.lower():
                return jsonify({
                    'success': False,
                    'errores': ['Ese archivo .xls no es la plantilla esperada. Descargue nuevamente el template .xlsx, complete las filas y vuelva a subirlo.']
                }), 400
            filas = filas_desde_html(contenido) if nombre.endswith('.xls') or '<table' in contenido.lower() else filas_desde_csv(contenido)

        datos = datos_desde_filas(filas)
        if not datos:
            return jsonify({'success': False, 'errores': ['No se encontraron filas para importar. Revise que el archivo tenga encabezados y datos.']}), 400

        confirmar_reemplazo = request.form.get('confirmar_reemplazo') == '1'
        try:
            resoluciones = json.loads(request.form.get('resoluciones') or '{}')
        except json.JSONDecodeError:
            resoluciones = {}
        creados = 0
        reemplazados = 0
        errores = []
        items_validos = []
        conflictos_por_clave = {}
        porcentajes_por_clave = {}
        resoluciones_pendientes = []
        for indice, item in enumerate(datos, start=2):
            if indice == 2 and (
                str(item.get('fecha', '')).startswith('YYYY')
                or str(item.get('cliente', '')).startswith('Debe existir en Datos Maestros')
            ):
                continue

            # Las planillas operativas pueden traer meses futuros ya armados con
            # nombres y tipos, pero sin datos. No deben convertirse en registros
            # de facturación en cero.
            campos_numericos_importacion = (
                'horas_objetivo', 'horas_facturadas', 'valor_hora',
                'facturado_horas_manual', 'tarifacion', 'unidad',
                'porcentaje_variable', 'variable_productivo',
                'ajuste_bono_penalizacion', 'netx_gen', 'otros',
                'total_facturado_informado',
            )
            if all(item.get(campo) in (None, '') for campo in campos_numericos_importacion):
                continue

            try:
                preparar_fila_migracion_facturacion(item)
            except (TypeError, ValueError) as exc:
                errores.append(f'Fila {indice}: valores inválidos: {exc}')
                continue

            item.setdefault('bonos', 0)
            item.setdefault('horas_penalizadas', 0)
            item['penalizaciones'] = normalizar_importe_ajuste('penalizaciones', item.get('penalizaciones'))
            item.setdefault('netx_gen', 0)
            item.setdefault('otros', 0)
            item.setdefault('tarifacion', None)
            item.setdefault('importe_fijo', None)
            item.setdefault('variable_objetivo', 0)
            item.setdefault('variable_productivo', 0)
            item.setdefault('valor_hora_objetivo', item.get('valor_hora'))

            cuenta_original = str(item.get('cliente') or '').strip()
            subcampania_original = str(item.get('subcampania') or '').strip()
            resolucion = resoluciones.get(str(indice)) or {}
            if item.get('_importacion_directa_personal'):
                _, error_asignacion = resolver_personal_importacion_directa(item)
            elif resolucion.get('accion') == 'vincular':
                asignacion_resuelta = db.session.get(AsignacionComercial, int(resolucion.get('asignacion_id') or 0))
                if not asignacion_resuelta or not asignacion_resuelta.activa:
                    errores.append(f'Fila {indice}: la asociación elegida ya no está disponible')
                    continue
                aplicar_asignacion_a_fila(item, asignacion_resuelta)
                error_asignacion = None
            elif resolucion.get('accion') == 'crear':
                gerente = str(resolucion.get('gerente') or '').strip()
                jefe_site = str(resolucion.get('jefe_site') or '').strip()
                if not gerente or not jefe_site:
                    errores.append(f'Fila {indice}: para crear el maestro debe indicar gerente y jefe de site')
                    continue
                cliente_nuevo = str(resolucion.get('cliente') or cuenta_original).strip()
                campania_nueva = str(resolucion.get('campania') or cuenta_original).strip()
                campos_nuevos = {
                    'cliente': cliente_nuevo,
                    'gerente': gerente,
                    'jefe_site': jefe_site,
                    'campania': campania_nueva,
                    'subcampania': subcampania_original,
                    'tipo_negocio': str(resolucion.get('tipo_negocio') or '').strip() or None,
                    'es_next_gen': bool(item.get('_es_next_gen_sin_horas')),
                }
                asignacion_resuelta, _ = obtener_o_crear_asignacion(campos_nuevos)
                aplicar_asignacion_a_fila(item, asignacion_resuelta)
                error_asignacion = None
            else:
                _, error_asignacion = resolver_asignacion_importacion(item)
            if error_asignacion:
                resoluciones_pendientes.append({
                    'fila': indice,
                    'cuenta': cuenta_original,
                    'subcampania': subcampania_original,
                    'mensaje': error_asignacion,
                    'opciones': opciones_maestro_para_fila(item),
                })
                continue

            if item.get('_es_next_gen_sin_horas'):
                item['es_next_gen'] = True
                item['tipo_jornada'] = 'NextGen'

            error_total = validar_total_fila_migracion(item)
            if error_total:
                errores.append(f'Fila {indice}: {error_total}')
                continue
            error_variable = validar_variable_fila_migracion(item)
            if error_variable:
                errores.append(f'Fila {indice}: {error_variable}')
                continue
            if item.get('porcentaje_variable') not in (None, ''):
                clave_porcentaje = (item['cliente'], item['campania'], item['mes'])
                porcentaje = parse_numero(item.get('porcentaje_variable'))
                anterior = porcentajes_por_clave.get(clave_porcentaje)
                if anterior is not None and abs(anterior - porcentaje) > 0.0001:
                    errores.append(
                        f'Fila {indice}: el % Variable {porcentaje} contradice otro valor {anterior} '
                        f'para {item["cliente"]} / {item["campania"]} / {item["mes"]}'
                    )
                    continue
                porcentajes_por_clave[clave_porcentaje] = porcentaje

            errores_fila = validar_payload_facturacion(
                item,
                exigir_configuracion_valor_hora=False,
                preservar_valores_importados=True,
            )
            if errores_fila:
                errores.append(f'Fila {indice}: ' + '; '.join(errores_fila))
                continue
            try:
                existentes = query_reemplazo_importacion(item).all()
                if existentes:
                    clave = clave_reemplazo_importacion(item)
                    conflictos_por_clave.setdefault(
                        clave,
                        detalle_conflicto_importacion(indice, item, existentes)
                    )
                items_validos.append((indice, item, existentes))
            except Exception as exc:
                errores.append(f'Fila {indice}: {exc}')

        if resoluciones_pendientes:
            db.session.rollback()
            return jsonify({
                'success': False,
                'requiere_resolucion_maestros': True,
                'mensaje': 'Hay servicios que necesitan vincularse o darse de alta.',
                'pendientes': resoluciones_pendientes,
                'errores': errores[:50],
            }), 409

        if errores:
            db.session.rollback()
            return jsonify({'success': False, 'creados': 0, 'errores': errores[:50]}), 400

        if conflictos_por_clave and not confirmar_reemplazo:
            return jsonify({
                'success': False,
                'requiere_confirmacion': True,
                'mensaje': 'La importacion reemplazaria datos existentes. Confirme para continuar.',
                'conflictos': list(conflictos_por_clave.values()),
                'errores': ['Hay datos existentes que coinciden por fecha, cliente, gerente, jefe de site, campania y sub campania.']
            }), 409

        grupos_reemplazo = {}
        items_nuevos = []
        for indice, item, existentes in items_validos:
            if existentes and confirmar_reemplazo:
                clave = clave_reemplazo_importacion(item)
                grupo = grupos_reemplazo.setdefault(clave, {
                    'items': [], 'existentes': {registro.id: registro for registro in existentes},
                })
                grupo['items'].append((indice, item))
                grupo['existentes'].update({registro.id: registro for registro in existentes})
            else:
                items_nuevos.append((indice, item))

        # Una clave puede representar varias líneas aditivas de la planilla.
        # Se reemplaza el conjunto completo de forma atómica para que una
        # reimportación de agosto no colapse dos líneas en una ni duplique total.
        for clave, grupo in grupos_reemplazo.items():
            existentes = list(grupo['existentes'].values())
            antes_grupo = [snapshot_modelo(registro) for registro in existentes]
            try:
                for registro in existentes:
                    db.session.delete(registro)
                db.session.flush()
                nuevos = []
                for indice, item in grupo['items']:
                    nuevo = crear_registro_facturacion(
                        item,
                        exigir_configuracion_valor_hora=False,
                        preservar_facturado_manual=True,
                    )
                    guardar_porcentaje_variable_importado(item)
                    nuevos.append(nuevo)
                db.session.flush()
                registrar_historial(
                    'edicion', 'facturacion', existentes[0].id,
                    f'Grupo de facturación reemplazado: {grupo["items"][0][1]["cliente"]} / {grupo["items"][0][1]["mes"]}',
                    antes={'filas': antes_grupo},
                    despues={'filas': [snapshot_modelo(registro) for registro in nuevos]},
                    detalle=f'{len(existentes)} fila(s) anteriores reemplazadas por {len(nuevos)} fila(s) importadas.',
                )
                reemplazados += len(nuevos)
            except Exception as exc:
                primera_fila = grupo['items'][0][0]
                errores.append(f'Fila {primera_fila}: {exc}')

        for indice, item in items_nuevos:
            try:
                crear_registro_facturacion(
                    item,
                    exigir_configuracion_valor_hora=False,
                    preservar_facturado_manual=True,
                )
                guardar_porcentaje_variable_importado(item)
                creados += 1
            except Exception as exc:
                errores.append(f'Fila {indice}: {exc}')

        if errores:
            db.session.rollback()
            return jsonify({'success': False, 'creados': 0, 'errores': errores[:50]}), 400

        if creados == 0 and reemplazados == 0:
            db.session.rollback()
            return jsonify({
                'success': False,
                'creados': 0,
                'errores': [
                    'No se importó ningún registro. La plantilla descargada viene vacía: complete los datos desde la primera fila debajo del encabezado y vuelva a subirla.'
                ]
            }), 400

        db.session.commit()
        partes_mensaje = []
        if creados:
            partes_mensaje.append(f'{creados} creado(s)')
        if reemplazados:
            partes_mensaje.append(f'{reemplazados} reemplazado(s)')
        return jsonify({
            'success': True,
            'creados': creados,
            'reemplazados': reemplazados,
            'mensaje': 'Importacion completada: ' + ', '.join(partes_mensaje)
        })
    except zipfile.BadZipFile:
        db.session.rollback()
        return jsonify({'success': False, 'errores': ['No pude leer el Excel. Descargue nuevamente el template .xlsx y vuelva a completarlo.']}), 400
    except Exception as exc:
        db.session.rollback()
        return jsonify({'success': False, 'errores': [str(exc)]}), 500


@main_bp.route('/api/excepciones-calculo', methods=['GET'])
@edicion_requerida
def api_excepciones_calculo():
    excepciones = ExcepcionCalculo.query.order_by(ExcepcionCalculo.cliente, ExcepcionCalculo.campania).all()
    return jsonify({'success': True, 'excepciones': [item.to_dict() for item in excepciones]})


@main_bp.route('/api/excepciones-calculo', methods=['POST'])
@edicion_requerida
def api_crear_excepcion_calculo():
    data = request.get_json() or {}
    cliente = str(data.get('cliente') or '').strip()
    campania = str(data.get('campania') or '').strip()
    tipo = str(data.get('tipo_calculo') or 'facturado_horas_manual').strip()
    try:
        ajuste_objetivo = parse_numero(data.get('ajuste_vh_objetivo_pct'))
        ajuste_alcanzado = parse_numero(data.get('ajuste_vh_alcanzado_pct'))
    except ValueError:
        return jsonify({'success': False, 'errores': ['Los porcentajes de ajuste no son válidos']}), 400
    if not cliente or not campania:
        return jsonify({'success': False, 'errores': ['Cliente y campaña son obligatorios']}), 400
    if tipo not in ('facturado_horas_manual', 'facturado_manual_ajustes_vh'):
        return jsonify({'success': False, 'errores': ['El tipo de cálculo no es válido']}), 400
    excepcion = ExcepcionCalculo.query.filter(
        func.lower(ExcepcionCalculo.cliente) == cliente.lower(),
        func.lower(ExcepcionCalculo.campania) == campania.lower(),
    ).first()
    es_nueva = excepcion is None
    if excepcion:
        excepcion.tipo_calculo = tipo
        excepcion.ajuste_vh_objetivo_pct = ajuste_objetivo
        excepcion.ajuste_vh_alcanzado_pct = ajuste_alcanzado
        excepcion.activa = True
    else:
        excepcion = ExcepcionCalculo(cliente=cliente, campania=campania, tipo_calculo=tipo, ajuste_vh_objetivo_pct=ajuste_objetivo, ajuste_vh_alcanzado_pct=ajuste_alcanzado)
        db.session.add(excepcion)
    db.session.flush()
    registrar_historial(
        'creacion' if es_nueva else 'edicion',
        'excepcion_calculo',
        excepcion.id,
        f'Excepción de cálculo: {cliente} / {campania}',
        despues=excepcion.to_dict(),
    )
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Excepción guardada', 'excepcion': excepcion.to_dict()})


@main_bp.route('/api/excepciones-calculo/<int:excepcion_id>', methods=['DELETE'])
@eliminacion_requerida
def api_eliminar_excepcion_calculo(excepcion_id):
    excepcion = db.session.get(ExcepcionCalculo, excepcion_id)
    if not excepcion:
        return jsonify({'success': False, 'errores': ['La excepción no existe']}), 404
    antes = excepcion.to_dict()
    db.session.delete(excepcion)
    registrar_historial(
        'eliminacion', 'excepcion_calculo', excepcion_id,
        f'Excepción eliminada: {antes["cliente"]} / {antes["campania"]}', antes=antes,
    )
    db.session.commit()
    return jsonify({'success': True, 'mensaje': 'Excepción eliminada'})


@main_bp.route('/api/asignaciones', methods=['GET'])
@carga_requerida
def api_asignaciones():
    """Lista de asociaciones predefinidas para carga y filtros."""
    solo_activas = request.args.get('activas', '1') != '0'
    query = AsignacionComercial.query
    if solo_activas:
        query = query.filter_by(activa=True)
    asignaciones = query.order_by(
        AsignacionComercial.cliente,
        AsignacionComercial.campania,
        AsignacionComercial.subcampania
    ).all()
    return jsonify({
        'success': True,
        'asignaciones': [asignacion.to_dict() for asignacion in asignaciones]
    })


@main_bp.route('/api/campanias', methods=['GET'])
@edicion_requerida
def api_campanias():
    """Lista el catálogo canónico y su ID estable por cliente + campaña."""
    solo_activas = request.args.get('activas', '1') != '0'
    query = Campania.query
    if solo_activas:
        query = query.filter_by(activa=True)
    campanias = query.order_by(Campania.cliente, Campania.nombre).all()
    return jsonify({
        'success': True,
        'campanias': [campania.to_dict() for campania in campanias],
    })


@main_bp.route('/api/campanias/<int:campania_id>/valor-hora', methods=['PATCH'])
@edicion_requerida
def api_configurar_valor_hora_campania(campania_id):
    campania = db.session.get(Campania, campania_id)
    if not campania:
        return jsonify({'success': False, 'errores': ['La campaña no existe']}), 404
    data = request.get_json() or {}
    valor_hora_variable = data.get('valor_hora_variable')
    if valor_hora_variable is not None and not isinstance(valor_hora_variable, bool):
        return jsonify({'success': False, 'errores': ['Debe indicar Sí, No o Sin configurar']}), 400
    antes = campania.to_dict()
    campania.valor_hora_variable = valor_hora_variable
    db.session.flush()
    despues = campania.to_dict()
    registrar_historial(
        'edicion', 'campania', campania.id,
        f'Configuración de valor hora: {campania.cliente} / {campania.nombre}',
        antes=antes, despues=despues,
        detalle=(
            'Variable' if campania.valor_hora_variable is True
            else 'Repite valor hora objetivo' if campania.valor_hora_variable is False
            else 'Sin configurar'
        ),
    )
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': 'Configuración de valor hora guardada',
        'campania': campania.to_dict(),
    })


@main_bp.route('/api/asignaciones/<int:asignacion_id>', methods=['GET'])
@edicion_requerida
def api_obtener_asignacion(asignacion_id):
    """Obtiene una asociacion puntual desde la base."""
    asignacion = AsignacionComercial.query.get_or_404(asignacion_id)
    return jsonify({'success': True, 'asignacion': asignacion.to_dict()})


@main_bp.route('/api/asignaciones', methods=['POST'])
@edicion_requerida
def api_crear_asignacion():
    """Alta de una asociacion cliente/gerente/campaña/sub campaña."""
    return crear_asignacion_desde_request()


@main_bp.route('/api/catalogos/asignaciones', methods=['POST'])
@edicion_requerida
def api_crear_asignacion_catalogos():
    """Alta de datos maestros desde la pantalla de catalogos."""
    return crear_asignacion_desde_request()


def crear_asignacion_desde_request():
    try:
        data = request.get_json(silent=True) or request.form.to_dict() or {}
        es_next_gen = str(data.get('es_next_gen') or '').lower() in ('1', 'true', 'on', 'si', 'sí')
        nombre_next_gen = str(data.get('nombre_next_gen') or data.get('cliente') or '').strip()
        campos = {
            'cliente': str(data.get('cliente') or '').strip(),
            'gerente': str(data.get('gerente') or '').strip(),
            'jefe_site': str(data.get('jefe_site') or '').strip(),
            'campania': str(data.get('campania') or '').strip(),
            'subcampania': str(data.get('subcampania') or '').strip(),
            'tipo_negocio': str(data.get('tipo_negocio') or '').strip() or None,
            'es_next_gen': es_next_gen,
        }
        if es_next_gen:
            campos.update(cliente=nombre_next_gen, gerente='NextGen', jefe_site='NextGen',
                          campania=nombre_next_gen, subcampania=nombre_next_gen, tipo_negocio='NextGen')
        errores = [f'{campo} es obligatorio' for campo, valor in campos.items() if campo not in ('tipo_negocio', 'es_next_gen') and not valor]
        if errores:
            return jsonify({'success': False, 'errores': errores}), 400

        existente = AsignacionComercial.query.filter_by(**campos).first()
        if existente:
            existente.activa = True
            if not existente.campania_id:
                existente.campania_catalogo = obtener_o_crear_campania(campos['cliente'], campos['campania'])
            db.session.commit()
            return jsonify({'success': True, 'mensaje': 'La asociacion ya existia y quedo activa', 'asignacion': existente.to_dict()})

        campania = obtener_o_crear_campania(campos['cliente'], campos['campania'])
        asignacion = AsignacionComercial(**campos, campania_catalogo=campania)
        db.session.add(asignacion)
        db.session.flush()
        try:
            registrar_historial(
                'creacion',
                'asignacion',
                asignacion.id,
                f'Asociacion creada: {asignacion.label}',
                despues=asignacion.to_dict(),
            )
        except Exception as historial_error:
            current_app.logger.exception('No se pudo registrar historial de asignacion: %s', historial_error)
        db.session.commit()
        return jsonify({'success': True, 'mensaje': 'Asociacion creada correctamente', 'asignacion': asignacion.to_dict()})
    except Exception as exc:
        db.session.rollback()
        current_app.logger.exception('Error creando asignacion')
        return jsonify({'success': False, 'errores': [f'No se pudo crear la asociacion: {exc}']}), 500


@main_bp.route('/api/asignaciones/<int:asignacion_id>', methods=['PATCH', 'PUT'])
@edicion_requerida
def api_actualizar_asignacion(asignacion_id):
    data = request.get_json() or {}
    if not validar_confirmacion_accion(data):
        return jsonify({'success': False, 'errores': ['La confirmacion no es valida']}), 403

    asignacion = AsignacionComercial.query.get_or_404(asignacion_id)
    if data.get('es_next_gen'):
        nombre_next_gen = str(data.get('nombre_next_gen') or data.get('cliente') or asignacion.cliente).strip()
        data.update(cliente=nombre_next_gen, gerente='NextGen', jefe_site='NextGen',
                    campania=nombre_next_gen, subcampania=nombre_next_gen, tipo_negocio='NextGen', es_next_gen=True)
    antes = snapshot_modelo(asignacion)
    valores_anteriores = filtros_asignacion(asignacion)
    nuevos_valores = {}
    campos_asociacion = ['cliente', 'gerente', 'jefe_site', 'campania', 'subcampania', 'tipo_negocio', 'es_next_gen']
    vigencia_desde_raw = data.get('vigencia_desde')
    vigencia_desde = mes_valido(vigencia_desde_raw) if vigencia_desde_raw else None

    for campo in campos_asociacion:
        if campo in data:
            if campo == 'es_next_gen':
                nuevos_valores[campo] = bool(data.get(campo))
                continue
            valor = str(data.get(campo) or '').strip()
            if campo == 'tipo_negocio':
                nuevos_valores[campo] = valor or None
                continue
            if not valor:
                return jsonify({'success': False, 'errores': [f'{campo} es obligatorio']}), 400
            nuevos_valores[campo] = valor

    if nuevos_valores:
        valores_finales = {**valores_anteriores, **nuevos_valores}
        if vigencia_desde_raw and not vigencia_desde:
            return jsonify({'success': False, 'errores': ['El mes de vigencia debe tener formato YYYY-MM']}), 400
        if vigencia_desde:
            asignacion_nueva, creada = obtener_o_crear_asignacion(valores_finales)
            anio_vigencia, mes_vigencia = (int(parte) for parte in vigencia_desde.split('-'))
            if mes_vigencia == 1:
                vigencia_anterior = f'{anio_vigencia - 1:04d}-12'
            else:
                vigencia_anterior = f'{anio_vigencia:04d}-{mes_vigencia - 1:02d}'

            # La identidad comercial permanece; gerente y jefe pueden cambiar.
            # Cerramos cualquier versión que cubra el mes anterior y abrimos la
            # versión seleccionada desde la fecha indicada.
            versiones = AsignacionComercial.query.filter_by(
                cliente=valores_finales['cliente'],
                campania=valores_finales['campania'],
                subcampania=valores_finales['subcampania'],
                tipo_negocio=valores_finales.get('tipo_negocio'),
                es_next_gen=bool(valores_finales.get('es_next_gen')),
                activa=True,
            ).all()
            for version in versiones:
                if version.id == asignacion_nueva.id:
                    continue
                if (not version.vigencia_desde or version.vigencia_desde <= vigencia_anterior) and (
                    not version.vigencia_hasta or version.vigencia_hasta >= vigencia_desde
                ):
                    version.vigencia_hasta = vigencia_anterior
            asignacion_nueva.vigencia_desde = vigencia_desde
            asignacion_nueva.vigencia_hasta = None
            # Se identifica el servicio por sus dimensiones estables. Así una
            # segunda corrección de vigencia también alcanza filas que ya habían
            # recibido otro gerente o jefe en una edición anterior.
            registros_asociados = Facturacion2026.query.filter_by(
                cliente=valores_finales['cliente'],
                campania=valores_finales['campania'],
                subcampania=valores_finales['subcampania'],
                tipo_negocio=valores_finales.get('tipo_negocio'),
                es_next_gen=bool(valores_finales.get('es_next_gen')),
            ).filter(Facturacion2026.mes >= vigencia_desde).all()
            for registro in registros_asociados:
                for campo, valor in valores_finales.items():
                    setattr(registro, campo, valor)

            despues = snapshot_modelo(asignacion_nueva)
            cambios = cambios_entre(antes, despues)
            detalle = {
                'vigencia_desde': vigencia_desde,
                'asignacion_origen_id': asignacion.id,
                'registros_actualizados': len(registros_asociados),
                'cambios': cambios,
            }
            registrar_historial(
                'creacion' if creada else 'edicion',
                'asignacion',
                asignacion_nueva.id,
                f'Asociacion vigente desde {vigencia_desde}: {asignacion_nueva.label}',
                antes=antes,
                despues=despues,
                detalle=serializar_json(detalle),
            )
            db.session.commit()
            db.session.refresh(asignacion_nueva)
            return jsonify({
                'success': True,
                'mensaje': f'Asociacion aplicada desde {vigencia_desde}. El historico anterior no se modifico.',
                'asignacion': asignacion_nueva.to_dict(),
                'asignacion_original': antes,
                'vigencia_desde': vigencia_desde,
                'registros_actualizados': len(registros_asociados),
                'cambios': cambios,
            })

        duplicada = AsignacionComercial.query.filter_by(**valores_finales).filter(
            AsignacionComercial.id != asignacion.id
        ).first()
        if duplicada:
            return jsonify({'success': False, 'errores': ['Ya existe otra asociación con esos datos']}), 400

        registros_asociados = Facturacion2026.query.filter_by(**valores_anteriores).all()
        for registro in registros_asociados:
            for campo, valor in valores_finales.items():
                setattr(registro, campo, valor)
        for campo, valor in valores_finales.items():
            setattr(asignacion, campo, valor)
        asignacion.campania_catalogo = obtener_o_crear_campania(
            valores_finales['cliente'], valores_finales['campania']
        )
    else:
        registros_asociados = []

    if 'activa' in data:
        asignacion.activa = bool(data['activa'])
    db.session.flush()
    despues = snapshot_modelo(asignacion)
    cambios = cambios_entre(antes, despues)
    if cambios:
        registrar_historial(
            'edicion',
            'asignacion',
            asignacion.id,
            f'Asociacion actualizada: {asignacion.label}',
            antes=antes,
            despues=despues,
            detalle=serializar_json(cambios),
        )
    db.session.commit()
    db.session.refresh(asignacion)
    return jsonify({
        'success': True,
        'mensaje': 'Asociacion actualizada correctamente',
        'asignacion': asignacion.to_dict(),
        'registros_actualizados': len(registros_asociados),
        'cambios': cambios,
    })


@main_bp.route('/api/asignaciones/<int:asignacion_id>', methods=['DELETE'])
@eliminacion_requerida
def api_eliminar_asignacion(asignacion_id):
    data = request.get_json() or {}
    if not validar_confirmacion_accion(data):
        return jsonify({'success': False, 'errores': ['La confirmacion no es valida']}), 403

    asignacion = AsignacionComercial.query.get_or_404(asignacion_id)
    antes = snapshot_modelo(asignacion)
    registros_asociados = query_por_asignacion(asignacion).all()
    registrar_historial(
        'eliminacion',
        'asignacion',
        asignacion.id,
        f'Asociacion eliminada: {asignacion.label}',
        antes=antes,
        detalle=f'Se elimino la asociacion y {len(registros_asociados)} carga(s) vinculadas.',
    )
    for registro in registros_asociados:
        db.session.delete(registro)
    db.session.delete(asignacion)
    db.session.commit()
    return jsonify({
        'success': True,
        'mensaje': 'Asociación eliminada',
        'registros_eliminados': len(registros_asociados)
    })


@main_bp.route('/api/seed', methods=['POST'])
@admin_requerido
def api_seed():
    """Endpoint para cargar datos de ejemplo"""
    from app.models import Facturacion2026
    
    # Verificar si ya hay datos
    if Facturacion2026.query.first():
        return jsonify({'success': False, 'mensaje': 'Ya existen datos en la base de datos'}), 400
    
    datos_seed = [
        # Enero 2026
        {'fecha': '2026-01-05', 'cliente': 'Acme Corp', 'tipo_jornada': 'Diurna', 
         'horas_objetivo': 160, 'horas_facturadas': 165, 'valor_hora': 45, 'tarifacion': None, 
         'bonos': 200, 'penalizaciones': 0, 'netx_gen': 50, 'otros': 0},
        {'fecha': '2026-01-12', 'cliente': 'Tech Solutions', 'tipo_jornada': 'Nocturna', 
         'horas_objetivo': 80, 'horas_facturadas': 75, 'valor_hora': 50, 'tarifacion': None, 
         'bonos': 0, 'penalizaciones': 100, 'netx_gen': 0, 'otros': 25},
        {'fecha': '2026-01-20', 'cliente': 'Global Services', 'tipo_jornada': 'Feriado', 
         'horas_objetivo': 160, 'horas_facturadas': 160, 'valor_hora': 42, 'tarifacion': 48, 
         'bonos': 150, 'penalizaciones': 0, 'netx_gen': 0, 'otros': 0},
        
        # Febrero 2026
        {'fecha': '2026-02-03', 'cliente': 'Acme Corp', 'tipo_jornada': 'Capacitación', 
         'horas_objetivo': 160, 'horas_facturadas': 170, 'valor_hora': 45, 'tarifacion': None, 
         'bonos': 350, 'penalizaciones': 0, 'netx_gen': 100, 'otros': 0},
        {'fecha': '2026-02-10', 'cliente': 'Tech Solutions', 'tipo_jornada': 'Diurnas Feriado', 
         'horas_objetivo': 80, 'horas_facturadas': 82, 'valor_hora': 50, 'tarifacion': None, 
         'bonos': 50, 'penalizaciones': 0, 'netx_gen': 0, 'otros': 0},
        {'fecha': '2026-02-18', 'cliente': 'Innovate Ltd', 'tipo_jornada': 'Nocturnas Feriado', 
         'horas_objetivo': 160, 'horas_facturadas': 155, 'valor_hora': 55, 'tarifacion': None, 
         'bonos': 0, 'penalizaciones': 150, 'netx_gen': 0, 'otros': 50},
        
        # Marzo 2026
        {'fecha': '2026-03-02', 'cliente': 'Global Services', 'tipo_jornada': 'horas líder', 
         'horas_objetivo': 160, 'horas_facturadas': 168, 'valor_hora': 42, 'tarifacion': 48, 
         'bonos': 280, 'penalizaciones': 0, 'netx_gen': 75, 'otros': 0},
        {'fecha': '2026-03-09', 'cliente': 'Acme Corp', 'tipo_jornada': 'radio', 
         'horas_objetivo': 160, 'horas_facturadas': 160, 'valor_hora': 45, 'tarifacion': None, 
         'bonos': 100, 'penalizaciones': 0, 'netx_gen': 0, 'otros': 0},
        {'fecha': '2026-03-16', 'cliente': 'Tech Solutions', 'tipo_jornada': 'Diurna', 
         'horas_objetivo': 80, 'horas_facturadas': 78, 'valor_hora': 50, 'tarifacion': None, 
         'bonos': 0, 'penalizaciones': 50, 'netx_gen': 0, 'otros': 0},
        {'fecha': '2026-03-23', 'cliente': 'Digital Dynamics', 'tipo_jornada': 'Nocturna', 
         'horas_objetivo': 160, 'horas_facturadas': 172, 'valor_hora': 52, 'tarifacion': None, 
         'bonos': 400, 'penalizaciones': 0, 'netx_gen': 150, 'otros': 0},
        
        # Abril 2026
        {'fecha': '2026-04-06', 'cliente': 'Innovate Ltd', 'tipo_jornada': 'Feriado', 
         'horas_objetivo': 160, 'horas_facturadas': 158, 'valor_hora': 55, 'tarifacion': None, 
         'bonos': 0, 'penalizaciones': 55, 'netx_gen': 0, 'otros': 0},
        {'fecha': '2026-04-13', 'cliente': 'Acme Corp', 'tipo_jornada': 'Capacitación', 
         'horas_objetivo': 160, 'horas_facturadas': 175, 'valor_hora': 45, 'tarifacion': None, 
         'bonos': 500, 'penalizaciones': 0, 'netx_gen': 200, 'otros': 0},
        {'fecha': '2026-04-20', 'cliente': 'Global Services', 'tipo_jornada': 'Diurnas Feriado', 
         'horas_objetivo': 160, 'horas_facturadas': 160, 'valor_hora': 42, 'tarifacion': 48, 
         'bonos': 150, 'penalizaciones': 0, 'netx_gen': 0, 'otros': 0},
    ]
    gerentes_por_cliente = {
        'Acme Corp': 'Laura Gomez',
        'Tech Solutions': 'Martin Perez',
        'Global Services': 'Sofia Alvarez',
        'Innovate Ltd': 'Diego Torres',
        'Digital Dynamics': 'Valeria Ruiz',
    }
    jefes_site_por_cliente = {
        'Acme Corp': 'Mariana Silva',
        'Tech Solutions': 'Roberto Diaz',
        'Global Services': 'Carolina Mendez',
        'Innovate Ltd': 'Pablo Rios',
        'Digital Dynamics': 'Natalia Castro',
    }
    campanias_por_cliente = {
        'Acme Corp': ('Retencion', 'Empresas premium'),
        'Tech Solutions': ('Soporte', 'Mesa tecnica'),
        'Global Services': ('Operacion 2026', 'Backoffice regional'),
        'Innovate Ltd': ('Crecimiento', 'Nuevas cuentas'),
        'Digital Dynamics': ('Transformacion', 'Automatizacion'),
    }
    
    try:
        for d in datos_seed:
            fecha = datetime.strptime(d['fecha'], '%Y-%m-%d').date()
            mes = fecha.strftime('%Y-%m')
            
            registro = Facturacion2026(
                fecha=fecha,
                mes=mes,
                cliente=d['cliente'],
                gerente=d.get('gerente') or gerentes_por_cliente.get(d['cliente'], 'Sin asignar'),
                jefe_site=d.get('jefe_site') or jefes_site_por_cliente.get(d['cliente'], 'Sin asignar'),
                campania=d.get('campania') or campanias_por_cliente.get(d['cliente'], ('Operacion 2026', 'General'))[0],
                subcampania=d.get('subcampania') or campanias_por_cliente.get(d['cliente'], ('Operacion 2026', 'General'))[1],
                tipo_jornada=d['tipo_jornada'],
                horas_objetivo=d['horas_objetivo'],
                horas_facturadas=d['horas_facturadas'],
                valor_hora_objetivo=d.get('valor_hora_objetivo', d['valor_hora']),
                valor_hora=d['valor_hora'],
                tarifacion=d['tarifacion'],
                importe_fijo=d.get('importe_fijo'),
                variable_productivo=d.get('variable_productivo', 0),
                bonos=d['bonos'],
                penalizaciones=d['penalizaciones'],
                netx_gen=d['netx_gen'],
                otros=d['otros']
            )
            db.session.add(registro)

            campania, subcampania = campanias_por_cliente.get(d['cliente'], ('Operacion 2026', 'General'))
            existe_asignacion = AsignacionComercial.query.filter_by(
                cliente=d['cliente'],
                gerente=gerentes_por_cliente.get(d['cliente'], 'Sin asignar'),
                jefe_site=jefes_site_por_cliente.get(d['cliente'], 'Sin asignar'),
                campania=campania,
                subcampania=subcampania
            ).first()
            if not existe_asignacion:
                campania_catalogo = obtener_o_crear_campania(d['cliente'], campania)
                db.session.add(AsignacionComercial(
                    campania_catalogo=campania_catalogo,
                    cliente=d['cliente'],
                    gerente=gerentes_por_cliente.get(d['cliente'], 'Sin asignar'),
                    jefe_site=jefes_site_por_cliente.get(d['cliente'], 'Sin asignar'),
                    campania=campania,
                    subcampania=subcampania
                ))
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'mensaje': f'Se cargaron {len(datos_seed)} registros de ejemplo'
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'errores': [str(e)]}), 500


@main_bp.route('/api/clientes', methods=['GET'])
@login_requerido
def api_clientes():
    """Endpoint para obtener lista de clientes únicos"""
    clientes = aplicar_filtros(Facturacion2026.query).with_entities(
        Facturacion2026.cliente
    ).distinct().order_by(Facturacion2026.cliente).all()
    return jsonify({
        'success': True,
        'clientes': [c[0] for c in clientes]
    })


# Endpoint para obtener lista de gerentes únicos
@main_bp.route('/api/gerentes', methods=['GET'])
@login_requerido
def api_gerentes():
    gerentes = aplicar_filtros(Facturacion2026.query).with_entities(
        Facturacion2026.gerente
    ).distinct().order_by(Facturacion2026.gerente).all()
    return jsonify({
        'success': True,
        'gerentes': [g[0] for g in gerentes if g[0]]
    })


@main_bp.route('/api/meses', methods=['GET'])
@login_requerido
def api_meses():
    """Endpoint para obtener lista de meses disponibles"""
    meses = aplicar_filtros(Facturacion2026.query).with_entities(
        Facturacion2026.mes
    ).distinct().order_by(Facturacion2026.mes).all()
    return jsonify({
        'success': True,
        'meses': [m[0] for m in meses]
    })
