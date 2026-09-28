"""Exportación anual de las vistas de Proyectados con sus columnas y agrupaciones."""
from collections import defaultdict
from copy import copy
from inspect import unwrap
from io import BytesIO
import unicodedata

from flask import current_app
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NUMBER = '#,##0.00;[Red]-#,##0.00'
MONEY = '$ #,##0.00;[Red]-$ #,##0.00'
PERCENT = '0.00%'


def key(value):
    return ''.join(c for c in unicodedata.normalize('NFD', str(value or '').lower().strip()) if not unicodedata.combining(c))


def table(book, title, headers, rows, formats=None, total_rows=()):
    ws = book.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row == 1 or cell.row in total_rows:
                cell.fill = PatternFill('solid', fgColor='1E3A5F' if cell.row == 1 else 'DBEAFE')
                cell.font = Font(bold=True, color='FFFFFF' if cell.row == 1 else '172554')
            if cell.row > 1:
                cell.number_format = (formats or {}).get(cell.column, NUMBER if isinstance(cell.value, (int, float)) else 'General')
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 28 if i <= 3 else 19
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = '1:1'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def projection_rows(items, plp=False):
    if plp:
        groups = defaultdict(list)
        for item in items:
            if item.get('tipo_plp'):
                groups[(item['mes'], item['tipo_plp'])].append(item)
        for (_, name), group in sorted(groups.items()):
            first = group[0]
            for night in (False, True):
                fraction = lambda x: (x.get('porcentaje_nocturnidad', 0) or 0) / 100 if night else 1 - (x.get('porcentaje_nocturnidad', 0) or 0) / 100
                total = lambda field: sum((x.get(field) or 0) * fraction(x) for x in group)
                hours, required = total('horas_proyectadas'), total('horas_requeridas')
                if hours <= 0:
                    continue
                loads = {x.get('carga_semanal') or 'L a V' for x in group}
                shifts = {x.get('carga_horaria') or 0 for x in group}
                yield ['Personal', name + (' nocturnidad' if night else ''), first['year'], first['mes_label'],
                       total('dotacion_requerida'), next(iter(loads)) if len(loads) == 1 else 'Varias',
                       f'{next(iter(shifts)):g} h' if len(shifts) == 1 else 'Varias',
                       max(x.get('dias_objetivo') or 0 for x in group), required, hours,
                       hours / required if required else 0]
    else:
        for item in items:
            if item.get('tipo_plp'):
                continue
            night = (item.get('porcentaje_nocturnidad') or 0) / 100
            splits = [(False, 1 - night), (True, night)] if item.get('tiene_nocturnidad') and night > 0 else [(False, 1)]
            shifts = item.get('jornadas') or [item]
            for nocturnal, fraction in splits:
                yield [item['cliente'], item['campania'] + (' nocturnidad' if nocturnal else ''), item['year'], item['mes_label'],
                       (item.get('dotacion_requerida') or 0) * fraction,
                       '\n'.join(f"{x.get('dotacion_requerida') or 0:g} p / {x.get('carga_semanal') or 'L a V'}" for x in shifts),
                       '\n'.join(f"{x.get('carga_horaria') or 0:g} h" for x in shifts), item.get('dias_objetivo') or 0,
                       (item.get('horas_requeridas') or 0) * fraction, (item.get('horas_proyectadas') or 0) * fraction,
                       (item.get('porcentaje_cumplimiento') or 0) / 100]


def crear_exportacion_proyectados(year):
    from app import routes as r
    app = current_app._get_current_object()

    def data(fn):
        with app.test_request_context('/exportacion', query_string={'year': year}):
            response = unwrap(fn)()
            result = response.get_json()
            if response.status_code != 200 or not result.get('success'):
                raise ValueError(f'No se pudo obtener {fn.__name__}')
            return result

    def existing_export(fn):
        with app.test_request_context('/exportacion', query_string={'year': year}):
            response = unwrap(fn)()
            response.direct_passthrough = False
            try:
                return load_workbook(BytesIO(response.get_data()))
            finally:
                response.close()

    book = Workbook()
    book.remove(book.active)
    projections = data(r.api_matriz_proyecciones)
    months = projections['meses']
    keys = [m['value'] for m in months]
    labels = [m['label'] for m in months]
    headers = ['Cliente', 'Campaña', 'Año', 'Mes', 'Dotación requerida', 'Carga', 'Hs. jornada', 'Q dias obj.', 'Hs. requeridas', 'Horas proyectadas', '% cumplimiento']
    for name, plp in [('Proyecciones - Todas', False), ('Proyecciones - PLP', True)]:
        table(book, name, headers, projection_rows(projections['proyecciones'], plp), {3: '0', 11: PERCENT})
    plp = [x for x in projections['proyecciones'] if x.get('tipo_plp')]
    table(book, 'PLP - Carga', ['Clasificación PLP', 'Mes', 'Campaña', 'Horas', 'Carga semanal', 'Horas jornada', '% cumplimiento'],
          [[x['tipo_plp'], x['mes_label'], x['campania'], x['horas_requeridas'], x.get('carga_semanal') or 'L a V',
            x.get('carga_horaria') or 6, (x.get('porcentaje_cumplimiento') or 0) / 100] for x in plp], {7: PERCENT})
    summary = []
    for month in months:
        amounts = {kind: sum(x['horas_requeridas'] for x in plp if x['mes'] == month['value'] and x['tipo_plp'] == kind)
                   for kind in ('Personal CX', 'Personal', 'Personal Soporte', 'Personal SMB')}
        cx, personal, support, smb = amounts.values()
        total = cx + personal + support
        summary.append([month['label'], cx, personal, support, total, support / total if total else 0, smb, total + smb])
    totals = [sum(row[i] for row in summary) for i in range(1, 8)]
    totals[4] = totals[2] / totals[3] if totals[3] else 0
    summary.append(['TOTAL'] + totals)
    table(book, 'PLP - Resumen mensual', ['Mes', 'Personal CX', 'Personal', 'Personal Soporte', 'Total Personal', '% Soporte', 'Personal SMB', 'Total'], summary, {6: PERCENT}, (len(summary) + 1,))
    distribution = data(r.api_personal_distribucion)
    shares = {(x['mes'], x['servicio']): x['porcentaje_diurno'] for x in distribution['distribuciones']}
    table(book, 'PLP - Distribución', ['Mes'] + [s + ' (% diurno)' for s in distribution['servicios']],
          [[m['label']] + [shares.get((m['value'], s), 100) / 100 for s in distribution['servicios']] for m in months],
          {i + 2: PERCENT for i in range(len(distribution['servicios']))})

    prices = data(r.api_matriz_precios)
    pairs = {}
    for item in prices['asociaciones'] + prices['precios']:
        pair = (key(item['cliente']), key(item['campania']))
        if pair not in pairs:
            pairs[pair] = dict(item)
        elif not pairs[pair].get('site') or pairs[pair]['site'] == 'Sin site':
            pairs[pair]['site'] = item.get('site') or 'Sin site'
    lookup = {(key(x['cliente']), key(x['campania']), x['mes']): x for x in prices['precios']}
    table(book, 'Precios', ['Cuenta', 'Site', 'Cliente'] + labels,
          [[x['cliente'], x.get('site') or 'Sin site', x['campania']] +
           [lookup.get((*pair, month), {}).get('precio_final', 0) for month in keys]
           for pair, x in sorted(pairs.items())], {i: MONEY for i in range(4, 16)})

    # Reutiliza la exportación que ya reproduce los grupos y totales de RESUMEN.
    def append_sheets(source, names):
        for original, name in zip(source.worksheets, names):
            target = book.create_sheet(name)
            for row in original:
                for cell in row:
                    out = target.cell(cell.row, cell.column, cell.value)
                    if isinstance(cell.value, str):
                        out.data_type = 's'
                    if cell.has_style:
                        out.font = copy(cell.font)
                        out.fill = copy(cell.fill)
                        out.border = copy(cell.border)
                        out.alignment = copy(cell.alignment)
                        out.protection = copy(cell.protection)
                        out.number_format = cell.number_format
            for col, dim in original.column_dimensions.items():
                target.column_dimensions[col] = copy(dim)
            for number, dim in original.row_dimensions.items():
                target.row_dimensions[number] = copy(dim)
            for merged in original.merged_cells.ranges:
                target.merge_cells(str(merged))
            target.freeze_panes = original.freeze_panes
            target.auto_filter = copy(original.auto_filter)

    append_sheets(existing_export(r.exportar_resumen_proyeccion), ['Facturación horas'])
    variables = data(r.api_variables)['filas']
    for estimation in (False, True):
        rows = []
        for x in variables:
            values = [(x.get('meses' if estimation else 'porcentajes') or {}).get(m, 0) for m in keys]
            rows.append([x.get('site', ''), x['cliente'], x['campania'], x.get('total', 0) if estimation else sum(values) / len(keys) / 100] +
                        (values if estimation else [v / 100 for v in values]))
        if estimation:
            rows.insert(0, ['TOTAL', '', 'GENERAL'] + [sum(row[i] for row in rows) for i in range(3, 16)])
        table(book, 'Estimación variable' if estimation else '% Variable',
              ['Site', 'Cliente', 'Campaña', 'Total' if estimation else 'Promedio'] + labels,
              rows, {i: MONEY if estimation else PERCENT for i in range(4, 17)}, (2,) if estimation else ())
    tariffs = data(r.api_tarifaciones)['filas']
    table(book, 'Tarifación', ['Site', 'Cliente / Campaña', 'Concepto', 'Total tarifación'] + labels,
          [[x.get('site', ''), x['cliente'] + '\n' + x['campania'], x['concepto'], x['total_monto']] +
           [x['montos'].get(m, 0) for m in keys] for x in tariffs], {i: MONEY for i in range(4, 17)})
    ng = data(r.api_next_gen)
    def local_number(value):
        return f'{value:,.2f}'.replace(',', '_').replace('.', ',').replace('_', '.')
    dual = lambda pesos, usd: f'$ {local_number(pesos)}\nUSD {local_number(usd)}'
    ngrows = [['TOTAL NEXT GEN', '', '', '', dual(ng['total_pesos']['total'], ng['total_usd']['total'])] +
              [dual(ng['total_pesos']['meses'][m], ng['total_usd']['meses'][m]) for m in keys]]
    ngrows += [[x.get('site', ''), x['cliente'], x['campania'], x['producto'], x['total_pesos']] +
               [dual(x['pesos'][m], x['usd'][m]) for m in keys] for x in ng['filas']]
    table(book, 'Next Gen', ['Site', 'Cliente', 'Campaña destino', 'Producto', 'Total pesos'] +
          [f"{m['label']}\nUSD {ng['dolares'][m['value']]}" for m in months], ngrows, {5: MONEY}, (2,))
    fixed = data(r.api_suma_fija)['filas']
    table(book, 'Suma fija', ['Site', 'Cliente', 'Campaña'] + labels,
          [[x.get('site') or 'Sin site', x['cliente'], x['campania']] +
           [x['meses'].get(m, {}).get('monto', 0) for m in keys] for x in fixed],
          {i: MONEY for i in range(4, 16)})
    control = data(r.api_control_proyecciones)['filas']
    for field, title in [('horas', 'Control - Horas'), ('dotaciones', 'Control - Dotaciones')]:
        rows = [[x.get('site', ''), x['cliente'], x['campania'], x['total_' + field]] + [x[field].get(m, 0) for m in keys] for x in control]
        rows.insert(0, ['TOTAL', '', 'GENERAL ' + field.upper()] + [sum(row[i] for row in rows) for i in range(3, 16)])
        table(book, title, ['Site', 'Cliente', 'Campaña', 'Total'] + labels, rows, total_rows=(2,))
    append_sheets(existing_export(r.exportar_notificaciones_proyecciones), ['Notificaciones', 'Notificaciones - Detalle'])
    holidays = data(r.api_calendario_operativo)['feriados']
    table(book, 'Calendario operativo', ['Fecha', 'Nombre', 'Tipo', 'Estado'],
          [[x['fecha'], x['nombre'], x['tipo'], 'Activo' if x['activo'] else 'Inactivo'] for x in holidays])
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output
