import io
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app, db
from app.models import HistorialCambio, ProformaPersonal
from app.routes import deshacer_item_historial


archivo = Path(r'C:\Users\jegil\Downloads\Proforma Proyectada Cable 08-2026.xlsx')
archivo_personal_pay = Path(r'C:\Users\jegil\Downloads\Proforma Proyectada Personal Pay 08-2026.xlsx')
contenido = archivo.read_bytes()
app = create_app()


def contenido_para_negocio(bytes_archivo, negocio):
    libro = load_workbook(io.BytesIO(bytes_archivo))
    hoja = libro[libro.sheetnames[0]]
    for numero in range(4, hoja.max_row + 1):
        if hoja.cell(numero, 1).value not in (None, ''):
            hoja.cell(numero, 3).value = negocio
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def subir(client, tipo, bytes_archivo, nombre='proforma.xlsx'):
    return client.post(
        '/api/proforma-personal/importar',
        data={'tipo_proforma': tipo, 'archivo': (io.BytesIO(bytes_archivo), nombre)},
        headers={'X-CSRF-Token': 'verificacion-tipos'},
        content_type='multipart/form-data',
    )


with app.test_client() as client, app.app_context():
    with client.session_transaction() as sesion:
        sesion['usuario_id'] = 1
        sesion['csrf_token'] = 'verificacion-tipos'
    original_commit = db.session.commit
    db.session.commit = db.session.flush
    try:
        resultados = {}
        contenidos = {
            'Masivo': contenido,
            'Personal Pay': contenido_para_negocio(contenido, 'PAY'),
            'Soporte': contenido_para_negocio(contenido, 'SOPORTE'),
        }
        for tipo in ('Masivo', 'Personal Pay', 'Soporte'):
            respuesta = subir(client, tipo, contenidos[tipo], f'{tipo}.xlsx')
            assert respuesta.status_code == 200, respuesta.get_data(as_text=True)
            resultados[tipo] = len(respuesta.get_json()['proformas'])
        assert len(set(resultados.values())) == 1 and resultados['Masivo'] > 2
        periodo = subir(client, 'Masivo', contenido).get_json()['periodos'][0]

        libro = load_workbook(io.BytesIO(contenido))
        hoja = libro[libro.sheetnames[0]]
        filas_validas = [
            numero for numero in range(4, hoja.max_row + 1)
            if any(hoja.cell(numero, columna).value not in (None, '') for columna in range(8, 17))
        ]
        conservar = set(filas_validas[:2])
        for numero in range(4, hoja.max_row + 1):
            if numero not in conservar:
                for columna in range(1, 17):
                    hoja.cell(numero, columna).value = None
        reducido = io.BytesIO()
        libro.save(reducido)
        respuesta_reemplazo = subir(client, 'Masivo', reducido.getvalue(), 'Masivo reducido.xlsx')
        assert respuesta_reemplazo.status_code == 200, respuesta_reemplazo.get_data(as_text=True)
        assert respuesta_reemplazo.get_json()['eliminadas'] == resultados['Masivo'] - 2

        conteos = {
            tipo: ProformaPersonal.query.filter_by(tipo_proforma=tipo, periodo=periodo).count()
            for tipo in ('Masivo', 'Personal Pay', 'Soporte')
        }
        assert conteos == {'Masivo': 2, 'Personal Pay': resultados['Personal Pay'], 'Soporte': resultados['Soporte']}

        historial = HistorialCambio.query.filter_by(
            entidad='proforma_personal', accion='importacion',
        ).order_by(HistorialCambio.id.desc()).first()
        resultado_deshacer, estado = deshacer_item_historial(historial)
        assert estado == 200, resultado_deshacer
        assert ProformaPersonal.query.filter_by(tipo_proforma='Masivo', periodo=periodo).count() == resultados['Masivo']
        tipo_incorrecto = subir(client, 'Personal Pay', contenido, 'Cable incorrecto.xlsx')
        assert tipo_incorrecto.status_code == 400 and 'NEGOCIO' in tipo_incorrecto.get_data(as_text=True)
        personal_pay = subir(client, 'Personal Pay', archivo_personal_pay.read_bytes(), archivo_personal_pay.name)
        assert personal_pay.status_code == 200, personal_pay.get_data(as_text=True)
        assert personal_pay.get_json()['periodos'] == ['202608']
        assert len(personal_pay.get_json()['proformas']) == 6
        print({'cargas': resultados, 'tras_reemplazo': conteos, 'deshacer': estado})
    finally:
        db.session.rollback()
        db.session.commit = original_commit
