import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from app import create_app

esperados = (
    'FDV', 'PERIODO', 'NEGOCIO', 'SITIO_PROVEEDOR', 'SEGMENTO', 'SUBSITIO',
    'TIPO_HORA', 'TOTAL_HORAS', 'PRECIO', 'MONTO_FIJO', 'PORCENTAJE_BONO_KPI_VS',
    'MONTO_VARIABLE_KPI_VS', 'PORCENTAJE_BONO_AC', 'MONTO_VARIABLE_AC',
    'MONTO_VARIABLE', 'TOTAL_PROYECCION',
)
app = create_app()
with app.test_client() as client:
    with client.session_transaction() as session:
        session['usuario_id'] = 1
    response = client.get('/api/proforma-personal/template')
    assert response.status_code == 200
    libro = load_workbook(io.BytesIO(response.data), data_only=False)
    assert libro.sheetnames == ['PIC_PROF_PAGO_CAT']
    hoja = libro.active
    assert tuple(hoja.cell(3, columna).value for columna in range(1, 17)) == esperados
    assert str(hoja['P2'].value).upper() == '=SUM(P4:P1000)'
    assert hoja['K4'].number_format == '0.00'
    assert hoja['M4'].number_format == '0.00'
    page = client.get('/carga-datos-personal').get_data(as_text=True)
    assert '/api/proforma-personal/template' in page
    assert 'tabHorasPersonal' not in page
    print({'status': response.status_code, 'hoja': hoja.title, 'columnas': hoja.max_column})
