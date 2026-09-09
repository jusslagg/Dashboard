import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from app import create_app

app = create_app()
with app.test_client() as client:
    with client.session_transaction() as test_session:
        test_session['usuario_id'] = 1

    page = client.get('/notificaciones-proyecciones')
    assert page.status_code == 200
    assert b'Notificaciones de cambios' in page.data

    response = client.get('/api/notificaciones-proyecciones?year=2026')
    assert response.status_code == 200
    data = response.get_json()
    items = data['notificaciones']
    assert items
    assert all(item['mes'].startswith('2026-') for item in items)
    assert all(item['detalle'] for item in items)
    assert all(item['comparacion'] and item['lectura'] for item in items)
    assert all(item['servicios'] and item['cantidad_servicios'] == len(item['servicios']) for item in items)
    assert all(item['estado'] in ('positiva', 'negativa', 'mixta') for item in items)
    assert any('horas_proyectadas' in item['metricas_cambiadas'] for item in items)
    assert any('precio_final' in item['metricas_cambiadas'] for item in items)
    assert any('total_proyectado' in item['metricas_cambiadas'] for item in items)

    filtered = client.get('/api/notificaciones-proyecciones?year=2026&estado=negativa&metrica=horas_proyectadas').get_json()['notificaciones']
    assert all(item['estado'] == 'negativa' for item in filtered)
    assert all('horas_proyectadas' in item['metricas_cambiadas'] for item in filtered)
    personal_septiembre = [item for item in items if item['mes'] == '2026-09' and item['cliente'] == 'Personal']
    assert len(personal_septiembre) == 1 and personal_septiembre[0]['cantidad_servicios'] > 1

    export = client.get('/api/notificaciones-proyecciones/exportar?year=2026&estado=negativa')
    assert export.status_code == 200
    assert export.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    workbook = load_workbook(io.BytesIO(export.data), read_only=True)
    sheet = workbook['REPORTE']
    detail_sheet = workbook['DETALLE']
    assert sheet.max_row == len([item for item in items if item['estado'] == 'negativa']) + 1
    assert sheet.cell(1, 11).value == 'Detalle completo'
    assert detail_sheet.max_row > sheet.max_row

    print(json.dumps({
        'total': len(items),
        'positivas': sum(item['estado'] == 'positiva' for item in items),
        'negativas': sum(item['estado'] == 'negativa' for item in items),
        'mixtas': sum(item['estado'] == 'mixta' for item in items),
        'filtradas_negativas_horas': len(filtered),
        'filas_excel_negativas': sheet.max_row - 1,
        'filas_detalle_excel': detail_sheet.max_row - 1,
    }, ensure_ascii=False))
